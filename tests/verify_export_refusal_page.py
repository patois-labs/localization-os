"""Issue #593: a browser export that a delivery check refuses gets a PAGE, not raw JSON.

The maintainer's AWS rehearsal (2026-10-07): clicking "Export txt/docx" on a held document showed
`{"ok":false,"gate_held":true,...}` in the browser. The checks are right; the presentation
was wrong. The web route `/documents/{id}/export` now answers a request that names
`Accept: text/html` with a styled 409 page (what is wrong, why, the first N affected
segments or paragraphs, and the action buttons). Every other client keeps the exact JSON:
`/api/documents/{id}/download`, the CLI, curl/requests (Accept `*/*`), fetch(application/json).

Pinned (anchored on element ids, never on copy):
  1. quality-gate hold (NOT bypassable): HTML 409, #export-act-review links the real editor,
     the affected segments are listed, NO #export-act-anyway; the JSON twin is unchanged.
  2. FULL_TEXT inline-tag quarantine ("tokens"): HTML 409 with #export-act-paragraphs and
     #export-act-anyway (allow_source_fallback=1); JSON twin unchanged.
  3. a quarantine kind with no way out ("count"): no anyway button.
  4. classic held locale: HTML offers the allow_held way out; allow_held=1 still exports.
  5. /api download: still JSON for the same states, even if the caller sends text/html.
  6. role rule: a reader sees the refusal but no override, and the override URL is refused for them.
  7. project single-file export from a browser lands on the document refusal page.
  8. more than N affected items: the page says how many are shown.

Hermetic: throwaway LOC_TMS_DATA_DIR. ASCII-only. Final line: "EXPORT REFUSAL PAGE OK".
"""

import os
import sys

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_exprefuse_")

import re
from io import BytesIO

from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app.db import engine, init_db
from app.main import app
from app.models import Project, Role, Segment, SegStatus, TargetLocale, TransMode, User
from app.services import apikeys, quality
from app.services.ingest import ingest_document
from tests._csrf import csrf_for

init_db()

LOC = "es-419"
HTML = {"Accept": "text/html,application/xhtml+xml,*/*;q=0.8"}
_SEQ = [0]

P0_RUNS = [("The ", "bold"), ("Acme Pro", "italic"), (" connector ships today.", "underline")]
BODY = [
    "Our platform helps studios launch their games worldwide.",
    "Localization is handled by a team of native speaking experts.",
]
T0_STRIPPED = "El conector Acme Pro se lanza hoy."
TBODY = [
    "Nuestra plataforma ayuda a los estudios a lanzar sus juegos.",
    "La localizacion la maneja un equipo de hablantes nativos.",
]


def _fail(msg):
    print("FAIL:", msg)
    sys.exit(1)


def _check(cond, msg):
    if not cond:
        _fail(msg)


def _seq():
    _SEQ[0] += 1
    return _SEQ[0]


def _project(source="en-US"):
    with Session(engine) as s:
        p = Project(
            name="exprefuse-%d" % _seq(),  # noqa: UP031
            org_id=1,
            source_locale=source,
            default_provider="mock",
        )
        s.add(p)
        s.commit()
        s.refresh(p)
        s.add(TargetLocale(project_id=p.id, locale=LOC))
        s.commit()
        return p.id


def _txt_doc(n_segments=3, status=SegStatus.NEEDS_REVIEW, gate_flag=True, target="traducido"):
    """A SEGMENTED doc with `n_segments` rows; each carries a gate-hold flag when asked."""
    pid = _project()
    text = " ".join("Sentence number %d is here." % i for i in range(n_segments)).encode()  # noqa: UP031
    with Session(engine) as s:
        p = s.get(Project, pid)
        doc = ingest_document(s, p, "plain.txt", text, TransMode.SEGMENTED)
        s.commit()
        s.refresh(doc)
        rows = s.exec(
            select(Segment).where(Segment.document_id == doc.id, Segment.target_locale == LOC)
        ).all()
        for sg in rows:
            sg.target_text = target
            sg.status = status
            if gate_flag:
                sg.qa_flags = [
                    {"type": quality.GATE_HOLD_FLAG_TYPE, "step": "Initial check", "cause": ""}
                ]
            s.add(sg)
        s.commit()
        return pid, doc.id, len(rows)


def _docx(p0_runs, plain):
    d = Docx()
    p = d.add_paragraph()
    for text, fmt in p0_runs:
        r = p.add_run(text)
        setattr(r, fmt, True)
    for t in plain:
        d.add_paragraph(t)
    b = BytesIO()
    d.save(b)
    return b.getvalue()


def _fulltext_doc(target_paras):
    pid = _project()
    with Session(engine) as s:
        p = s.get(Project, pid)
        doc = ingest_document(s, p, "Marketing.docx", _docx(P0_RUNS, BODY), TransMode.FULL_TEXT)
        s.commit()
        s.refresh(doc)
        seg = s.exec(
            select(Segment).where(Segment.document_id == doc.id, Segment.target_locale == LOC)
        ).one()
        seg.target_text = "\n\n".join(target_paras)
        seg.status = SegStatus.MT
        s.add(seg)
        s.commit()
        return pid, doc.id


def _api_key():
    with Session(engine) as s:
        u = s.exec(select(User).where(User.email == "refuse-admin@x.io")).first()
        if not u:
            u = User(email="refuse-admin@x.io", name="A", role=Role.ADMIN, is_active=True)
            s.add(u)
            s.commit()
            s.refresh(u)
        uid = u.id
    _row, raw = apikeys.create_key(uid, "refuse-api", scopes="api docs")
    return raw


client = TestClient(app)
KEY = _api_key()
API = {"X-Loc-Token": KEY}


def _ids(html):
    return set(re.findall(r'id="([^"]+)"', html))


def _href(html, ident):
    m = re.search(r'<a[^>]*id="%s"[^>]*href="([^"]*)"' % re.escape(ident), html)  # noqa: UP031
    if not m:
        m = re.search(r'<a[^>]*href="([^"]*)"[^>]*id="%s"' % re.escape(ident), html)  # noqa: UP031
    return m.group(1).replace("&amp;", "&") if m else None


def test_gate_hold():
    _pid, did, n = _txt_doc(3)
    url = "/documents/%d/export?locale=%s" % (did, LOC)  # noqa: UP031
    # JSON twin: unchanged shape.
    j = client.get(url)
    _check(j.status_code == 409, j.status_code)
    _check(j.headers["content-type"].startswith("application/json"), j.headers["content-type"])
    _check(j.json().get("gate_held") is True, j.json())
    # Browser: a page.
    r = client.get(url, headers=HTML)
    _check(r.status_code == 409, r.status_code)
    _check(r.headers["content-type"].startswith("text/html"), r.headers["content-type"])
    ids = _ids(r.text)
    _check({"export-refusal", "export-refusal-why", "export-act-review"} <= ids, sorted(ids))
    _check("export-act-anyway" not in ids, "a non-bypassable gate must never offer Export anyway")
    _check("export-act-paragraphs" not in ids, "no paragraphs button on a gate hold")
    review = _href(r.text, "export-act-review")
    _check(review == "/documents/%d/editor?locale=%s" % (did, LOC), review)  # noqa: UP031
    _check(client.get(review).status_code == 200, "the Review link must open the real editor")
    _check(len(re.findall(r'id="export-refusal-item-\d+"', r.text)) == n, "all segments listed")
    # The raw message stays available for support.
    _check("export-refusal-raw" in ids, "the exact check message must be on the page")
    # The JSON twin's review_url (API clients and MCP follow it) must be a LIVE page: the
    # post editor, not the old /documents/N/segments that no web route serves.
    api = client.get("/api/documents/%d/download?locale=%s" % (did, LOC), headers=API).json()  # noqa: UP031
    _check(api["review_url"] == review == j.json()["review_url"], (api["review_url"], review))
    _check(client.get(api["review_url"]).status_code == 200, "API review_url must resolve")
    # An /api caller that happens to send text/html still gets JSON.
    a = client.get("/api/documents/%d/download?locale=%s" % (did, LOC), headers={**API, **HTML})  # noqa: UP031
    _check(
        a.status_code == 409 and a.json().get("gate_held") is True, (a.status_code, a.text[:200])
    )
    print(
        "  1. gate hold: HTML 409 + editor Review link + segments, NO anyway; JSON/API unchanged OK"
    )


def test_more_than_n():
    _pid, did, n = _txt_doc(14)
    r = client.get("/documents/%d/export?locale=%s" % (did, LOC), headers=HTML)  # noqa: UP031
    _check(r.status_code == 409, r.status_code)
    _check(len(re.findall(r'id="export-refusal-item-\d+"', r.text)) == 10, "first 10 only")
    m = re.search(r'id="export-refusal-items-more"[^>]*>\s*([^<]*)<', r.text)
    _check(m and "10" in m.group(1) and str(n) in m.group(1), m and m.group(1))
    print("  8. 14 held segments: first 10 listed with an 'of 14' note OK")


def test_quarantine_tokens():
    _pid, did = _fulltext_doc([T0_STRIPPED] + TBODY)  # noqa: RUF005
    url = "/documents/%d/export?locale=%s" % (did, LOC)  # noqa: UP031
    j = client.get(url)
    _check(j.status_code == 409 and j.json().get("quarantined") is True, j.text[:200])
    _check("export_anyway_url" in j.json() and "paragraphs_url" in j.json(), j.json())
    r = client.get(url, headers=HTML)
    _check(r.status_code == 409, r.status_code)
    _check(r.headers["content-type"].startswith("text/html"), r.headers["content-type"])
    ids = _ids(r.text)
    _check({"export-act-paragraphs", "export-act-anyway"} <= ids, sorted(ids))
    _check(_href(r.text, "export-act-paragraphs") == j.json()["paragraphs_url"], "paragraphs link")
    _check(_href(r.text, "export-act-anyway") == j.json()["export_anyway_url"], "anyway link")
    _check("export-refusal-item-0" in r.text, "the damaged paragraph is listed")
    pg = client.get(_href(r.text, "export-act-paragraphs"))
    _check(pg.status_code == 200, "View paragraphs must open")
    ok = client.get(_href(r.text, "export-act-anyway"), headers=HTML)
    _check(ok.status_code == 200 and ok.headers.get("x-export-substituted") == "1", ok.status_code)
    a = client.get("/api/documents/%d/download?locale=%s" % (did, LOC), headers={**API, **HTML})  # noqa: UP031
    _check(a.status_code == 409 and a.json().get("quarantined") is True, a.text[:200])
    print(
        "  2. tokens quarantine: HTML 409, paragraphs + anyway links match the API; anyway works OK"
    )


def test_quarantine_count_has_no_way_out():
    _pid, did = _fulltext_doc([T0_STRIPPED, *TBODY[:1], "extra one", "extra two"])
    r = client.get("/documents/%d/export?locale=%s" % (did, LOC), headers=HTML)  # noqa: UP031
    _check(r.status_code == 409, r.status_code)
    ids = _ids(r.text)
    _check("export-act-anyway" not in ids, "count desync offers no override")
    _check("export-nobypass-note" in ids, "says it cannot be bypassed")
    print("  3. count/untranslated quarantine: no Export anyway button OK")


def test_classic_held():
    pid = _project()
    with Session(engine) as s:
        doc = ingest_document(
            s, s.get(Project, pid), "h.txt", b"Hello world. This is a test.", TransMode.SEGMENTED
        )
        s.commit()
        s.refresh(doc)
        did = doc.id
    url = "/documents/%d/export?locale=%s" % (did, LOC)  # noqa: UP031
    j = client.get(url)
    _check(j.status_code == 409 and j.json().get("held") is True, j.text[:200])
    r = client.get(url, headers=HTML)
    _check(r.status_code == 409, r.status_code)
    _check("export-act-anyway" in _ids(r.text), "classic hold is bypassable: button offered")
    _check(_href(r.text, "export-act-anyway") == j.json()["export_anyway_url"], "anyway link")
    _check(client.get(_href(r.text, "export-act-anyway"), headers=HTML).status_code == 200, "ships")
    print("  4. classic held locale: Export anyway offered, still ships with allow_held OK")


def test_unknown_locale_and_ok_unchanged():
    _pid, did, _n = _txt_doc(1, status=SegStatus.APPROVED, gate_flag=False)
    ok = client.get("/documents/%d/export?locale=%s" % (did, LOC), headers=HTML)  # noqa: UP031
    _check(
        ok.status_code == 200 and "attachment" in ok.headers.get("content-disposition", ""),
        ok.status_code,
    )
    nf = client.get("/documents/%d/export?locale=xx-XX" % did, headers=HTML)  # noqa: UP031
    _check(
        nf.status_code == 404 and nf.headers["content-type"].startswith("text/html"), nf.status_code
    )
    nj = client.get("/documents/%d/export?locale=xx-XX" % did)  # noqa: UP031
    _check(nj.status_code == 404 and nj.json().get("ok") is False, nj.text[:100])
    print(
        "  5. a deliverable export still downloads; unknown locale: HTML for browsers, JSON else OK"
    )


def test_project_export_redirects_browser():
    pid, did, _n = _txt_doc(2)
    url = "/projects/%d/export?locales=%s" % (pid, LOC)  # noqa: UP031
    r = client.get(url, headers=HTML, follow_redirects=False)
    _check(r.status_code == 303, r.status_code)
    _check(
        r.headers["location"] == f"/documents/{did}/export?locale={LOC}",
        r.headers["location"],
    )
    page = client.get(r.headers["location"], headers=HTML)
    _check(page.status_code == 409 and "export-act-review" in _ids(page.text), page.status_code)
    plain = client.get(url, follow_redirects=False)
    _check(
        plain.status_code == 303 and "err=export_held" in plain.headers["location"], plain.headers
    )
    print(
        "  7. project single-file export: browser lands on the document page;"
        " others keep the flash OK"
    )


def test_reader_cannot_override():
    """Auth ON: a viewer sees the refusal with no override control, and typing the override
    URL by hand does not ship the file either (the flags are ignored for a read-only role);
    an admin still gets the file."""
    from app.services import security, settings

    FIXTURE_PW = "fixture-pass-1234!!"  # pregate: allow-secret test fixture
    _pid, did = _fulltext_doc([T0_STRIPPED] + TBODY)  # noqa: RUF005

    def _owned_by_viewer(pid):
        # project scoping: a viewer sees only projects it created or is a member of
        with Session(engine) as s:
            pr = s.get(Project, pid)
            pr.created_by = "viewer593@x.io"
            s.add(pr)
            s.commit()

    _owned_by_viewer(_pid)
    security.create_user("viewer593@x.io", "Vi", Role.VIEWER, FIXTURE_PW)
    security.create_user("admin593@x.io", "Ad", Role.ADMIN, FIXTURE_PW)
    settings.set_many({"security_auth_required": "1"})
    settings.invalidate_cache()
    try:

        def _session(who):
            c = TestClient(app)
            csrf_for(c)
            r = c.post(
                "/auth/login",
                data={"email": who, "password": FIXTURE_PW, "next": "/"},
                follow_redirects=False,
            )
            _check(r.status_code == 303, (who, r.status_code))
            csrf_for(c)
            return c

        base = "/documents/%d/export?locale=%s" % (did, LOC)  # noqa: UP031
        dlg = {"X-Export-Dialog": "1", "Accept": "application/json"}
        v = _session("viewer593@x.io")
        _r0 = v.get(base, headers=dlg)
        _check(_r0.status_code == 409, ("viewer dialog refusal", _r0.status_code))
        ref = _r0.json()["refusal"]
        _check(not ref["export_anyway_url"] and ref["can_override"] is False, ref)
        page = v.get(base, headers=HTML)
        _check(page.status_code == 409 and "export-act-anyway" not in _ids(page.text), "no button")
        _check("export-anyway-note" in _ids(page.text), "viewer is told why")
        forced = v.get(base + "&allow_source_fallback=1", headers=HTML)
        _check(forced.status_code == 409, "viewer typing the override URL must still be refused")
        a = _session("admin593@x.io")
        ok = a.get(base + "&allow_source_fallback=1", headers=HTML)
        _check(
            ok.status_code == 200 and ok.headers.get("x-export-substituted") == "1", ok.status_code
        )
        aref = a.get(base, headers=dlg).json()["refusal"]
        _check(aref["export_anyway_url"] and aref["can_override"] is True, aref)
        # The classic held locale override obeys the same rule.
        hp = _project()
        _owned_by_viewer(hp)
        with Session(engine) as s:
            hd = ingest_document(
                s, s.get(Project, hp), "h2.txt", b"Hello world. A test.", TransMode.SEGMENTED
            )
            s.commit()
            s.refresh(hd)
            hid = hd.id
        hurl = "/documents/%d/export?locale=%s&allow_held=1" % (hid, LOC)  # noqa: UP031
        _check(v.get(hurl, headers=HTML).status_code == 409, "viewer allow_held refused")
        _check(a.get(hurl, headers=HTML).status_code == 200, "admin allow_held ships")

        # Gate (#593): the /api twins (what the remote CLI and MCP call) apply the same
        # rule. A viewer's docs key asking for either override is a 403, never the file;
        # without the flag it still gets the ordinary refusal; an admin key ships.
        def _key_for(email):
            with Session(engine) as s:
                uid = s.exec(select(User).where(User.email == email)).first().id
            return {"X-Loc-Token": apikeys.create_key(uid, "k593", scopes="api docs")[1]}

        vk, ak = _key_for("viewer593@x.io"), _key_for("admin593@x.io")
        dl = "/api/documents/%d/download?locale=%s" % (did, LOC)  # noqa: UP031
        r1 = client.get(dl + "&allow_source_fallback=1", headers=vk)
        _check(
            r1.status_code == 403 and r1.json().get("ok") is False,
            ("viewer /api fallback", r1.status_code),
        )
        _check(client.get(dl, headers=vk).status_code == 409, "viewer without a flag: refusal")
        r2 = client.get(dl + "&allow_source_fallback=1", headers=ak)
        _check(
            r2.status_code == 200 and r2.headers.get("x-export-substituted") == "1", r2.status_code
        )
        hdl = "/api/documents/%d/download?locale=%s&allow_held=1" % (hid, LOC)  # noqa: UP031
        r3 = client.get(hdl, headers=vk)
        _check(r3.status_code == 403, ("viewer /api allow_held", r3.status_code))
        _check(client.get(hdl, headers=ak).status_code == 200, "admin /api allow_held ships")
        pex = "/api/projects/%d/export?locales=%s&allow_held=1" % (hp, LOC)  # noqa: UP031
        r4 = client.get(pex, headers=vk)
        _check(r4.status_code == 403, ("viewer /api project allow_held", r4.status_code))
        _check(client.get(pex, headers=ak).status_code == 200, "admin /api project allow_held")
    finally:
        settings.set_many({"security_auth_required": "0"})
        settings.invalidate_cache()
    print(
        "  6. read-only role: no override button, hand-typed override URL refused on web and"
        " /api (403); admin ships OK"
    )


def test_dialog_payload_and_wiring():
    """The in-page dialog asks with X-Export-Dialog: 1 and gets the refusal as data."""
    DLG = {"X-Export-Dialog": "1", "Accept": "application/json"}
    _pid, did, _n = _txt_doc(3)
    url = "/documents/%d/export?locale=%s" % (did, LOC)  # noqa: UP031
    r = client.get(url, headers=DLG)
    _check(r.status_code == 409 and r.json().get("ok") is False, r.text[:200])
    ref = r.json()["refusal"]
    _check(ref["review_url"] == "/documents/%d/editor?locale=%s" % (did, LOC), ref)  # noqa: UP031
    _check(not ref["export_anyway_url"] and ref["not_bypassable"], "gate hold: never anyway")
    _check(len(ref["items"]) == 3 and ref["steps"], ref)
    # The plain JSON (no dialog header) is the exact legacy body.
    _check(client.get(url).json().get("gate_held") is True, "legacy JSON unchanged")
    _pid2, did2 = _fulltext_doc([T0_STRIPPED] + TBODY)  # noqa: RUF005
    q = client.get("/documents/%d/export?locale=%s" % (did2, LOC), headers=DLG).json()["refusal"]  # noqa: UP031
    _check(q["paragraphs_url"] and "allow_source_fallback=1" in q["export_anyway_url"], q)
    # A deliverable export is still the file, even with the dialog header.
    _p3, did3, _n3 = _txt_doc(1, status=SegStatus.APPROVED, gate_flag=False)
    ok = client.get("/documents/%d/export?locale=%s" % (did3, LOC), headers=DLG)  # noqa: UP031
    _check(
        ok.status_code == 200 and "attachment" in ok.headers["content-disposition"], ok.status_code
    )
    # Wiring: the document page marks its export link, loads the script, renames the button.
    page = client.get("/documents/%d?locale=%s" % (did3, LOC)).text  # noqa: UP031
    _check('id="export-doc-link"' in page and "data-export-dialog" in page, "link marked")
    _check("export_dialog.js" in page, "script loaded")
    _check("post-editing editor" not in page, "old button label gone")
    _check(
        _href(page, "export-doc-link") is not None
        or 'href="/documents/%d/editor?locale=%s"' % (did3, LOC) in page,  # noqa: UP031
        "editor link present",
    )
    with open(os.path.join(_ROOT, "app", "static", "export_dialog.js"), encoding="utf-8") as fh:
        js = fh.read()
    _check("showModal" in js and "innerHTML" not in js, "native modal, no innerHTML")
    _check("export-page-live" in js and "aria-live" in js, "announces progress and result")
    _check("btn-secondary" in js, "Export anyway is the quiet button")
    # Issue #564: one Close control in the refusal dialog (footer), not a second in the header.
    _check(js.count("'Close'") == 1 and "export-dialog-close-x" not in js, "single Close button")
    # A project single-file export asked by the dialog follows to the document refusal.
    proj = client.get("/projects/%d/export?locales=%s" % (_pid, LOC), headers=DLG)  # noqa: UP031
    _check(
        proj.status_code == 409 and "refusal" in proj.json(), (proj.status_code, proj.text[:100])
    )
    print("  9. dialog payload: editor link, no anyway on the gate, anyway on tokens; wiring OK")


def main():
    test_dialog_payload_and_wiring()
    test_gate_hold()
    test_quarantine_tokens()
    test_quarantine_count_has_no_way_out()
    test_classic_held()
    test_unknown_locale_and_ok_unchanged()
    test_reader_cannot_override()
    test_project_export_redirects_browser()
    test_more_than_n()
    print("EXPORT REFUSAL PAGE OK")


if __name__ == "__main__":
    main()
