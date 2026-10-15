"""Chunk 5 verification — multi-agent workflow builder + orchestrator, offline.

- prompt-authoring assistant drafts a prompt with placeholders
- build a workflow via HTTP; steps persist
- run it: each step's output feeds the next (chaining), run + step traces saved
- reuse the same global workflow on a second document
- human-in-the-loop: a human_review step PAUSES; resume runs remaining steps
- auto-chunking: a large document is split into >1 chunk for a step
- failover: a step whose primary provider errors falls over to the backup
- the SSE /runs/stream endpoint emits run_started + step + done events
"""

import os
import sys
from io import BytesIO

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_test_")

from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app.cli.__main__ import main as cli_main
from app.db import engine, init_db
from app.main import app
from app.models import StepRun, StepType, Workflow, WorkflowRun, WorkflowStep
from app.services import engines as engines_svc
from app.services import orchestrator
from app.services.providers import (
    DeadEngineError,
    Provider,
    ProviderTruncated,
    register_provider,
)


class Failer(Provider):
    name = "failer"

    def complete(self, prompt, system=None, task=None, payload=None, **kw):
        raise RuntimeError("down")


class Truncator(Provider):
    """Mimics a local/keyed engine (e.g. LM Studio) with a tight output-token
    window: the model hits max_tokens before finishing, so complete() raises
    ProviderTruncated -- exactly like LocalEngine.complete() does on a real
    finish_reason=length response."""

    name = "truncator"

    def complete(self, prompt, system=None, task=None, payload=None, **kw):
        raise ProviderTruncated("local finish_reason=length - output truncated")


def docx_from(paragraphs) -> bytes:
    d = Docx()
    for p in paragraphs:
        d.add_paragraph(p)
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def new_project(client, name):
    r = client.post(
        "/projects",
        data={
            "name": name,
            "source_locale": "en-US",
            "target_locales": "fr-FR",
            "default_provider": "claude",
        },
        follow_redirects=False,
    )
    return int(r.headers["location"].rsplit("/", 1)[-1])


def upload(client, pid, paragraphs, mode="segmented"):
    r = client.post(
        f"/projects/{pid}/documents",
        data={"mode": mode},
        files={
            "file": (
                "d.docx",
                docx_from(paragraphs),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    return int(r.headers["location"].rsplit("/", 1)[-1])


def _drain_stream(client, run_id):
    """Read the read-only SSE observer to completion. observe_run() blocks the
    stream until the run reaches a terminal/paused state, so the TestClient
    returns the full event log in one GET."""
    return client.get(f"/runs/stream?run_id={run_id}").text


def add_step(
    session,
    wf_id,
    order,
    name,
    step_type=StepType.LLM_AGENT,
    provider=None,
    fallback_provider=None,
    prompt_text="",
    chunk_policy="auto",
    backoff_base=0.0,
):
    session.add(
        WorkflowStep(
            workflow_id=wf_id,
            order_index=order,
            name=name,
            step_type=step_type,
            provider=provider,
            fallback_provider=fallback_provider,
            prompt_text=prompt_text,
            chunk_policy=chunk_policy,
            retry_max=1,
            backoff_base=backoff_base,
        )
    )
    session.commit()


def main() -> None:
    init_db()
    client = TestClient(app)

    # --- prompt-authoring assistant ---
    draft = orchestrator.author_prompt("a strict technical reviewer")
    assert "{text}" in draft and "{target_locale}" in draft, draft
    print("  author_prompt drafted a prompt with placeholders")

    # --- HARDENING (review finding 2): author_prompt with a dead engine:<id> ref ---
    # The builder's provider dropdown offers engine:<id> refs; a dead one (deleted/
    # disabled/keyless/undecryptable) must not 500 the web route or crash the CLI —
    # author_prompt raises DeadEngineError (unchanged, NOT swallowed into a draft
    # string), and both callers surface it as a clean, non-500/non-traceback failure.
    dead_prof = engines_svc.create_profile("AuthorDeadEngine", "deepl", api_key="k1")
    engines_svc.update_profile(dead_prof.id, enabled=False)
    dead_ref = f"engine:{dead_prof.id}"
    try:
        orchestrator.author_prompt("a translator", dead_ref)
        raise AssertionError("author_prompt with a dead engine:<id> must raise DeadEngineError")
    except DeadEngineError as e:
        assert dead_ref in str(e), str(e)

    wf_for_author = Workflow(name="Author dead-engine wf", scope="global")
    with Session(engine) as s:
        s.add(wf_for_author)
        s.commit()
        s.refresh(wf_for_author)
        wf_for_author_id = wf_for_author.id
    r = client.post(
        f"/workflows/{wf_for_author_id}/author",
        data={"description": "a translator", "provider": dead_ref},
    )
    assert r.status_code == 400, (
        "web author route must return a clean error, not 500",
        r.status_code,
    )
    assert dead_ref in r.text, "web error page must name the dead ref"
    print(
        "  finding-2 web: POST /workflows/{id}/author with a dead engine:<id> -> clean 400, not 500"
    )

    import io as _io
    from contextlib import redirect_stderr as _redirect_stderr

    errbuf = _io.StringIO()
    with _redirect_stderr(errbuf):
        rc = cli_main(
            [
                "workflow",
                "author-prompt",
                str(wf_for_author_id),
                "a translator",
                "--provider",
                dead_ref,
            ]
        )
    assert rc != 0, ("CLI author-prompt with a dead engine:<id> must exit non-zero", rc)
    assert dead_ref in errbuf.getvalue(), ("CLI stderr must name the dead ref", errbuf.getvalue())
    print(
        "  finding-2 CLI: `workflow author-prompt --provider <dead ref>` exits non-zero, "
        "clean stderr message (no traceback)"
    )

    # --- REGRESSION (live-use bug): author_prompt with a provider that raises
    # ProviderTruncated (local/keyed engine with a tight output-token window
    # hitting finish_reason=length) must NOT surface as a raw 500 / traceback.
    # orchestrator.author_prompt() calls prov.complete() directly (bypassing
    # run_completion's try/except), so this exception reaches the two callers
    # unmodified -- both must catch it and respond cleanly, mirroring the
    # DeadEngineError handling right above.
    register_provider("truncator", Truncator())
    try:
        orchestrator.author_prompt("a translator", "truncator")
        raise AssertionError(
            "author_prompt with a truncating provider must raise ProviderTruncated"
        )
    except ProviderTruncated:
        pass

    r = client.post(
        f"/workflows/{wf_for_author_id}/author",
        data={"description": "a translator", "provider": "truncator"},
    )
    assert r.status_code == 400, (
        "web author route must return a clean 400 on truncation, not 500",
        r.status_code,
    )
    assert "ran out of output space" in r.text, "web error page must explain the truncation"
    print(
        "  regression web: POST /workflows/{id}/author with a truncating provider "
        "-> clean 400, not 500"
    )

    errbuf2 = _io.StringIO()
    with _redirect_stderr(errbuf2):
        rc2 = cli_main(
            [
                "workflow",
                "author-prompt",
                str(wf_for_author_id),
                "a translator",
                "--provider",
                "truncator",
            ]
        )
    assert rc2 != 0, ("CLI author-prompt with a truncating provider must exit non-zero", rc2)
    assert "ran out of output space" in errbuf2.getvalue(), (
        "CLI stderr must explain the truncation",
        errbuf2.getvalue(),
    )
    print(
        "  regression CLI: `workflow author-prompt --provider truncator` exits non-zero, "
        "clean stderr message (no traceback)"
    )

    # --- build a workflow via HTTP ---
    r = client.post(
        "/workflows", data={"name": "EN-FR 2-step", "scope": "global"}, follow_redirects=False
    )
    wf_id = int(r.headers["location"].rsplit("/", 1)[-1])
    client.post(
        f"/workflows/{wf_id}/steps",
        data={
            "name": "Translator",
            "step_type": "llm_agent",
            "provider": "claude",
            "chunk_policy": "auto",
            "prompt_text": "Translate to {target_locale}: {text}",
        },
    )
    client.post(
        f"/workflows/{wf_id}/steps",
        data={
            "name": "Reviewer",
            "step_type": "llm_agent",
            "provider": "claude",
            "chunk_policy": "auto",
            "prompt_text": "Review and improve: {text}",
        },
    )
    page = client.get(f"/workflows/{wf_id}").text
    assert "Translator" in page and "Reviewer" in page, "steps not shown in builder"
    with Session(engine) as s:
        steps = orchestrator.get_steps(s, wf_id)
    assert len(steps) == 2, len(steps)
    print(f"  built workflow with {len(steps)} steps via HTTP")

    # --- run it: chaining + traces ---
    pid = new_project(client, "WF Test")
    doc1 = upload(client, pid, ["Adventure awaits.", "Save big this weekend."])
    with Session(engine) as s:
        summary = orchestrator.run_workflow(s, wf_id, doc1, "fr-FR")
    assert summary["status"] == "completed", summary
    assert len(summary["steps"]) == 2, summary
    # two agent passes => the mock prefix appears twice (chaining proof)
    assert summary["text"].count("[fr-FR]") >= 2, summary["text"][:120]
    with Session(engine) as s:
        runs = s.exec(select(StepRun)).all()
    assert len([sr for sr in runs if sr.output_text]) >= 2
    print(f"  ran workflow: {len(summary['steps'])} steps chained; final has nested prefixes")

    # --- reuse the same global workflow on a second document ---
    doc2 = upload(client, pid, ["A different sentence entirely."])
    with Session(engine) as s:
        s2 = orchestrator.run_workflow(s, wf_id, doc2, "fr-FR")
    assert s2["status"] == "completed", s2
    print("  reused the same global workflow on a second document")

    # --- human-in-the-loop pause/resume ---
    with Session(engine) as s:
        wf2 = Workflow(name="With human review", scope="global")
        s.add(wf2)
        s.commit()
        s.refresh(wf2)
        wf2_id = wf2.id
        add_step(s, wf2_id, 0, "Translator", prompt_text="Translate: {text}")
        add_step(s, wf2_id, 1, "Human review", step_type=StepType.HUMAN_REVIEW)
        add_step(s, wf2_id, 2, "Finalizer", prompt_text="Finalize: {text}")
    doc3 = upload(client, pid, ["Pause here please."])
    with Session(engine) as s:
        paused = orchestrator.run_workflow(s, wf2_id, doc3, "fr-FR")
    assert paused["status"] == "paused" and paused["paused_step"] == 1, paused
    run_id = paused["run_id"]
    with Session(engine) as s:
        resumed = orchestrator.resume_workflow(s, run_id, "fr-FR", "HUMAN EDITED TEXT")
    assert resumed["status"] == "completed", resumed
    assert "HUMAN EDITED TEXT" in resumed["text"], resumed["text"]
    print(
        f"  human-in-the-loop: paused at step {paused['paused_step']}, resumed to completion with edited text"  # noqa: E501
    )

    # --- Finding 3 regression: observer must surface the CURRENT pause, not a
    #     stale resolved one. A workflow with TWO human_review steps leaves the
    #     first step's StepRun at status=PAUSED forever after resume; a fresh
    #     observer reconnecting while paused at step 2 used to replay step 1's
    #     (stale) pause because it iterates StepRuns id-order and yielded the
    #     FIRST PAUSED row. observe_run now only treats a PAUSED StepRun as live
    #     when the RUN is paused AND the step's order_index == current_step_index. ---
    import json as _json

    from app.models import RunStatus as _RunStatus

    with Session(engine) as s:
        wf5 = Workflow(name="Two human steps", scope="global")
        s.add(wf5)
        s.commit()
        s.refresh(wf5)
        wf5_id = wf5.id
        add_step(s, wf5_id, 0, "Translator", prompt_text="Translate: {text}")
        add_step(s, wf5_id, 1, "First review", step_type=StepType.HUMAN_REVIEW)
        add_step(s, wf5_id, 2, "Second review", step_type=StepType.HUMAN_REVIEW)
        add_step(s, wf5_id, 3, "Finalizer", prompt_text="Finalize: {text}")
    doc7 = upload(client, pid, ["Two pauses please."])
    # start the run (POST) -> background thread pauses at step 1 (index 1)
    r = client.post(
        f"/documents/{doc7}/run",
        data={"workflow_id": wf5_id, "locale": "fr-FR"},
        follow_redirects=False,
    )
    run7 = int(r.headers["location"].split("/runs/", 1)[1].split("?")[0])
    body1 = _drain_stream(client, run7)
    ev1 = [_json.loads(ln[len("data: ") :]) for ln in body1.splitlines() if ln.startswith("data: ")]
    paused1 = next(e for e in ev1 if e["type"] == "paused")
    assert paused1["step_index"] == 1, ("first pause must be step 1", paused1)
    # resume past step 1 -> background thread continues and pauses at step 2 (index 2)
    client.post(
        f"/runs/{run7}/resume",
        data={"locale": "fr-FR", "edited_text": "EDIT ONE"},
        follow_redirects=False,
    )
    body2 = _drain_stream(client, run7)  # fresh observer reconnect while paused at step 2
    ev2 = [_json.loads(ln[len("data: ") :]) for ln in body2.splitlines() if ln.startswith("data: ")]
    paused2 = next(e for e in ev2 if e["type"] == "paused")
    assert paused2["step_index"] == 2, (
        "reconnect must surface step 2's pause, not the stale step-1 pause",
        paused2,
    )
    assert paused2["step_name"] == "Second review", paused2
    with Session(engine) as s:
        run_now = s.get(WorkflowRun, run7)
        assert run_now.status == _RunStatus.PAUSED and run_now.current_step_index == 2, (
            "run should be paused at step 2",
            run_now.status,
            run_now.current_step_index,
        )
    print(
        "  finding-3: reconnect observer surfaces the CURRENT (step-2) pause, not the stale step-1 one"  # noqa: E501
    )

    # --- Finding 6 regression: double start/resume is guarded. A second start
    #     POST for the same (workflow, document) re-attaches to the live run
    #     (one WorkflowRun); a second resume advances only once (one extra step). ---
    doc8 = upload(client, pid, ["Guard me."])
    r_a = client.post(
        f"/documents/{doc8}/run",
        data={"workflow_id": wf2_id, "locale": "fr-FR"},
        follow_redirects=False,
    )
    run8 = int(r_a.headers["location"].split("/runs/", 1)[1].split("?")[0])
    _drain_stream(client, run8)  # let it reach its (single) human pause
    r_b = client.post(
        f"/documents/{doc8}/run",
        data={"workflow_id": wf2_id, "locale": "fr-FR"},
        follow_redirects=False,
    )
    run8_again = int(r_b.headers["location"].split("/runs/", 1)[1].split("?")[0])
    assert run8_again == run8, ("double start must re-attach to the same run", run8, run8_again)
    with Session(engine) as s:
        n_runs = len(
            s.exec(
                select(WorkflowRun).where(
                    WorkflowRun.workflow_id == wf2_id, WorkflowRun.document_id == doc8
                )
            ).all()
        )
    assert n_runs == 1, ("double start created a second run", n_runs)
    # double resume: the second resume must be a no-op (run already advanced).
    client.post(
        f"/runs/{run8}/resume",
        data={"locale": "fr-FR", "edited_text": "R1"},
        follow_redirects=False,
    )
    _drain_stream(client, run8)
    with Session(engine) as s:
        idx_after_first = s.get(WorkflowRun, run8).current_step_index
    client.post(
        f"/runs/{run8}/resume",
        data={"locale": "fr-FR", "edited_text": "R2"},
        follow_redirects=False,
    )
    with Session(engine) as s:
        idx_after_second = s.get(WorkflowRun, run8).current_step_index
    assert idx_after_second == idx_after_first, (
        "a second resume of a non-paused run must be a no-op",
        idx_after_first,
        idx_after_second,
    )
    print("  finding-6: double start re-attaches (one run); double resume is a no-op")

    # --- auto-chunking on a large document ---
    with Session(engine) as s:
        wf3 = Workflow(name="Chunker", scope="global")
        s.add(wf3)
        s.commit()
        s.refresh(wf3)
        wf3_id = wf3.id
        add_step(s, wf3_id, 0, "Translator", prompt_text="Translate: {text}", chunk_policy="auto")
    big = [("Localization workflow paragraph %d. " % i) * 12 for i in range(60)]  # noqa: UP031
    doc4 = upload(client, pid, big, mode="full_text")
    with Session(engine) as s:
        st = orchestrator.run_workflow(s, wf3_id, doc4, "fr-FR")
    assert st["steps"][0]["chunks"] > 1, f"expected >1 chunk, got {st['steps'][0]['chunks']}"
    print(f"  auto-chunking: large doc split into {st['steps'][0]['chunks']} chunks")

    # --- failover within a step ---
    register_provider("failer", Failer())
    with Session(engine) as s:
        wf4 = Workflow(name="Failover step", scope="global")
        s.add(wf4)
        s.commit()
        s.refresh(wf4)
        wf4_id = wf4.id
        add_step(
            s,
            wf4_id,
            0,
            "Translator",
            provider="failer",
            fallback_provider="mock",
            prompt_text="Translate: {text}",
            backoff_base=0.0,
        )
    doc5 = upload(client, pid, ["Failover sentence."])
    with Session(engine) as s:
        stf = orchestrator.run_workflow(s, wf4_id, doc5, "fr-FR")
    assert stf["steps"][0]["provider"] == "mock", stf["steps"][0]
    print("  failover within a step: primary 'failer' fell over to 'mock'")

    # --- POST starts a run; SSE stream is a read-only observer ---
    doc6 = upload(client, pid, ["Stream me."])
    r = client.post(
        f"/documents/{doc6}/run",
        data={"workflow_id": wf_id, "locale": "fr-FR"},
        follow_redirects=False,
    )
    assert r.status_code == 303, r.status_code
    run6 = int(r.headers["location"].split("/runs/", 1)[1].split("?")[0])
    body = _drain_stream(client, run6)
    assert (
        '"type": "run_started"' in body and '"type": "step"' in body and '"type": "done"' in body
    ), body[:300]
    print("  POST starts a run; SSE observer emits run_started + step + done")

    # --- the observer does NOT execute: opening the stream twice for one run
    #     must not add StepRuns; a GET for a run never starts a second run ---
    with Session(engine) as s:
        n_steps_1 = len(s.exec(select(StepRun).where(StepRun.workflow_run_id == run6)).all())
        n_runs_1 = len(s.exec(select(WorkflowRun)).all())
    _drain_stream(client, run6)  # reconnect / second tab
    _drain_stream(client, run6)
    with Session(engine) as s:
        n_steps_2 = len(s.exec(select(StepRun).where(StepRun.workflow_run_id == run6)).all())
        n_runs_2 = len(s.exec(select(WorkflowRun)).all())
    assert n_steps_2 == n_steps_1, (n_steps_1, n_steps_2)
    assert n_runs_2 == n_runs_1, (n_runs_1, n_runs_2)
    print("  re-opening the stream re-attaches to the same run (no re-execution)")

    # --- multi-locale run remembers its OWN locale (dual-surface QA finding) ---
    #     A run launched for a locale that is NOT the alphabetically-first target
    #     must persist that locale, and the run view + resume form must use it —
    #     not the first locale found on the document. Regression for the bug where
    #     WorkflowRun had no target_locale and run_detail inferred sorted()[0].
    ml_pid = int(
        client.post(
            "/projects",
            data={
                "name": "Multi-locale",
                "source_locale": "en-US",
                "target_locales": "de-DE, fr-FR",
                "default_provider": "claude",
            },
            follow_redirects=False,
        )
        .headers["location"]
        .rsplit("/", 1)[-1]
    )
    ml_doc = upload(client, ml_pid, ["Locale sentence."])
    # a workflow with a human_review step so the run PAUSES and renders a resume form
    with Session(engine) as s:
        ml_wf = Workflow(name="ML review", scope="global")
        s.add(ml_wf)
        s.commit()
        s.refresh(ml_wf)
        ml_wf_id = ml_wf.id
        add_step(s, ml_wf_id, 0, "Translator", prompt_text="Translate: {text}")
        add_step(s, ml_wf_id, 1, "Human", step_type=StepType.HUMAN_REVIEW)
    # start the run for fr-FR (NOT de-DE, which sorts first)
    r = client.post(
        f"/documents/{ml_doc}/run",
        data={"workflow_id": ml_wf_id, "locale": "fr-FR"},
        follow_redirects=False,
    )
    assert r.status_code == 303, r.status_code
    ml_run = int(r.headers["location"].split("/runs/", 1)[1].split("?")[0])
    _drain_stream(client, ml_run)  # let it pause at the human step
    with Session(engine) as s:
        row = s.get(WorkflowRun, ml_run)
        assert row.target_locale == "fr-FR", (
            "run must persist its launch locale",
            row.target_locale,
        )
    # the run view's resume form must carry fr-FR, never the sorted-first de-DE
    view = client.get(f"/runs/{ml_run}").text
    import re as _re

    m = _re.search(
        r'action="/runs/%d/resume".*?name="locale"[^>]*value="([^"]*)"' % ml_run,  # noqa: UP031
        view,
        _re.S,  # noqa: RUF100, UP031
    )
    assert m and m.group(1) == "fr-FR", (
        "resume form must submit the run's own locale, not sorted()[0]",
        m.group(1) if m else None,
    )
    print("  multi-locale run persists + renders its own locale (not sorted-first)")

    print("CHUNK 5 OK")


if __name__ == "__main__":
    main()
