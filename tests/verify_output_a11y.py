"""Translated-output accessibility (VPAT 504.2, 504.2.1, 504.2.2, 504.3; EN 301 549 5.4,
11.8.2, 11.8.3).

Audit finding 2026-09-24: translated files kept the SOURCE language declaration and did
not translate image alternative text, so a screen reader read a Spanish document with an
English voice and heard English alt text. This suite pins the fix per format:

  1. HTML: alt / title / aria-label (img, area, a) are extracted as `a11y:*` units
     APPENDED after the body units, written back escaped; `<html lang>` (+ xml:lang,
     content-language meta) declares the target and `dir="rtl"` is set for an RTL
     target; build() with no mapping stays byte-identical; a document ingested before
     alt extraction (mapping without the new indices) keeps its alt text untouched.
  2. DOCX: wp:docPr descr/title extracted and written back; w:lang in docDefaults
     carries the target in the right slot (val / eastAsia / bidi); dc:language set;
     <w:bidi/> on translated paragraphs for an RTL target.
  3. Markdown: image lines marked `a11y:md-image`, empty alt reported, build unchanged.
  4. PDF: /Lang is the target (typst `set text(lang, region)`), untagged path unchanged.
  5. XLSX dc:language, XLIFF target-language / trgLang, PO Language header.
  6. End to end: upload -> segment notes on alt units -> Editor shows the `a11y_alt`
     prompt (and the save response keeps it) -> export declares the target and carries
     the translated alt text.
  7. Formats without images report nothing (the prompt stays off).

ASCII output only.
"""

import io
import os
import re
import sys
import zipfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_outa11y_")

from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app.db import engine, init_db
from app.filters import get_filter
from app.main import app
from app.models import Document, Segment

HTML_SRC = (
    b"<!DOCTYPE html>\n"
    b'<html lang="en" xml:lang="en"><head>'
    b'<meta http-equiv="Content-Language" content="en"><title>Guide</title></head>\n'
    b'<body><p>Hello <b>world</b> <img src="a/cat.png" alt="A cat &amp; a dog"></p>\n'
    b'<img src="x.png"><img src="deco.png" alt="">\n'
    b'<p><a href="/r" title="Open the report" aria-label="Report link">Report</a></p>\n'
    b"<script>var s='<img alt=\"skip me\">';</script></body></html>\n"
)


def _html_units():
    return get_filter("html").parse(HTML_SRC).units


# --------------------------------------------------------------------------- #
# 1. HTML
# --------------------------------------------------------------------------- #
def test_html_alt_units_and_writeback():
    f = get_filter("html")
    units = _html_units()
    body = [u for u in units if not u.style]
    alts = [u for u in units if u.style]
    assert [u.style for u in alts] == [
        "a11y:img-alt",
        "a11y:link-title",
        "a11y:link-aria-label",
    ], [u.style for u in alts]
    assert [u.text for u in alts] == ["A cat & a dog", "Open the report", "Report link"]
    # Appended AFTER every body unit: body indices are exactly 0..n-1.
    assert [u.p_index for u in body] == list(range(len(body)))
    assert [u.p_index for u in alts] == list(range(len(body), len(body) + 3))
    # The script-embedded <img alt> is never extracted.
    assert all("skip me" not in u.text for u in units)
    # No mapping -> byte-identical.
    assert f.build(HTML_SRC, {}) == HTML_SRC
    # Old document: a mapping with ONLY the body indices leaves the alt text alone.
    old = {str(u.p_index): "ES " + u.text for u in body}
    out_old = f.build(HTML_SRC, old).decode()
    assert 'alt="A cat &amp; a dog"' in out_old and 'title="Open the report"' in out_old
    # Write-back, escaped: quotes and angle brackets cannot break out of the attribute.
    m = dict(old)
    m[str(alts[0].p_index)] = 'Un gato & un "perro" <b>'
    m[str(alts[1].p_index)] = "Abrir el informe"
    out = f.build(HTML_SRC, m).decode()
    assert 'alt="Un gato &amp; un &quot;perro&quot; &lt;b&gt;"' in out, out
    assert 'title="Abrir el informe"' in out and 'aria-label="Report link"' in out
    # Re-parse the translated file: the alt unit reads back as the translation.
    re_alts = [u.text for u in f.parse(out.encode()).units if u.style == "a11y:img-alt"]
    assert re_alts == ['Un gato & un "perro" <b>'], re_alts
    print("  1a) html: alt/title/aria-label units appended, escaped write-back, old docs safe")


def test_html_language_and_direction():
    f = get_filter("html")
    out = f.build_localized(HTML_SRC, {}, "es_ES").decode()
    assert '<html lang="es-ES" xml:lang="es-ES">' in out, out[:120]
    assert 'content="es-ES"' in out
    assert "dir=" not in out.split(">", 2)[1]  # LTR target: no dir added
    ar = f.build_localized(HTML_SRC, {}, "ar-EG").decode()
    assert re.search(r'<html [^>]*dir="rtl"', ar) and 'lang="ar-EG"' in ar, ar[:120]
    # An existing dir is flipped for an LTR target.
    rtl_src = b'<html lang="he" dir="rtl"><body><p>x y</p></body></html>'
    assert b'dir="ltr"' in f.build_localized(rtl_src, {}, "fr")
    # A fragment (no <html>) gets no declaration and stays exactly as build() made it.
    frag = b"<p>Hello <img src='a.png' alt='A cat'></p>"
    assert f.build_localized(frag, {"0": "Hola"}, "es") == f.build(frag, {"0": "Hola"})
    # The placeholder locale and junk declare nothing.
    for loc in ("\u2014", "", None, 'x" onload="y'):
        assert f.build_localized(HTML_SRC, {}, loc) == HTML_SRC, loc
    print("  1b) html: <html lang/xml:lang>, content-language, dir=rtl, fragments untouched")


def test_html_missing_alt_findings():
    sk = get_filter("html").parse(HTML_SRC).skeleton
    info = sk.get("a11y_images")
    assert info and info["images"] == 3, info
    kinds = [(f["kind"], f["where"]) for f in info["missing_alt"]]
    assert kinds == [("missing", "x.png"), ("empty", "deco.png")], kinds
    none = get_filter("html").parse(b"<p>No images here.</p>").skeleton
    assert "a11y_images" not in none
    print("  1c) html: missing and empty alt reported; no key when there are no images")


# --------------------------------------------------------------------------- #
# 2. DOCX
# --------------------------------------------------------------------------- #
def _docx_with_images() -> bytes:
    from docx import Document as Docx
    from docx.oxml.ns import qn
    from PIL import Image

    png = io.BytesIO()
    Image.new("RGB", (4, 4), "red").save(png, "PNG")
    d = Docx()
    d.add_paragraph("Hello world.")
    d.add_paragraph().add_run().add_picture(io.BytesIO(png.getvalue()))
    d.add_paragraph("A second paragraph.").add_run().add_picture(io.BytesIO(png.getvalue()))
    dps = list(d.element.body.iter(qn("wp:docPr")))
    dps[0].set("descr", "A red square")
    dps[0].set("title", "Square")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def _part(data: bytes, name: str) -> str:
    return zipfile.ZipFile(io.BytesIO(data)).read(name).decode("utf-8")


def test_docx():
    f = get_filter("docx")
    src = _docx_with_images()
    r = f.parse(src)
    alts = [(u.text, u.style) for u in r.units if u.style.startswith("a11y:")]
    assert alts == [("A red square", "a11y:img-alt"), ("Square", "a11y:img-title")], alts
    body_idx = [u.p_index for u in r.units if not u.style.startswith("a11y:")]
    assert body_idx == [0, 2], body_idx  # paragraph indices unchanged
    miss = r.skeleton["a11y_images"]["missing_alt"]
    assert [m["kind"] for m in miss] == ["missing"] and miss[0]["unit"] == 2, miss
    m = {str(u.p_index): "ES " + u.text for u in r.units}
    out = f.build_localized(src, m, "es-ES")
    doc_xml = _part(out, "word/document.xml")
    assert 'descr="ES A red square"' in doc_xml and 'title="ES Square"' in doc_xml
    styles = _part(out, "word/styles.xml")
    assert re.search(r'<w:lang [^>]*w:val="es-ES"', styles), styles[:400]
    assert "<dc:language>es-ES</dc:language>" in _part(out, "docProps/core.xml")
    assert "<w:bidi/>" not in doc_xml
    ar = f.build_localized(src, m, "ar")
    assert re.search(r'<w:lang [^>]*w:bidi="ar"', _part(ar, "word/styles.xml"))
    assert _part(ar, "word/document.xml").count("<w:bidi/>") == 2  # the 2 translated paras
    ja = f.build_localized(src, m, "ja-JP")
    assert re.search(r'<w:lang [^>]*w:eastAsia="ja-JP"', _part(ja, "word/styles.xml"))
    # #730: Thai / Hindi letters are complex-script characters: the target goes into the
    # w:bidi slot (it used to land in w:val, leaving the template's w:bidi="ar-SA" over Thai).
    for loc in ("th-TH", "hi-IN", "km"):
        cs = _part(f.build_localized(src, m, loc), "word/styles.xml")
        assert re.search(rf'<w:lang [^>]*w:bidi="{loc}"', cs), (loc, cs[:600])
        assert not re.search(rf'w:val="{loc}"', cs), (loc, "Latin slot must keep its language")
        assert "<w:bidi/>" not in _part(f.build_localized(src, m, loc), "word/document.xml")
    # gate: a romanized complex-script target is written in Latin letters: Latin slot
    from app.filters import a11y

    for loc, want in (
        ("hi-Latn", "latin"),
        ("sa-Latn-IN", "latin"),
        ("hi-Deva-IN", "bidi"),
        ("th", "bidi"),
        ("ar", "bidi"),
        ("ja-JP", "eastAsia"),
        ("es-ES", "latin"),
    ):
        assert a11y.script_class(loc) == want, (loc, a11y.script_class(loc))
    # Plain build() declares nothing new (the round-trip ratchets pin build()).
    assert "es-ES" not in _part(f.build(src, m), "word/styles.xml")
    print("  2)  docx: docPr descr/title round-trip, w:lang slot, dc:language, w:bidi")


def _strip_parts(data: bytes, drop: set) -> bytes:
    """A copy of a .docx package without the parts in `drop` (and their relationships
    and content-type overrides): a real-world file some producers write."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            if info.filename in drop:
                continue
            body = zin.read(info.filename)
            if info.filename.endswith(".rels") or info.filename == "[Content_Types].xml":
                text = body.decode("utf-8")
                for name in drop:
                    leaf = name.rsplit("/", 1)[-1]
                    text = re.sub(
                        r"<Relationship [^>]*Target=\"[^\"]*" + re.escape(leaf) + r"\"[^>]*/>",
                        "",
                        text,
                    )
                    text = re.sub(
                        r"<Override [^>]*PartName=\"/" + re.escape(name) + r"\"[^>]*/>", "", text
                    )
                body = text.encode("utf-8")
            zout.writestr(info, body)
    return out.getvalue()


def test_docx_without_core_or_styles():
    """Gate round 1 (M2): a package with no docProps/core.xml and no word/styles.xml
    must not gain python-docx's invented parts (a core.xml titled "Word Document", a
    default styles.xml). Only the run-level w:lang of translated paragraphs is set."""
    from docx import Document as Docx

    d = Docx()
    d.add_paragraph("Hello world.")
    d.add_paragraph("Second line.")
    b = io.BytesIO()
    d.save(b)
    bare = _strip_parts(b.getvalue(), {"docProps/core.xml", "word/styles.xml"})
    names_in = sorted(zipfile.ZipFile(io.BytesIO(bare)).namelist())
    assert "docProps/core.xml" not in names_in and "word/styles.xml" not in names_in
    f = get_filter("docx")
    m = {"0": "Hola mundo.", "1": "Segunda linea."}
    for loc, attr in (
        ("es-ES", "w:val"),
        ("ar", "w:bidi"),
        ("ja", "w:eastAsia"),
        ("th", "w:bidi"),  # #730
    ):
        out = f.build_localized(bare, m, loc)
        z = zipfile.ZipFile(io.BytesIO(out))
        assert sorted(z.namelist()) == names_in, (loc, sorted(z.namelist()), names_in)
        for name in z.namelist():
            assert b"Word Document" not in z.read(name), (loc, name)
            assert b"python-docx" not in z.read(name) or name == "docProps/app.xml", (loc, name)
        doc_xml = z.read("word/document.xml").decode("utf-8")
        tag = {"es-ES": "es-ES", "ar": "ar", "ja": "ja", "th": "th"}[loc]
        assert doc_xml.count(f'<w:lang {attr}="{tag}"/>') == 2, (loc, doc_xml[-600:])
        # Same entry list and the same bytes for every part but the document itself.
        zin = zipfile.ZipFile(io.BytesIO(bare))
        changed = [n for n in names_in if zin.read(n) != z.read(n)]
        assert changed == ["word/document.xml"], (loc, changed)
    print("  2b) docx without core.xml/styles.xml: no invented parts, run-level w:lang only")


# --------------------------------------------------------------------------- #
# 3. Markdown
# --------------------------------------------------------------------------- #
def test_markdown():
    f = get_filter("md")
    src = b"# Title\n\nSee ![A diagram](d.png) here.\n\n![](empty.png)\n\n```\n![](code.png)\n```\n"
    r = f.parse(src)
    styles = {u.text: u.style for u in r.units}
    assert styles["See ![A diagram](d.png) here."] == "a11y:md-image", styles
    assert styles["Title"] == ""
    info = r.skeleton["a11y_images"]
    assert info["images"] == 2, info  # the fenced one is code, not an image
    assert [(x["kind"], x["where"]) for x in info["missing_alt"]] == [("empty", "empty.png")]
    assert f.build(src, {}) == src  # #677: a no-op build is the original bytes
    print("  3)  markdown: image lines marked, empty alt reported, fenced code ignored")


# --------------------------------------------------------------------------- #
# 4. PDF
# --------------------------------------------------------------------------- #
def test_pdf_lang():
    f = get_filter("pdf")
    texts = {"0": "Hola mundo", "1": "Adios"}
    for loc, want in (("es_ES", b"/Lang(es-ES)"), ("ar", b"/Lang(ar)"), ("es-419", b"/Lang(es)")):
        out = f.build_localized(b"", texts, loc)
        assert want in out, (loc, re.findall(rb"/Lang ?\([^)]*\)", out))
        assert b"/StructTreeRoot" in out  # typst tags the PDF by default
        assert b"pdfuaid" not in out  # and we never claim PDF/UA
    from app.filters import pdf_filter

    assert pdf_filter._build_source(["x"]) == pdf_filter._build_source(["x"], "")
    assert "lang" not in pdf_filter._build_source(["x"])
    print("  4)  pdf: /Lang is the target locale; tagged; no PDF/UA claim")


# --------------------------------------------------------------------------- #
# 5. XLSX / XLIFF / PO
# --------------------------------------------------------------------------- #
def test_other_declarations():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = "Hello"
    b = io.BytesIO()
    wb.save(b)
    out = get_filter("xlsx").build_localized(b.getvalue(), {"0": "Hola"}, "es-ES")
    assert "<dc:language" in _part(out, "docProps/core.xml")
    assert ">es-ES</dc:language>" in _part(out, "docProps/core.xml")

    x12 = (
        b'<?xml version="1.0" encoding="UTF-8"?><xliff version="1.2" '
        b'xmlns="urn:oasis:names:tc:xliff:document:1.2"><file source-language="en" '
        b'target-language="de" datatype="plaintext" original="a"><body><trans-unit id="1">'
        b"<source>Hello</source></trans-unit></body></file></xliff>"
    )
    x = get_filter("xliff")
    assert b'target-language="es-ES"' in x.build_localized(x12, {"0": "Hola"}, "es_ES")
    x20 = (
        b'<?xml version="1.0"?><xliff version="2.0" xmlns="urn:oasis:names:tc:xliff:'
        b'document:2.0" srcLang="en"><file id="f"><unit id="u"><segment><source>Hello'
        b"</source></segment></unit></file></xliff>"
    )
    assert b'trgLang="ar"' in x.build_localized(x20, {"0": "x"}, "ar")

    po = (
        b'# header\nmsgid ""\nmsgstr ""\n"Content-Type: text/plain; charset=UTF-8\\n"\n'
        b'"Language: en\\n"\n\nmsgid "Hello"\nmsgstr ""\n'
    )
    p = get_filter("po")
    assert b"Language: es_ES\\n" in p.build_localized(po, {"0": "Hola"}, "es-ES")
    assert p.build_localized(po, {"0": "Hola"}, "") == p.build(po, {"0": "Hola"})
    print("  5)  xlsx dc:language, xliff target-language/trgLang, po Language header")


def test_svg_language():
    """#691: an SVG export declares the target on its root element."""
    f = get_filter("svg")
    svg = (
        b'<?xml version="1.0"?>\n<!-- c -->\n<svg xmlns="http://www.w3.org/2000/svg" '
        b'xml:lang="en"><text x="1" y="2">Hello</text><text xml:lang="la">Lorem</text></svg>'
    )
    out = f.build_localized(svg, {"0": "Hallo"}, "de-DE")
    assert b'<svg xmlns="http://www.w3.org/2000/svg" xml:lang="de-DE">' in out, out
    assert b">Hallo<" in out and b'<text xml:lang="la">' in out, out  # nested kept
    assert f.build_localized(svg, {"0": "Hallo"}, "") == f.build(svg, {"0": "Hallo"})
    assert f.build(svg, {}) == svg
    bare = b'<svg xmlns="http://www.w3.org/2000/svg"><text>Hi</text></svg>'
    assert f.build_localized(bare, {}, "ar").startswith(b'<svg xml:lang="ar" xmlns='), bare
    plain = b"<svg lang='en' xmlns=\"http://www.w3.org/2000/svg\"><text>Hi</text></svg>"
    out = f.build_localized(plain, {}, "fr-FR")
    assert b'lang="fr-FR"' in out and b"xml:lang" not in out, out
    both = b'<svg xml:lang="en" lang="en"><text>Hi</text></svg>'
    want = b'<svg xml:lang="ja" lang="ja"><text>Hi</text></svg>'
    assert f.build_localized(both, {}, "ja") == want
    # gate: `lang=` text inside another attribute's quoted value is not an attribute
    tricky = (
        b'<svg aria-label="set lang=en" data-n=\' xml:lang=x\' xml:lang="en"><text>Hi</text></svg>'
    )
    want = (
        b'<svg aria-label="set lang=en" data-n=\' xml:lang=x\' xml:lang="de"><text>Hi</text></svg>'
    )
    assert f.build_localized(tricky, {}, "de") == want, f.build_localized(tricky, {}, "de")
    print("  5b) svg root xml:lang / lang")


# --------------------------------------------------------------------------- #
# 6. End to end: ingest -> editor -> export
# --------------------------------------------------------------------------- #
def test_end_to_end():
    client = TestClient(app)
    r = client.post(
        "/projects",
        data={
            "name": "Output a11y",
            "source_locale": "en-US",
            "target_locales": "es-ES",
            "default_provider": "claude",
        },
        follow_redirects=False,
    )
    assert r.status_code == 303, r.text
    pid = int(r.headers["location"].rsplit("/", 1)[-1])
    r = client.post(
        f"/projects/{pid}/documents",
        data={"mode": "segmented"},
        files={"file": ("guide.html", HTML_SRC, "text/html")},
        follow_redirects=False,
    )
    assert r.status_code == 303, r.text
    doc_id = int(r.headers["location"].rsplit("/", 1)[-1])
    with Session(engine) as s:
        segs = sorted(
            s.exec(
                select(Segment).where(
                    Segment.document_id == doc_id, Segment.target_locale == "es-ES"
                )
            ).all(),
            key=lambda x: x.order_index,
        )
        rows = [(x.id, x.order_index, x.source_text, x.meta) for x in segs]
        flags = (s.get(Document, doc_id).skeleton or {}).get("a11y_alt_flags")
    noted = [(src, (meta or {}).get("notes", "")) for _i, _o, src, meta in rows if meta]
    assert [n[0] for n in noted] == ["A cat & a dog", "Open the report", "Report link"], noted
    assert noted[0][1].startswith("Image alternative text"), noted
    assert all(meta is None for _i, _o, src, meta in rows if src not in dict(noted)), rows
    assert flags and [f["kind"] for f in flags] == ["missing", "empty"], flags

    page = client.get(f"/documents/{doc_id}/editor?locale=es-ES")
    assert page.status_code == 200, page.status_code
    assert "a11y_alt" in page.text and "no alternative text in the source" in page.text
    assert "Image alternative text (read aloud by screen readers)" in page.text

    for sid, _o, src, _m in rows:
        r = client.post(f"/segments/{sid}/save", data={"target_text": "ES " + src, "tte_ms": "0"})
        assert r.status_code == 200, r.text
    flagged_order = flags[0]["order_index"]
    sid_flagged, src_flagged = next((i, src) for i, o, src, _m in rows if o == flagged_order)
    r = client.post(
        f"/segments/{sid_flagged}/save",
        data={"target_text": "ES2 " + src_flagged, "tte_ms": "0"},
    )
    assert any(f.get("type") == "a11y_alt" for f in r.json()["qa"]), r.json()

    resp = client.get(f"/documents/{doc_id}/export?locale=es-ES")
    assert resp.status_code == 200, resp.text
    out = resp.content.decode("utf-8")
    assert '<html lang="es-ES" xml:lang="es-ES">' in out, out[:160]
    assert 'alt="ES A cat &amp; a dog"' in out, out
    assert 'title="ES Open the report"' in out and 'aria-label="ES Report link"' in out
    print("  6)  end to end: notes on alt segments, editor prompt, export lang + alt")
    return doc_id


def test_headless_surfaces(doc_id: int):
    """The same prompt on the API and the CLI (four-surface rule)."""
    import subprocess

    from app.services import apikeys, security

    user = security.create_user("oa@test.invalid", "OA admin", "admin", "oa-password-1234")
    _, raw = apikeys.create_key(user.id, "oa", scopes="api docs admin")
    client = TestClient(app)
    r = client.get(f"/api/documents/{doc_id}/segments", headers={"X-Loc-Token": raw})
    assert r.status_code == 200, r.text
    summ = r.json().get("a11y_images")
    assert summ and summ["images"] == 3, r.json()
    assert [m["kind"] for m in summ["missing_alt"]] == ["missing", "empty"], summ
    proc = subprocess.run(
        [sys.executable, "-m", "app.cli", "doc", "show", str(doc_id)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=120,
        env=os.environ.copy(),
        check=False,
    )
    assert "Images: 3, 2 without alt text in the source" in proc.stdout, (
        proc.stdout[-800:] + proc.stderr[-800:]
    )
    print("  6b) api a11y_images + cli `doc show` line report the same prompt")


# --------------------------------------------------------------------------- #
# 7. Formats without images: prompt stays off
# --------------------------------------------------------------------------- #
def test_full_text_mode_keeps_alt_out():
    """Gate round 1 (L5): a FULL_TEXT document carries no alt units (the one body would
    otherwise end in context-free alt paragraphs); its alt text is left as it is, the
    language is still declared, and the export paragraph guard is not tripped."""
    client = TestClient(app)
    r = client.post(
        "/projects",
        data={
            "name": "Output a11y FT",
            "source_locale": "en-US",
            "target_locales": "es-ES",
            "default_provider": "claude",
        },
        follow_redirects=False,
    )
    pid = int(r.headers["location"].rsplit("/", 1)[-1])
    r = client.post(
        f"/projects/{pid}/documents",
        data={"mode": "full_text"},
        files={"file": ("ft.html", HTML_SRC, "text/html")},
        follow_redirects=False,
    )
    assert r.status_code == 303, r.text
    doc_id = int(r.headers["location"].rsplit("/", 1)[-1])
    with Session(engine) as s:
        doc = s.get(Document, doc_id)
        units = (doc.skeleton or {}).get("units") or []
        segs = s.exec(
            select(Segment).where(Segment.document_id == doc_id, Segment.target_locale == "es-ES")
        ).all()
    assert units and not any((u.get("style") or "").startswith("a11y:") for u in units), units
    assert len(segs) == 1 and "A cat & a dog" not in segs[0].source_text, segs[0].source_text
    parts = segs[0].source_text.split("\n\n")
    r = client.post(
        f"/segments/{segs[0].id}/save",
        data={"target_text": "\n\n".join("ES " + p for p in parts), "tte_ms": "0"},
    )
    assert r.status_code == 200, r.text
    resp = client.get(f"/documents/{doc_id}/export?locale=es-ES")
    assert resp.status_code == 200, resp.text
    out = resp.content.decode("utf-8")
    assert '<html lang="es-ES"' in out and 'alt="A cat &amp; a dog"' in out, out[:300]
    print("  6c) full-text mode: no alt units, alt left as is, language declared, export ok")


def test_html_edge_alts():
    """Gate round 1 (L7): a numeric-only alt is not a segment; an alt holding a raw `>`
    (the tokenizer splits that tag) or an entity-encoded one is not reported as empty."""
    f = get_filter("html")
    units = f.parse(b'<p>Text here.</p><img src="n.png" alt="12345">').units
    assert [u.text for u in units] == ["Text here."], [u.text for u in units]
    for src in (
        b'<p>Text here.</p><img src="g.png" alt="a > b">',
        b'<p>Text here.</p><img src="h.png" alt="a &gt; b">',
        b"<p>Text here.</p><img src='i.png' alt='x > y'>",
    ):
        info = f.parse(src).skeleton.get("a11y_images")
        assert info and info["missing_alt"] == [], (src, info)
    info = f.parse(b'<p>T.</p><img src="e.png" alt="">').skeleton["a11y_images"]
    assert [m["kind"] for m in info["missing_alt"]] == ["empty"], info
    # HubSpot's same-shape check counts BODY units only (gate round 1, L3): one paragraph
    # more and one alt fewer is NOT the same shape.
    from app.services.connectors.hubspot import _body_unit_count

    a = f.parse(b'<p>One.</p><img src="a.png" alt="Alt one">').units
    b = f.parse(b"<p>One.</p><p>Two.</p>").units
    assert len(a) == len(b) and _body_unit_count(a) == 1 and _body_unit_count(b) == 2
    print(
        "  1d) html: numeric-only alt not a segment; raw '>' / entity alt not flagged empty; "
        "hubspot shape = body units"
    )


def test_prompt_off_without_images():
    for fmt, data in (
        ("txt", b"Hello world.\n"),
        ("json", b'{"a": "Hello world"}'),
        ("srt", b"1\n00:00:01,000 --> 00:00:02,000\nHello\n"),
    ):
        filt = get_filter(fmt)
        if filt is None:
            continue
        sk = filt.parse(data).skeleton
        assert "a11y_images" not in sk, (fmt, sk)
        assert filt.build_localized(data, {}, "es-ES") == filt.build(data, {}), fmt
    print("  7)  txt/json/srt: no image prompt, no declaration change")


def main() -> None:
    init_db()
    test_html_alt_units_and_writeback()
    test_html_language_and_direction()
    test_html_missing_alt_findings()
    test_html_edge_alts()
    test_docx()
    test_docx_without_core_or_styles()
    test_markdown()
    test_pdf_lang()
    test_other_declarations()
    test_svg_language()
    doc_id = test_end_to_end()
    test_full_text_mode_keeps_alt_out()
    test_prompt_off_without_images()
    test_headless_surfaces(doc_id)
    print("OUTPUT A11Y OK")


if __name__ == "__main__":
    main()
