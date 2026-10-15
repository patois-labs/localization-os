"""Lazy third-party imports stay lazy (test speed, 2026-10-08; offline).

Every verify suite, every CLI call and every server start pays the app's import time. The
heaviest third-party packages that only some requests need are imported where they are used,
not at module top level:

  * openpyxl (and the numpy it pulls in): app/filters/xlsx_filter.py, app/filters/xlsx_guard.py
  * python-docx: app/filters/docx_filter.py
  * pypdfium2: app/filters/pdf_filter.py
  * PyJWT: app/services/oidc.py

This suite pins that, so a new top-level import of one of them is caught here instead of
silently adding back up to a second per process:

  [A] `import app.main` loads none of them.
  [B] `python -m app.cli --help` loads none of them.
  [C] docx_filter.qn (a local copy so the module-level tag constants need no python-docx)
      matches docx.oxml.ns.qn for every prefix the module uses.
  [D] the lazy entry points still reach the real libraries: xlsx_filter.load_workbook reads a
      workbook, DocxFilter parses a .docx, PdfFilter parses a PDF, oidc._signing_key resolves
      a JWKS key.
  [E] server start (app.main._startup) still refuses an install missing any of the four,
      as the old top-level imports did at import time.
"""

import io
import os
import subprocess
import sys

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)

from tests._bootstrap import init_env

init_env("loctms_lazyimp_")

HEAVY = ("openpyxl", "docx", "pypdfium2", "jwt", "numpy")

_PROBE = rf"""
import runpy, sys
mode = sys.argv[1]
if mode == "main":
    import app.main  # noqa: F401
else:
    sys.argv = ["app.cli", "--help"]
    try:
        runpy.run_module("app.cli", run_name="__main__")
    except SystemExit:
        pass
heavy = {HEAVY!r}
print("LOADED=" + ",".join(sorted(m for m in heavy if m in sys.modules)))
"""


def _loaded(mode):
    proc = subprocess.run(
        [sys.executable, "-c", _PROBE, mode],
        cwd=_ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=300,
        env=os.environ.copy(),
    )
    assert proc.returncode == 0, proc.stderr[-2000:]
    line = [ln for ln in proc.stdout.splitlines() if ln.startswith("LOADED=")]
    assert line, proc.stdout[-2000:]
    return [m for m in line[-1][len("LOADED=") :].split(",") if m]


def test_a_app_import_is_lazy():
    loaded = _loaded("main")
    assert loaded == [], f"import app.main loaded {loaded}: import them where they are used"
    print("  [A] import app.main loads none of", ", ".join(HEAVY), "OK")


def test_b_cli_help_is_lazy():
    loaded = _loaded("cli")
    assert loaded == [], f"python -m app.cli --help loaded {loaded}"
    print("  [B] python -m app.cli --help loads none of", ", ".join(HEAVY), "OK")


def test_c_docx_qn_matches_python_docx():
    import docx.oxml.ns as ns

    from app.filters import docx_filter

    for prefix, uri in docx_filter._NSMAP.items():
        assert ns.nsmap[prefix] == uri, prefix
        assert docx_filter.qn(f"{prefix}:x") == ns.qn(f"{prefix}:x"), prefix
    for tag in ("w:p", "w:rPr", "w:tbl", "wp:inline", "pic:pic"):
        assert docx_filter.qn(tag) == ns.qn(tag), tag
    print("  [C] docx_filter.qn matches docx.oxml.ns.qn for w, wp and pic OK")


def test_d_lazy_entry_points_reach_the_libraries():
    from docx import Document
    from openpyxl import Workbook

    from app.filters.docx_filter import DocxFilter
    from app.filters.pdf_filter import PdfFilter
    from app.filters.xlsx_filter import XlsxFilter, load_workbook

    wb = Workbook()
    wb.active["A1"] = "Hello world"
    buf = io.BytesIO()
    wb.save(buf)
    assert load_workbook(io.BytesIO(buf.getvalue())).active["A1"].value == "Hello world"
    units = XlsxFilter().parse(buf.getvalue()).units
    assert [u.text for u in units] == ["Hello world"], units

    d = Document()
    d.add_paragraph("Hello docx")
    buf = io.BytesIO()
    d.save(buf)
    units = DocxFilter().parse(buf.getvalue()).units
    assert [u.text for u in units] == ["Hello docx"], units

    pdf = PdfFilter().build(b"", {})
    units = PdfFilter().parse(pdf).units
    assert isinstance(units, list)

    import json

    from cryptography.hazmat.primitives.asymmetric import rsa
    from jwt.algorithms import RSAAlgorithm

    from app.db import init_db
    from app.services import oidc

    init_db()  # the egress guard reads its settings from the store
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    jwk = json.loads(RSAAlgorithm.to_jwk(key.public_key()))
    jwk["kid"] = "k1"
    import jwt

    token = jwt.encode({"sub": "x"}, key, algorithm="RS256", headers={"kid": "k1"})

    import httpx

    def handler(request):
        return httpx.Response(200, json={"keys": [jwk]})

    got = oidc._signing_key(
        token,
        {"jwks_uri": "https://idp.test/jwks"},
        transport=httpx.MockTransport(handler),
        iss="https://idp.test",
    )
    assert got.public_numbers() == key.public_key().public_numbers()
    print("  [D] xlsx/docx/pdf filters and the OIDC JWKS key reach their lazy libraries OK")


def test_e_server_start_still_requires_the_lazy_libraries():
    """A broken install missing one of them used to fail at import time; it must still fail
    at server start (app.main._require_lazy_libraries), not at the first upload or sign-in."""
    from unittest.mock import patch

    from app import main

    assert set(main._LAZY_REQUIRED_LIBS) == {"openpyxl", "docx", "pypdfium2", "jwt"}
    main._require_lazy_libraries()  # the real install has all four
    real = __import__("importlib.util").util.find_spec

    def without_jwt(name, *a, **k):
        return None if name == "jwt" else real(name, *a, **k)

    with patch("importlib.util.find_spec", side_effect=without_jwt):
        try:
            main._require_lazy_libraries()
        except RuntimeError as exc:
            assert "PyJWT" in str(exc), exc
        else:
            raise AssertionError("server start must refuse an install without PyJWT")
    with open(main.__file__, encoding="utf-8") as fh:
        src = fh.read()
    body = src[src.index("def _startup() -> None:") :]
    assert body.split("\n", 2)[1].strip() == "_require_lazy_libraries()", body[:200]
    print("  [E] server start refuses an install missing a lazily imported library OK")


if __name__ == "__main__":
    test_a_app_import_is_lazy()
    test_b_cli_help_is_lazy()
    test_c_docx_qn_matches_python_docx()
    test_d_lazy_entry_points_reach_the_libraries()
    test_e_server_start_still_requires_the_lazy_libraries()
    print("LAZY IMPORTS OK")
