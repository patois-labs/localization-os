"""Microsoft 365 Files connector (SharePoint libraries + OneDrive) - fully offline.

FakeGraph is an httpx.MockTransport tenant that models the parts of Microsoft Graph this
connector depends on (taxonomy G): the app-only token endpoint, drive / site / library lookup,
folder children with @odata.nextLink paging, driveItem metadata ($select honoured) with a
content hash AND a content tag that move on a content edit but NOT on a metadata touch, path
addressed create with conflictBehavior=fail (parent folders auto-created, 409 on a clash),
replace-in-place with If-Match (412), a 302 download redirect to a file host that must never see
the bearer token, a 4 MiB simple-upload cap with chunked upload sessions (320 KiB alignment),
Sites.Selected style per-drive denial (403), 401 on an expired token, 429 + Retry-After, and
injected 5xx. The fake would fail the suite if a write landed outside the target root.

Covers the taxonomy A-H checklist (the connector build blueprint section 3):
  A  caller article_ids and operator paths cannot escape the configured drive / source folder
     (hostile id table, foreign folder, target root, other drive, prefix-of-folder sibling,
     folder id, hostile nextLink / redirect / upload URL), zero writes and zero downloads.
  B  a held-back unit never creates or changes a copy (all-or-nothing per file); a copy edited
     by hand is never overwritten (zero MT spend, If-Match race, delete-to-regenerate).
  C  steady state: run 2 lists metadata only, downloads nothing, writes nothing; a metadata-only
     touch of the source is not a change; a content edit re-translates just that file.
  D  supported formats selected; .ppt / pdf / lock files / empty / oversize / bad names skipped
     WITH a reason; a skip is never silent.
  E  the Max files cap round-robins across sub-folders (distribution asserted, not the total).
  F  a held or failed pair records no content hash; a healthy re-run completes it.
  G  the fake enforces Graph's own rules (name characters, 4 MiB simple cap, chunk alignment).
  H  WriteBudget ceiling, egress guard, token never in a message or log, four parity surfaces.
Also: 429 backoff, 401 remint, upload session for large files, interrupted-build resume,
pre-existing human copy fails closed (needs_link) and adoption, credential pin on scope edits.

ASCII-only prints; final line is 'verify_msgraph OK'.
"""

import base64
import hashlib
import logging
import os
import re
import sys
from io import BytesIO
from urllib.parse import parse_qs, quote, unquote

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env
from tests._piece_patch import set_pieces

init_env("loctms_msgraph_")

import docx
import httpx
import openpyxl
from fastapi.testclient import TestClient
from sqlmodel import Session

from app import config as app_config
from app.db import engine, init_db
from app.main import app
from app.models import Connection
from app.services.connectors import base as cbase
from app.services.connectors import idmap
from app.services.connectors import store as cstore
from app.services.connectors.msgraph import MsGraphConnector
from app.services.connectors.msgraph_graph import copy_changed, validate_item_id
from app.services.connectors.store_part2 import _has_credential

TENANT = "contoso.onmicrosoft.com"
CLIENT = "11111111-2222-3333-4444-555555555555"
FIXTURE_SECRET = "fixture-app-credential-9f8e7d6c"  # pregate: allow-secret offline fixture
SITE_URL = "https://contoso.sharepoint.com/sites/Loc"
SITE_ID = "contoso.sharepoint.com,11111111-aaaa-bbbb-cccc-000000000001,11111111-aaaa-bbbb-cccc-000000000002"  # noqa: E501


# ============================================================================
# Fake Microsoft Graph tenant.
# ============================================================================
def _b64(data: bytes) -> str:
    return base64.b64encode(hashlib.sha256(data).digest()[:20]).decode()


class FakeGraph:
    GRAPH = "graph.microsoft.com"
    LOGIN = "login.microsoftonline.com"
    FILES = "contoso.sharepoint.com"

    def __init__(self, *, page_size=200, hashes=True, encode_paths=False):
        self.drive = "b!DriveOne"
        self.other_drive = "b!DriveTwo"
        self.page_size = page_size
        self.hashes = hashes
        self.encode_paths = encode_paths
        self.items: dict = {}
        self._n = 0
        self.root = self._mk("", None, True, "ROOT")
        self.root_other = None
        self.tokens: set = set()
        self.mints = 0
        self.token_ttl = 3600
        self.log: list = []
        self.writes: list = []  # (kind, path) of every SUCCESSFUL mutating request
        self.creates: list = []  # conflictBehavior values seen on path creates
        self.downloads: list = []
        self.leaked: list = []  # requests to the file host that carried a bearer token
        self.stray_hosts: list = []
        self.denied_drives: set = set()
        self.fail: list = []  # injected failures
        self.sessions: dict = {}
        self.chunk_unit = 320 * 1024
        self.simple_max = 4 * 1024 * 1024
        self.redirect_to = None
        self.evil_next = False
        self.evil_upload = False
        self.upload_host = None  # override the uploadUrl host (another tenant)
        self.web_url = f"https://{self.FILES}/sites/Loc/Shared%20Documents"
        self.on_put = None  # hook(item) run just before a replace-in-place is applied
        self.libraries = [
            {"id": self.drive, "name": "Documents"},
            {"id": self.other_drive, "name": "Archive"},
        ]
        self.sites = {f"{self.FILES}:/sites/Loc": SITE_ID}
        self.users = {"alex@contoso.com": self.drive}
        self.bad_secret = False

    # ---- model --------------------------------------------------------------
    def _mk(self, name, parent, folder, iid=None, data=b""):
        self._n += 1
        iid = iid or f"01ITEM{self._n:04d}"
        it = {
            "id": iid,
            "name": name,
            "parent": parent,
            "folder": folder,
            "data": data,
            "cv": 1,
            "ev": 1,
        }
        self.items[iid] = it
        return it

    def _children(self, pid):
        return [i for i in self.items.values() if i["parent"] == pid]

    def _child(self, pid, name):
        for i in self._children(pid):
            if i["name"].lower() == name.lower():
                return i
        return None

    def _find(self, parts):
        cur = self.root
        for p in parts:
            cur = self._child(cur["id"], p)
            if cur is None:
                return None
        return cur

    def _ensure_folder(self, parts):
        cur = self.root
        for p in parts:
            nxt = self._child(cur["id"], p)
            if nxt is None:
                nxt = self._mk(p, cur["id"], True)
            cur = nxt
        return cur

    def _path_parts(self, it):
        out = []
        while it and it["id"] != "ROOT":
            out.append(it["name"])
            it = self.items.get(it["parent"])
        return list(reversed(out))

    def put_file(self, path, data):
        parts = path.split("/")
        parent = self._ensure_folder(parts[:-1])
        cur = self._child(parent["id"], parts[-1])
        if cur is None:
            cur = self._mk(parts[-1], parent["id"], False, data=data)
        else:
            cur["data"], cur["cv"], cur["ev"] = data, cur["cv"] + 1, cur["ev"] + 1
        return cur

    def read(self, path):
        it = self._find(path.split("/"))
        return None if it is None or it["folder"] else it["data"]

    def exists(self, path):
        return self._find(path.split("/")) is not None

    def item(self, path):
        return self._find(path.split("/"))

    def edit(self, path, data):
        """A content edit (a person saves the file): content hash, cTag and eTag all move."""
        it = self._find(path.split("/"))
        it["data"], it["cv"], it["ev"] = data, it["cv"] + 1, it["ev"] + 1

    def touch(self, path):
        """A metadata-only change (rename of a column, a label): eTag moves, content does not."""
        it = self._find(path.split("/"))
        it["ev"] += 1

    def delete(self, path):
        it = self._find(path.split("/"))
        del self.items[it["id"]]

    def files_under(self, prefix):
        out = []
        for it in self.items.values():
            if not it["folder"]:
                p = "/".join(self._path_parts(it))
                if p.startswith(prefix):
                    out.append(p)
        return sorted(out)

    def transport(self):
        return httpx.MockTransport(self._handle)

    def invalidate_tokens(self):
        self.tokens.clear()

    def inject(self, method, contains, status, *, times=1, json_body=None, headers=None):
        self.fail.append(
            {
                "method": method,
                "contains": contains,
                "status": status,
                "times": times,
                "json": json_body,
                "headers": headers or {},
            }
        )

    # ---- rendering ----------------------------------------------------------
    def _json(self, it, select=None):
        parent = self.items.get(it["parent"]) if it["parent"] else None
        ppath = "/".join(
            quote(p, safe="") if self.encode_paths else p
            for p in (self._path_parts(parent) if parent else [])
        )
        d = {
            "id": it["id"],
            "name": it["name"],
            "cTag": f'"c:{it["id"]},{it["cv"]}"',
            "eTag": f'"{it["id"]},{it["ev"]}"',
            "lastModifiedDateTime": f"2026-10-10T10:00:{it['ev'] % 60:02d}Z",
            "parentReference": {
                "driveId": self.drive,
                "id": parent["id"] if parent else None,
                "path": f"/drives/{self.drive}/root:" + (f"/{ppath}" if ppath else ""),
            },
        }
        if it.get("remote"):
            d["remoteItem"] = {"id": "REMOTE", "parentReference": {"driveId": "b!Elsewhere"}}
        if it["folder"]:
            d["folder"] = {"childCount": len(self._children(it["id"]))}
        else:
            d["size"] = len(it["data"])
            d["file"] = {"mimeType": "application/octet-stream"}
            if self.hashes:
                d["file"]["hashes"] = {"quickXorHash": _b64(it["data"])}
        if select:
            keep = {s.strip() for s in select.split(",")} | {"id"}
            d = {k: v for k, v in d.items() if k in keep}
        return d

    @staticmethod
    def _resp(status, body=None, headers=None):
        return httpx.Response(status, json=body if body is not None else {}, headers=headers or {})

    def _err(self, status, code, msg="x"):
        return self._resp(
            status, {"error": {"code": code, "message": msg}}, {"request-id": "req-0001"}
        )

    # ---- dispatch -----------------------------------------------------------
    def _handle(self, req):
        host = req.url.host
        self.log.append((req.method, host, unquote(req.url.path), req.url.query.decode()))
        for f in self.fail:
            if f["times"] > 0 and f["method"] == req.method and f["contains"] in str(req.url):
                f["times"] -= 1
                return httpx.Response(
                    f["status"],
                    json=f["json"] or {"error": {"code": "injected"}},
                    headers=f["headers"],
                )
        if host == self.LOGIN:
            return self._token(req)
        if host == self.FILES:
            return self._files(req)
        if host == self.GRAPH:
            return self._graph(req)
        self.stray_hosts.append(host)
        return httpx.Response(599)

    def _token(self, req):
        form = {k: v[0] for k, v in parse_qs(req.content.decode()).items()}
        if (
            form.get("client_id") != CLIENT
            or form.get("grant_type") != "client_credentials"
            or form.get("scope") != "https://graph.microsoft.com/.default"
            or not req.url.path.startswith(f"/{TENANT}/")
        ):
            return self._resp(400, {"error": "invalid_request", "error_description": "bad form"})
        if self.bad_secret or form.get("client_secret") != FIXTURE_SECRET:
            # a real AAD error never echoes the secret; this one does, to prove it is scrubbed
            return self._resp(
                401,
                {
                    "error": "invalid_client",
                    "error_description": f"AADSTS7000215: Invalid client secret provided {form.get('client_secret')}\r\nTrace ID: x",  # noqa: E501
                },
            )
        self.mints += 1
        tok = f"tok-{self.mints}-fixture"
        self.tokens.add(tok)
        return self._resp(
            200, {"access_token": tok, "expires_in": self.token_ttl, "token_type": "Bearer"}
        )

    def _files(self, req):
        if req.headers.get("authorization"):
            self.leaked.append(str(req.url))
        path = unquote(req.url.path)
        m = re.fullmatch(r"/download/(?P<id>[^/]+)", path)
        if m and req.method == "GET":
            it = self.items.get(m["id"])
            if it is None:
                return self._err(404, "itemNotFound")
            self.downloads.append(it["id"])
            return httpx.Response(200, content=it["data"])
        m = re.fullmatch(r"/upload/(?P<sid>[^/]+)", path)
        if m and req.method == "PUT":
            return self._chunk(m["sid"], req)
        if m and req.method == "DELETE":
            self.sessions.pop(m["sid"], None)
            return httpx.Response(204)
        return self._err(404, "notFound")

    def _chunk(self, sid, req):
        s = self.sessions.get(sid)
        if s is None:
            return self._err(404, "itemNotFound")
        mm = re.fullmatch(r"bytes (\d+)-(\d+)/(\d+)", req.headers.get("content-range", ""))
        if not mm:
            return self._err(400, "invalidRange")
        a, b, total = map(int, mm.groups())
        if (
            a != len(s["buf"])
            or len(req.content) != b - a + 1
            or (total != s["total"] and s["total"])
        ):
            return self._err(416, "invalidRange", "chunk out of order")
        s["total"] = total
        if b + 1 < total and (b - a + 1) % self.chunk_unit != 0:
            return self._err(400, "invalidRange", "chunk size must be a multiple of 320 KiB")
        s["buf"] += req.content
        if len(s["buf"]) < total:
            return self._resp(202, {"nextExpectedRanges": [f"{len(s['buf'])}-"]})
        it = self.items[s["id"]]
        it["data"], it["cv"], it["ev"] = bytes(s["buf"]), it["cv"] + 1, it["ev"] + 1
        self.writes.append(("session", "/".join(self._path_parts(it))))
        return self._resp(200, self._json(it))

    def _graph(self, req):
        auth = req.headers.get("authorization", "")
        if not auth.startswith("Bearer ") or auth[7:] not in self.tokens:
            return self._err(401, "InvalidAuthenticationToken", "token expired")
        path = unquote(req.url.path)
        if not path.startswith("/v1.0/"):
            return self._err(404, "notFound")
        rest = path[5:]
        q = parse_qs(req.url.query.decode())
        select = (q.get("$select") or [""])[0]
        m = re.fullmatch(r"/sites/(?P<host>[^:/]+):(?P<p>/.*)", rest)
        if m:
            sid = self.sites.get(f"{m['host']}:{m['p']}")
            return self._resp(200, {"id": sid}) if sid else self._err(404, "itemNotFound")
        m = re.fullmatch(r"/sites/(?P<sid>[^/]+)/drive", rest)
        if m:
            return (
                self._resp(200, {"id": self.libraries[0]["id"], "name": "Documents"})
                if m["sid"] == SITE_ID
                else self._err(404, "itemNotFound")
            )
        m = re.fullmatch(r"/sites/(?P<sid>[^/]+)/drives", rest)
        if m:
            return (
                self._resp(200, {"value": self.libraries})
                if m["sid"] == SITE_ID
                else self._err(404, "itemNotFound")
            )
        m = re.fullmatch(r"/users/(?P<u>[^/]+)/drive", rest)
        if m:
            d = self.users.get(m["u"])
            return self._resp(200, {"id": d}) if d else self._err(404, "itemNotFound")
        m = re.fullmatch(r"/drives/(?P<d>[^/]+)(?P<rest>/.*)?", rest)
        if not m:
            return self._err(404, "notFound")
        d, sub = m["d"], m["rest"] or ""
        if d in self.denied_drives:
            return self._err(403, "accessDenied", "Access denied")
        if d != self.drive:
            return self._err(404, "itemNotFound")
        return self._drive(req, sub, q, select)

    def _drive(self, req, sub, q, select):
        meth = req.method
        if sub == "" and meth == "GET":
            body = {"id": self.drive}
            if self.web_url:
                body["webUrl"] = self.web_url
            return self._resp(200, body)
        if sub == "/root" and meth == "GET":
            return self._resp(200, self._json(self.root, select))
        m = re.fullmatch(r"/root:/(?P<p>.+?)(?P<content>:/content)?", sub)
        if m:
            parts = m["p"].split("/")
            if meth == "GET" and not m["content"]:
                it = self._find(parts)
                return (
                    self._resp(200, self._json(it, select))
                    if it
                    else self._err(404, "itemNotFound")
                )
            if meth == "PUT" and m["content"]:
                return self._create(req, parts)
        m = re.fullmatch(r"/items/(?P<id>[^/]+)(?P<tail>/.*)?", sub)
        if m:
            it = self.items.get(m["id"])
            tail = m["tail"] or ""
            if it is None:
                return self._err(404, "itemNotFound")
            if not tail and meth == "GET":
                return self._resp(200, self._json(it, select))
            if tail == "/children" and meth == "GET":
                if not it["folder"]:
                    return self._err(400, "invalidRequest", "not a folder")
                kids = sorted(self._children(it["id"]), key=lambda i: i["name"].lower())
                top = int((q.get("$top") or [self.page_size])[0])
                size = min(top, self.page_size)
                start = int((q.get("$skiptoken") or [0])[0])
                body = {"value": [self._json(k, select) for k in kids[start : start + size]]}
                if start + size < len(kids):
                    nxt = f"https://{self.GRAPH}/v1.0/drives/{self.drive}/items/{it['id']}/children?$skiptoken={start + size}&$select={select}"  # noqa: E501
                    body["@odata.nextLink"] = (
                        "https://evil.example/steal?x=1" if self.evil_next else nxt
                    )
                return self._resp(200, body)
            if tail == "/content" and meth == "GET":
                if it["folder"]:
                    return self._err(400, "invalidRequest", "folder")
                loc = self.redirect_to or f"https://{self.FILES}/download/{it['id']}?tempauth=abc"
                return httpx.Response(302, headers={"Location": loc})
            if tail == "/content" and meth == "PUT":
                return self._replace(req, it)
            if tail == "/createUploadSession" and meth == "POST":
                if it["folder"]:
                    return self._err(400, "invalidRequest", "folder")
                if self.on_put:
                    hook, self.on_put = self.on_put, None
                    hook(it)
                want = req.headers.get("if-match")
                if want and want != f'"{it["id"]},{it["ev"]}"':
                    return self._err(412, "preconditionFailed")
                sid = f"S{len(self.sessions) + 1}"
                self.sessions[sid] = {"id": it["id"], "buf": bytearray(), "total": 0}
                url = (
                    "https://evil.example/upload/x"
                    if self.evil_upload
                    else f"https://{self.upload_host or self.FILES}/upload/{sid}"
                )
                return self._resp(200, {"uploadUrl": url})
        return self._err(404, "notFound", f"unhandled {meth} {sub}")

    def _valid_name(self, name):
        return not (
            not name
            or any(c in name for c in '"*:<>?\\|')
            or name != name.strip()
            or name.endswith(".")
            or len(name) > 255
        )

    def _create(self, req, parts):
        if not all(self._valid_name(p) for p in parts):
            return self._err(400, "invalidRequest", "invalid name")
        qs = req.url.query.decode()
        mode = re.search(r"conflictBehavior=(\w+)", qs)
        mode = mode.group(1) if mode else "replace"
        self.creates.append(mode)
        if len(req.content) > self.simple_max:
            return self._err(413, "requestEntityTooLarge")
        existing = self._find(parts)
        if existing is not None and mode == "fail":
            return self._err(409, "nameAlreadyExists")
        if existing is not None and existing["folder"]:
            return self._err(409, "nameAlreadyExists")
        parent = self._ensure_folder(parts[:-1])
        if existing is None:
            existing = self._mk(parts[-1], parent["id"], False, data=req.content)
            self.writes.append(("create", "/".join(parts)))
            return self._resp(201, self._json(existing))
        existing["data"], existing["cv"], existing["ev"] = (
            req.content,
            existing["cv"] + 1,
            existing["ev"] + 1,
        )
        self.writes.append(("replace", "/".join(parts)))
        return self._resp(200, self._json(existing))

    def _replace(self, req, it):
        if it["folder"]:
            return self._err(400, "invalidRequest", "folder")
        if len(req.content) > self.simple_max:
            return self._err(413, "requestEntityTooLarge")
        if self.on_put:
            hook, self.on_put = self.on_put, None
            hook(it)
        want = req.headers.get("if-match")
        if want and want != f'"{it["id"]},{it["ev"]}"':
            return self._err(412, "preconditionFailed")
        it["data"], it["cv"], it["ev"] = req.content, it["cv"] + 1, it["ev"] + 1
        self.writes.append(("replace", "/".join(self._path_parts(it))))
        return self._resp(200, self._json(it))


# ============================================================================
# helpers
# ============================================================================
def mk_docx(*paras):
    d = docx.Document()
    for p in paras:
        d.add_paragraph(p)
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def docx_texts(data):
    return [p.text for p in docx.Document(BytesIO(data)).paragraphs if p.text.strip()]


def mk_xlsx(*rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(list(r))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def xlsx_texts(data):
    wb = openpyxl.load_workbook(BytesIO(data))
    return [c.value for row in wb.active.iter_rows() for c in row if c.value is not None]


BASE_CFG = {
    "tenant_id": TENANT,
    "client_id": CLIENT,
    "site_url": SITE_URL,
    "source_folder": "Source",
    "target_root": "Translations",
    "source_locale": "en-US",
    "target_locales": "fr-FR",
}


def _cfg(extra=None):
    return {**BASE_CFG, **(extra or {})}


def _conn(name, extra=None, secret=None):
    return cstore.create(
        "msgraph", name, _cfg(extra), secret or {"client_secret": FIXTURE_SECRET}, actor="admin"
    )


def _mc(extra=None, fake=None, secret=None):
    fake = fake or FakeGraph()
    c = MsGraphConnector(
        _cfg(extra), secret or {"client_secret": FIXTURE_SECRET}, transport=fake.transport()
    )
    return c, fake


def _run(cid, fake, **kw):
    return cstore.msgraph_translate(cid, actor="auto", transport=fake.transport(), **kw)


def _hashes(cid):
    with Session(engine) as s:
        return dict((s.get(Connection, cid).config or {}).get("msgraph_hashes") or {})


def _raises(fn, *a, **k):
    try:
        fn(*a, **k)
    except cbase.ConnectorError:
        return True
    return False


def _needs_review(mod, only_sources=None):
    """Wrap pretranslate_document so some (or all) segments come back NEEDS_REVIEW."""
    from sqlmodel import select

    from app.models import Segment, SegStatus

    orig = mod.pretranslate_document

    def patched(s, doc_id, loc, *a, **kw):
        out = orig(s, doc_id, loc, *a, **kw)
        for sg in s.exec(select(Segment).where(Segment.document_id == doc_id)).all():
            if only_sources is None or sg.source_text in only_sources:
                sg.status = SegStatus.NEEDS_REVIEW
                s.add(sg)
        s.commit()
        return out

    return orig, patched


class _Capture(logging.Handler):
    def __init__(self):
        super().__init__(level=logging.DEBUG)
        self.lines = []

    def emit(self, record):
        self.lines.append(record.getMessage())


def _outside_target(fake):
    return [w for w in fake.writes if not w[1].startswith("Translations/")]


def main() -> None:
    from app.web._routes import walk_routes

    init_db()
    cap = _Capture()
    logging.getLogger().addHandler(cap)
    logging.getLogger().setLevel(logging.DEBUG)

    # ---- (1) capability, test(), drive resolution in every mode ---------------------
    assert MsGraphConnector.capability.topology == "T3"
    mc, fk = _mc()
    fk.put_file("Source/a.docx", mk_docx("Hello"))
    ok, msg = mc.test()
    assert ok, msg
    assert mc._base() == "https://graph.microsoft.com/v1.0"
    assert fk.mints == 1 and not fk.stray_hosts
    # library by name, default library, drive id, onedrive user, site id
    mcl, fkl = _mc({"library": "documents"})
    fkl.put_file("Source/a.docx", mk_docx("Hi"))
    assert mcl.test()[0], "library name match is case-insensitive"
    mcl2, _ = _mc({"library": "Nope"}, fake=fkl)
    ok, msg = mcl2.test()
    assert not ok and "Documents" in msg and "Archive" in msg, ("lists the libraries", msg)
    for extra in (
        {"site_url": "", "drive_id": fk.drive},
        {"site_url": "", "onedrive_user": "alex@contoso.com"},
        {"site_url": "", "site_id": SITE_ID},
    ):
        mcx, _ = _mc(extra, fake=fk)
        assert mcx.test()[0], extra
    # exactly one drive choice
    for extra in ({"drive_id": fk.drive}, {"site_url": ""}):
        ok, msg = _mc(extra, fake=fk)[0].test()
        assert not ok and "exactly one" in msg, (extra, msg)
    print(
        "  connector: T3, test(); site URL / library / default / drive id / OneDrive user / site id; exactly-one rule"  # noqa: E501
    )

    # ---- (2) operator config strings are validated fail-closed (taxonomy A) ---------
    bad_cfgs = [
        {"source_folder": ".."},
        {"source_folder": "a/../b"},
        {"source_folder": "a\\b"},
        {"source_folder": "a//b"},
        {"source_folder": "%2e%2e/x"},
        {"source_folder": "a#b"},
        {"source_folder": "a\x00b"},
        {"source_folder": "a\r\nb"},
        {"source_folder": "/"},
        {"source_folder": ""},
        {"source_folder": "a?b"},
        {"source_folder": "a:b"},
        {"target_root": ".."},
        {"target_root": "../x"},
        {"target_root": "x/./y"},
        {"target_root": "Source"},
        {"target_root": "source/Sub"},
        {"source_folder": "Translations/in"},
        {"source_folder": "Source/Deep", "target_root": "Source"},
        {"tenant_id": "../x"},
        {"tenant_id": "a b"},
        {"tenant_id": "x?y"},
        {"client_id": "a/b"},
        {"site_url": "http://contoso.sharepoint.com/sites/Loc"},
        {"site_url": "https://evil.example/sites/Loc"},
        {"site_url": "https://contoso.sharepoint.com.evil.example/sites/Loc"},
        {"site_url": "https://user:pw@contoso.sharepoint.com/sites/Loc"},
        {"site_url": "https://contoso.sharepoint.com/sites/Loc?x=1"},
        {"site_url": "https://contoso.sharepoint.com/sites/../x"},
        {"site_url": "https://contoso.sharepoint.com:8443/sites/Loc"},
        {"site_url": "", "site_id": "../x"},
        {"site_url": "", "drive_id": "a/b"},
        {"site_url": "", "drive_id": "b!x?y=1"},
        {"site_url": "", "onedrive_user": "../x"},
        {"extensions": "ppt"},
        {"extensions": "docx, ../x"},
    ]
    for extra in bad_cfgs:
        mcb, fkb = _mc(extra)
        ok, msg = mcb.test()
        assert not ok, ("refused", extra, msg)
        assert not fkb.writes, extra
        r = False
        try:
            mcb.discover_articles()
        except cbase.ConnectorError:
            r = True
        assert r, ("discovery refuses", extra)
        assert not [x for x in fkb.log if x[1] == fkb.GRAPH and "/drives/" in x[2]], (
            "no drive request",
            extra,
        )
        if "tenant_id" in extra or "client_id" in extra:
            assert not fkb.log, (
                "a hostile tenant / client id sends nothing at all",
                extra,
                fkb.log,
            )
    print(f"  A: {len(bad_cfgs)} hostile operator configs refused before any drive request")

    # ---- (3) D: classification (both directions) + pagination + recursion ----------
    fk3 = FakeGraph(page_size=2)
    for i in range(5):
        fk3.put_file(f"Source/f{i}.docx", mk_docx(f"Text {i}"))
    fk3.put_file("Source/Sub/Deep/x.html", b"<p>Hello deep</p>")
    fk3.put_file("Source/notes.md", b"# Title\n\nSome text\n")
    fk3.put_file("Source/deck.ppt", b"legacy-binary-deck")
    fk3.put_file("Source/file.pdf", b"%PDF-1.4")
    fk3.put_file("Source/~$lock.docx", b"x")
    fk3.put_file("Source/empty.txt", b"")
    fk3.put_file("Source/big.docx", mk_docx("Big") + b"\0" * (1100 * 1024))
    fk3.put_file("Source/we#ird.docx", mk_docx("Hash name"))
    fk3.put_file("Source/data.bin", b"zz")
    fk3.put_file("Translations/fr-FR/old.docx", mk_docx("Not source"))
    fk3.put_file("Other/private.docx", mk_docx("Private"))
    mc3, _ = _mc({"max_file_mb": "1"}, fake=fk3)
    arts = mc3.discover_articles()
    rels = sorted(a["relpath"] for a in arts)
    assert rels == sorted([f"f{i}.docx" for i in range(5)] + ["Sub/Deep/x.html", "notes.md"]), rels
    assert not fk3.downloads, "discovery downloads nothing"
    sk = mc3.sel.skipped_dict()
    for name, why in [
        ("deck.ppt", "legacy .ppt"),
        ("file.pdf", "PDF"),
        ("~$lock.docx", "lock"),
        ("empty.txt", "empty"),
        ("big.docx", "limit"),
        ("we#ird.docx", "cannot carry"),
        ("data.bin", "no file filter"),
    ]:
        assert name in sk and why in sk[name], (name, sk.get(name))
    assert not any("Translations" in a["relpath"] or "private" in a["relpath"] for a in arts)
    assert "no file filter" in mc3.sel.selection_summary(), "loud skip leads the summary"
    top, _ = _mc({"include_subfolders": "0"}, fake=fk3)
    assert all("/" not in a["relpath"] for a in top.discover_articles())
    only, _ = _mc({"extensions": "docx"}, fake=fk3)
    assert {a["fmt"] for a in only.discover_articles()} == {"docx"}
    print(
        "  D: formats selected; ppt/pdf/lock/empty/oversize/bad-name/unknown skipped WITH reasons; paging + recursion; nothing downloaded"  # noqa: E501
    )

    # ---- (4) E: the Max files cap round-robins across sub-folders -------------------
    fk4 = FakeGraph()
    for i in range(20):
        fk4.put_file(f"Source/A/a{i:02d}.docx", mk_docx(f"a{i}"))
    for i in range(3):
        fk4.put_file(f"Source/B/b{i}.docx", mk_docx(f"b{i}"))
        fk4.put_file(f"Source/C/c{i}.docx", mk_docx(f"c{i}"))
    mc4, _ = _mc({"max_files": "9"}, fake=fk4)
    got = mc4.discover_articles()
    per = {}
    for a in got:
        per[a["rel"][0]] = per.get(a["rel"][0], 0) + 1
    assert len(got) == 9 and per == {"A": 3, "B": 3, "C": 3}, (
        "E: distribution, not just total",
        per,
    )
    print(f"  E: cap 9 over folders of 20/3/3 -> {per} (no folder starved)")

    # ---- (5) full flow: create, mirrored tree, nothing outside the target root -----
    fk5 = FakeGraph()
    d_doc = mk_docx("Hello world", "Second paragraph")
    x_book = mk_xlsx(("Name", "Price"), ("Widget", 5))
    fk5.put_file("Source/doc.docx", d_doc)
    fk5.put_file("Source/Sub/book.xlsx", x_book)
    fk5.put_file("Source/Sub/page.html", b"<html><body><p>Hello page</p></body></html>")
    fk5.put_file("Source/readme.md", b"# Title\n\nBody text here\n")
    src_snapshot = {p: fk5.read(p) for p in fk5.files_under("Source")}
    src_tags = {p: fk5.item(p)["cv"] for p in fk5.files_under("Source")}
    c5 = _conn("flow", {"target_locales": "fr-FR, pt-br"})
    ok, msg, s5 = _run(c5.id, fk5)
    assert ok, msg
    assert s5["created"] == 8 and s5["updated"] == 0 and not s5["failures"], s5
    tree = fk5.files_under("Translations")
    assert tree == sorted(
        f"Translations/{loc}/{p}"
        for loc in ("fr-FR", "pt-BR")
        for p in ("doc.docx", "readme.md", "Sub/book.xlsx", "Sub/page.html")
    ), tree
    fr = docx_texts(fk5.read("Translations/fr-FR/doc.docx"))
    assert len(fr) == 2 and all(t.lower().startswith("[fr") for t in fr), fr
    assert all(
        t.lower().startswith("[pt") for t in docx_texts(fk5.read("Translations/pt-BR/doc.docx"))
    )
    assert any(
        str(t).lower().startswith("[fr")
        for t in xlsx_texts(fk5.read("Translations/fr-FR/Sub/book.xlsx"))
    )
    assert b"[fr" in fk5.read("Translations/fr-FR/Sub/page.html").lower()
    assert {p: fk5.read(p) for p in fk5.files_under("Source")} == src_snapshot, (
        "source bytes untouched"
    )
    assert {p: fk5.item(p)["cv"] for p in fk5.files_under("Source")} == src_tags, (
        "source never written"
    )
    assert not _outside_target(fk5), ("every write is under the target root", _outside_target(fk5))
    assert set(fk5.creates) == {"fail"}, ("creates never replace silently", fk5.creates)
    assert not fk5.leaked and not fk5.stray_hosts, (fk5.leaked, fk5.stray_hosts)
    assert len(_hashes(c5.id)) == 8
    rows = idmap.list_for_connection(c5.id)
    assert len(rows) == 8 and all(r.status == idmap.DONE and r.target_id for r in rows), rows
    print(
        "  flow: docx/xlsx/html/md x 2 locales -> mirrored per-locale copies, source untouched, id map done, no token to the file host"  # noqa: E501
    )

    # ---- (6) C: second run is metadata-only and writes nothing ---------------------
    fk5.downloads.clear()
    fk5.writes.clear()
    n_log = len(fk5.log)
    ok, msg, s6 = _run(c5.id, fk5)
    assert ok and s6["created"] == 0 and s6["updated"] == 0 and s6["skipped"] == 8, s6
    assert not fk5.downloads and not fk5.writes, (
        "C: nothing downloaded or written",
        fk5.downloads,
        fk5.writes,
    )
    assert not [x for x in fk5.log[n_log:] if x[0] in ("PUT", "POST") and x[1] != fk5.LOGIN], (
        "no mutating request"
    )
    print("  C: run 2 over unchanged source -> 8 skipped, 0 downloads, 0 writes")

    # metadata-only touch is not a change; a content edit re-translates just that file
    fk5.touch("Source/doc.docx")
    ok, msg, s6b = _run(c5.id, fk5)
    assert s6b["skipped"] == 8 and not fk5.downloads, ("metadata touch is not a change", s6b)
    fk5.edit("Source/doc.docx", mk_docx("Hello world CHANGED", "Second paragraph"))
    ok, msg, s6c = _run(c5.id, fk5)
    assert ok and s6c["updated"] == 2 and s6c["skipped"] == 6 and s6c["created"] == 0, s6c
    assert fk5.downloads == [fk5.item("Source/doc.docx")["id"]], (
        "only the changed file downloaded",
        fk5.downloads,
    )
    assert "CHANGED" in " ".join(docx_texts(fk5.read("Translations/fr-FR/doc.docx")))
    print(
        "  C: metadata touch ignored; a content edit re-translates only that file (2 copies updated)"  # noqa: E501
    )

    # no content hash from Graph: the cTag carries the same guarantees
    fk6 = FakeGraph(hashes=False)
    fk6.put_file("Source/a.docx", mk_docx("Alpha"))
    c6 = _conn("ctag")
    assert _run(c6.id, fk6)[2]["created"] == 1
    fk6.touch("Source/a.docx")
    r = _run(c6.id, fk6)[2]
    assert r["skipped"] == 1 and r["updated"] == 0, ("cTag identity: touch ignored", r)
    fk6.edit("Source/a.docx", mk_docx("Alpha two"))
    r = _run(c6.id, fk6)[2]
    assert r["updated"] == 1, ("cTag identity: edit detected", r)
    print("  C: with no content hash from Graph the cTag identity behaves the same")

    # ---- (7) B/F: held-back never creates or changes a copy -------------------------
    import app.services.translate as T

    fk7 = FakeGraph()
    fk7.put_file("Source/h.docx", mk_docx("Keep me", "And me"))
    c7 = _conn("held")
    orig, patched = _needs_review(T)  # EVERY segment held
    set_pieces("app.services.translate", "pretranslate_document", patched)
    try:
        ok, msg, s7 = _run(c7.id, fk7)
    finally:
        set_pieces("app.services.translate", "pretranslate_document", orig)
    assert ok is not None and s7["created"] == 0 and s7["updated"] == 0, s7
    assert not fk7.files_under("Translations"), (
        "B: no placeholder or copy created",
        fk7.files_under("Translations"),
    )
    assert not fk7.writes, fk7.writes
    assert s7["held_back"] == 1 and not s7["failures"], s7
    assert not _hashes(c7.id), ("F: no hash recorded for a held pair", _hashes(c7.id))
    ok, msg, s7b = _run(c7.id, fk7)
    assert ok and s7b["created"] == 1, ("F: a healthy re-run completes it", s7b)
    assert len(docx_texts(fk7.read("Translations/fr-FR/h.docx"))) == 2
    # PARTIAL hold on an EXISTING copy: the copy keeps every stored translation
    fk7.edit("Source/h.docx", mk_docx("Keep me NEW", "And me"))
    before = fk7.read("Translations/fr-FR/h.docx")
    h_before = dict(_hashes(c7.id))
    orig, patched = _needs_review(T, only_sources={"And me"})
    set_pieces("app.services.translate", "pretranslate_document", patched)
    try:
        ok, msg, s7c = _run(c7.id, fk7)
    finally:
        set_pieces("app.services.translate", "pretranslate_document", orig)
    assert fk7.read("Translations/fr-FR/h.docx") == before, (
        "B: partial hold leaves the stored copy exactly as it was"
    )
    assert s7c["updated"] == 0 and s7c["held_back"] == 1, s7c
    assert _hashes(c7.id) == h_before, ("F: hash NOT advanced on a held pair", _hashes(c7.id))
    ok, msg, s7d = _run(c7.id, fk7)
    assert (
        ok
        and s7d["updated"] == 1
        and "NEW" in " ".join(docx_texts(fk7.read("Translations/fr-FR/h.docx")))
    ), s7d
    # a NEW locale variant is not created while a unit is held either
    c7b = _conn("held2", {"target_locales": "fr-FR, de-DE"})
    orig, patched = _needs_review(T, only_sources={"And me"})
    set_pieces("app.services.translate", "pretranslate_document", patched)
    try:
        _run(c7b.id, fk7)
    finally:
        set_pieces("app.services.translate", "pretranslate_document", orig)
    assert not fk7.exists("Translations/de-DE/h.docx"), (
        "B: no partially translated variant is created"
    )
    print(
        "  B/F: all-held -> nothing created; partial hold on new AND existing copy -> not written, copy intact, no hash; healthy re-run completes"  # noqa: E501
    )

    # ---- (8) B: a copy edited by hand is never overwritten --------------------------
    fk8 = FakeGraph()
    fk8.put_file("Source/e.docx", mk_docx("Edit target"))
    c8 = _conn("edited")
    assert _run(c8.id, fk8)[2]["created"] == 1
    human = mk_docx("Version corrigee par un humain")
    fk8.edit("Translations/fr-FR/e.docx", human)
    fk8.edit("Source/e.docx", mk_docx("Edit target v2"))
    calls = []
    real_pre = T.pretranslate_document

    def counting(*a, **kw):
        calls.append(1)
        return real_pre(*a, **kw)

    set_pieces("app.services.translate", "pretranslate_document", counting)
    try:
        fk8.writes.clear()
        ok, msg, s8 = _run(c8.id, fk8)
    finally:
        set_pieces("app.services.translate", "pretranslate_document", real_pre)
    assert fk8.read("Translations/fr-FR/e.docx") == human, "B: the human's copy is untouched"
    assert not fk8.writes and not calls, ("zero writes and ZERO MT spend", fk8.writes, calls)
    assert s8["edited_copies"] == [
        {"article_id": fk8.item("Source/e.docx")["id"], "locale": "fr-fr"}
    ], s8
    assert "edited in SharePoint" in msg and ok, msg
    # a person deleting the copy regenerates it
    fk8.delete("Translations/fr-FR/e.docx")
    ok, msg, s8b = _run(c8.id, fk8)
    assert (
        ok
        and fk8.exists("Translations/fr-FR/e.docx")
        and "v2" in " ".join(docx_texts(fk8.read("Translations/fr-FR/e.docx")))
    ), s8b
    # unchanged source + deleted copy is also regenerated (the probe reads the copy)
    fk8.delete("Translations/fr-FR/e.docx")
    ok, msg, s8c = _run(c8.id, fk8)
    assert fk8.exists("Translations/fr-FR/e.docx") and s8c["skipped"] == 0, (
        "deleted copy is not suppressed",
        s8c,
    )
    # the explicit retry path skips the pre-check, so the write itself must still refuse
    fk8.edit("Translations/fr-FR/e.docx", human)
    fk8.edit("Source/e.docx", mk_docx("Edit target v2b"))
    sid8 = fk8.item("Source/e.docx")["id"]
    ok, msg, s8r = _run(c8.id, fk8, only_pairs={(sid8, "fr-fr")})
    assert fk8.read("Translations/fr-FR/e.docx") == human, (
        "the write-time check protects a retried pair too"
    )
    assert any("edited-copy" in f["error"] for f in s8r["failures"]), s8r
    # a probe that could not tell (503s) must not turn into permission: the write re-checks
    copy_id = fk8.item("Translations/fr-FR/e.docx")["id"]
    fk8.inject("GET", f"/items/{copy_id}?", 503, times=3)
    saved_sleep = cbase._sleep
    cbase._sleep = lambda s: None
    try:
        fk8.writes.clear()
        ok, msg, s8u = _run(c8.id, fk8)
    finally:
        cbase._sleep = saved_sleep
    assert fk8.read("Translations/fr-FR/e.docx") == human and not fk8.writes, (
        "unknown probe still protects",
        s8u,
    )
    assert any("edited-copy" in f["error"] for f in s8u["failures"]), s8u
    fk8.delete("Translations/fr-FR/e.docx")
    ok, msg, s8b2 = _run(c8.id, fk8)
    assert fk8.exists("Translations/fr-FR/e.docx"), s8b2
    # an edit that lands BETWEEN our check and our write loses the race (If-Match 412)
    fk8.edit("Source/e.docx", mk_docx("Edit target v3"))
    racer = mk_docx("Raced human edit")
    fk8.on_put = lambda it: it.update(data=racer, ev=it["ev"] + 1, cv=it["cv"] + 1)
    ok, msg, s8d = _run(c8.id, fk8)
    assert fk8.read("Translations/fr-FR/e.docx") == racer, (
        "412 protects an edit made after the check"
    )
    assert any("edited-copy" in f["error"] for f in s8d["failures"]), s8d
    # the protection can be switched off (live QA fallback)
    c8b = _conn("unprotected", {"protect_edited_copies": "0"})
    fk8b = FakeGraph()
    fk8b.put_file("Source/e.docx", mk_docx("One"))
    _run(c8b.id, fk8b)
    fk8b.edit("Translations/fr-FR/e.docx", human)
    fk8b.edit("Source/e.docx", mk_docx("Two"))
    ok, msg, s8e = _run(c8b.id, fk8b)
    assert s8e["updated"] == 1 and fk8b.read("Translations/fr-FR/e.docx") != human
    print(
        "  B: hand-edited copy untouched (0 writes, 0 MT calls), delete regenerates, If-Match 412 race safe, switch-off works"  # noqa: E501
    )

    # the same protection when Graph returns no content hash (the cTag decides)
    fk8c = FakeGraph(hashes=False)
    fk8c.put_file("Source/e.docx", mk_docx("One"))
    c8c = _conn("edited-ctag")
    _run(c8c.id, fk8c)
    fk8c.edit("Translations/fr-FR/e.docx", human)
    fk8c.edit("Source/e.docx", mk_docx("Two"))
    ok, msg, s8f = _run(c8c.id, fk8c)
    assert fk8c.read("Translations/fr-FR/e.docx") == human and s8f["updated"] == 0, s8f
    assert copy_changed("q:A|c:1", "q:A|c:2") is False, "a content hash outranks the cTag"
    assert copy_changed("q:A|c:1", "q:B|c:1") is True
    assert copy_changed("c:1", "c:2") is True
    assert copy_changed("c:1", "q:B") is False and copy_changed("", "c:1") is False, (
        "unknown is not an edit"
    )
    print(
        "  B: the edited-copy check also works on the cTag alone; unknown identities are never an edit"  # noqa: E501
    )

    # ---- (9) interrupted builds resume; a human file at the path fails closed -------
    fk9 = FakeGraph()
    fk9.put_file("Source/r.docx", mk_docx("Resume me"))
    c9 = _conn("resume")
    fk9.inject(
        "PUT", "/content", 503, times=3
    )  # the fill (not the create) fails: placeholder remains
    # the first PUT is the create (":/content?..."), so let it through and fail the fill only
    fk9.fail.clear()
    fk9.inject("PUT", "/items/", 503, times=3)
    ok, msg, s9 = _run(c9.id, fk9)
    ph = fk9.item("Translations/fr-FR/r.docx")
    assert ph is not None and ph["data"] == b"" and s9["failures"], (
        "placeholder left, pair failed",
        s9,
    )
    assert not _hashes(c9.id), "F: a failed pair records no hash"
    row = idmap.list_for_connection(c9.id)[0]
    assert row.status == idmap.BUILDING
    n_creates = len(fk9.creates)
    ok, msg, s9b = _run(c9.id, fk9)
    assert ok and s9b["created"] == 1 and len(fk9.creates) == n_creates, (
        "resumed, not re-created",
        s9b,
    )
    assert (
        fk9.read("Translations/fr-FR/r.docx") != b"" and len(fk9.files_under("Translations")) == 1
    )
    # residue with NO map row (a crash before the reserve): the 0-byte file is resumed
    fk9b = FakeGraph()
    fk9b.put_file("Source/r.docx", mk_docx("Resume me too"))
    fk9b.put_file("Translations/fr-FR/r.docx", b"")
    c9b = _conn("residue")
    ok, msg, s9c = _run(c9b.id, fk9b)
    assert ok and s9c["created"] == 1 and fk9b.read("Translations/fr-FR/r.docx"), s9c
    # a HUMAN-made copy at the path, no map row: fail closed, never overwritten
    fk9c = FakeGraph()
    fk9c.put_file("Source/n.docx", mk_docx("Needs link"))
    theirs = mk_docx("Une traduction humaine existante")
    fk9c.put_file("Translations/fr-FR/n.docx", theirs)
    c9c = _conn("needslink")
    ok, msg, s9d = _run(c9c.id, fk9c)
    assert fk9c.read("Translations/fr-FR/n.docx") == theirs and s9d["created"] == 0, (
        "never overwritten"
    )
    assert any("needs_link" in f["error"] for f in s9d["failures"]), s9d
    sid = fk9c.item("Source/n.docx")["id"]
    assert idmap.resolve_needs_link(c9c.id, sid, "fr-fr", "adopt")[0]
    ok, msg, s9e = _run(c9c.id, fk9c)
    assert s9e["updated"] == 1 and fk9c.read("Translations/fr-FR/n.docx") != theirs, (
        "adopt then update in place",
        s9e,
    )
    print(
        "  resume: failed fill leaves a 0-byte placeholder -> next run resumes it (no duplicate create); crash residue resumed; human copy fails closed (needs_link) and adopt updates it"  # noqa: E501
    )

    # ---- (10) A: caller article_ids cannot escape the source folder ------------------
    fk10 = FakeGraph()
    fk10.put_file("Source/ok.docx", mk_docx("In scope"))
    fk10.put_file("Source/Sub/ok2.docx", mk_docx("In scope sub"))
    fk10.put_file("SourceOld/sib.docx", mk_docx("Prefix sibling"))
    fk10.put_file("Other/private.docx", mk_docx("Private"))
    fk10.put_file("Translations/fr-FR/t.docx", mk_docx("A copy"))
    sub_id = fk10.item("Source/Sub/ok2.docx")["id"]
    outside = {
        "private": fk10.item("Other/private.docx")["id"],
        "prefix sibling": fk10.item("SourceOld/sib.docx")["id"],
        "a target copy": fk10.item("Translations/fr-FR/t.docx")["id"],
        "a folder": fk10.item("Source/Sub")["id"],
        "the root": "ROOT",
    }
    hostile = [
        "..",
        "../x",
        "%2e%2e%2f",
        "a/b",
        "a?x=1",
        "a#f",
        "a\x00b",
        "a\r\nb",
        "a b",
        "",
        "ïïd",
        "x" * 201,
        "ok.docx",
        "Source/ok.docx",
        "01ITEM0001/../x",
        "ITEM%00",
        "a;b",
    ]
    c10 = _conn("scope")
    for label, bad in [*outside.items(), *[(repr(h), h) for h in hostile]]:
        fk10.writes.clear()
        fk10.downloads.clear()
        ok, msg, s10 = _run(c10.id, fk10, article_ids=[bad])
        assert not ok, (label, msg)
        assert not fk10.writes and not fk10.downloads, (label, "ZERO writes, ZERO downloads")
        assert not fk10.stray_hosts
        assert "\x00" not in msg and "\r" not in msg, ("message is log-safe", repr(msg))
    for bad in hostile:
        assert _raises(validate_item_id, bad), repr(bad)
    # a foreign DRIVE id: the item answers like a missing one
    ok, msg, _s = _run(c10.id, fk10, article_ids=["b!OtherDriveItem"])
    assert not fk10.writes
    ok, msg, s10 = _run(c10.id, fk10, article_ids=[sub_id])
    assert ok and s10["created"] == 1 and fk10.exists("Translations/fr-FR/Sub/ok2.docx"), (msg, s10)
    assert not fk10.exists("Translations/fr-FR/ok.docx"), "only the asked-for file was translated"
    print(
        f"  A: {len(hostile) + len(outside)} hostile / out-of-scope ids refused with zero writes and zero downloads; a valid nested id works and mirrors"  # noqa: E501
    )

    # ---- (10b) names with spaces / accents / ampersands, paths percent-encoded by Graph ---
    fk10b = FakeGraph(encode_paths=True)
    fk10b.put_file("Source Docs/Sub Folder/Résumé & notes.docx", mk_docx("Accented name"))
    fk10b.put_file("Source Docs Old/x.docx", mk_docx("Sibling with a shared prefix"))
    c10b = _conn("names", {"source_folder": "Source Docs"})
    ok, msg, s10b = _run(c10b.id, fk10b)
    assert ok and s10b["created"] == 1, (msg, s10b)
    assert fk10b.exists("Translations/fr-FR/Sub Folder/Résumé & notes.docx"), fk10b.files_under(
        "Translations"
    )
    rid = fk10b.item("Source Docs/Sub Folder/Résumé & notes.docx")["id"]
    sib = fk10b.item("Source Docs Old/x.docx")["id"]
    ok, msg, _s = _run(c10b.id, fk10b, article_ids=[rid])
    assert ok, ("encoded parentReference.path still scope-checks correctly", msg)
    n_w = len(fk10b.writes)
    ok, msg, _s = _run(c10b.id, fk10b, article_ids=[sib])
    assert not ok and len(fk10b.writes) == n_w, (
        "a sibling folder sharing the prefix is out of scope",
        msg,
    )
    print(
        "  names: spaces, accents and & round-trip; a percent-encoded parent path is scope-checked; a prefix-sharing sibling folder is refused"  # noqa: E501
    )

    # ---- (11) A: hostile links from Graph itself -----------------------------------
    fk11 = FakeGraph(page_size=1)
    fk11.put_file("Source/a.docx", mk_docx("A"))
    fk11.put_file("Source/b.docx", mk_docx("B"))
    fk11.evil_next = True
    mc11, _ = _mc(fake=fk11)
    assert _raises(mc11.discover_articles), "nextLink outside Graph is refused"
    assert not fk11.stray_hosts and not any(x[1] == "evil.example" for x in fk11.log)
    fk12 = FakeGraph()
    fk12.put_file("Source/a.docx", mk_docx("A"))
    fk12.redirect_to = "https://evil.example/exfil?f=1"
    c12 = _conn("redirect")
    ok, msg, s12 = _run(c12.id, fk12)
    assert not fk12.stray_hosts and s12["created"] == 0 and s12["unreadable"] == 1, (msg, s12)
    assert not fk12.leaked, "no bearer token ever reaches a file host"
    fk13 = FakeGraph()
    fk13.put_file("Source/a.docx", mk_docx("A"))
    set_pieces("app.services.connectors.msgraph", "SIMPLE_UPLOAD_MAX", 10)
    fk13.evil_upload = True
    c13 = _conn("evilupload")
    try:
        ok, msg, s13 = _run(c13.id, fk13)
    finally:
        set_pieces("app.services.connectors.msgraph", "SIMPLE_UPLOAD_MAX", 4 * 1024 * 1024)
    assert not fk13.stray_hosts and s13["failures"], (msg, s13)
    print(
        "  A: hostile nextLink / download redirect / upload URL refused, no request to the stray host, token never sent to a file host"  # noqa: E501
    )

    # ---- (12) token handling: remint on 401, bad secret, never logged ----------------
    fk14 = FakeGraph()
    fk14.put_file("Source/a.docx", mk_docx("A"))
    c14 = _conn("token")
    fk14.invalidate_tokens()
    ok, msg, s14 = _run(c14.id, fk14)
    assert ok and s14["created"] == 1, (msg, s14)
    mc14, fk14b = _mc(fake=fk14)
    mc14._token(mc14._client().__enter__())
    fk14b.invalidate_tokens()
    assert mc14.test()[0], "an expired token is minted again once"
    assert fk14.mints >= 2
    badk = FakeGraph()
    badk.put_file("Source/a.docx", mk_docx("A"))
    badk.bad_secret = True
    cbad = _conn("badsecret")
    ok, msg, _sbad = _run(cbad.id, badk)
    assert not ok and "sign-in failed" in msg and FIXTURE_SECRET not in msg, (
        "H: secret scrubbed from the error",
        msg,
    )
    ok, m2 = _mc(fake=badk)[0].test()
    assert not ok and FIXTURE_SECRET not in m2 and "AADSTS7000215" in m2, m2
    blob = "\n".join(cap.lines)
    assert FIXTURE_SECRET not in blob, "H: the client secret is in no log line"
    assert not re.search(r"tok-\d+-fixture", blob), "H: no access token in any log line"
    for fk_ in (fk5, fk14):
        assert not fk_.leaked
    print(
        "  H: 401 -> one remint; wrong secret -> clean message with the secret scrubbed; secret and token in no log line"  # noqa: E501
    )

    # ---- (13) vendor errors, rate limits, retries --------------------------------------
    fk15 = FakeGraph()
    fk15.put_file("Source/a.docx", mk_docx("A"))
    fk15.denied_drives.add(fk15.drive)
    ok, msg = _mc(fake=fk15)[0].test()
    assert not ok and "Sites.Selected" in msg and "accessDenied" in msg and "request-id" in msg, msg
    ok, msg, _s = _run(_conn("denied").id, fk15)
    assert not ok and "Sites.Selected" in msg, msg
    fk16 = FakeGraph()
    ok, msg = _mc(fake=fk16)[0].test()
    assert not ok and "not found" in msg, ("missing source folder is named", msg)
    slept = []
    real_sleep = cbase._sleep
    cbase._sleep = lambda s: slept.append(s)
    try:
        fk17 = FakeGraph()
        fk17.put_file("Source/a.docx", mk_docx("A"))
        fk17.inject("GET", "/children", 429, times=2, headers={"Retry-After": "3"})
        fk17.inject("GET", "/download/", 503, times=1)
        ok, msg, s17 = _run(_conn("throttle").id, fk17)
        assert ok and s17["created"] == 1, (msg, s17)
        assert slept and 3.0 in slept, ("Retry-After honoured through base._request", slept)
        fk18 = FakeGraph()
        fk18.put_file("Source/a.docx", mk_docx("A"))
        fk18.inject("PUT", "conflictBehavior=fail", 503, times=1)
        ok, msg, s18 = _run(_conn("nocreateretry").id, fk18)
        assert fk18.creates == [] and s18["failures"], (
            "a create is NOT blindly retried on 5xx",
            s18,
        )
        fk19 = FakeGraph()
        fk19.put_file("Source/a.docx", mk_docx("A"))
        fk19.inject(
            "PUT",
            "conflictBehavior=fail",
            507,
            times=1,
            json_body={"error": {"code": "quotaLimitReached", "message": "full"}},
        )
        ok, msg, s19 = _run(_conn("quota").id, fk19)
        assert any("quotaLimitReached" in f["error"] for f in s19["failures"]), s19
    finally:
        cbase._sleep = real_sleep
    print(
        "  vendor errors: 403 names Sites.Selected + request-id, missing folder named, 429 Retry-After + 5xx retried, create not blindly retried, 507 reported"  # noqa: E501
    )

    # ---- (14) large files use an upload session in aligned chunks --------------------
    fk20 = FakeGraph()
    fk20.put_file(
        "Source/a.docx", mk_docx(*[f"Paragraph number {i} with some words" for i in range(60)])
    )
    fk20.chunk_unit = 1024
    set_pieces("app.services.connectors.msgraph", "SIMPLE_UPLOAD_MAX", 1000)
    set_pieces("app.services.connectors.msgraph", "CHUNK", 4096)
    fk20.simple_max = 1000
    try:
        c20 = _conn("session")
        ok, msg, s20 = _run(c20.id, fk20)
    finally:
        set_pieces("app.services.connectors.msgraph", "SIMPLE_UPLOAD_MAX", 4 * 1024 * 1024)
        set_pieces("app.services.connectors.msgraph", "CHUNK", 320 * 1024 * 10)
    assert ok and s20["created"] == 1, (msg, s20)
    assert any(w[0] == "session" for w in fk20.writes), (
        "session used above the simple cap",
        fk20.writes,
    )
    got20 = fk20.read("Translations/fr-FR/a.docx")
    assert len(docx_texts(got20)) == 60 and not fk20.leaked, (
        "multi-chunk upload reassembled correctly"
    )
    print(
        "  upload: a file over the simple cap goes through an upload session in aligned chunks, bytes intact"  # noqa: E501
    )

    # ---- (14b) review fix: the upload session carries If-Match (lost race on a large file) ----
    fk20.edit(
        "Source/a.docx", mk_docx(*[f"Edited paragraph {i} with some words" for i in range(60)])
    )
    racer20 = mk_docx("Raced human edit of a large copy")
    fk20.on_put = lambda it: it.update(data=racer20, ev=it["ev"] + 1, cv=it["cv"] + 1)
    set_pieces("app.services.connectors.msgraph", "SIMPLE_UPLOAD_MAX", 1000)
    set_pieces("app.services.connectors.msgraph", "CHUNK", 4096)
    n_sess = len([w for w in fk20.writes if w[0] == "session"])
    try:
        ok, msg, s20b = _run(c20.id, fk20)
    finally:
        set_pieces("app.services.connectors.msgraph", "SIMPLE_UPLOAD_MAX", 4 * 1024 * 1024)
        set_pieces("app.services.connectors.msgraph", "CHUNK", 320 * 1024 * 10)
    assert fk20.read("Translations/fr-FR/a.docx") == racer20, "session 412 kept the human edit"
    assert len([w for w in fk20.writes if w[0] == "session"]) == n_sess, "no session write landed"
    assert any("edited-copy" in f["error"] for f in s20b["failures"]), s20b
    print("  upload: createUploadSession carries If-Match; a race on a large copy loses safely")

    # ---- (14c) review fix: source files are never written, whatever the id map or tree says ----
    fk22 = FakeGraph()
    fk22.put_file("Source/a.docx", mk_docx("Source text"))
    c22 = _conn("srcguard", {"protect_edited_copies": "0"})
    ok, msg, s22 = _run(c22.id, fk22)
    assert ok and s22["created"] == 1
    src_before = fk22.read("Source/a.docx")
    # a corrupted / hostile id map row pointing the copy at the SOURCE file itself
    rows = idmap.get(c22.id, fk22.item("Source/a.docx")["id"], "fr-fr")
    assert rows is not None
    with Session(engine) as sx:
        from app.models import ConnectorIdMap

        r = sx.get(ConnectorIdMap, rows.id)
        r.target_id = fk22.item("Source/a.docx")["id"]
        sx.add(r)
        sx.commit()
    fk22.edit("Source/a.docx", mk_docx("Source text v2"))
    src_before = fk22.read("Source/a.docx")
    fk22.writes.clear()
    ok, msg, s22b = _run(c22.id, fk22)
    assert fk22.read("Source/a.docx") == src_before and not fk22.writes, (
        "an id map target inside the source tree is never written",
        fk22.writes,
    )
    assert any("not inside the target root" in f["error"] for f in s22b["failures"]), s22b
    # a pre-made shortcut folder in the copy's path is refused before any create
    fk22b = FakeGraph()
    fk22b.put_file("Source/b.docx", mk_docx("Shortcut text"))
    fk22b.put_file("Translations/placeholder.txt", b"x")
    sc = fk22b._ensure_folder(["Translations", "fr-FR"])
    sc["remote"] = True
    c22b = _conn("shortcut")
    ok, msg, s22c = _run(c22b.id, fk22b)
    assert not fk22b.writes and s22c["created"] == 0, (fk22b.writes, s22c)
    # shortcut folders under the source are not walked
    fk22c = FakeGraph()
    fk22c.put_file("Source/ok.docx", mk_docx("fine"))
    fk22c._ensure_folder(["Source", "Linked"])["remote"] = True
    fk22c.put_file("Source/Linked/inside.docx", mk_docx("never read"))
    c22c = _conn("noshortcutwalk")
    ok, msg, s22d = _run(c22c.id, fk22c)
    assert s22d["created"] == 1 and not fk22c.exists("Translations/fr-FR/Linked/inside.docx"), s22d
    print(
        "  A: id-map target in the source tree, shortcut ancestors and shortcut sources: zero source writes"  # noqa: E501
    )

    # ---- (14d) Review gate: file URLs pinned to the drive's own tenant SharePoint hosts ----
    from app.services.connectors.msgraph_graph import sharepoint_hosts, trusted_file_url

    hs = sharepoint_hosts("https://contoso.sharepoint.com/sites/Loc/Shared%20Documents")
    assert hs == {"contoso.sharepoint.com", "contoso-my.sharepoint.com"}, hs
    assert sharepoint_hosts("https://contoso-my.sharepoint.com/personal/alex") == hs
    assert sharepoint_hosts("https://contosoeur.sharepoint.com/sites/X") == {
        "contosoeur.sharepoint.com",
        "contosoeur-my.sharepoint.com",
    }, "a multi-geo satellite pins its own hosts"
    for bad in (
        "",
        "http://contoso.sharepoint.com/x",
        "https://contoso.sharepoint.com.evil.example/x",
        "https://contoso.sharepoint.us/x",
        "https://onedrive.live.com/x",
        "https://u:p@contoso.sharepoint.com/x",
        "https://-my.sharepoint.com/x",
    ):
        assert sharepoint_hosts(bad) == frozenset(), bad
    for good in (
        "https://contoso.sharepoint.com/download?tempauth=1",
        "https://contoso-my.sharepoint.com:443/up/x",
    ):
        assert trusted_file_url(good, hs), good
    for bad in (
        "https://fabrikam.sharepoint.com/download?tempauth=1",
        "https://fabrikam-my.sharepoint.com/up/x",
        "https://b0mpua-by3301.files.1drv.com/y23",
        "https://sn3302.up.1drv.com/up/x",
        "https://contoso.sharepoint.com.evil.example/x",
        "https://contoso.sharepoint.com:8443/x",
        "https://u@contoso.sharepoint.com/x",
        "http://contoso.sharepoint.com/x",
        "https://contoso.sharepoint.com:99999/x",
    ):
        assert not trusted_file_url(bad, hs), bad
    assert not trusted_file_url("https://contoso.sharepoint.com/x", frozenset())
    # another tenant's SharePoint as the download redirect: refused, no request, nothing written
    fk23 = FakeGraph()
    fk23.put_file("Source/a.docx", mk_docx("A"))
    fk23.redirect_to = "https://fabrikam.sharepoint.com/download/x?tempauth=abc"
    ok, msg, s23 = _run(_conn("othertenant").id, fk23)
    assert fk23.stray_hosts == [] and not fk23.writes and s23["unreadable"] == 1, (msg, s23)
    assert not any(x[1] == "fabrikam.sharepoint.com" for x in fk23.log)
    # another tenant's OneDrive as the uploadUrl: refused before any chunk is sent
    fk24 = FakeGraph()
    fk24.put_file("Source/a.docx", mk_docx("A"))
    c24 = _conn("othertenantup", {"protect_edited_copies": "0"})
    ok, msg, s24 = _run(c24.id, fk24)
    assert ok and s24["created"] == 1, s24
    fk24.edit("Source/a.docx", mk_docx("A changed"))
    fk24.upload_host = "fabrikam-my.sharepoint.com"
    before24 = fk24.read("Translations/fr-FR/a.docx")
    set_pieces("app.services.connectors.msgraph", "SIMPLE_UPLOAD_MAX", 10)
    try:
        ok, msg, s24b = _run(c24.id, fk24)
    finally:
        set_pieces("app.services.connectors.msgraph", "SIMPLE_UPLOAD_MAX", 4 * 1024 * 1024)
    assert s24b["failures"] and fk24.read("Translations/fr-FR/a.docx") == before24, s24b
    assert not any(x[1] == "fabrikam-my.sharepoint.com" for x in fk24.log), "no chunk sent"
    # a drive whose webUrl is not SharePoint trusts no file host at all (fail closed)
    fk25 = FakeGraph()
    fk25.put_file("Source/a.docx", mk_docx("A"))
    fk25.web_url = "https://evil.example/sites/Loc"
    ok, msg, s25 = _run(_conn("badweburl").id, fk25)
    assert not fk25.downloads and not fk25.writes and s25["unreadable"] == 1, (msg, s25)
    fk25.web_url = None
    ok, msg, s25b = _run(_conn("noweburl").id, fk25)
    assert not fk25.downloads and not fk25.writes, (msg, s25b)
    print(
        "  A: download/upload URLs pinned to the drive's own <tenant>[-my].sharepoint.com; another tenant, 1drv.com, no webUrl: refused"  # noqa: E501
    )

    # ---- (14e) Review gate: folder cap and Max files ceiling (a whole site as the source) ----
    fk26 = FakeGraph()
    for k in range(4):
        fk26.put_file(f"Source/F{k}/a{k}.docx", mk_docx(f"Folder {k}"))
    set_pieces("app.services.connectors.msgraph", "MAX_FOLDERS", 4)
    try:
        ok, msg, _s26 = _run(_conn("manyfolders").id, fk26)
    finally:
        set_pieces("app.services.connectors.msgraph", "MAX_FOLDERS", 500)
    assert not ok and "more than 4 folders" in msg and "Narrow the Source folder" in msg, msg
    assert not fk26.downloads and not fk26.writes, "refused before any download or MT"
    fk26.put_file("Source/top.docx", mk_docx("Top level"))
    set_pieces("app.services.connectors.msgraph", "MAX_FOLDERS", 4)
    try:
        ok, msg, s26b = _run(_conn("manyfolders2", {"include_subfolders": "0"}).id, fk26)
    finally:
        set_pieces("app.services.connectors.msgraph", "MAX_FOLDERS", 500)
    assert ok and s26b["created"] == 1 and "folders" not in msg, ("top level only", msg, s26b)
    fk27 = FakeGraph()
    for k in range(4):
        fk27.put_file(f"Source/b{k}.docx", mk_docx(f"File {k}"))
    set_pieces("app.services.connectors.msgraph", "MAX_FILES_CEILING", 2)
    try:
        ok, msg, s27 = _run(_conn("ceiling", {"max_files": "100000"}).id, fk27)
    finally:
        set_pieces("app.services.connectors.msgraph", "MAX_FILES_CEILING", 500)
    assert s27["created"] == 2 and len(fk27.downloads) == 2, s27
    print(
        "  E: more than MAX_FOLDERS folders refused before any download; Max files clamped to the 500 ceiling"  # noqa: E501
    )

    # ---- (15) unreadable / empty files do not stop the run --------------------------
    fk21 = FakeGraph()
    fk21.put_file("Source/good.docx", mk_docx("Good file"))
    fk21.put_file("Source/corrupt.docx", b"this is not a zip archive")
    fk21.put_file("Source/blank.docx", mk_docx(" "))
    c21 = _conn("mixed")
    ok, msg, s21 = _run(c21.id, fk21)
    assert ok and s21["created"] == 1 and s21["unreadable"] == 1 and s21["empty_skipped"] == 1, s21
    assert "corrupt.docx" in msg and "could not be read" in msg, msg
    assert fk21.exists("Translations/fr-FR/good.docx") and not fk21.exists(
        "Translations/fr-FR/corrupt.docx"
    )
    fk21b = FakeGraph()
    fk21b.put_file("Source/corrupt.docx", b"nope")
    ok, msg, s21b = _run(_conn("onlybad").id, fk21b)
    assert not ok and s21b["unreadable"] == 1, ("only unreadable files -> not green", ok, s21b)
    print(
        "  a corrupt or empty file is reported and skipped; the other files are still translated; all-unreadable is not green"  # noqa: E501
    )

    # ---- (16) locales: canonical folder names, rejection, source dropped ---------------
    fk22 = FakeGraph()
    fk22.put_file("Source/a.docx", mk_docx("A"))
    c22 = _conn("locales", {"target_locales": "fr-FR, ../x, en-US, de-DE"})
    ok, msg, s22 = _run(c22.id, fk22)
    assert ok and set(s22["locales"]) == {"fr-fr", "de-de"} and s22["rejected"] == ["../x"], s22
    assert s22["dropped_source"] == ["en-us"], s22
    assert fk22.files_under("Translations") == [
        "Translations/de-DE/a.docx",
        "Translations/fr-FR/a.docx",
    ]
    nosrc = _conn("nosrc", {"source_locale": ""})
    ok, msg, _s = _run(nosrc.id, fk22)
    assert not ok and "Source locale" in msg, msg
    print(
        "  locales: folders use canonical names, a hostile locale is rejected, the source locale is dropped, a missing source locale fails loud"  # noqa: E501
    )

    # ---- (17) H: WriteBudget ceiling, egress guard -----------------------------------
    saved = getattr(app_config, "MAX_WRITES_PER_RUN", None)
    app_config.MAX_WRITES_PER_RUN = 1
    try:
        fk23 = FakeGraph()
        fk23.put_file("Source/a.docx", mk_docx("A"))
        fk23.put_file("Source/b.docx", mk_docx("B"))
        ok, msg, s23 = _run(_conn("budget", {"target_locales": "fr-FR, de-DE"}).id, fk23)
        assert ok and s23["ceiling_hit"] and s23["created"] + s23["updated"] == 1, s23
        assert len(fk23.files_under("Translations")) == 1, (
            "exactly one copy written",
            fk23.files_under("Translations"),
        )
    finally:
        app_config.MAX_WRITES_PER_RUN = saved
    from app.services.connectors import egress_guard

    fk24 = FakeGraph()
    fk24.put_file("Source/a.docx", mk_docx("A"))
    real = egress_guard.egress_is_real
    egress_guard.egress_is_real = lambda t: True
    try:
        ok, msg, _s = _run(_conn("egress").id, fk24)
    finally:
        egress_guard.egress_is_real = real
    assert not ok and not fk24.writes and not fk24.downloads, (
        "mock MT + real egress blocked, nothing touched",
        msg,
    )
    print(
        "  H: WriteBudget MAX_WRITES_PER_RUN=1 -> exactly one copy; egress guard refuses mock output and touches nothing"  # noqa: E501
    )

    # ---- (18) credential pin: scope edits clear the secret, benign edits do not ---------
    c25 = _conn("pin")
    cfg25 = dict(cstore.get(c25.id).config)
    cstore.update(c25.id, "pin", {**cfg25, "max_files": "10"}, {"client_secret": ""})
    assert cstore.get(c25.id).secret_enc and _has_credential(cstore.get(c25.id)), (
        "benign edit keeps the secret"
    )
    cstore.update(
        c25.id, "pin", {**cfg25, "max_files": "10", "source_folder": "Other"}, {"client_secret": ""}
    )
    assert not _has_credential(cstore.get(c25.id)), (
        "moving the source folder clears the stored secret"
    )
    print(
        "  credential pin: editing the source folder clears the stored secret; editing max_files does not"  # noqa: E501
    )

    # ---- (19) pull() and retry ---------------------------------------------------------
    fk26 = FakeGraph()
    fk26.put_file("Source/a.docx", mk_docx("Pull me", "And me"))
    mc26, _ = _mc(fake=fk26)
    items = mc26.pull()
    assert [i.text for i in items] == ["Pull me", "And me"] and all(
        i.field.startswith("body:") for i in items
    ), items
    assert _raises(mc26.push, items)
    fk27 = FakeGraph()
    fk27.put_file("Source/a.docx", mk_docx("Retry me"))
    c27 = _conn("retry")
    fk27.inject("PUT", "/items/", 503, times=3)
    ok, msg, s27 = _run(c27.id, fk27)
    assert s27["failures"], s27
    ok, msg, s27b = cstore.retry_failed(
        c27.id, s27["failures"], actor="auto", transport=fk27.transport()
    )
    assert ok and fk27.read("Translations/fr-FR/a.docx"), (msg, s27b)
    print("  pull() lists units; push refuses; retry_failed re-attempts exactly the failed pair")

    # ---- (20) H: four-surface parity ---------------------------------------------------
    from app.services.connectors import REGISTRY

    assert "msgraph" in REGISTRY
    tc = TestClient(app)
    assert "Microsoft 365 Files" in tc.get("/connectors").text
    typ = tc.get("/connectors/msgraph").text
    assert "Source folder" in typ and "Client secret" in typ, "type form renders"
    from app.cli.__main__ import build_parser

    assert "connector msgraph-translate" in _leaf_paths(build_parser())
    from starlette.routing import Route

    api_paths = {r.path for r in walk_routes(app) if isinstance(r, Route)}
    assert "/api/connectors/{cid}/msgraph/translate" in api_paths
    from app.services.connectors import harvest as _hv

    assert _hv._map_msgraph({"title": "doc.docx"})["source"].get("title") == "doc.docx"
    print(
        "  H parity: REGISTRY + /connectors tile + type form + /api route + CLI leaf + harvest title"  # noqa: E501
    )

    blob = "\n".join(cap.lines)
    assert FIXTURE_SECRET not in blob and not re.search(r"tok-\d+-fixture", blob), (
        "H: secret/token in no log line (whole suite)"
    )
    print("verify_msgraph OK")


def _leaf_paths(parser, prefix=""):
    import argparse

    out = set()
    subs = [a for a in parser._actions if isinstance(a, argparse._SubParsersAction)]
    if not subs:
        return {prefix} if prefix else set()
    for sp in subs:
        for name, subparser in sp.choices.items():
            out |= _leaf_paths(subparser, f"{prefix} {name}".strip())
    return out


if __name__ == "__main__":
    main()
