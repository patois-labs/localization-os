"""Loud/refusing FT export source-fallback, TAGLESS arm (honest-failure, 2026-08-13).

THE BUG (big-doc flood): a sentence-tuned model emitted far fewer paragraphs than the
source for a FULL_TEXT doc. In the TAGLESS export arm `full.split("\\n\\n")` then yielded
fewer parts than source units, and `paragraph_texts_for` SILENTLY substituted the ENGLISH
source for every missing paragraph via `parts[idx] if idx < len(parts) else u["text"]`.
Observed live: a 36KB "es" deliverable holding 851 ENGLISH sentences vs 7 Spanish, served
under a target-locale filename with ZERO warning. (The TAGGED arm was already protected:
GUARD A quarantines on a count desync.)

THE FIX (export.py, tagless arm), after the adversarial review folded in:
  * a count desync leaves NO trustworthy positional map, so - exactly like guard A on the
    tagged arm - it is an UNCONDITIONAL quarantine. `allow_source_fallback` does NOT open a
    door here (removed: an interior drop shifts every later paragraph, and a tail-share
    number measures the wrong quantity);
  * comparison + mapping run on NON-BLANK parts only, so trailing/leading/interior blank
    "\\n\\n" split artifacts export byte-correct while a real paragraph deletion still refuses;
  * parts < units -> kind="untranslated"; parts > units -> kind="count"; a blank stored
    body -> kind="untranslated" ("stored translated body is blank"); units == [] -> guard
    skipped (skeleton-less/fixture docs keep main's behavior).
The TAGGED arm and the sanctioned held-export path (untranslated target_text "" -> source,
the NOT-NULL default) are untouched.

Hermetic + offline (throwaway LOC_TMS_DATA_DIR, provider keys blanked). ASCII-only.
Final line: "FT EXPORT UNTRANSLATED OK".
"""

import os
import sys

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
for _s in (sys.stdout, sys.stderr):
    try:  # noqa: SIM105
        _s.reconfigure(encoding="utf-8", errors="backslashreplace")
    except Exception:
        pass
from tests._bootstrap import init_env

init_env("loctms_ftuntrans_")

import tempfile
from io import BytesIO

from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app.db import engine, init_db
from app.main import app
from app.models import (
    Document,
    Project,
    Segment,
    SegStatus,
    TargetLocale,
    TransMode,
)
from app.services import export, segment
from app.services.ingest import ingest_document
from tests._cli_runner import run_cli

init_db()

LOC = "es-ES"
_SEQ = [0]


def _seq():
    _SEQ[0] += 1
    return _SEQ[0]


def _build_docx(paras, first_runs=None):
    d = Docx()
    if first_runs is not None:
        p = d.add_paragraph()
        for text, fmt in first_runs:
            r = p.add_run(text)
            setattr(r, fmt, True)
    for t in paras:
        d.add_paragraph(t)
    b = BytesIO()
    d.save(b)
    return b.getvalue()


def _docx_text(data):
    return "\n".join(p.text for p in Docx(BytesIO(data)).paragraphs)


def _fixture(source_paras, target_text, first_runs=None, locale=LOC):
    """An INGESTED FULL_TEXT docx; `target_text` set verbatim on the body segment (None
    leaves it untranslated). Returns {doc_id, seg_id, units, source_text}."""
    with Session(engine) as s:
        p = Project(
            name="ftut-%d" % _seq(),  # noqa: UP031
            org_id=1,
            source_locale="en-US",
            default_provider="mock",  # noqa: RUF100, UP031
        )
        s.add(p)
        s.commit()
        s.refresh(p)
        s.add(TargetLocale(project_id=p.id, locale=locale))
        s.commit()
        doc = ingest_document(
            s, p, "doc.docx", _build_docx(source_paras, first_runs), TransMode.FULL_TEXT
        )
        s.commit()
        s.refresh(doc)
        seg = s.exec(
            select(Segment).where(Segment.document_id == doc.id, Segment.target_locale == locale)
        ).one()
        if target_text is not None:
            seg.target_text = target_text
            seg.status = SegStatus.MT
            s.add(seg)
            s.commit()
        return {
            "doc_id": doc.id,
            "seg_id": seg.id,
            "units": list((doc.skeleton or {})["units"]),
            "source_text": seg.source_text,
        }


def _docs_key():
    from app.services import apikeys, security

    u = security.get_by_email("ftut@example.com")
    if u is None:
        u = security.create_user("ftut@example.com", "FT UT", "admin", "Pw!12345xyz-ftut")
    _row, raw = apikeys.create_key(u.id, "ftut-%d" % _seq(), scopes="api docs")  # noqa: UP031
    return {"X-Loc-Token": raw}


def _quarantines(doc_id, allow_source_fallback, kind):
    with Session(engine) as s:
        d = s.get(Document, doc_id)
        try:
            export.build_translated_file(s, d, LOC, allow_source_fallback=allow_source_fallback)
        except export.TagExportQuarantine as exc:
            assert exc.kind == kind, (kind, exc.kind, str(exc))
            assert str(exc).isascii(), str(exc)  # repo law: ASCII-safe CLI output
            return str(exc)
    raise AssertionError(f"expected a TagExportQuarantine (kind={kind}), got a file")


# ===========================================================================
# 1. parts < units -> REFUSE (kind="untranslated"); the flag does NOT open it.
# ===========================================================================
def test_parts_lt_units_refused_and_flag_does_not_open():
    # counts DIFFER (2 real parts, 3 units) -> the "English source would ship" shape.
    fx = _fixture(["Source one.", "Source two.", "Source three."], "Trad uno.\n\nTrad dos.")
    msg = _quarantines(fx["doc_id"], allow_source_fallback=False, kind="untranslated")
    assert "2 real paragraph(s), the source has 3" in msg, msg
    assert "silently substitute ENGLISH source" in msg, msg
    assert "1 missing paragraph" in msg, msg
    assert "allow_source_fallback" not in msg, msg  # the door is gone
    # WITH the flag: still refused (a count desync has no trustworthy positional map).
    _quarantines(fx["doc_id"], allow_source_fallback=True, kind="untranslated")
    print("  1. parts<units REFUSED (kind=untranslated); allow_source_fallback does NOT open it OK")


# ===========================================================================
# 2. parts > units -> REFUSE (kind="count") unconditionally, even WITH the flag.
# ===========================================================================
def test_parts_gt_units_refused_unconditionally():
    fx = _fixture(["Source one.", "Source two."], "Trad uno.\n\nTrad dos.\n\nTrad tres.")
    for flag in (False, True):
        msg = _quarantines(fx["doc_id"], allow_source_fallback=flag, kind="count")
        assert "silently dropped" in msg, msg
    print("  2. parts>units REFUSED (kind=count) unconditionally, flag does NOT open it OK")


# ===========================================================================
# 3. a blank STORED target body -> REFUSE (kind="untranslated"); the None-target held
#    path still source-echoes (untouched).
# ===========================================================================
def test_blank_stored_body_refused_but_held_untouched():
    # a whitespace-only stored target (the #10 hole) refuses loud.
    fxb = _fixture(["Source one.", "Source two."], "   ")
    msg = _quarantines(fxb["doc_id"], allow_source_fallback=True, kind="untranslated")
    assert "stored translated body is blank" in msg, msg
    # the SANCTIONED held path (untranslated target_text "" -> source) still exports source,
    # never this quarantine: build source-echoes exactly as main / verify_export_held_locale.
    fxh = _fixture(["Hello world.", "Second line."], None)
    with Session(engine) as s:
        d = s.get(Document, fxh["doc_id"])
        data = export.build_translated_file(s, d, LOC)  # must NOT quarantine
    assert "Hello world." in _docx_text(data), _docx_text(data)
    print(
        "  3. blank STORED body REFUSED (kind=untranslated); None-target held export source-echoes OK"  # noqa: E501
    )


# ===========================================================================
# 4. blank "\n\n" split artifacts are NOT paragraphs: the four working shapes export
#    byte-correct; an interior blank that DELETES a real paragraph still refuses.
# ===========================================================================
def test_blank_part_shapes():
    working = {
        "trailing \\n\\n": "Trad uno.\n\nTrad dos.\n\n",
        "leading blank": "\n\nTrad uno.\n\nTrad dos.",
        "doubled delimiter": "Trad uno.\n\n\n\nTrad dos.",
        "trailing whitespace": "Trad uno.\n\nTrad dos.\n\n   ",
    }
    from app.services import fulltext_paragraphs as ftp

    for label, body in working.items():
        fx = _fixture(["Source one.", "Source two."], body)
        rep = []
        with Session(engine) as s:
            d = s.get(Document, fx["doc_id"])
            # ENFORCE the "cannot disagree on structure" property: split_view OK AND
            # export clean, TOGETHER, on the same artifact body (item 2c).
            v = ftp.split_view(s, d, LOC)
            m = export.paragraph_texts_for(s, d, LOC, report=rep)
        assert v["ok"] is True, (label, v)
        assert v["damaged_count"] == 0, (label, v)
        assert m == {"0": "Trad uno.", "1": "Trad dos."}, (label, m)
        assert rep == [], (label, rep)  # no substitution
    # C2: 3 units, an INTERIOR blank -> real position 1 is blank in the translated body.
    # Must REFUSE (was previously a silent positional shift), naming the blank paragraph -
    # and NOT claiming English source would ship (counts coincide, so it ships blank).
    fxc = _fixture(["Source one.", "Source two.", "Source three."], "Trad uno.\n\n\n\nTrad dos.")
    msg = _quarantines(fxc["doc_id"], allow_source_fallback=False, kind="untranslated")
    assert "paragraph(s) 1 of 3 are blank" in msg, msg
    assert "set-paragraph" in msg and "revert-paragraph" in msg, msg
    assert "ENGLISH source" not in msg, msg  # not the source-substitution shape
    print("  4. blank-artifact shapes: split_view OK + export clean together; interior-blank")
    print("     names the blank paragraph and refuses (no false English-source claim) OK")


# ===========================================================================
# 5. a clean count-aligned tagless body is byte-unchanged.
# ===========================================================================
def test_clean_aligned_body_unchanged():
    fx = _fixture(["Source one.", "Source two."], "Trad uno.\n\nTrad dos.")
    rep = []
    with Session(engine) as s:
        d = s.get(Document, fx["doc_id"])
        m = export.paragraph_texts_for(s, d, LOC, report=rep)
    assert m == {"0": "Trad uno.", "1": "Trad dos."}, m
    assert rep == [], rep
    print("  5. clean count-aligned tagless body: mapping unchanged, no report OK")


# ===========================================================================
# 6. the TAGGED arm is untouched: a count desync is still GUARD A (kind="count"),
#    NOT the new "untranslated" kind.
# ===========================================================================
def test_tagged_arm_unchanged():
    fx = _fixture(
        ["Body one.", "Body two."],
        target_text=None,
        first_runs=[("Bold ", "bold"), ("title.", "italic")],
    )
    with Session(engine) as s:
        seg = s.get(Segment, fx["seg_id"])
        assert segment.find_placeholders(seg.source_text), "fixture must be TAGGED"
        seg.target_text = "Solo un parrafo traducido."  # 1 part, 3 units
        seg.status = SegStatus.MT
        s.add(seg)
        s.commit()
    _quarantines(fx["doc_id"], allow_source_fallback=False, kind="count")  # guard A
    print("  6. TAGGED count desync still owned by GUARD A (kind=count), untouched OK")


# ===========================================================================
# 7. units == [] (skeleton-less/fixture doc) -> guard skipped, main behavior preserved.
# ===========================================================================
def test_no_units_short_circuit():
    with Session(engine) as s:
        p = Project(
            name="ftut-nounits-%d" % _seq(),  # noqa: UP031
            org_id=1,
            source_locale="en-US",
            default_provider="mock",
        )
        s.add(p)
        s.commit()
        s.refresh(p)
        s.add(TargetLocale(project_id=p.id, locale=LOC))
        s.commit()
        doc = Document(
            project_id=p.id,
            filename="d.docx",
            fmt="docx",
            mode=TransMode.FULL_TEXT,
            skeleton={"units": [], "seg_locale": "en-US"},
        )
        s.add(doc)
        s.commit()
        s.refresh(doc)
        s.add(
            Segment(
                document_id=doc.id,
                target_locale=LOC,
                order_index=0,
                source_text="",
                target_text="Algo",
                status=SegStatus.MT,
            )
        )
        s.commit()
        d = s.get(Document, doc.id)
        m = export.paragraph_texts_for(s, d, LOC)  # must NOT raise
    assert m == {}, m
    print("  7. units==[] short-circuits the tagless guard (main behavior preserved) OK")


# ===========================================================================
# 8. the surfaces: web + /api + CLI all refuse (both flag states), no way-out URL.
# ===========================================================================
def test_surfaces():
    H = _docs_key()
    client = TestClient(app)
    fx = _fixture(["Source one.", "Source two.", "Source three."], "Trad uno.\n\nTrad dos.")
    did = fx["doc_id"]

    # web: 409 both with and without the flag; NO export_anyway_url (no override), but a
    # paragraphs_url DIAGNOSIS link is offered.
    r = client.get("/documents/%d/export?locale=%s" % (did, LOC))  # noqa: UP031
    assert r.status_code == 409, (r.status_code, r.text[:300])
    assert r.json()["kind"] == "untranslated", r.json()
    assert "export_anyway_url" not in r.json(), r.json()
    assert "paragraphs" in r.json().get("paragraphs_url", ""), r.json()
    r2 = client.get("/documents/%d/export?locale=%s&allow_source_fallback=1" % (did, LOC))  # noqa: UP031
    assert r2.status_code == 409, (r2.status_code, r2.text[:300])
    assert r2.json()["kind"] == "untranslated", r2.json()

    # /api: same, both flag states.
    ra = client.get("/api/documents/%d/download?locale=%s" % (did, LOC), headers=H)  # noqa: UP031
    assert ra.status_code == 409, (ra.status_code, ra.text[:300])
    assert ra.json()["kind"] == "untranslated", ra.json()
    ra2 = client.get(
        "/api/documents/%d/download?locale=%s&allow_source_fallback=1" % (did, LOC),  # noqa: UP031
        headers=H,  # noqa: RUF100, UP031
    )
    assert ra2.status_code == 409, (ra2.status_code, ra2.text[:300])

    # CLI: fails both with and without the flag; the hint points at the diagnosis path
    # (doc paragraphs -> set-paragraph / revert-paragraph), NOT the source-fallback flag.
    tmp = tempfile.mkdtemp(prefix="ftut_")
    code, out, err = run_cli(
        ["--json", "doc", "export", str(did), "--locale", LOC, "-o", os.path.join(tmp, "no.docx")]
    )
    assert code != 0, (code, out, err)
    assert "doc paragraphs" in (out + err) and "revert-paragraph" in (out + err), (out, err)
    code, out, err = run_cli(
        [
            "--json",
            "doc",
            "export",
            str(did),
            "--locale",
            LOC,
            "-o",
            os.path.join(tmp, "yes.docx"),
            "--allow-source-fallback",
        ]
    )
    assert code != 0, (code, out, err)  # the flag does not rescue a count desync
    print("  8. web + /api + CLI: refuse both flag states, no source-fallback way-out OK")


# ===========================================================================
# 9. BLOCKER (delta review): a legit XLIFF with an empty <source/> ingested full_text
#    exports clean in BOTH the translation and the source-only paths. Layer 1 (the
#    xliff filter now skips the blank source) makes it a clean 2-unit doc.
# ===========================================================================
_XLIFF_BLANK_MID = (
    b'<?xml version="1.0"?>\n'
    b'<xliff version="1.2"><file source-language="en" target-language="es"><body>'
    b'<trans-unit id="1"><source>Alpha source paragraph here.</source></trans-unit>'
    b'<trans-unit id="2"><source>   </source></trans-unit>'
    b'<trans-unit id="3"><source>Gamma source paragraph here.</source></trans-unit>'
    b"</body></file></xliff>"
)


def test_xliff_empty_source_roundtrips():
    with Session(engine) as s:
        p = Project(
            name="ftut-xliff-%d" % _seq(),  # noqa: UP031
            org_id=1,
            source_locale="en-US",
            default_provider="mock",  # noqa: RUF100, UP031
        )
        s.add(p)
        s.commit()
        s.refresh(p)
        s.add(TargetLocale(project_id=p.id, locale=LOC))
        s.commit()
        doc = ingest_document(s, p, "blank.xlf", _XLIFF_BLANK_MID, TransMode.FULL_TEXT)
        s.commit()
        s.refresh(doc)
        units = list((doc.skeleton or {})["units"])
        # layer 1: the empty <source/> is skipped -> a clean 2-unit doc.
        assert len(units) == 2, units
        seg = s.exec(select(Segment).where(Segment.document_id == doc.id)).one()
        did = doc.id
    # source-only / held path (target "" -> source) exports clean, no quarantine.
    with Session(engine) as s:
        d = s.get(Document, did)
        out_src = export.build_translated_file(s, d, LOC)
    assert b"Alpha source paragraph here." in out_src, out_src[:200]
    # a perfect 2-paragraph translation exports clean too.
    with Session(engine) as s:
        seg = s.exec(select(Segment).where(Segment.document_id == did)).one()
        seg.target_text = "Alfa parrafo traducido.\n\nGamma parrafo traducido."
        seg.status = SegStatus.MT
        s.add(seg)
        s.commit()
        d = s.get(Document, did)
        out_tr = export.build_translated_file(s, d, LOC)
    assert b"Alfa parrafo traducido." in out_tr and b"Gamma parrafo traducido." in out_tr
    print("  9. XLIFF empty <source/> full_text: filter skips it -> translation AND")
    print("     source-only export both clean (blocker) OK")


# ===========================================================================
# 10. BLOCKER (defense-in-depth): a LEGACY row whose skeleton carries a whitespace-only
#     unit (bypass the filter, build units directly) still exports via the symmetric
#     predicate - a blank unit maps to its own source, real counts match.
# ===========================================================================
def test_legacy_blank_unit_exports_via_symmetric_predicate():
    units = [
        {"p_index": 0, "text": "Alpha."},
        {"p_index": 1, "text": "   "},  # legacy whitespace-only source unit
        {"p_index": 2, "text": "Gamma."},
    ]
    src_body = "\n\n".join(u["text"] for u in units)
    with Session(engine) as s:
        p = Project(
            name="ftut-legacy-%d" % _seq(),  # noqa: UP031
            org_id=1,
            source_locale="en-US",
            default_provider="mock",  # noqa: RUF100, UP031
        )
        s.add(p)
        s.commit()
        s.refresh(p)
        s.add(TargetLocale(project_id=p.id, locale=LOC))
        s.commit()
        doc = Document(
            project_id=p.id,
            filename="legacy.docx",
            fmt="docx",
            mode=TransMode.FULL_TEXT,
            skeleton={"units": units, "seg_locale": "en-US"},
        )
        s.add(doc)
        s.commit()
        s.refresh(doc)
        # a perfect translation: the blank unit's slot is blank in the body too.
        s.add(
            Segment(
                document_id=doc.id,
                target_locale=LOC,
                order_index=0,
                source_text=src_body,
                target_text="Alfa.\n\n\n\nGamma.",
                status=SegStatus.MT,
            )
        )
        s.commit()
        d = s.get(Document, doc.id)
        m = export.paragraph_texts_for(s, d, LOC)  # must NOT raise
    # non-blank parts (2) match non-blank units (2); the blank unit maps to its own source.
    assert m == {"0": "Alfa.", "1": "   ", "2": "Gamma."}, m
    # source-only path (blank target) also aligns instead of quarantining.
    with Session(engine) as s:
        seg = s.exec(select(Segment).where(Segment.document_id == d.id)).one()
        seg.target_text = ""
        s.add(seg)
        s.commit()
        d = s.get(Document, d.id)
        m2 = export.paragraph_texts_for(s, d, LOC)
    assert m2 == {"0": "Alpha.", "1": "   ", "2": "Gamma."}, m2
    print("  10. legacy whitespace-only source unit exports via the symmetric predicate")
    print("      (translation AND source-only) - defense-in-depth OK")


# ===========================================================================
# 11. BLOCKER-2a: a PO with a whitespace-only `msgid " "` (deliberately NOT filtered)
#     ingested full_text exports clean via the symmetric predicate (translation AND
#     source-only); an under-translation still refuses.
# ===========================================================================
_PO_BLANK_MID = (
    b'msgid ""\nmsgstr ""\n\n'
    b'msgid "Alpha source paragraph."\nmsgstr ""\n\n'
    b'msgid " "\nmsgstr ""\n\n'
    b'msgid "Gamma source paragraph."\nmsgstr ""\n'
)


def test_po_blank_msgid_roundtrips():
    with Session(engine) as s:
        p = Project(
            name="ftut-po-%d" % _seq(),  # noqa: UP031
            org_id=1,
            source_locale="en-US",
            default_provider="mock",  # noqa: RUF100, UP031
        )
        s.add(p)
        s.commit()
        s.refresh(p)
        s.add(TargetLocale(project_id=p.id, locale=LOC))
        s.commit()
        doc = ingest_document(s, p, "blank.po", _PO_BLANK_MID, TransMode.FULL_TEXT)
        s.commit()
        s.refresh(doc)
        units = list((doc.skeleton or {})["units"])
        # PO keeps the whitespace msgid (no filter guard) -> a 3-unit doc with a blank unit.
        assert len(units) == 3, units
        assert not (units[1]["text"] or "").strip(), units
        did = doc.id
    # source-only path exports clean (blank unit -> its own source, real counts match).
    with Session(engine) as s:
        d = s.get(Document, did)
        out_src = export.build_translated_file(s, d, LOC)
    assert b"Alpha source paragraph." in out_src and b"Gamma source paragraph." in out_src
    # a perfect translation (blank unit slot blank) exports clean.
    with Session(engine) as s:
        seg = s.exec(select(Segment).where(Segment.document_id == did)).one()
        seg.target_text = "Alfa traducido.\n\n\n\nGamma traducido."
        seg.status = SegStatus.MT
        s.add(seg)
        s.commit()
        d = s.get(Document, did)
        out_tr = export.build_translated_file(s, d, LOC)
    assert b"Alfa traducido." in out_tr and b"Gamma traducido." in out_tr
    # an under-translation (a real paragraph dropped) still refuses.
    with Session(engine) as s:
        seg = s.exec(select(Segment).where(Segment.document_id == did)).one()
        seg.target_text = "Alfa traducido."
        s.add(seg)
        s.commit()
    _quarantines(did, allow_source_fallback=False, kind="untranslated")
    print(
        '  11. PO msgid " " full_text: translation + source-only export clean, '
        "under-translation refuses (blocker-2a) OK"
    )


# ===========================================================================
# 12. BLOCKER-2b: direct align_parts coverage - a legacy blank unit at 0 / middle / last,
#     under, over, all-blank.
# ===========================================================================
def test_align_parts_unit_coverage():
    from app.services import fulltext_paragraphs as ftp

    A, B = "tA", "tB"
    body = "tA\n\ntB"
    # blank unit at index 0 / middle / last -> the blank unit maps to its own source,
    # the two real parts fill the two real units in order.
    at0 = ftp.align_parts(body, [{"text": " "}, {"text": "s0"}, {"text": "s1"}])
    assert at0 == ([" ", A, B], ""), at0
    mid = ftp.align_parts(body, [{"text": "s0"}, {"text": " "}, {"text": "s1"}])
    assert mid == ([A, " ", B], ""), mid
    last = ftp.align_parts(body, [{"text": "s0"}, {"text": "s1"}, {"text": " "}])
    assert last == ([A, B, " "], ""), last
    # under: fewer real parts than real units.
    assert ftp.align_parts(body, [{"text": "s0"}, {"text": "s1"}, {"text": "s2"}]) == (
        None,
        "under",
    )
    # over: more real parts than real units.
    assert ftp.align_parts("tA\n\ntB\n\ntC", [{"text": "s0"}, {"text": "s1"}]) == (None, "over")
    # all-blank: no real parts, no real units -> tolerated, all-blank aligned (not a desync).
    allblank = ftp.align_parts("  \n\n\t", [{"text": " "}, {"text": "\t"}])
    assert allblank == ([" ", "\t"], ""), allblank
    print(
        "  12. align_parts: legacy blank unit at 0/mid/last maps to own source; under/over/"
        "all-blank classified (blocker-2b) OK"
    )


# ===========================================================================
# 13. BLOCKER-2d: LOCK the PO no-guard decision - the whitespace `msgid " "` unit is KEPT
#     and build() aligns positionally. A future "symmetric cleanup" of po_filter that
#     skips it (renumbering re-exports of legacy docs) breaks THIS test loudly.
# ===========================================================================
def test_po_filter_keeps_whitespace_msgid():
    from app.filters.po_filter import PoFilter

    filt = PoFilter()
    parsed = filt.parse(_PO_BLANK_MID)
    texts = [u.text for u in parsed.units]
    # the whitespace msgid is KEPT as a unit (only a FALSY msgid is skipped).
    assert texts == ["Alpha source paragraph.", " ", "Gamma source paragraph."], texts
    # build() keys on the running unit counter: unit 1 (the whitespace msgid) keeps its
    # slot, so the mapping lands positionally on the right msgids.
    out = filt.build(_PO_BLANK_MID, {"0": "Alfa", "1": "MID", "2": "Gamma"}).decode()
    assert 'msgid "Alpha source paragraph."\nmsgstr "Alfa"' in out, out
    assert 'msgid " "\nmsgstr "MID"' in out, out
    assert 'msgid "Gamma source paragraph."\nmsgstr "Gamma"' in out, out
    print(
        "  13. po_filter KEEPS a whitespace msgid + aligns build positionally "
        "(locks the no-guard decision, blocker-2d) OK"
    )


if __name__ == "__main__":
    test_parts_lt_units_refused_and_flag_does_not_open()
    test_parts_gt_units_refused_unconditionally()
    test_blank_stored_body_refused_but_held_untouched()
    test_blank_part_shapes()
    test_clean_aligned_body_unchanged()
    test_tagged_arm_unchanged()
    test_no_units_short_circuit()
    test_surfaces()
    test_xliff_empty_source_roundtrips()
    test_legacy_blank_unit_exports_via_symmetric_predicate()
    test_po_blank_msgid_roundtrips()
    test_align_parts_unit_coverage()
    test_po_filter_keeps_whitespace_msgid()
    print("FT EXPORT UNTRANSLATED OK")
