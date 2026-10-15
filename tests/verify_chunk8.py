"""Chunk 8 verification — export round-trip + dashboard, offline.

- segmented export: translated sentences are re-joined into their paragraphs;
  untranslated paragraphs fall back to source; structure preserved
- full-text export: the whole-doc translation is split back across paragraphs
- the export endpoint returns a .docx attachment
- the project dashboard renders rollup stats
"""

import os
import sys
from io import BytesIO

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_test_")

from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app.db import engine, init_db
from app.main import app
from app.models import Segment


def docx_from(paragraphs):
    d = Docx()
    for p in paragraphs:
        d.add_paragraph(p)
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def upload(client, pid, paragraphs, mode):
    r = client.post(
        f"/projects/{pid}/documents",
        data={"mode": mode},
        files={
            "file": (
                "d.docx",
                docx_from(paragraphs),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    return int(r.headers["location"].rsplit("/", 1)[-1])


def main() -> None:
    init_db()
    client = TestClient(app)
    r = client.post(
        "/projects",
        data={
            "name": "Export Test",
            "source_locale": "en-US",
            "target_locales": "fr-FR",
            "default_provider": "claude",
        },
        follow_redirects=False,
    )
    pid = int(r.headers["location"].rsplit("/", 1)[-1])

    # --- segmented round-trip ---
    doc = upload(
        client,
        pid,
        ["Adventure awaits. Save big.", "Second paragraph here.", "Leave me alone."],
        "segmented",
    )
    with Session(engine) as s:
        segs = sorted(
            s.exec(
                select(Segment).where(Segment.document_id == doc, Segment.target_locale == "fr-FR")
            ).all(),
            key=lambda x: x.order_index,
        )
        ids = [x.id for x in segs]
    assert len(ids) == 4, ids  # 2 + 1 + 1 sentences
    for sid, tgt in zip(ids[:3], ["A0", "A1", "B0"], strict=False):
        client.post(f"/segments/{sid}/save", data={"target_text": tgt, "tte_ms": "0"})
    # ids[3] (in paragraph 3) left untranslated -> should fall back to source

    resp = client.get(f"/documents/{doc}/export?locale=fr-FR")
    assert resp.status_code == 200
    assert "attachment" in resp.headers["content-disposition"]
    assert "wordprocessingml" in resp.headers["content-type"]
    texts = [p.text for p in Docx(BytesIO(resp.content)).paragraphs]
    assert "A0 A1" in texts, texts  # two sentences re-joined into para 1
    assert "B0" in texts, texts
    assert "Leave me alone." in texts, texts  # untranslated paragraph kept its source
    print(f"  segmented export: paragraphs rebuilt + untranslated fallback ({len(texts)} paras)")

    # --- full-text round-trip ---
    ft = upload(client, pid, ["First para.", "Second para."], "full_text")
    with Session(engine) as s:
        seg = s.exec(
            select(Segment).where(Segment.document_id == ft, Segment.target_locale == "fr-FR")
        ).first()
        seg_id = seg.id
    client.post(
        f"/segments/{seg_id}/save", data={"target_text": "Première.\n\nDeuxième.", "tte_ms": "0"}
    )
    resp = client.get(f"/documents/{ft}/export?locale=fr-FR")
    texts = [p.text for p in Docx(BytesIO(resp.content)).paragraphs]
    assert "Première." in texts and "Deuxième." in texts, texts
    print("  full-text export: whole-doc translation split back across paragraphs")

    # --- project Overview stats block ---
    # Key on the "Overview" section heading + a stats-block-specific marker
    # ("segment TM" tile label) + the approved-% legend, NOT the word "Dashboard"
    # (which the sidebar nav renders on every page and does not prove this block).
    page = client.get(f"/projects/{pid}").text
    assert "Overview" in page and "segment TM" in page and "approved" in page, (
        "project overview stats not rendered"
    )
    print("  project Overview renders rollup stats")

    print("CHUNK 8 OK")


if __name__ == "__main__":
    main()
