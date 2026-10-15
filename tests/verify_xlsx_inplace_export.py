"""#685: the Excel export writes translations into the ORIGINAL package.

The old export rebuilt the workbook through openpyxl, which drops every part it does not
model (text boxes, shapes, EMF/WMF pictures, printer settings, the calc chain). The new one
patches the worksheet / shared-string / drawing XML in place and copies every other member
byte for byte, as the Word and PowerPoint filters do. Text-box paragraphs are offered for
translation under walk 2, numbered after the cells, so cell segment ids never move.

  [A] untouched members are byte-identical; a no-op build changes nothing
  [B] translated cells read back (shared, rich, inline, str, implicit address), numbers and
      formulas untouched, a literal "=..." translation stays a string
  [C] shared strings: one translation for every use rewrites the item in place; different
      translations append a new item and repoint only that cell; rich formatting kept
      when uniform, phonetic runs dropped
  [D] walk 2: text-box paragraphs after the cells (Fallback copy skipped), written back
      with the target language; walk 1 (an old import) leaves the drawing alone; the cell
      units are the same list the pre-#685 parse produced
  [E] column mode writes a target column into new cells and rows, in address order
  [F] .xlsm keeps its VBA project and content type; Strict in, Strict out; a package
      without docProps/core.xml gets one for the language
  [G] end to end: upload, translate, export through export.build_translated_file; an
      unknown walk stamp is refused with "Re-import"

Fixtures: tests/fixtures/formats/xlsx/make_xlsx_objects.py (hand-written XML).
Hermetic + offline. Final line: "XLSX INPLACE EXPORT OK".
"""

import io
import os
import sys
import warnings
import zipfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_xlsxinplace_")

warnings.filterwarnings("ignore", module="openpyxl")

from app.filters import get_filter
from app.filters.xlsx_filter import XlsxFilter
from tests.fixtures.formats.ooxml_strict import make_strict
from tests.fixtures.formats.xlsx import make_xlsx_objects as M

KEPT = (
    "xl/drawings/vmlDrawing1.vml",
    "xl/media/image1.emf",
    "xl/charts/chart1.xml",
    "xl/printerSettings/printerSettings1.bin",
    "xl/calcChain.xml",
    "xl/comments1.xml",
    "xl/styles.xml",
    "xl/workbook.xml",
    "[Content_Types].xml",
    "xl/worksheets/_rels/sheet1.xml.rels",
    "xl/drawings/_rels/drawing1.xml.rels",
)


def _members(data: bytes) -> dict:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return {i.filename: zf.read(i) for i in zf.infolist()}


def _wb(data: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(data))


def _all(prefix="DE "):
    return {str(i): prefix + t for i, t in enumerate(M.TEXTS)}


def test_a_parts_survive():
    f = XlsxFilter()
    data = M.make_workbook()
    before = _members(data)
    same = _members(f.build(data, {}))
    assert list(same) == list(before), "member list and order kept"
    changed = [k for k in before if before[k] != same[k]]
    assert changed == [], changed
    out = _members(f.build_localized(data, _all(), "de-DE"))
    assert list(out) == list(before)
    for name in KEPT:
        assert out[name] == before[name], f"{name} must come back byte for byte"
    assert b">Box line one<" not in out["xl/drawings/drawing1.xml"]
    print("  [A] untouched members byte-identical; no-op build changes nothing")


def test_b_cells():
    f = XlsxFilter()
    data = M.make_workbook()
    tr = _all()
    tr["3"] = "=1+1"  # a translation that looks like a formula stays text
    wb = _wb(f.build_localized(data, tr, "de-DE"))
    ws = wb["Report"]
    assert ws["A1"].value == "DE Quarterly report", ws["A1"].value
    assert ws["B1"].value == "DE Bold title"
    assert ws["A2"].value == "DE Mixed run text"
    assert ws["B2"].value == "=1+1" and ws["B2"].data_type == "s"
    assert ws["A3"].value == "DE Plain str cell"
    assert ws["B3"].value == "---"
    assert ws["B4"].value == "DE Quarterly report"
    assert ws["A5"].value == "DE Implicit row cell", ws["A5"].value
    assert ws["C1"].value == 42 and ws["C2"].value == "=C1*2"
    assert wb["Notes"]["A1"].value == "DE Second sheet text"
    assert wb.properties.language == "de-DE"
    assert len(ws.data_validations.dataValidation) == 1
    print("  [B] cells translated in place; numbers, formulas, validation kept")


def test_c_shared_strings():
    f = XlsxFilter()
    data = M.make_workbook()
    # same translation for both uses of shared item 0: rewritten in place
    out = _members(f.build(data, {"0": "Bericht", "5": "Bericht"}))
    sst = out["xl/sharedStrings.xml"].decode()
    assert sst.count("<si>") == 5 and 'uniqueCount="5"' in sst, sst
    assert "<si><t>Bericht</t></si>" in sst
    # different translations: a new item for the second, A1's item untouched elsewhere
    out = _members(f.build(data, {"0": "Bericht"}))
    sst = out["xl/sharedStrings.xml"].decode()
    assert 'uniqueCount="6"' in sst and "<si><t>Quarterly report</t></si>" in sst, sst
    ws = _wb(f.build(data, {"0": "Bericht"}))["Report"]
    assert ws["A1"].value == "Bericht" and ws["B4"].value == "Quarterly report"
    # rich text: uniform runs keep bold/size, mixed runs become plain, rPh dropped
    out = _members(f.build(data, {"1": "Fetter Titel", "2": "Gemischt"}))
    sst = out["xl/sharedStrings.xml"].decode()
    assert '<si><r><rPr><b/><sz val="14"/></rPr><t>Fetter Titel</t></r></si>' in sst, sst
    assert '<si><t>Gemischt</t><phoneticPr fontId="1"/></si>' in sst, sst
    assert "mikusu" not in sst
    # leading/trailing whitespace is preserved
    ws = _wb(f.build(data, {"0": " x ", "5": " x "}))["Report"]
    assert ws["A1"].value == " x ", repr(ws["A1"].value)
    print("  [C] shared strings: in place when shared alike, appended otherwise; runs ok")


def test_d_walk2_textbox():
    from app.filters.xlsx_filter import _iter_cells, _load_checked, walk_version

    f = XlsxFilter()
    data = M.make_workbook()
    res = f.parse(data)
    assert [u.text for u in res.units] == M.TEXTS, [u.text for u in res.units]
    assert res.skeleton["xlsx_walk"] == 2
    assert [u.style for u in res.units][M.CELL_COUNT :] == ["text box", "text box"]
    # the cell units are exactly the pre-#685 walk (openpyxl order), ids unchanged
    old = [c.value for _s, c in _iter_cells(_load_checked(data))]
    assert [u.text for u in res.units[: M.CELL_COUNT]] == old
    assert walk_version({"units": []}) == 1 and walk_version(res.skeleton) == 2
    try:
        walk_version({"xlsx_walk": 99})
        raise AssertionError("unknown walk accepted")
    except ValueError as exc:
        assert "Re-import" in str(exc)
    out = _members(f.build_localized(data, _all(), "ar-EG"))
    drawing = out["xl/drawings/drawing1.xml"].decode()
    assert "DE Box line one" in drawing and "DE Box line two" in drawing
    assert "Fallback copy" in drawing, "the mc:Fallback copy is not walked"
    assert 'lang="ar-EG"' in drawing and 'rtl="1"' in drawing
    # B-rtl hook on the patched package: every worksheet view right to left (a sheet
    # without sheetViews gets one), never for a left-to-right target
    for name in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
        assert b'rightToLeft="1"' in out[name], name
    ltr = _members(f.build_localized(data, _all(), "de-DE"))
    assert b"rightToLeft" not in ltr["xl/worksheets/sheet1.xml"]
    assert _wb(f.build_localized(data, _all(), "ar-EG"))["Notes"].sheet_view.rightToLeft
    again = f.parse(f.build(data, _all()))
    assert [u.text for u in again.units] == ["DE " + t for t in M.TEXTS]
    # walk 1: an old import never touches the drawing, even with extra keys
    out1 = _members(f.build_localized(data, _all(), "de-DE", walk=1))
    assert out1["xl/drawings/drawing1.xml"] == _members(data)["xl/drawings/drawing1.xml"]
    print("  [D] walk 2 text boxes after the cells; walk 1 leaves the drawing alone")


def test_e_column_mode():
    f = XlsxFilter()
    sheet = (
        f'{M.DECL}<worksheet xmlns="{M.SML}"><dimension ref="A1:B3"/><sheetData>'
        '<row r="1"><c r="A1" t="inlineStr"><is><t>en-US</t></is></c>'
        '<c r="B1" t="inlineStr"><is><t>de-DE</t></is></c></row>'
        '<row r="2"><c r="A2" t="inlineStr"><is><t>Hello</t></is></c></row>'
        '<row r="4" spans="1:1"><c r="A4" t="inlineStr"><is><t>Bye</t></is></c></row>'
        "</sheetData></worksheet>"
    )
    m = _members(M.make_workbook())
    m["xl/worksheets/sheet1.xml"] = sheet.encode()
    data = M._zip(m)
    res = f.parse(
        data,
        columns="source=en-US,target.de-DE=de-DE",
        source_locale="en-US",
        target_locales=["de-DE"],
    )
    assert [u.text for u in res.units] == ["Hello", "Bye"], [u.text for u in res.units]
    out = f.build_localized(
        data, {"0": "Hallo", "1": "Tschuess"}, "de-DE", table=res.skeleton["table"]
    )
    ws = _wb(out)["Report"]
    assert (ws["B2"].value, ws["B4"].value, ws["A2"].value) == ("Hallo", "Tschuess", "Hello")
    xml = _members(out)["xl/worksheets/sheet1.xml"].decode()
    assert xml.index('r="A4"') < xml.index('r="B4"') and 'ref="A1:B4"' in xml, xml
    assert _members(out)["xl/drawings/drawing1.xml"] == m["xl/drawings/drawing1.xml"]
    print("  [E] column mode creates target cells in order; drawing untouched")


def test_f_packages():
    f = XlsxFilter()
    data = M.make_workbook(macro=True)
    out = _members(f.build_localized(data, _all(), "de-DE"))
    assert out["xl/vbaProject.bin"] == M.VBA
    assert M.XLSM_MAIN.encode() in out["[Content_Types].xml"]
    # Strict in, Strict out: untouched Strict parts keep their bytes
    strict = make_strict.make_xlsx()
    texts = [u.text for u in f.parse(strict).units]
    sout = f.build_localized(strict, {str(i): "DE " + t for i, t in enumerate(texts)}, "de-DE")
    sm, om = _members(sout), _members(strict)
    assert b"purl.oclc.org/ooxml/spreadsheetml/main" in sm["xl/worksheets/sheet1.xml"]
    assert (
        sm["xl/styles.xml"] == om["xl/styles.xml"]
        and sm["xl/workbook.xml"] == om["xl/workbook.xml"]
    )
    assert [u.text for u in f.parse(sout).units] == ["DE " + t for t in texts]
    # no docProps/core.xml: one is created and related for the language
    bare = f.build_localized(M.make_workbook(core=False), {}, "fr-FR")
    assert _wb(bare).properties.language == "fr-FR"
    bm = _members(bare)
    assert (
        b"core-properties" in bm["_rels/.rels"]
        and b"/docProps/core.xml" in bm["[Content_Types].xml"]
    )
    print("  [F] xlsm VBA kept; Strict round trip stays Strict; core.xml created when absent")


def test_g_end_to_end():
    from sqlmodel import Session, select

    from app.db import engine, init_db
    from app.models import Document, Project, Segment, SegStatus, TargetLocale
    from app.services import export as export_svc
    from app.services.export import ExportUnreadableBlob
    from app.services.ingest import ingest_document

    init_db()
    with Session(engine) as s:
        p = Project(name="xlsx685", source_locale="en", default_provider="mock")
        s.add(p)
        s.commit()
        s.refresh(p)
        s.add(TargetLocale(project_id=p.id, locale="de-DE"))
        s.commit()
        doc = ingest_document(s, p, "objects.xlsx", M.make_workbook(), "segmented")
        segs = s.exec(
            select(Segment).where(Segment.document_id == doc.id, Segment.target_locale == "de-DE")
        ).all()
        assert sorted(x.source_text for x in segs) == sorted(M.TEXTS), [x.source_text for x in segs]
        for seg in segs:
            seg.target_text = "DE " + seg.source_text
            seg.status = SegStatus.MT
            s.add(seg)
        s.commit()
        out = export_svc.build_translated_file(s, doc, "de-DE")
        assert [u.text for u in get_filter("xlsx").parse(out).units] == ["DE " + t for t in M.TEXTS]
        assert _members(out)["xl/media/image1.emf"] == M.EMF
        # an old (walk 1) import: the drawing stays as uploaded
        doc = s.get(Document, doc.id)
        sk = dict(doc.skeleton)
        sk.pop("xlsx_walk")
        doc.skeleton = sk
        s.add(doc)
        s.commit()
        out1 = export_svc.build_translated_file(s, doc, "de-DE")
        assert b">Box line one<" in _members(out1)["xl/drawings/drawing1.xml"]
        sk = dict(doc.skeleton)
        sk["xlsx_walk"] = 7
        doc.skeleton = sk
        s.add(doc)
        s.commit()
        try:
            export_svc.build_translated_file(s, doc, "de-DE")
            raise AssertionError("an unknown walk stamp must be refused")
        except ExportUnreadableBlob as exc:
            assert "Re-import" in str(exc)
    print("  [G] upload -> export keeps the picture; walk 1 and unknown stamps handled")


def _column_book(rows_xml: str) -> bytes:
    sheet = (
        f'{M.DECL}<worksheet xmlns="{M.SML}"><sheetData>'
        '<row r="1"><c r="A1" t="inlineStr"><is><t>en-US</t></is></c>'
        '<c r="B1" t="inlineStr"><is><t>de-DE</t></is></c></row>'
        f"{rows_xml}</sheetData></worksheet>"
    )
    m = _members(M.make_workbook())
    m["xl/worksheets/sheet1.xml"] = sheet.encode()
    return M._zip(m)


def _column_parse(f, data):
    return f.parse(
        data,
        columns="source=en-US,target.de-DE=de-DE",
        source_locale="en-US",
        target_locales=["de-DE"],
    )


def test_h_gate_formula_targets_and_scale():
    """Gate: a target cell that held a formula. A shared-formula master gives its
    dependents their own formula before it becomes text (else they point at a group
    with no formula: a corrupt file), and the calc chain (which would name a cell with
    no formula) is dropped with its content type and relationship. Column mode stays
    linear in the row count (the row lookup was a scan per created cell)."""
    import time

    f = XlsxFilter()
    data = _column_book(
        '<row r="2"><c r="A2" t="inlineStr"><is><t>Hello</t></is></c>'
        '<c r="B2" t="str"><f t="shared" ref="B2:B4" si="0">UPPER(A2)</f><v>HELLO</v></c></row>'
        '<row r="3"><c r="A3" t="inlineStr"><is><t>Bye</t></is></c>'
        '<c r="B3" t="str"><f t="shared" si="0"/><v>BYE</v></c></row>'
        '<row r="4"><c r="A4" t="inlineStr"><is><t>Yes</t></is></c>'
        '<c r="B4" t="str"><f t="shared" si="0"/><v>YES</v></c></row>'
    )
    res = _column_parse(f, data)
    out = f.build_localized(data, {"0": "Hallo"}, "de-DE", table=res.skeleton["table"])
    om = _members(out)
    xml = om["xl/worksheets/sheet1.xml"].decode()
    assert '<c r="B2" t="inlineStr"><is><t>Hallo</t></is></c>' in xml, xml
    assert "<f>UPPER(A3)</f>" in xml and "<f>UPPER(A4)</f>" in xml, xml
    assert 'si="0"' not in xml, xml
    assert "xl/calcChain.xml" not in om
    assert b"calcChain" not in om["[Content_Types].xml"]
    assert b"calcChain" not in om["xl/_rels/workbook.xml.rels"]
    ws = _wb(out)["Report"]
    assert (ws["B2"].value, ws["B3"].value, ws["B4"].value) == ("Hallo", "=UPPER(A3)", "=UPPER(A4)")
    # no formula touched: the calc chain stays, byte for byte
    plain = _column_book('<row r="2"><c r="A2" t="inlineStr"><is><t>Hello</t></is></c></row>')
    res = _column_parse(f, plain)
    out = _members(f.build_localized(plain, {"0": "Hallo"}, "de-DE", table=res.skeleton["table"]))
    assert out["xl/calcChain.xml"] == _members(plain)["xl/calcChain.xml"]
    # scale: 20000 new target cells (a scan per cell took about 21 s for 5000 rows)
    n = 20000
    big = _column_book(
        "".join(
            f'<row r="{i}"><c r="A{i}" t="inlineStr"><is><t>Text {i}</t></is></c></row>'
            for i in range(2, n + 2)
        )
    )
    res = _column_parse(f, big)
    t0 = time.monotonic()
    out = f.build_localized(
        big, {str(i): f"DE {i}" for i in range(n)}, "de-DE", table=res.skeleton["table"]
    )
    took = time.monotonic() - t0
    assert took < 60, f"column mode took {took:.1f}s for {n} rows"
    xml = _members(out)["xl/worksheets/sheet1.xml"].decode()
    assert xml.index(f'r="A{n + 1}"') < xml.index(f'r="B{n + 1}"')
    print(f"  [H] formula targets unshared, calc chain dropped; {n} column rows in {took:.1f}s")


if __name__ == "__main__":
    test_a_parts_survive()
    test_b_cells()
    test_c_shared_strings()
    test_d_walk2_textbox()
    test_e_column_mode()
    test_f_packages()
    test_g_end_to_end()
    test_h_gate_formula_targets_and_scale()
    print("XLSX INPLACE EXPORT OK")
