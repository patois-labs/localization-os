"""Workflow segmented write-back verification (#43, offline).

Phase 1 of the LLM-workflow write-back plan: the multi-agent orchestrator now
runs over SEGMENTED / per-row content and writes each Segment.target_text (it was
FULL_TEXT-only before). This suite proves:

  (1) a SEGMENTED doc with 2+ segments run through a 2-step workflow writes EVERY
      seg.target_text (all non-empty, correctly chained per segment);
  (2) the FULL_TEXT path still writes its single segment (regression — the sacred
      gate: the FULL_TEXT body is byte-unchanged, so its behavior must be too);
  (3) human_review over a multi-segment doc pauses ONCE at run level, resume seeds
      the edited blob back + re-splits by index, run completes, every seg written;
  (4) TM short-circuit: a 100% segment-TM match on step 0 fills that segment WITHOUT
      a step-0 provider call (proven by a counting provider) while the other segment
      still calls the provider.

Hermetic: throwaway LOC_TMS_DATA_DIR + cleared provider keys (incl LOCAL_BASE_URL)
force MockProvider. ASCII prints only.
"""

import os
import sys
from io import BytesIO

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_wfseg_")

from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app.db import engine, init_db
from app.main import app
from app.models import (
    FullTextTM,
    RunStatus,
    Segment,
    SegmentTM,
    StepRun,
    StepType,
    Workflow,
    WorkflowRun,
    WorkflowStep,
)
from app.services import orchestrator, settings, tm
from app.services.providers import Provider, register_provider


class Counter(Provider):
    """A translate provider that counts calls (to prove the TM short-circuit
    skips step-0 provider calls for TM-matched segments)."""

    name = "counter"
    calls = 0

    def complete(self, prompt, system=None, task=None, payload=None, **kw):
        Counter.calls += 1
        return "[counter] " + (payload or {}).get("text", "")


class Capture(Provider):
    """A translate provider that RECORDS the exact text (and full prompt) the
    engine was asked to translate, so a test can prove the engine only ever saw
    PII-redacted text (placeholders, never the raw email/SSN). It ECHOES the text
    back (like MockProvider) so the orchestrator's pii.restore() genuinely
    reinstates the PII in the delivered target — exercising the full
    redact -> engine -> restore round-trip."""

    name = "capture"

    def __init__(self):
        self.seen_texts = []
        self.seen_prompts = []

    def complete(self, prompt, system=None, task=None, payload=None, **kw):
        text = (payload or {}).get("text", "")
        self.seen_texts.append(text)
        self.seen_prompts.append(prompt or "")
        return "[fr-FR] " + text


class Boom(Provider):
    """A provider whose completions always raise — run_completion exhausts retries
    and returns text='' with an error, so the orchestrator sees an empty step
    output (proves failed-step integrity: FAILED status, NO blank target written)."""

    name = "boom"

    def complete(self, prompt, system=None, task=None, payload=None, **kw):
        raise RuntimeError("provider outage")


class LqaCapture(Provider):
    """Captures the exact source+target blobs an LQA step handed the engine, so a
    test can prove redact_pii scrubs PII on the LQA path too. Returns a minimal
    scorecard shape lqa.score_text expects to build its card from."""

    name = "lqacap"

    def __init__(self):
        self.seen = []

    def complete(self, prompt, system=None, task=None, payload=None, **kw):
        self.seen.append(prompt or "")
        p = payload or {}
        self.seen.append(str(p.get("source", "")) + "\n" + str(p.get("target", "")))
        # empty error list -> lqa builds a clean scorecard
        return '{"score": 100, "errors": []}'


# raw PII planted in segment sources — must NEVER reach the engine when redacting
PII_EMAIL = "john.doe@example.com"
PII_SSN = "123-45-6789"


def docx_from(paragraphs) -> bytes:
    d = Docx()
    for p in paragraphs:
        d.add_paragraph(p)
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def new_project(client, name, targets="fr-FR"):
    r = client.post(
        "/projects",
        data={
            "name": name,
            "source_locale": "en-US",
            "target_locales": targets,
            "default_provider": "claude",
        },
        follow_redirects=False,
    )
    return int(r.headers["location"].rsplit("/", 1)[-1])


def upload(client, pid, paragraphs, mode="segmented"):
    r = client.post(
        f"/projects/{pid}/documents",
        data={"mode": mode},
        files={
            "file": (
                "d.docx",
                docx_from(paragraphs),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    return int(r.headers["location"].rsplit("/", 1)[-1])


def add_step(
    session,
    wf_id,
    order,
    name,
    step_type=StepType.LLM_AGENT,
    provider="claude",
    prompt_text="",
    chunk_policy="auto",
):
    session.add(
        WorkflowStep(
            workflow_id=wf_id,
            order_index=order,
            name=name,
            step_type=step_type,
            provider=provider,
            prompt_text=prompt_text,
            chunk_policy=chunk_policy,
            retry_max=1,
            backoff_base=0.0,
        )
    )
    session.commit()


def segments_for(doc_id, locale="fr-FR"):
    with Session(engine) as s:
        return sorted(
            s.exec(
                select(Segment).where(
                    Segment.document_id == doc_id, Segment.target_locale == locale
                )
            ).all(),
            key=lambda x: x.order_index,
        )


def main() -> None:
    init_db()
    client = TestClient(app)

    # =====================================================================
    # (1) SEGMENTED doc, 2 segments, 2-step workflow -> EVERY seg written
    # =====================================================================
    with Session(engine) as s:
        wf = Workflow(name="Seg 2-step", scope="global")
        s.add(wf)
        s.commit()
        s.refresh(wf)
        wf_id = wf.id
        add_step(s, wf_id, 0, "Translator", prompt_text="Translate: {text}")
        add_step(s, wf_id, 1, "Reviewer", prompt_text="Review: {text}")

    pid = new_project(client, "Seg WF")
    doc = upload(client, pid, ["Adventure awaits.", "Save big this weekend."])
    with Session(engine) as s:
        summary = orchestrator.run_workflow(s, wf_id, doc, "fr-FR")
    assert summary["status"] == "completed", summary
    assert len(summary["steps"]) == 2, summary

    segs = segments_for(doc)
    assert len(segs) == 2, ("expected 2 segments", len(segs))
    for sg in segs:
        assert (sg.target_text or "").strip(), ("segment left untranslated", sg.order_index)
        # two agent passes => the mock prefix nests twice (per-segment chaining)
        assert sg.target_text.count("[fr-FR]") >= 2, (
            "per-segment chaining missing",
            sg.target_text,
        )
        assert sg.status == "mt" and sg.mt_confidence == 0.9, (sg.status, sg.mt_confidence)
    # each segment carries its OWN source, not a merged blob
    assert "Adventure" in segs[0].target_text and "Save big" in segs[1].target_text, (
        "segments must translate their own source, not a merged blob",
        segs[0].target_text,
        segs[1].target_text,
    )
    # ONE StepRun per STEP (aggregate) so observe_run/SSE are unchanged
    with Session(engine) as s:
        srs = s.exec(
            select(StepRun).where(StepRun.workflow_run_id == summary["run_id"]).order_by(StepRun.id)
        ).all()
    assert len(srs) == 2, ("expected one StepRun per step (aggregate)", len(srs))
    print(
        "  (1) segmented 2-seg doc: both seg.target_text written, per-segment chained, "
        "2 aggregate StepRuns"
    )

    # =====================================================================
    # (2) FULL_TEXT path still writes its single segment (regression)
    # =====================================================================
    with Session(engine) as s:
        wff = Workflow(name="Full 1-step", scope="global")
        s.add(wff)
        s.commit()
        s.refresh(wff)
        wff_id = wff.id
        add_step(s, wff_id, 0, "Translator", prompt_text="Translate: {text}")
    ftdoc = upload(client, pid, ["Whole document one.", "Whole document two."], mode="full_text")
    with Session(engine) as s:
        ft = orchestrator.run_workflow(s, wff_id, ftdoc, "fr-FR")
    assert ft["status"] == "completed", ft
    ftsegs = segments_for(ftdoc)
    # full-text mode = one segment holding the whole doc; it must be filled
    filled = [sg for sg in ftsegs if (sg.target_text or "").strip()]
    assert filled, "FULL_TEXT run wrote no target text (regression)"
    assert "[fr-FR]" in filled[0].target_text, filled[0].target_text
    print(f"  (2) full_text path unchanged: single segment written ({len(filled)} filled)")

    # =====================================================================
    # (3) human_review over multi-segment doc -> ONE run-level pause -> resume
    # =====================================================================
    with Session(engine) as s:
        wfh = Workflow(name="Seg human", scope="global")
        s.add(wfh)
        s.commit()
        s.refresh(wfh)
        wfh_id = wfh.id
        add_step(s, wfh_id, 0, "Translator", prompt_text="Translate: {text}")
        add_step(s, wfh_id, 1, "Human", step_type=StepType.HUMAN_REVIEW)
        add_step(s, wfh_id, 2, "Finalizer", prompt_text="Finalize: {text}")
    hdoc = upload(client, pid, ["First para here.", "Second para here.", "Third para here."])
    with Session(engine) as s:
        paused = orchestrator.run_workflow(s, wfh_id, hdoc, "fr-FR")
    assert paused["status"] == "paused" and paused["paused_step"] == 1, paused
    blob = paused["paused_text"]
    # the pause blob assembles ALL segment targets joined by the known joiner
    parts = blob.split(orchestrator.SEG_JOINER)
    assert len(parts) == 3, ("pause blob must assemble all 3 segment targets", len(parts))
    for p in parts:
        assert "[fr-FR]" in p, ("each assembled part is a translated segment", p)
    run_id = paused["run_id"]
    with Session(engine) as s:
        run = s.get(WorkflowRun, run_id)
        assert run.status == RunStatus.PAUSED and run.current_step_index == 1, (
            run.status,
            run.current_step_index,
        )
    # operator edits the blob (keeping the joiner) and resumes
    edited = orchestrator.SEG_JOINER.join(["EDITED ONE", "EDITED TWO", "EDITED THREE"])
    with Session(engine) as s:
        resumed = orchestrator.resume_workflow(s, run_id, "fr-FR", edited)
    assert resumed["status"] == "completed", resumed
    hsegs = segments_for(hdoc)
    assert len(hsegs) == 3, len(hsegs)
    for i, sg in enumerate(hsegs):
        assert (sg.target_text or "").strip(), ("segment left untranslated after resume", i)
        # the finalizer ran on the edited per-segment text
        assert "EDITED" in sg.target_text, (
            "edited blob must re-split by index to segments",
            sg.target_text,
        )
    assert "EDITED ONE" in hsegs[0].target_text and "EDITED THREE" in hsegs[2].target_text, (
        "re-split must align edited parts to segments by index",
        hsegs[0].target_text,
        hsegs[2].target_text,
    )
    print(
        "  (3) human_review: ONE run-level pause over 3 segments, resume re-splits "
        "edited blob by index, all written"
    )

    # =====================================================================
    # (4) TM short-circuit on step 0 skips the provider call for a 100% match
    # =====================================================================
    register_provider("counter", Counter())
    tmpid = new_project(client, "TM SC")
    tmdoc = upload(client, tmpid, ["Known sentence.", "Unknown sentence."])
    # seed a 100% segment-TM entry for the FIRST segment's source only
    with Session(engine) as s:
        tm.write_back(s, tmpid, "Known sentence.", "PHRASE CONNUE", "en-US", "fr-FR")
    with Session(engine) as s:
        wftm = Workflow(name="TM SC wf", scope="global")
        s.add(wftm)
        s.commit()
        s.refresh(wftm)
        wftm_id = wftm.id
        # step 0 uses the counting provider so we can observe skipped calls
        add_step(s, wftm_id, 0, "Translator", provider="counter", prompt_text="Translate: {text}")
    Counter.calls = 0
    with Session(engine) as s:
        tmrun = orchestrator.run_workflow(s, wftm_id, tmdoc, "fr-FR")
    assert tmrun["status"] == "completed", tmrun
    # only the UNKNOWN segment should have hit the provider (1 call, not 2)
    assert Counter.calls == 1, (
        "TM 100% hit must skip the step-0 provider call for that segment",
        Counter.calls,
    )
    tmsegs = segments_for(tmdoc)
    assert tmsegs[0].target_text == "PHRASE CONNUE", (
        "TM-matched segment target must come from TM verbatim",
        tmsegs[0].target_text,
    )
    assert "[counter]" in tmsegs[1].target_text or "Unknown" in tmsegs[1].target_text, (
        "non-matched segment must be provider-translated",
        tmsegs[1].target_text,
    )
    print(
        "  (4) TM short-circuit: 100% step-0 match skips provider (1 call, not 2); "
        "TM target written verbatim"
    )

    # =====================================================================
    # (5) PII per-step redaction (segmented path)
    #   (a) step.redact_pii=True -> engine sees REDACTED text (placeholders, no
    #       raw email/SSN); the raw PII never lands in the ENGINE input,
    #       segment-TM, full-text-TM, or StepRun.output_text as RAW *unredacted*
    #       leakage; it IS correctly restored in the final target (mirror
    #       translate.py: the delivered target legitimately carries the PII back),
    #       and no orphan «PII_...» placeholder tokens remain.
    #   (b) step WITHOUT the flag + global pii OFF -> text passes through raw.
    #   (c) global pii_redaction ON + step flag OFF -> still redacts (the
    #       `step.redact_pii or pii.enabled()` path).
    # =====================================================================
    def _no_placeholder(s):
        return "«PII_" not in (s or "")

    # --- (5a) per-step redact_pii=True ---------------------------------------
    cap_a = Capture()
    register_provider("capture", cap_a)
    settings.set_many({"pii_redaction": "0"})  # global OFF: the STEP opts in
    pid_a = new_project(client, "PII step on")
    doc_a = upload(
        client, pid_a, [f"Email me at {PII_EMAIL} today.", f"My SSN is {PII_SSN} exactly."]
    )
    with Session(engine) as s:
        wf_a = Workflow(name="PII redact wf", scope="global")
        s.add(wf_a)
        s.commit()
        s.refresh(wf_a)
        wf_a_id = wf_a.id
        s.add(
            WorkflowStep(
                workflow_id=wf_a_id,
                order_index=0,
                name="Translator",
                step_type=StepType.LLM_AGENT,
                provider="capture",
                prompt_text="Translate: {text}\nSource: {source}",
                chunk_policy="auto",
                retry_max=1,
                backoff_base=0.0,
                redact_pii=True,
            )
        )
        s.commit()
    with Session(engine) as s:
        run_a = orchestrator.run_workflow(s, wf_a_id, doc_a, "fr-FR")
    assert run_a["status"] == "completed", run_a

    # the engine NEVER saw the raw email or SSN — only placeholders
    all_seen = "\n".join(cap_a.seen_texts) + "\n" + "\n".join(cap_a.seen_prompts)
    assert cap_a.seen_texts, "capture engine was never called"
    assert PII_EMAIL not in all_seen, ("raw email reached the engine", all_seen)
    assert PII_SSN not in all_seen, ("raw SSN reached the engine", all_seen)
    assert "«PII_EMAIL" in all_seen and "«PII_SSN" in all_seen, (
        "engine should have seen PII placeholders",
        all_seen,
    )

    # the final target correctly RESTORES the PII (delivered file is complete),
    # and carries NO orphan placeholder tokens (restore fully completed)
    segs_a = segments_for(doc_a)
    assert PII_EMAIL in segs_a[0].target_text, (
        "email must be restored in target",
        segs_a[0].target_text,
    )
    assert PII_SSN in segs_a[1].target_text, (
        "ssn must be restored in target",
        segs_a[1].target_text,
    )
    for sg in segs_a:
        assert _no_placeholder(sg.target_text), (
            "orphan PII placeholder left in target",
            sg.target_text,
        )

    # StepRun.output_text carries NO orphan placeholder (restored blob), and
    # StepRun.input_text (the step's incoming working text) likewise has no stray
    # placeholder token left over.
    with Session(engine) as s:
        srs_a = s.exec(
            select(StepRun).where(StepRun.workflow_run_id == run_a["run_id"]).order_by(StepRun.id)
        ).all()
    for sr in srs_a:
        assert _no_placeholder(sr.output_text), (
            "orphan placeholder in StepRun.output_text",
            sr.output_text,
        )
        assert _no_placeholder(sr.input_text), (
            "orphan placeholder in StepRun.input_text",
            sr.input_text,
        )

    # the RAW PII must never have leaked into the TM tables in *redacted* form,
    # and no placeholder token was persisted to any TM row (orchestrator doesn't
    # write TM, but assert the tables are clean of placeholder leakage regardless)
    with Session(engine) as s:
        seg_tm_rows = s.exec(select(SegmentTM)).all()
        ft_tm_rows = s.exec(select(FullTextTM)).all()
    for row in seg_tm_rows + ft_tm_rows:
        assert _no_placeholder(row.source_text) and _no_placeholder(row.target_text), (
            "PII placeholder leaked into a TM row",
            row.source_text,
            row.target_text,
        )
    print(
        "  (5a) per-step redact_pii=True: engine saw placeholders (no raw email/SSN); "
        "final target restores PII; no orphan placeholders in target/StepRun/TM"
    )

    # --- (5b) no flag + global OFF -> raw passthrough ------------------------
    cap_b = Capture()
    register_provider("capture", cap_b)
    settings.set_many({"pii_redaction": "0"})
    pid_b = new_project(client, "PII off")
    doc_b = upload(client, pid_b, [f"Reach me at {PII_EMAIL} now."])
    with Session(engine) as s:
        wf_b = Workflow(name="PII off wf", scope="global")
        s.add(wf_b)
        s.commit()
        s.refresh(wf_b)
        wf_b_id = wf_b.id
        s.add(
            WorkflowStep(
                workflow_id=wf_b_id,
                order_index=0,
                name="Translator",
                step_type=StepType.LLM_AGENT,
                provider="capture",
                prompt_text="Translate: {text}",
                chunk_policy="auto",
                retry_max=1,
                backoff_base=0.0,
                redact_pii=False,
            )
        )
        s.commit()
    with Session(engine) as s:
        run_b = orchestrator.run_workflow(s, wf_b_id, doc_b, "fr-FR")
    assert run_b["status"] == "completed", run_b
    # no redaction: the engine sees the RAW email (proves the flag actually gates)
    assert any(PII_EMAIL in t for t in cap_b.seen_texts), (
        "without the flag + global off, engine must see raw text",
        cap_b.seen_texts,
    )
    print("  (5b) no flag + global OFF: text passes through unredacted (engine sees raw)")

    # --- (5c) global pii_redaction ON + step flag OFF -> still redacts --------
    cap_c = Capture()
    register_provider("capture", cap_c)
    settings.set_many({"pii_redaction": "1"})  # global ON
    try:
        pid_c = new_project(client, "PII global on")
        doc_c = upload(client, pid_c, [f"Contact {PII_EMAIL} please."])
        with Session(engine) as s:
            wf_c = Workflow(name="PII global wf", scope="global")
            s.add(wf_c)
            s.commit()
            s.refresh(wf_c)
            wf_c_id = wf_c.id
            s.add(
                WorkflowStep(
                    workflow_id=wf_c_id,
                    order_index=0,
                    name="Translator",
                    step_type=StepType.LLM_AGENT,
                    provider="capture",
                    prompt_text="Translate: {text}",
                    chunk_policy="auto",
                    retry_max=1,
                    backoff_base=0.0,
                    redact_pii=False,
                )
            )
            s.commit()
        with Session(engine) as s:
            run_c = orchestrator.run_workflow(s, wf_c_id, doc_c, "fr-FR")
        assert run_c["status"] == "completed", run_c
        # global ON drives redaction even though the step flag is OFF
        seen_c = "\n".join(cap_c.seen_texts)
        assert PII_EMAIL not in seen_c, ("global pii ON must redact (or path)", seen_c)
        assert "«PII_EMAIL" in seen_c, ("global pii ON must send placeholders", seen_c)
        segs_c = segments_for(doc_c)
        assert PII_EMAIL in segs_c[0].target_text, (
            "global-on run must restore PII in target",
            segs_c[0].target_text,
        )
        print(
            "  (5c) global pii ON + step flag OFF: still redacts via `or pii.enabled()` "
            "(engine sees placeholders; PII restored in target)"
        )
    finally:
        settings.set_many({"pii_redaction": "0"})  # never leak the global on into later suites

    # =====================================================================
    # (6) F1: human_review as the LAST step -> resume persists operator edits
    #   [translate, human_review]. last step is human, so on resume start==2 and
    #   the step loop is empty; the post-loop write-back must STILL write the
    #   operator's edited blob to EVERY seg.target_text (not the pre-review MT).
    # =====================================================================
    with Session(engine) as s:
        wf6 = Workflow(name="Human last", scope="global")
        s.add(wf6)
        s.commit()
        s.refresh(wf6)
        wf6_id = wf6.id
        add_step(s, wf6_id, 0, "Translator", prompt_text="Translate: {text}")
        add_step(s, wf6_id, 1, "Human", step_type=StepType.HUMAN_REVIEW)
    doc6 = upload(client, pid, ["Alpha source.", "Beta source."])
    with Session(engine) as s:
        paused6 = orchestrator.run_workflow(s, wf6_id, doc6, "fr-FR")
    assert paused6["status"] == "paused" and paused6["paused_step"] == 1, paused6
    # pre-review targets are the MT ([fr-FR] ...), NOT yet the operator edit
    edited6 = orchestrator.SEG_JOINER.join(["OPERATOR ALPHA", "OPERATOR BETA"])
    with Session(engine) as s:
        resumed6 = orchestrator.resume_workflow(s, paused6["run_id"], "fr-FR", edited6)
    assert resumed6["status"] == "completed", resumed6
    segs6 = segments_for(doc6)
    assert len(segs6) == 2, len(segs6)
    assert segs6[0].target_text == "OPERATOR ALPHA", (
        "human-last resume must persist the edited blob, not the pre-review MT",
        segs6[0].target_text,
    )
    assert segs6[1].target_text == "OPERATOR BETA", segs6[1].target_text
    assert "[fr-FR]" not in segs6[0].target_text, (
        "edited target must NOT retain the pre-review MT",
        segs6[0].target_text,
    )
    print(
        "  (6) F1 human-review-LAST: resume writes the EDITED blob to every "
        "seg.target_text (not the discarded pre-review MT)"
    )

    # =====================================================================
    # (7) F2: a failing provider -> StepRun FAILED, WorkflowRun FAILED, and NO
    #   seg.target_text written (targets stay empty/prior, NOT blank-MT).
    # =====================================================================
    register_provider("boom", Boom())
    pid7 = new_project(client, "Boom fail")
    doc7 = upload(client, pid7, ["Will fail one.", "Will fail two."])
    with Session(engine) as s:
        wf7 = Workflow(name="Boom wf", scope="global")
        s.add(wf7)
        s.commit()
        s.refresh(wf7)
        wf7_id = wf7.id
        add_step(s, wf7_id, 0, "Translator", provider="boom", prompt_text="Translate: {text}")
    # capture prior target state (fresh doc -> empty targets)
    pre7 = segments_for(doc7)
    assert all(not (sg.target_text or "").strip() for sg in pre7), "fresh targets should be empty"
    with Session(engine) as s:
        run7 = orchestrator.run_workflow(s, wf7_id, doc7, "fr-FR")
    # summary is NOT 'completed' (the run failed)
    assert run7["status"] != "completed", ("failed run must not report completed", run7)
    with Session(engine) as s:
        wr7 = s.get(WorkflowRun, run7["run_id"])
        assert wr7.status == RunStatus.FAILED, ("WorkflowRun must be FAILED", wr7.status)
        srs7 = s.exec(
            select(StepRun).where(StepRun.workflow_run_id == run7["run_id"]).order_by(StepRun.id)
        ).all()
    assert srs7 and srs7[-1].status == RunStatus.FAILED, (
        "failed step's StepRun must be FAILED",
        [sr.status for sr in srs7],
    )
    post7 = segments_for(doc7)
    for sg in post7:
        assert not (sg.target_text or "").strip(), (
            "failed run must leave NO blank/partial target written",
            sg.order_index,
            sg.target_text,
        )
    print(
        "  (7) F2 failed step: StepRun FAILED + WorkflowRun FAILED + NO target "
        "written (no blank-MT)"
    )

    # =====================================================================
    # (7b) HARDENING: a workflow STEP whose provider is a dead engine:<id> ref
    #      (disabled/deleted/keyless) FAILS the run LOUD — StepRun + WorkflowRun
    #      FAILED, the error names the dead ref, and NO mock output is written.
    #      (Mirrors (7) but the failure source is a dead engine ref, not an outage.)
    # =====================================================================
    from app.services import engines as _eng

    prof_dead = _eng.create_profile("WFDeadDeepL", "deepl", api_key="k1")
    _eng.update_profile(prof_dead.id, enabled=False)  # dead AFTER it was configured
    dead_ref = f"engine:{prof_dead.id}"
    pid7b = new_project(client, "WF dead engine")
    doc7b = upload(client, pid7b, ["Segment one.", "Segment two."])
    with Session(engine) as s:
        wf7b = Workflow(name="WF dead engine wf", scope="global")
        s.add(wf7b)
        s.commit()
        s.refresh(wf7b)
        wf7b_id = wf7b.id
        add_step(s, wf7b_id, 0, "Translator", provider=dead_ref, prompt_text="Translate: {text}")
    with Session(engine) as s:
        run7b = orchestrator.run_workflow(s, wf7b_id, doc7b, "fr-FR")
    assert run7b["status"] != "completed", ("dead-engine run must not report completed", run7b)
    assert dead_ref in (run7b.get("error") or ""), (
        "workflow error must name the dead ref",
        run7b.get("error"),
    )
    with Session(engine) as s:
        wr7b = s.get(WorkflowRun, run7b["run_id"])
        assert wr7b.status == RunStatus.FAILED, ("WorkflowRun must be FAILED", wr7b.status)
    post7b = segments_for(doc7b)
    for sg in post7b:
        assert not (sg.target_text or "").strip(), (
            "dead-engine run must leave NO mock target written",
            sg.order_index,
            sg.target_text,
        )
    print(
        "  (7b) dead engine:<id> step: WorkflowRun FAILED, error names the ref, "
        "NO mock target written"
    )

    # =====================================================================
    # (8) Issue #23 (b): resume with an altered SEG_JOINER count must REFUSE
    #   cleanly (run stays PAUSED + re-resumable), NOT terminally FAIL.
    #   (Behavior CHANGED from the old F4 terminal-FAIL: the count check moved to a
    #   graceful re-pause; the old assertion asserted the bug this issue fixes.)
    # =====================================================================
    with Session(engine) as s:
        wf8 = Workflow(name="Count mismatch", scope="global")
        s.add(wf8)
        s.commit()
        s.refresh(wf8)
        wf8_id = wf8.id
        add_step(s, wf8_id, 0, "Translator", prompt_text="Translate: {text}")
        add_step(s, wf8_id, 1, "Human", step_type=StepType.HUMAN_REVIEW)
        add_step(s, wf8_id, 2, "Finalizer", prompt_text="Finalize: {text}")
    doc8 = upload(client, pid, ["Gamma one.", "Gamma two."])
    with Session(engine) as s:
        paused8 = orchestrator.run_workflow(s, wf8_id, doc8, "fr-FR")
    assert paused8["status"] == "paused", paused8
    run8_id = paused8["run_id"]
    pre8 = segments_for(doc8)
    pre8_targets = [sg.target_text for sg in pre8]
    # operator deleted a delimiter: 2 segments but only 1 part now
    bad_blob = "ONLY ONE PART NO DELIMITER"
    with Session(engine) as s:
        resumed8 = orchestrator.resume_workflow(s, run8_id, "fr-FR", bad_blob)
    # NEW: the generator refuses via a distinct 'rejected' status; the run is NOT failed.
    assert resumed8["status"] == "rejected", (
        "count-mismatch resume must REFUSE (status=='rejected'), not fail",
        resumed8,
    )
    assert "segment count" in (resumed8.get("error") or ""), (
        "rejected summary must carry the realignment reason",
        resumed8,
    )
    with Session(engine) as s:
        wr8 = s.get(WorkflowRun, run8_id)
        assert wr8.status == RunStatus.PAUSED, (
            "count-mismatch must leave the run PAUSED, not FAILED",
            wr8.status,
        )
        assert wr8.current_step_index == 1, (
            "cursor must be restored to the human_review step (re-resumable)",
            wr8.current_step_index,
        )
        assert wr8.finished_at is None, ("a still-paused run must not be finished", wr8)
    post8 = segments_for(doc8)
    assert [sg.target_text for sg in post8] == pre8_targets, (
        "count-mismatch must not corrupt/overwrite targets",
        pre8_targets,
        [sg.target_text for sg in post8],
    )
    # RE-RESUMABLE: a subsequent CORRECT resume completes the run.
    good_blob = orchestrator.SEG_JOINER.join(["FIX ONE", "FIX TWO"])
    with Session(engine) as s:
        resumed8b = orchestrator.resume_workflow(s, run8_id, "fr-FR", good_blob)
    assert resumed8b["status"] == "completed", (
        "run must be re-resumable after a refused mismatch",
        resumed8b,
    )
    post8b = segments_for(doc8)
    assert all("FIX" in (sg.target_text or "") for sg in post8b), (
        "the corrected edit must land after re-resume",
        [s.target_text for s in post8b],
    )
    print(
        "  (8) issue #23(b) resume count-mismatch: REFUSED cleanly (status=='rejected'), "
        "run stays PAUSED @ human step, targets intact, then re-resumes to completion"
    )

    # =====================================================================
    # (8b) Issue #23 (a) show-run + (b) --edited-segments round-trip
    # =====================================================================
    with Session(engine) as s:
        wf8c = Workflow(name="Show + segments", scope="global")
        s.add(wf8c)
        s.commit()
        s.refresh(wf8c)
        wf8c_id = wf8c.id
        add_step(s, wf8c_id, 0, "Translator", prompt_text="Translate: {text}")
        add_step(s, wf8c_id, 1, "Human", step_type=StepType.HUMAN_REVIEW)
        add_step(s, wf8c_id, 2, "Finalizer", prompt_text="Finalize: {text}")
    doc8c = upload(client, pid, ["Delta one.", "Delta two.", "Delta three."])
    with Session(engine) as s:
        paused8c = orchestrator.run_workflow(s, wf8c_id, doc8c, "fr-FR")
        run8c_id = paused8c["run_id"]
        # (a) show-run exposes the segments (count + content) for a paused run
        view = orchestrator.paused_review(s, run8c_id)
    assert view is not None and view["segment_count"] == 3, (
        "show-run must expose all 3 segments for the paused run",
        view,
    )
    assert all("Delta" in seg for seg in view["segments"]), (
        "each exposed segment must carry its own reviewed text",
        view["segments"],
    )
    assert view["paused_step"]["type"] == StepType.HUMAN_REVIEW, view["paused_step"]
    assert view["delimiter"] == orchestrator.SEG_JOINER, view
    # (b) wrong-count --edited-segments is REFUSED before any state change
    with Session(engine) as s:
        try:
            orchestrator.prepare_resume_seed(s, run8c_id, edited_segments=["only", "two"])
            raise AssertionError("wrong-count edited_segments must raise ResumeMismatch")
        except orchestrator.ResumeMismatch:
            pass
        wr8c = s.get(WorkflowRun, run8c_id)
        assert wr8c.status == RunStatus.PAUSED, (
            "a refused edited_segments must not touch run state",
            wr8c.status,
        )
    # (b) correct-count --edited-segments round-trips to completion (no hand delimiters)
    with Session(engine) as s:
        seed = orchestrator.prepare_resume_seed(
            s, run8c_id, edited_segments=["SEG A", "SEG B", "SEG C"]
        )
        resumed8c = orchestrator.resume_workflow(s, run8c_id, "fr-FR", seed)
    assert resumed8c["status"] == "completed", resumed8c
    post8c = segments_for(doc8c)
    got = [sg.target_text or "" for sg in post8c]
    assert "SEG A" in got[0] and "SEG B" in got[1] and "SEG C" in got[2], (
        "edited_segments must map 1:1 to segments by index",
        got,
    )
    print(
        "  (8c) issue #23(a/b): show-run exposes 3 segments; --edited-segments "
        "wrong-count refused (run untouched), correct-count round-trips 1:1"
    )

    # =====================================================================
    # (9) F3: LQA step honors redact_pii -> engine sees ONLY placeholders
    #   [translate(redact_pii=True), lqa] on a doc with email+SSN in source.
    # =====================================================================
    lqa_cap = LqaCapture()
    register_provider("lqacap", lqa_cap)
    cap9 = Capture()
    register_provider("capture", cap9)
    settings.set_many({"pii_redaction": "0"})
    pid9 = new_project(client, "LQA pii")
    doc9 = upload(client, pid9, [f"Mail {PII_EMAIL} here.", f"SSN {PII_SSN} here."])
    with Session(engine) as s:
        wf9 = Workflow(name="LQA pii wf", scope="global")
        s.add(wf9)
        s.commit()
        s.refresh(wf9)
        wf9_id = wf9.id
        s.add(
            WorkflowStep(
                workflow_id=wf9_id,
                order_index=0,
                name="Translator",
                step_type=StepType.LLM_AGENT,
                provider="capture",
                prompt_text="Translate: {text}",
                chunk_policy="auto",
                retry_max=1,
                backoff_base=0.0,
                redact_pii=True,
            )
        )
        s.add(
            WorkflowStep(
                workflow_id=wf9_id,
                order_index=1,
                name="LQA",
                step_type=StepType.LQA,
                provider="lqacap",
                chunk_policy="auto",
                retry_max=1,
                backoff_base=0.0,
                redact_pii=True,
            )
        )
        s.commit()
    with Session(engine) as s:
        run9 = orchestrator.run_workflow(s, wf9_id, doc9, "fr-FR")
    assert run9["status"] == "completed", run9
    lqa_seen = "\n".join(lqa_cap.seen)
    assert lqa_seen, "LQA provider was never called"
    assert PII_EMAIL not in lqa_seen, ("raw email reached the LQA engine", lqa_seen)
    assert PII_SSN not in lqa_seen, ("raw SSN reached the LQA engine", lqa_seen)
    assert "«PII_" in lqa_seen, ("LQA engine should have seen placeholders", lqa_seen)
    print("  (9) F3 LQA honors redact_pii: LQA engine saw ONLY placeholders (no raw email/SSN)")

    # =====================================================================
    # (10) F5: TM-origin segment keeps TM provenance; MT segment stays MT/0.9
    # =====================================================================
    tmpid2 = new_project(client, "TM prov")
    tmdoc2 = upload(client, tmpid2, ["Provenance known.", "Provenance unknown."])
    with Session(engine) as s:
        tm.write_back(s, tmpid2, "Provenance known.", "TM PROVENANCE", "en-US", "fr-FR")
    with Session(engine) as s:
        wf10 = Workflow(name="TM prov wf", scope="global")
        s.add(wf10)
        s.commit()
        s.refresh(wf10)
        wf10_id = wf10.id
        add_step(s, wf10_id, 0, "Translator", prompt_text="Translate: {text}")
    with Session(engine) as s:
        run10 = orchestrator.run_workflow(s, wf10_id, tmdoc2, "fr-FR")
    assert run10["status"] == "completed", run10
    segs10 = segments_for(tmdoc2)
    assert segs10[0].target_text == "TM PROVENANCE", segs10[0].target_text
    assert segs10[0].tm_match_pct == 100, (
        "TM-origin segment must carry TM provenance (tm_match_pct=100)",
        segs10[0].tm_match_pct,
    )
    # MT segment: MT status, confidence 0.9, no TM match pct
    assert segs10[1].status == "mt" and segs10[1].mt_confidence == 0.9, (
        segs10[1].status,
        segs10[1].mt_confidence,
    )
    print("  (10) F5 TM provenance: TM-origin seg -> tm_match_pct=100; MT seg -> mt/0.9")

    # =====================================================================
    # (11) HARDENING (review finding 3): FULL_TEXT run with a dead engine:<id>
    #   step must NOT report a false "completed" with a blank target. Before the
    #   fix, the FULL_TEXT loop's DeadEngineError catch set the StepRun FAILED +
    #   err, but run.status was set COMPLETED unconditionally after the loop --
    #   the API/CLI got {"ok": true, "status": "completed"} with empty text.
    #   Mirrors test (7b) (the segmented-path equivalent) but for FULL_TEXT mode.
    # =====================================================================
    from app.services import engines as _eng11

    prof_dead11 = _eng11.create_profile("WFDeadFullText", "deepl", api_key="k1")
    _eng11.update_profile(prof_dead11.id, enabled=False)
    dead_ref11 = f"engine:{prof_dead11.id}"
    pid11 = new_project(client, "FT dead engine")
    doc11 = upload(client, pid11, ["A sentence to translate."], mode="full_text")
    with Session(engine) as s:
        wf11 = Workflow(name="FT dead engine wf", scope="global")
        s.add(wf11)
        s.commit()
        s.refresh(wf11)
        wf11_id = wf11.id
        add_step(s, wf11_id, 0, "Translator", provider=dead_ref11, prompt_text="Translate: {text}")
    pre11 = segments_for(doc11)
    assert all(not (sg.target_text or "").strip() for sg in pre11), "fresh targets should be empty"
    with Session(engine) as s:
        run11 = orchestrator.run_workflow(s, wf11_id, doc11, "fr-FR")
    assert run11["status"] != "completed", (
        "FULL_TEXT dead-engine run must not report completed",
        run11,
    )
    assert run11["status"] == "failed", run11
    assert dead_ref11 in (run11.get("error") or ""), (
        "FULL_TEXT run error must name the dead ref",
        run11.get("error"),
    )
    with Session(engine) as s:
        wr11 = s.get(WorkflowRun, run11["run_id"])
        assert wr11.status == RunStatus.FAILED, ("WorkflowRun must be FAILED", wr11.status)
        srs11 = s.exec(
            select(StepRun).where(StepRun.workflow_run_id == run11["run_id"]).order_by(StepRun.id)
        ).all()
    assert srs11 and srs11[-1].status == RunStatus.FAILED, (
        "failed step's StepRun must be FAILED",
        [sr.status for sr in srs11],
    )
    post11 = segments_for(doc11)
    for sg in post11:
        assert not (sg.target_text or "").strip(), (
            "FULL_TEXT dead-engine run must leave NO mock target written",
            sg.order_index,
            sg.target_text,
        )
    print(
        "  (11) F3 FULL_TEXT dead engine:<id> step: run FAILED (not a false "
        "'completed'), error names the ref, NO blank target written"
    )

    # A healthy FULL_TEXT run must still report COMPLETED (regression, unaffected
    # by the any_step_failed gate added for (11)).
    doc11b = upload(client, pid11, ["Another healthy sentence."], mode="full_text")
    with Session(engine) as s:
        wf11b = Workflow(name="FT healthy wf", scope="global")
        s.add(wf11b)
        s.commit()
        s.refresh(wf11b)
        wf11b_id = wf11b.id
        add_step(s, wf11b_id, 0, "Translator", prompt_text="Translate: {text}")
    with Session(engine) as s:
        run11b = orchestrator.run_workflow(s, wf11b_id, doc11b, "fr-FR")
    assert run11b["status"] == "completed", run11b
    post11b = segments_for(doc11b)
    assert any((sg.target_text or "").strip() for sg in post11b), (
        "healthy FULL_TEXT run must still write its target"
    )
    print("  (11b) healthy FULL_TEXT run still reports COMPLETED + writes its target")

    # =====================================================================
    # (12) HARDENING (review finding 4): DeadEngineError messages carry step/pair
    #   context (f"step '{name}' ({src}->{locale}): {exc}"), not the bare provider
    #   message, so a multi-step run's error names WHICH step/pair died.
    # =====================================================================
    assert run11["error"].startswith("step 'Translator' (en-US->fr-FR): "), (
        "FULL_TEXT dead-engine error must carry step/pair context",
        run11["error"],
    )
    assert run7b.get("error", "").startswith("step 'Translator' (en-US->fr-FR): "), (
        "segmented dead-engine error must carry step/pair context",
        run7b.get("error"),
    )
    print("  (12) F4 error messages carry step/pair context on both FULL_TEXT and segmented paths")

    print("WORKFLOW SEGMENTS OK")


if __name__ == "__main__":
    main()
