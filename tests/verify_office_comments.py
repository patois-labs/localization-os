"""#690 + decision 8 (2026-10-11): comments in Office files, and SVG <title>/<desc>.

  A  OPTION: app/services/comment_option. A project created through project_create starts
     with comments OFF in every format (Word included); a project with no row (created
     before the option) keeps the old split (Word on, the rest off); switching it applies
     to every format; deleting the project removes its row; the switch is audited.
  B  FILTERS, comments off: the unit lists are exactly the pre-#690 ones (no comment
     text, same indices), Word drops only its comments-part units.
  C  FILTERS, comments on: Word comments; Excel notes (the bold "Author:" run kept out)
     and threaded comments (Excel's placeholder note skipped); PowerPoint classic and
     modern comments with replies; OpenDocument annotations (ODT, and a comment on a
     NUMBER cell in ODS); Visio comment entries. Authors, initials and dates never units.
  D  EXPORT: each comment comes back translated; authors and dates byte-for-byte; the
     body still translates; a document imported WITHOUT comments exports them unchanged.
     Excel: notes and threaded comments are patched into the package-patching export (B-xlsx);
  E  INGEST: the option reaches the filter at upload (off, then on) and at re-import.
  F  SURFACES: web toggle (manager+, 403 for a translator), /api route + detail fields,
     CLI `project set-comments` + `project show`, MCP configure_project action=comments,
     help + llms.txt + CLI docs.
  G  SVG: <title> and <desc> are a11y:svg-title / a11y:svg-desc units appended after the
     <text> units; nested markup, <metadata> and <style> are left alone; export writes them
     back escaped; an SVG imported before keeps its indices.

Hermetic + offline. Final line: "verify_office_comments OK".
"""

import os
import sys

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_comments_")

import importlib.util
import io
import json
import zipfile
from contextlib import redirect_stdout
from types import SimpleNamespace

from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app.cli.__main__ import main as cli_main
from app.db import engine, init_db
from app.filters import get_filter
from app.filters.office_comments import COMMENT_BASE
from app.main import app
from app.models import AuditLog, Document, Project, Role, Segment, Setting, TransMode
from app.services import comment_option, content, ingest, security, settings
from app.services.export import build_translated_file
from tests._csrf import csrf_for
from tests.fixtures.formats.pptx.make_pptx import make_pptx

PW = "Str0ng-Passw0rd!x"
_SEQ = [0]


def _seq():
    _SEQ[0] += 1
    return _SEQ[0]


def _members(data: bytes) -> dict:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return {i.filename: zf.read(i) for i in zf.infolist()}


def _texts(fmt, data, **kw):
    return [(u.p_index, u.text, u.style) for u in get_filter(fmt).parse(data, **kw).units]


def _comment_units(fmt, data, **kw):
    return [t for t in _texts(fmt, data, **kw) if t[0] >= COMMENT_BASE]


# ------------------------------------------------------------------------------ fixtures
def docx_bytes() -> bytes:
    from docx import Document as DocxDocument

    d = DocxDocument()
    p = d.add_paragraph("The body paragraph.")
    d.add_comment(p.runs, text="Please check this sentence.", author="Ann Reviewer", initials="AR")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


_XL_TC_NS = "http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments"
_TC_REL = "http://schemas.microsoft.com/office/2017/10/relationships/threadedComment"
_TC_CT = "application/vnd.ms-excel.threadedcomments+xml"
THREAD_ID = "{11111111-2222-3333-4444-555555555555}"
REPLY_ID = "{66666666-2222-3333-4444-555555555555}"
PLACEHOLDER = (
    "[Threaded comment]\n\nYour version of Excel allows you to read this threaded comment."
)


def xlsx_bytes(threaded: bool = False) -> bytes:
    """A workbook with one classic note (bold "Ann:" label run + body run) on B2, and,
    with `threaded`, a threaded comment + reply on C3 whose legacy placeholder note Excel
    writes is also present."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Quarterly budget"
    ws["B2"] = 42
    ws["B2"].comment = Comment("Ann:\nCheck the total.", "Ann")
    if threaded:
        ws["C3"] = "Owner"
        ws["C3"].comment = Comment(PLACEHOLDER, "tc=" + THREAD_ID)
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    parts = _members(data)
    # Rewrite the B2 note the way Excel writes it: a bold "Ann:" run, then the body run.
    cname = next(n for n in parts if n.startswith("xl/comments") and n.endswith(".xml"))
    xml = parts[cname].decode("utf-8")
    start = xml.index('<comment ref="B2"')
    end = xml.index("</comment>", start) + len("</comment>")
    note = (
        '<comment ref="B2" authorId="0"><text>'
        '<r><rPr><b/><sz val="9"/></rPr><t>Ann:</t></r>'
        '<r><rPr><sz val="9"/></rPr><t xml:space="preserve">\nCheck the total.</t></r>'
        "</text></comment>"
    )
    parts[cname] = (xml[:start] + note + xml[end:]).encode("utf-8")
    if threaded:
        sheet = "xl/worksheets/sheet1.xml"
        rels = "xl/worksheets/_rels/sheet1.xml.rels"
        parts["xl/threadedComments/threadedComment1.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<ThreadedComments xmlns="{_XL_TC_NS}">'
            f'<threadedComment ref="C3" dT="2026-10-01T10:00:00.00" '
            f'personId="{{AAAA0000-0000-0000-0000-000000000000}}" id="{THREAD_ID}">'
            "<text>Who owns this line?</text></threadedComment>"
            f'<threadedComment ref="C3" dT="2026-10-02T10:00:00.00" '
            f'personId="{{BBBB0000-0000-0000-0000-000000000000}}" id="{REPLY_ID}" '
            f'parentId="{THREAD_ID}"><text>Finance does.</text></threadedComment>'
            "</ThreadedComments>"
        ).encode()
        r = parts[rels].decode("utf-8")
        parts[rels] = r.replace(
            "</Relationships>",
            f'<Relationship Id="rIdTc1" Type="{_TC_REL}" '
            'Target="../threadedComments/threadedComment1.xml"/></Relationships>',
        ).encode("utf-8")
        ct = parts["[Content_Types].xml"].decode("utf-8")
        parts["[Content_Types].xml"] = ct.replace(
            "</Types>",
            '<Override PartName="/xl/threadedComments/threadedComment1.xml" '
            f'ContentType="{_TC_CT}"/></Types>',
        ).encode("utf-8")
        assert sheet in parts
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, raw in parts.items():
            zf.writestr(name, raw)
    return out.getvalue()


_P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_P188 = "http://schemas.microsoft.com/office/powerpoint/2018/8/main"


def pptx_bytes() -> bytes:
    classic = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<p:cmLst xmlns:p="{_P_NS}"><p:cm authorId="0" dt="2026-10-01T10:00:00.000" idx="1">'
        '<p:pos x="10" y="10"/><p:text>Shorten this title.</p:text></p:cm></p:cmLst>'
    )
    modern = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<p188:cmLst xmlns:a="{_A_NS}" xmlns:r="http://schemas.openxmlformats.org/'
        f'officeDocument/2006/relationships" xmlns:p188="{_P188}">'
        '<p188:cm id="{C0000000-0000-0000-0000-000000000001}" '
        'authorId="{A0000000-0000-0000-0000-000000000001}" created="2026-10-01T10:00:00.000">'
        "<p188:txBody><a:bodyPr/><a:p><a:r><a:t>Is this chart current?</a:t></a:r></a:p>"
        "</p188:txBody><p188:replyLst>"
        '<p188:reply id="{C0000000-0000-0000-0000-000000000002}" '
        'authorId="{A0000000-0000-0000-0000-000000000002}" created="2026-10-02T10:00:00.000">'
        "<p188:txBody><a:bodyPr/><a:p><a:r><a:t>Yes, updated today.</a:t></a:r></a:p>"
        "</p188:txBody></p188:reply></p188:replyLst></p188:cm></p188:cmLst>"
    )
    rt = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
    rt_modern = "http://schemas.microsoft.com/office/2018/10/relationships/comments"
    return make_pptx(
        extra_parts={
            "ppt/comments/comment1.xml": classic,
            "ppt/comments/modernComment_100_1.xml": modern,
        },
        extra_slide_rels=(
            ("rIdC1", rt, "../comments/comment1.xml", ""),
            ("rIdC2", rt_modern, "../comments/modernComment_100_1.xml", ""),
        ),
    )


_ODF_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" '
    'xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0" office:version="1.3"'
)


def odf_bytes(body: str, mimetype: str) -> bytes:
    content = (
        f'<?xml version="1.0" encoding="UTF-8"?>\n<office:document-content {_ODF_NS}>'
        f"<office:body>{body}</office:body></office:document-content>"
    )
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
        f'manifest:version="1.3"><manifest:file-entry manifest:full-path="/" '
        f'manifest:media-type="{mimetype}"/><manifest:file-entry '
        'manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>'
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), mimetype)
        z.writestr("content.xml", content, zipfile.ZIP_DEFLATED)
        z.writestr("META-INF/manifest.xml", manifest, zipfile.ZIP_DEFLATED)
    return buf.getvalue()


ANNOTATION = (
    "<office:annotation><dc:creator>Ann Reviewer</dc:creator>"
    "<dc:date>2026-10-01T10:00:00</dc:date><meta:creator-initials>AR</meta:creator-initials>"
    "<text:p>Rephrase this, please.</text:p></office:annotation>"
)
ODT_BODY = (
    "<office:text><text:p>First paragraph" + ANNOTATION + " of the text.</text:p>"
    "<text:p>Second paragraph.</text:p></office:text>"
)
ODS_BODY = (
    '<office:spreadsheet><table:table table:name="S"><table:table-row>'
    '<table:table-cell office:value-type="string"><text:p>Revenue</text:p></table:table-cell>'
    '<table:table-cell office:value-type="float" office:value="12">'
    "<office:annotation><dc:creator>Bo</dc:creator><text:p>Rounded figure.</text:p>"
    "</office:annotation><text:p>12</text:p></table:table-cell>"
    "</table:table-row></table:table></office:spreadsheet>"
)

_VNS = "http://schemas.microsoft.com/office/visio/2012/main"


def vsdx_bytes(with_comments: bool = True) -> bytes:
    page_rel = "http://schemas.microsoft.com/visio/2010/relationships/page"
    parts = {
        "[Content_Types].xml": (
            "<?xml version='1.0' encoding='utf-8'?>"
            "<Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'>"
            "<Default Extension='rels' ContentType='application/vnd.openxmlformats-package."
            "relationships+xml'/><Default Extension='xml' ContentType='application/xml'/>"
            "<Override PartName='/visio/document.xml' "
            "ContentType='application/vnd.ms-visio.drawing.main+xml'/></Types>"
        ),
        "_rels/.rels": (
            "<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'/>"
        ),
        "visio/document.xml": f"<VisioDocument xmlns='{_VNS}'/>",
        "visio/pages/pages.xml": (
            f"<Pages xmlns='{_VNS}' xmlns:r='http://schemas.openxmlformats.org/officeDocument/"
            "2006/relationships'><Page ID='0' NameU='Page-1' Name='Process'>"
            "<PageSheet/><Rel r:id='rId1'/></Page></Pages>"
        ),
        "visio/pages/_rels/pages.xml.rels": (
            "<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>"
            f"<Relationship Id='rId1' Type='{page_rel}' Target='page1.xml'/></Relationships>"
        ),
        "visio/pages/page1.xml": (
            f"<PageContents xmlns='{_VNS}'><Shapes><Shape ID='1' Type='Shape'>"
            "<Text>Start here</Text></Shape></Shapes></PageContents>"
        ),
    }
    if with_comments:
        parts["visio/_rels/document.xml.rels"] = (
            "<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>"
            "<Relationship Id='rId9' Type='http://schemas.microsoft.com/visio/2010/"
            "relationships/comments' Target='comments.xml'/></Relationships>"
        )
        parts["visio/comments.xml"] = (
            "<?xml version='1.0' encoding='utf-8' ?>"
            f"<Comments xmlns='{_VNS}'><AuthorList>"
            "<AuthorEntry Name='Ann Reviewer' Initials='AR' ResolutionID='' ID='0'/>"
            "</AuthorList><CommentList>"
            "<CommentEntry AuthorID='0' PageID='0' ShapeID='1' Date='2026-10-01T10:00:00' "
            "CommentID='0'>Move this shape left.</CommentEntry>"
            "<CommentEntry AuthorID='0' PageID='0' Date='2026-10-01T11:00:00' "
            "CommentID='1'>  </CommentEntry>"
            "</CommentList></Comments>"
        )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, raw in parts.items():
            zf.writestr(name, raw)
    return out.getvalue()


SVG = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10">'
    "<title>Sales &amp; returns</title>"
    "<desc> Bars per quarter </desc>"
    "<metadata><title>Not announced</title></metadata>"
    "<g><title>Q1 <tspan>bar</tspan></title><desc>---</desc></g>"
    "<text x='1' y='1'>Q1</text><text x='1' y='5'>Q2 total</text>"
    "</svg>"
)


# ------------------------------------------------------------------------------ A
def _mk_project(**kw):
    with Session(engine) as s:
        p = content.project_create(s, "Comments %d" % _seq(), target_locales="de-DE", **kw)  # noqa: UP031
        return SimpleNamespace(id=p.id, customer_id=None)


def _legacy_project():
    """A project created without project_create (as before the option existed)."""
    with Session(engine) as s:
        from app.models import TargetLocale

        p = Project(name="Legacy %d" % _seq())  # noqa: UP031
        s.add(p)
        s.commit()
        s.refresh(p)
        s.add(TargetLocale(project_id=p.id, locale="de-DE"))
        s.commit()
        return SimpleNamespace(id=p.id, customer_id=None)


def test_A_option():
    new = _mk_project()
    assert comment_option.stored(new) == "0", "a new project must start with comments off"
    for fmt in comment_option.COMMENT_FORMATS:
        assert comment_option.enabled_for(new, fmt) is False, fmt
    # gate: the paths that create a project without project_create (a department inbox,
    # a connector's sync project) stamp it too, so "no row" means only "pre-option".
    from app.models import Connection
    from app.services import routing
    from app.services.connectors import store_part2

    with Session(engine) as s:
        inbox = routing.ensure_department_inbox(s, "Comments Dept %d" % _seq(), "a@example.com")  # noqa: UP031
        assert comment_option.stored(inbox.id) == "0", "a new inbox project starts off"
        conn = Connection(kind="jira", name="Comments conn %d" % _seq())  # noqa: UP031
        s.add(conn)
        s.commit()
        s.refresh(conn)
        sync = store_part2._ensure_project(s, conn)
        assert comment_option.stored(sync.id) == "0", "a new connector project starts off"
        assert comment_option.enabled_for(sync, "docx") is False
    legacy = _legacy_project()
    assert comment_option.stored(legacy) is None
    assert comment_option.enabled_for(legacy, "docx") is True, "pre-option Word stays on"
    for fmt in ("xlsx", "pptx", "odt", "ods", "odp", "odg", "vsdx"):
        assert comment_option.enabled_for(legacy, fmt) is False, fmt
    assert comment_option.status(legacy) == {
        "translate_comments": False,
        "translate_comments_word": True,
        "translate_comments_source": "legacy",
    }
    st = comment_option.set_for(legacy, True, actor="tester@example.com")
    assert st["translate_comments"] and st["translate_comments_word"], st
    assert all(comment_option.enabled_for(legacy, f) for f in comment_option.COMMENT_FORMATS)
    comment_option.set_for(legacy, False, actor="tester@example.com")
    assert comment_option.enabled_for(legacy, "docx") is False, "an explicit off covers Word"
    for value, expect in (("on", True), ("0", False), (True, True), ("maybe", None), (2, None)):
        assert comment_option.parse_switch(value) is expect, value
    with Session(engine) as s:
        rows = s.exec(select(AuditLog).where(AuditLog.target == f"project {legacy.id}")).all()
        acts = {r.action for r in rows}
    assert {"project.translate_comments_on", "project.translate_comments_off"} <= acts, acts
    # delete removes the row
    with Session(engine) as s:
        content.delete_project(s, new.id, actor="tester")
        s.commit()
    with Session(engine) as s:
        assert s.get(Setting, f"translate_comments:{new.id}") is None
    settings.invalidate_cache()
    assert comment_option.stored(new.id) is None
    print("A ok: new projects off, legacy Word on, explicit switch, audit, delete cleanup")


def test_A2_stamp_in_caller_transaction():
    """Release 2026-10-12 sweep: the stamp is written through the caller's session.

    A caller that runs project_create inside its own outer transaction (a savepoint
    session on one connection, as verify_key_paging_scale and
    verify_glossary_paging_scale do) hit "database is locked" on the SECOND create: the
    stamp opened another write session while the outer connection held SQLite's writer
    lock. The row now joins the caller's session, so it is visible there at once and no
    second writer exists. Nothing is asserted after the outer rollback: with pysqlite a
    released outermost savepoint may already have committed, so that state is driver
    detail, not product behaviour."""
    with engine.connect() as connection:
        outer = connection.begin()
        try:
            with Session(connection, join_transaction_mode="create_savepoint") as s:
                for _ in range(2):
                    p = content.project_create(s, "Comments tx %d" % _seq(), target_locales="de-DE")  # noqa: UP031
                    row = s.get(Setting, f"translate_comments:{p.id}")
                    assert row is not None and row.value == "0", "stamp joins the caller's session"
        finally:
            outer.rollback()
    settings.invalidate_cache()
    # The returned Project stays readable after the caller's session closes.
    with Session(engine) as s:
        made = content.project_create(s, "Comments detached %d" % _seq(), target_locales="de-DE")  # noqa: UP031
    assert made.name.startswith("Comments detached"), made.name
    assert comment_option.stored(made.id) == "0"
    print("A2 ok: two creates inside one outer transaction, each stamped in the caller's session")


# ------------------------------------------------------------------------------ B + C
def test_BC_filters():
    # DOCX: comments in by default (legacy), out with comments=False, indices kept
    d = docx_bytes()
    on = _texts("docx", d)
    off = _texts("docx", d, comments=False)
    assert any("Please check this sentence." in t for _i, t, _s in on), on
    assert not any("Please check" in t for _i, t, _s in off), off
    assert off == [u for u in on if "Please check" not in u[1]], (on, off)
    assert _texts("docx", d, comments=True) == on

    # XLSX
    x = xlsx_bytes(threaded=True)
    assert _comment_units("xlsx", x) == [], "comments off by default"
    cells = _texts("xlsx", x)
    with_c = _texts("xlsx", x, comments=True)
    assert with_c[: len(cells)] == cells, "cell units unchanged"
    got = [(t, s) for _i, t, s in with_c[len(cells) :]]
    assert got == [
        ("Who owns this line?", "comment"),
        ("Finance does.", "comment"),
        ("Check the total.", "comment"),
    ], got
    assert not any("Ann:" in t or "Threaded comment" in t for _i, t, _s in with_c), with_c

    # PPTX
    p = pptx_bytes()
    base = _texts("pptx", p)
    assert _comment_units("pptx", p) == []
    pc = _texts("pptx", p, comments=True)
    assert pc[: len(base)] == base
    assert [t for _i, t, _s in pc[len(base) :]] == [
        "Shorten this title.",
        "Is this chart current?",
        "Yes, updated today.",
    ], pc[len(base) :]
    assert all(s == "slide 1 comment" for _i, _t, s in pc[len(base) :])

    # ODT / ODS
    odt = odf_bytes(ODT_BODY, "application/vnd.oasis.opendocument.text")
    ob = _texts("odt", odt)
    assert not any("Rephrase" in t or "Ann Reviewer" in t for _i, t, _s in ob), ob
    oc = _texts("odt", odt, comments=True)
    assert oc[: len(ob)] == ob
    assert [(t, s) for _i, t, s in oc[len(ob) :]] == [("Rephrase this, please.", "comment")]
    ods = odf_bytes(ODS_BODY, "application/vnd.oasis.opendocument.spreadsheet")
    assert [t for _i, t, _s in _comment_units("ods", ods, comments=True)] == ["Rounded figure."]
    assert _comment_units("ods", ods) == []

    # VSDX
    v = vsdx_bytes()
    vb = _texts("vsdx", v)
    assert _comment_units("vsdx", v) == []
    vc = _texts("vsdx", v, comments=True)
    assert vc[: len(vb)] == vb
    assert [(t, s) for _i, t, s in vc[len(vb) :]] == [("Move this shape left.", "comment")], vc
    assert _texts("vsdx", vsdx_bytes(with_comments=False), comments=True) == vb
    print("B+C ok: comments off = pre-#690 units; on = every Office family, no authors/dates")


# ------------------------------------------------------------------------------ D
def _de(units):
    return {str(i): "DE " + t for i, t, _s in units}


def test_D_export():
    # DOCX
    d = docx_bytes()

    def _word(units):  # the comment unit carries the annotationRef token: swap the words
        return {
            str(i): t.replace("Please check this sentence.", "Bitte Satz pruefen.")
            for i, t, _s in units
        }

    out = get_filter("docx").build(d, _word(_texts("docx", d)))
    com = _members(out)["word/comments.xml"].decode("utf-8")
    assert "Bitte Satz pruefen." in com and 'w:author="Ann Reviewer"' in com, com
    assert "w:annotationRef" in com and 'w:initials="AR"' in com
    # imported without comments: the comment stays as it is
    out = get_filter("docx").build(d, _word(_texts("docx", d, comments=False)))
    assert "Bitte" not in _members(out)["word/comments.xml"].decode("utf-8")

    # XLSX: the package-patching export (B-xlsx) keeps every comment part; the hook
    # translates notes and threaded comments in it
    x = xlsx_bytes(threaded=True)
    units = _texts("xlsx", x, comments=True)
    built = get_filter("xlsx").build_localized(x, _de(units), "de-DE")
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(built))["Budget"]
    assert ws["A1"].value == "DE Quarterly budget"
    assert ws["B2"].comment.text == "Ann:\nDE Check the total.", repr(ws["B2"].comment.text)
    assert ws["C3"].comment.text.startswith("[Threaded comment]"), "placeholder untouched"
    from app.filters import office_comments

    patched = office_comments.xlsx_apply(x, x, _de(units))
    tc = _members(patched)["xl/threadedComments/threadedComment1.xml"].decode("utf-8")
    assert "<text>DE Who owns this line?</text>" in tc and "<text>DE Finance does.</text>" in tc
    assert 'dT="2026-10-01T10:00:00.00"' in tc and THREAD_ID in tc, tc
    m0, m1 = _members(x), _members(patched)
    cname = next(n for n in m0 if n.startswith("xl/comments") and n.endswith(".xml"))
    note = m1[cname].decode("utf-8")
    assert "<t>Ann:</t>" in note and "\nDE Check the total.</t>" in note, note
    unchanged = [n for n in m0 if n not in (cname, "xl/threadedComments/threadedComment1.xml")]
    assert all(m0[n] == m1[n] for n in unchanged), "untouched parts must keep their bytes"
    no_c = get_filter("xlsx").build(x, _de(_texts("xlsx", x)))
    ws = load_workbook(io.BytesIO(no_c))["Budget"]
    assert ws["B2"].comment.text == "Ann:\nCheck the total."
    # gate (after B-xlsx #685): the REAL export patches the original package, so the
    # threaded comment and its reply translate end to end, and with the option off every
    # comment part (notes, threaded comments, persons, VML) comes back byte-identical.
    comment_parts = [
        n
        for n in m0
        if n.startswith(("xl/comments", "xl/threadedComments/", "xl/persons/", "xl/drawings/vml"))
    ]
    assert "xl/threadedComments/threadedComment1.xml" in comment_parts, comment_parts
    m2 = _members(built)
    tc = m2["xl/threadedComments/threadedComment1.xml"].decode("utf-8")
    assert "<text>DE Who owns this line?</text>" in tc and "<text>DE Finance does.</text>" in tc
    assert 'dT="2026-10-01T10:00:00.00"' in tc and THREAD_ID in tc and REPLY_ID in tc, tc
    assert "\nDE Check the total.</t>" in m2[cname].decode("utf-8")
    m3 = _members(no_c)
    for n in comment_parts:
        assert m3.get(n) == m0[n], f"option off: {n} must keep its bytes"
    m4 = _members(get_filter("xlsx").build_localized(x, _de(_texts("xlsx", x)), "de-DE"))
    for n in comment_parts:
        assert m4.get(n) == m0[n], f"option off (localized): {n} must keep its bytes"

    # PPTX
    p = pptx_bytes()
    out = _members(
        get_filter("pptx").build_localized(p, _de(_texts("pptx", p, comments=True)), "de-DE")
    )
    cl = out["ppt/comments/comment1.xml"].decode("utf-8")
    assert "<p:text>DE Shorten this title.</p:text>" in cl and 'dt="2026-10-01T10:00:00.000"' in cl
    mo = out["ppt/comments/modernComment_100_1.xml"].decode("utf-8")
    assert "DE Is this chart current?" in mo and "DE Yes, updated today." in mo
    assert 'authorId="{A0000000-0000-0000-0000-000000000002}"' in mo
    plain = _members(get_filter("pptx").build(p, _de(_texts("pptx", p))))
    assert plain["ppt/comments/comment1.xml"] == _members(p)["ppt/comments/comment1.xml"]

    # ODT
    odt = odf_bytes(ODT_BODY, "application/vnd.oasis.opendocument.text")
    out = _members(get_filter("odt").build(odt, _de(_texts("odt", odt, comments=True))))
    c = out["content.xml"].decode("utf-8")
    assert "DE Rephrase this, please." in c and "<dc:creator>Ann Reviewer</dc:creator>" in c, c
    assert "<dc:date>2026-10-01T10:00:00</dc:date>" in c and "DE Second paragraph." in c
    assert c.index("DE First paragraph") < c.index("DE Rephrase") < c.index("of the text")
    rebuilt = get_filter("odt").build(odt, _de(_texts("odt", odt, comments=True)))
    again = [t for _i, t, _s in _texts("odt", rebuilt, comments=True)]
    assert len(again) == 3 and again[2] == "DE Rephrase this, please.", again

    # VSDX
    v = vsdx_bytes()
    out = _members(get_filter("vsdx").build(v, _de(_texts("vsdx", v, comments=True))))
    vc = out["visio/comments.xml"].decode("utf-8")
    assert "DE Move this shape left." in vc and 'Name="Ann Reviewer"' in vc.replace("'", '"')
    assert "2026-10-01T10:00:00" in vc
    plain = _members(get_filter("vsdx").build(v, _de(_texts("vsdx", v))))
    assert plain["visio/comments.xml"] == _members(v)["visio/comments.xml"]
    print("D ok: every family exports translated comments with authors/dates kept")


# ------------------------------------------------------------------------------ E
def _ingest(project, name, data):
    with Session(engine) as s:
        proj = s.get(Project, project.id)
        doc = ingest.ingest_document(s, proj, name, data, TransMode.SEGMENTED)
        return doc.id


def _sources(doc_id) -> str:
    """Every source text of the document, joined (a Word comment unit carries tokens)."""
    with Session(engine) as s:
        segs = s.exec(select(Segment).where(Segment.document_id == doc_id)).all()
        return " | ".join(seg.source_text for seg in segs)


def test_E_ingest():
    proj = _mk_project()
    x = xlsx_bytes()
    off_id = _ingest(proj, "off.xlsx", x)
    assert "Check the total." not in _sources(off_id)
    d_off = _ingest(proj, "off.docx", docx_bytes())
    assert "Please check this sentence." not in _sources(d_off), "new project: Word off too"
    comment_option.set_for(proj, True)
    on_id = _ingest(proj, "on.xlsx", x)
    assert "Check the total." in _sources(on_id)
    with Session(engine) as s:
        doc = s.get(Document, on_id)
        assert doc.skeleton.get("comments") is True
        for seg in s.exec(select(Segment).where(Segment.document_id == on_id)).all():
            seg.target_text = "DE " + seg.source_text
            seg.status = "translated"
            s.add(seg)
        s.commit()
        doc = s.get(Document, on_id)
        built = build_translated_file(s, doc, "de-DE")
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(built))["Budget"]
    assert ws["B2"].comment.text == "Ann:\nDE Check the total.", ws["B2"].comment.text
    # re-import of the comments-off document follows the option now
    assert ingest.reingest_document(off_id) is True
    assert "Check the total." in _sources(off_id), "re-import follows the option now"
    legacy = _legacy_project()
    d_legacy = _ingest(legacy, "legacy.docx", docx_bytes())
    assert "Please check this sentence." in _sources(d_legacy), "pre-option Word stays on"
    x_legacy = _ingest(legacy, "legacy.xlsx", x)
    assert "Check the total." not in _sources(x_legacy)
    print("E ok: the option reaches the filter at upload; export writes the comment")


# ------------------------------------------------------------------------------ F
def _read(rel: str) -> str:
    with open(os.path.join(_ROOT, rel), encoding="utf-8") as fh:
        return fh.read()


def _json_out(text):
    text = text.strip()
    try:
        return json.loads(text)
    except ValueError:
        return json.loads(text[text.index("{") :])


def _login(email, role):
    security.create_user(email, email, role, PW)
    c = TestClient(app)
    csrf_for(c)
    r = c.post(
        "/auth/login", data={"email": email, "password": PW, "next": "/"}, follow_redirects=False
    )
    assert r.status_code == 303, (email, r.status_code)
    csrf_for(c)
    return c


def test_F_surfaces():
    proj = _mk_project()
    pid = proj.id
    c = TestClient(app)
    tok = csrf_for(c)
    r = c.get(f"/projects/{pid}")
    assert r.status_code == 200 and "Translate comments in Office files" in r.text, r.status_code
    assert "kept in the source language" in r.text
    r = c.post(
        f"/projects/{pid}/translate-comments",
        data={"translate_comments": "1", "csrf_token": tok},
        follow_redirects=False,
    )
    assert r.status_code == 303, (r.status_code, r.text[:200])
    assert comment_option.stored(pid) == "1"
    r = c.get(f"/projects/{pid}")
    assert "Stop translating comments" in r.text
    # /api
    from app.services import apikeys

    admin = security.create_user("cadm-%d@example.com" % _seq(), "a", Role.ADMIN, PW)  # noqa: UP031
    _row, raw = apikeys.create_key(admin.id, "cm", scopes="docs")
    c.headers["X-Loc-Token"] = raw
    j = c.get(f"/api/projects/{pid}").json()
    assert j["translate_comments"] is True and j["translate_comments_source"] == "project", j
    r = c.post(f"/api/projects/{pid}/translate-comments", json={"translate_comments": False})
    assert r.status_code == 200 and r.json()["translate_comments"] is False, r.text
    r = c.post(f"/api/projects/{pid}/translate-comments", json={"translate_comments": "maybe"})
    assert r.status_code == 400, r.text
    r = c.post("/api/projects/999999/translate-comments", json={"translate_comments": True})
    assert r.status_code == 404, r.text
    del c.headers["X-Loc-Token"]
    # CLI
    buf = io.StringIO()
    with redirect_stdout(buf):
        rc = cli_main(["--json", "project", "set-comments", str(pid), "on"])
    assert rc == 0, buf.getvalue()
    assert _json_out(buf.getvalue())["translate_comments"] is True
    buf = io.StringIO()
    with redirect_stdout(buf):
        rc = cli_main(["--json", "project", "show", str(pid)])
    assert rc == 0 and _json_out(buf.getvalue())["translate_comments"] is True
    buf = io.StringIO()
    with redirect_stdout(buf):
        rc = cli_main(["project", "show", str(pid)])
    assert rc == 0 and "comments in Office files: translated" in buf.getvalue(), buf.getvalue()
    # MCP: configure_project action=comments, routed to the /api route
    spec = importlib.util.spec_from_file_location(
        "mcp_registry_c", os.path.join(_ROOT, "integrations", "mcp-server", "registry.py")
    )
    reg = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = reg  # dataclasses resolve through sys.modules
    spec.loader.exec_module(reg)
    tool = reg.tool_by_name("configure_project")
    assert ("POST", "/api/projects/{pid}/translate-comments") in tool.routes
    assert "comments" in tool.input_schema["properties"]["action"]["enum"]
    assert tool.input_schema["properties"]["translate_comments"]["type"] == "boolean"
    server_src = _read(os.path.join("integrations", "mcp-server", "server.py"))
    assert '"/api/projects/%d/translate-comments"' in server_src
    # docs
    for path, needle in (
        ("llms.txt", "/api/projects/{pid}/translate-comments"),
        ("docs/CLI.md", "project set-comments"),
        ("app/templates/help.html", "/api/projects/{pid}/translate-comments"),
    ):
        assert needle in _read(path), path
    # auth ON: a translator cannot switch it, a manager can
    settings.set_many({"security_auth_required": "1"})
    settings.invalidate_cache()
    try:
        tr = _login("ctr-%d@example.com" % _seq(), Role.TRANSLATOR)  # noqa: UP031
        r = tr.post(
            f"/projects/{pid}/translate-comments",
            data={"translate_comments": "0"},
            follow_redirects=False,
        )
        assert r.status_code == 403, (r.status_code, r.text[:200])
        assert comment_option.stored(pid) == "1"
        mg = _login("cmg-%d@example.com" % _seq(), Role.MANAGER)  # noqa: UP031
        r = mg.post(
            f"/projects/{pid}/translate-comments",
            data={"translate_comments": "0"},
            follow_redirects=False,
        )
        assert r.status_code == 303, (r.status_code, r.text[:200])
        assert comment_option.stored(pid) == "0"
    finally:
        settings.set_many({"security_auth_required": "0"})
        settings.invalidate_cache()
    print("F ok: web toggle (role gated), /api route + detail, CLI, MCP, docs")


# ------------------------------------------------------------------------------ G
def test_G_svg():
    from app.filters.image_filter import ImageFilter

    f = ImageFilter()
    data = SVG.encode("utf-8")
    units = [(u.p_index, u.text, u.style) for u in f._parse_svg(data).units]
    assert units == [
        (0, "Q1", ""),
        (1, "Q2 total", ""),
        (2, "Sales & returns", "a11y:svg-title"),
        (3, "Bars per quarter", "a11y:svg-desc"),
    ], units
    out = f._build_svg(data, {"0": "T1", "1": "Q2 gesamt", "2": "Umsatz & Retouren", "3": "Balken"})
    text = out.decode("utf-8")
    assert "<title>Umsatz &amp; Retouren</title>" in text, text
    assert "<desc> Balken </desc>" in text
    assert "<metadata><title>Not announced</title></metadata>" in text
    assert "<title>Q1 <tspan>bar</tspan></title>" in text
    # an SVG imported before (text units only): its mapping never touches title/desc
    old = f._build_svg(data, {"0": "T1", "1": "Q2 gesamt"}).decode("utf-8")
    assert "<title>Sales &amp; returns</title>" in old
    print("G ok: SVG title/desc units appended, nested/metadata skipped, escaped on export")


def main():
    init_db()
    test_A_option()
    test_A2_stamp_in_caller_transaction()
    test_BC_filters()
    test_D_export()
    test_E_ingest()
    test_F_surfaces()
    test_G_svg()
    print("verify_office_comments OK")


if __name__ == "__main__":
    main()
