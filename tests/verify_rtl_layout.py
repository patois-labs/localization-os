"""Right-to-left layout on export (B-rtl, the maintainer's decisions 17 and 20, 2026-10-11).

For a right-to-left target (the app's own list: a11y.target_direction):

  xlsx   every worksheet <sheetView> gets rightToLeft="1" (column A on the right).
  ods    every sheet's table style gets style:writing-mode="rl-tb"; a sheet with no
         automatic table style gets one. Text documents' tables are not touched.
  idml   the ParagraphStyleRange of every translated paragraph gets
         ParagraphDirection="RightToLeftDirection" and the World-Ready composer
         ("HL Composer Optyca", single-line stays single-line); an explicit LTR
         CharacterDirection on a translated range becomes DefaultDirection; untranslated
         paragraphs (a footnote nobody translated) are left alone.
  icml   the same through the shared IDML helper.
  svg    every translated <text> gets direction="rtl" and a mirrored text-anchor
         (start <-> end, middle kept, inherited and inline-style anchors honoured); a
         tspan with its own anchor is mirrored too; untranslated text is untouched.

A left-to-right target changes none of this. Hermetic + offline.
Final line: "RTL LAYOUT OK".
"""

# ruff: noqa: E501  (hand-written XML fixtures are long literals)

import io
import os
import sys
import zipfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_rtl_layout_")

from lxml import etree

from app.filters import get_filter

AR = "ar-SA"
HE = "he-IL"


def member(data: bytes, name: str) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read(name)


# ---------------------------------------------------------------------- xlsx
def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "One"
    ws["A1"] = "Name"
    ws["B1"] = "Price"
    ws["A2"] = "Coffee"
    ws["B2"] = 3
    ws2 = wb.create_sheet("Two")
    ws2["A1"] = "Second sheet"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


SML = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _sheet_views(data: bytes) -> list:
    views = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in sorted(z.namelist()):
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                root = etree.fromstring(z.read(name))
                views += list(root.iter(SML + "sheetView"))
    return views


def test_xlsx_rtl():
    filt = get_filter("xlsx")
    data = _xlsx()
    us = filt.parse(data).units
    mapping = {str(u.p_index): "نص " + u.text for u in us}
    out = filt.build_localized(data, mapping, AR)
    views = _sheet_views(out)
    assert len(views) == 2, len(views)
    for v in views:
        assert v.get("rightToLeft") in ("1", "true"), etree.tostring(v)
    for locale in ("de-DE", "ja-JP"):
        out = filt.build_localized(data, mapping, locale)
        assert all(v.get("rightToLeft") in (None, "0", "false") for v in _sheet_views(out)), locale
    # plain build (no target) never flips
    assert all(
        v.get("rightToLeft") in (None, "0", "false")
        for v in _sheet_views(filt.build(data, mapping))
    )
    print("  [xlsx] every sheet right to left for ar-SA only OK")


# ---------------------------------------------------------------------- ods
ODF_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" '
    'xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0" '
)
ODS_CONTENT = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    f'<office:document-content {ODF_NS}office:version="1.3">'
    "<office:automatic-styles>"
    '<style:style style:name="ta1" style:family="table" style:master-page-name="Default">'
    '<style:table-properties table:display="true" style:writing-mode="lr-tb"/></style:style>'
    '<style:style style:name="co1" style:family="table-column"/>'
    "</office:automatic-styles>"
    "<office:body><office:spreadsheet>"
    '<table:table table:name="One" table:style-name="ta1"><table:table-column table:style-name="co1"/>'
    '<table:table-row><table:table-cell office:value-type="string"><text:p>Hello sheet</text:p></table:table-cell>'
    '<table:table-cell office:value-type="float" office:value="3"><text:p>3</text:p></table:table-cell></table:table-row>'
    "</table:table>"
    '<table:table table:name="Two"><table:table-row><table:table-cell office:value-type="string">'
    "<text:p>Second sheet</text:p></table:table-cell></table:table-row></table:table>"
    '<table:table table:name="Three" table:style-name="ta1"><table:table-row><table:table-cell office:value-type="string">'
    "<text:p>Third sheet</text:p></table:table-cell></table:table-row></table:table>"
    "</office:spreadsheet></office:body></office:document-content>"
)
ODT_CONTENT = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    f'<office:document-content {ODF_NS}office:version="1.3">'
    "<office:automatic-styles>"
    '<style:style style:name="Table1" style:family="table"><style:table-properties style:width="6in"/></style:style>'
    "</office:automatic-styles>"
    '<office:body><office:text><text:p>Body text</text:p><table:table table:name="T" table:style-name="Table1">'
    "<table:table-row><table:table-cell><text:p>Cell text</text:p></table:table-cell></table:table-row>"
    "</table:table></office:text></office:body></office:document-content>"
)


def _odf(content: str, mimetype: str) -> bytes:
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">'
        f'<manifest:file-entry manifest:full-path="/" manifest:media-type="{mimetype}"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), mimetype)
        z.writestr("content.xml", content, zipfile.ZIP_DEFLATED)
        z.writestr("META-INF/manifest.xml", manifest, zipfile.ZIP_DEFLATED)
    return buf.getvalue()


ODS = _odf(ODS_CONTENT, "application/vnd.oasis.opendocument.spreadsheet")
ODT = _odf(ODT_CONTENT, "application/vnd.oasis.opendocument.text")
OFFICE = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
STYLE = "{urn:oasis:names:tc:opendocument:xmlns:style:1.0}"
TABLE = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"


def _table_modes(data: bytes) -> dict:
    """{sheet name: writing-mode of its table style or None}."""
    root = etree.fromstring(member(data, "content.xml"))
    styles = {
        st.get(STYLE + "name"): st
        for st in root.iter(STYLE + "style")
        if st.get(STYLE + "family") == "table"
    }
    out = {}
    for t in root.iter(TABLE + "table"):
        st = styles.get(t.get(TABLE + "style-name"))
        tp = st.find(STYLE + "table-properties") if st is not None else None
        out[t.get(TABLE + "name")] = tp.get(STYLE + "writing-mode") if tp is not None else None
    return out


def test_ods_rtl():
    filt = get_filter("ods")
    us = filt.parse(ODS).units
    mapping = {str(u.p_index): "שלום " + u.text for u in us}
    out = filt.build_localized(ODS, mapping, HE)
    modes = _table_modes(out)
    assert modes == {"One": "rl-tb", "Two": "rl-tb", "Three": "rl-tb"}, modes
    root = etree.fromstring(member(out, "content.xml"))
    # display="true" and the master page survive on the existing style; a new style is made
    # for the unstyled sheet; column styles are untouched
    ta1 = next(s for s in root.iter(STYLE + "style") if s.get(STYLE + "name") == "ta1")
    assert ta1.get(STYLE + "master-page-name") == "Default"
    assert ta1.find(STYLE + "table-properties").get(TABLE + "display") == "true"
    names = [s.get(STYLE + "name") for s in root.iter(STYLE + "style")]
    assert len(names) == len(set(names)) == 3, names
    # Arabic too (the app's list, not a hard-coded pair)
    assert set(_table_modes(filt.build_localized(ODS, mapping, AR)).values()) == {"rl-tb"}
    for locale in ("fr-FR", "zh-CN"):
        modes = _table_modes(filt.build_localized(ODS, mapping, locale))
        assert modes == {"One": "lr-tb", "Two": None, "Three": "lr-tb"}, (locale, modes)
    # a text document's table keeps its direction (paragraph direction is #724's job)
    t = get_filter("odt")
    tu = t.parse(ODT).units
    out = t.build_localized(ODT, {str(u.p_index): "ש " + u.text for u in tu}, HE)
    assert _table_modes(out) == {"T": None}, _table_modes(out)
    print("  [ods] every sheet right to left for he-IL / ar-SA, ODT tables untouched OK")


# ---------------------------------------------------------------------- idml / icml
PKG_NS = "http://ns.adobe.com/AdobeInDesign/idml/1.0/packaging"
DATE = (2026, 10, 11, 9, 0, 0)
NORMAL_CS = "CharacterStyle/$ID/[No character style]"


def psr(inner: str, extra: str = ' Composer="HL Composer"') -> str:
    return f'<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/$ID/NormalParagraphStyle"{extra}>{inner}</ParagraphStyleRange>'


def csr(inner: str, extra: str = "") -> str:
    return f'<CharacterStyleRange AppliedCharacterStyle="{NORMAL_CS}"{extra}>{inner}</CharacterStyleRange>'


STORY_BODY = (
    # 0: body paragraph with a footnote (footnote = its own paragraph, NOT translated)
    psr(
        csr("<Content>Hello </Content>", ' CharacterDirection="LeftToRightDirection"')
        + csr(
            "<Content>world</Content><Footnote>"
            + psr(
                csr("<Content>Footnote text.</Content>"),
                ' ParagraphDirection="LeftToRightDirection"',
            )
            + "</Footnote>"
        )
    )
    # 1: single-line composer, link text in a hyperlink wrapper
    + psr(
        '<HyperlinkTextSource Self="h1" Name="l">'
        + csr("<Content>Link text here</Content>", ' CharacterDirection="LeftToRightDirection"')
        + "</HyperlinkTextSource>",
        ' Composer="HL Single" ParagraphDirection="LeftToRightDirection"',
    )
    # 2: left untranslated
    + psr(csr("<Content>Stays English.</Content>"), "")
)


def _idml(body: str) -> bytes:
    story = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<idPkg:Story xmlns:idPkg="{PKG_NS}" DOMVersion="20.0">'
        f'<Story Self="u1" AppliedTOCStyle="n" TrackChanges="false"><StoryPreference OpticalMarginAlignment="false" StoryDirection="LeftToRightDirection"/>'
        f"{body}</Story></idPkg:Story>\n"
    ).encode()
    designmap = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<?aid style="50" type="document" readerVersion="6.0" featureSet="257" product="20.0(73)" ?>\n'
        f'<Document xmlns:idPkg="{PKG_NS}" DOMVersion="20.0" Self="d">'
        '<Language Self="Language/$ID/Arabic" Name="$ID/Arabic" Id="300"/>'
        '<idPkg:Story src="Stories/Story_u1.xml"/></Document>\n'
    ).encode()
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        zi = zipfile.ZipInfo("mimetype", date_time=DATE)
        zf.writestr(zi, b"application/vnd.adobe.indesign-idml-package")
        for name, data in (
            ("designmap.xml", designmap),
            ("META-INF/container.xml", b'<?xml version="1.0"?><container/>'),
            ("Stories/Story_u1.xml", story),
        ):
            zi = zipfile.ZipInfo(name, date_time=DATE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zf.writestr(zi, data)
    return out.getvalue()


def _check_story(story_el, rtl: bool) -> None:
    top = [c for c in story_el if c.tag == "ParagraphStyleRange"]
    assert len(top) == 3, len(top)
    fn = story_el.find(".//Footnote/ParagraphStyleRange")
    if rtl:
        assert top[0].get("ParagraphDirection") == "RightToLeftDirection", top[0].attrib
        assert top[0].get("Composer") == "HL Composer Optyca", top[0].attrib
        assert top[1].get("ParagraphDirection") == "RightToLeftDirection", top[1].attrib
        assert top[1].get("Composer") == "HL Single Optyca", top[1].attrib
        dirs = {
            r.get("CharacterDirection")
            for r in top[0].iter("CharacterStyleRange")
            if fn not in r.iterancestors()
        }
        dirs |= {r.get("CharacterDirection") for r in top[1].iter("CharacterStyleRange")}
        assert "LeftToRightDirection" not in dirs, dirs
    else:
        assert top[0].get("ParagraphDirection") is None and top[0].get("Composer") == "HL Composer"
        assert top[1].get("ParagraphDirection") == "LeftToRightDirection"
        assert top[1].get("Composer") == "HL Single"
    # never touched: the untranslated footnote and the untranslated paragraph
    assert fn.get("ParagraphDirection") == "LeftToRightDirection", fn.attrib
    assert fn.get("Composer") is None, fn.attrib
    assert top[2].get("ParagraphDirection") is None and top[2].get("Composer") is None
    sp = story_el.find("StoryPreference")
    assert sp.get("StoryDirection") == "LeftToRightDirection"


def _mapping(filt, data):
    us = filt.parse(data).units
    texts = {u.text: u for u in us}
    want = {}
    for u in us:
        if "Footnote" in u.text or "Stays" in u.text:
            continue
        want[str(u.p_index)] = u.text.replace("Hello", "مرحبا").replace(
            "Link text here", "نص الرابط"
        )
    assert len(want) == 2, texts
    return want


def test_idml_rtl():
    filt = get_filter("idml")
    data = _idml(STORY_BODY)
    mapping = _mapping(filt, data)
    for locale, rtl in ((AR, True), (HE, True), ("fa-IR", True), ("de-DE", False)):
        out = filt.build_localized(data, mapping, locale)
        root = etree.fromstring(member(out, "Stories/Story_u1.xml"))
        _check_story(root.find("Story"), rtl)
    root = etree.fromstring(member(filt.build(data, mapping), "Stories/Story_u1.xml"))
    _check_story(root.find("Story"), False)
    print("  [idml] RTL paragraph direction + World-Ready composer on translated paragraphs OK")


def test_icml_rtl():
    filt = get_filter("icml")
    data = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<?aid style="50" type="snippet" readerVersion="6.0" featureSet="513" product="20.0(73)" ?>\n'
        '<?aid SnippetType="InCopyInterchange"?>\n'
        '<Document DOMVersion="20.0" Self="d">'
        '<Story Self="u1" AppliedTOCStyle="n" TrackChanges="false"><StoryPreference OpticalMarginAlignment="false" StoryDirection="LeftToRightDirection"/>'
        f"{STORY_BODY}</Story></Document>\n"
    ).encode()
    mapping = _mapping(filt, data)
    for locale, rtl in ((AR, True), ("ur-PK", True), ("en-GB", False)):
        out = filt.build_localized(data, mapping, locale)
        _check_story(etree.fromstring(out).find("Story"), rtl)
    print("  [icml] the same through the shared IDML helper OK")


# ---------------------------------------------------------------------- svg
SVG = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    '<svg xmlns="http://www.w3.org/2000/svg" width="400" height="200">'
    '<text id="a" x="10" y="20">Left label</text>'
    '<text id="b" x="200" y="40" text-anchor="middle">Centered</text>'
    '<text id="c" x="390" y="60" style="font-size:12px; text-anchor: end">Right label</text>'
    '<g text-anchor="end"><text id="d" x="390" y="80">Inherited end</text></g>'
    '<text id="e" x="10" y="100"><tspan x="10" text-anchor="start">First line</tspan>'
    '<tspan x="10" dy="14">Second line</tspan></text>'
    '<text id="f" x="10" y="140" direction="ltr">Explicit ltr</text>'
    '<text id="g" x="10" y="160">Not translated</text>'
    "</svg>\n"
)


def _svg_text(root, tid):
    return next(t for t in root.iter("{http://www.w3.org/2000/svg}text") if t.get("id") == tid)


def test_svg_rtl():
    filt = get_filter("image")
    from app.filters.image_filter import _split_tags, _svg_units

    n = len(_svg_units(_split_tags(SVG)))
    assert n == 8, n
    mapping = {str(k): "טקסט " + str(k) for k in range(n - 1)}  # last stays
    out = filt.build_localized(SVG.encode(), mapping, HE)
    root = etree.fromstring(out)
    a, b, c, d, e, f, g = (_svg_text(root, x) for x in "abcdefg")
    for t in (a, b, c, d, e, f):
        assert t.get("direction") == "rtl", (t.get("id"), t.attrib)
    assert a.get("text-anchor") == "end", a.attrib
    assert b.get("text-anchor") == "middle", b.attrib
    assert "text-anchor: start" in c.get("style") and c.get("text-anchor") is None, c.attrib
    assert c.get("style").startswith("font-size:12px"), c.get("style")
    assert d.get("text-anchor") == "start", d.attrib
    assert e.get("text-anchor") == "end", e.attrib
    spans = list(e)
    assert spans[0].get("text-anchor") == "end" and spans[1].get("text-anchor") is None, [
        s.attrib for s in spans
    ]
    assert g.get("direction") is None and g.get("text-anchor") is None, g.attrib
    assert root.get("{http://www.w3.org/XML/1998/namespace}lang") == HE
    # left-to-right targets and plain build: no direction, no anchor change
    for data in (
        filt.build_localized(SVG.encode(), mapping, "de-DE"),
        filt.build(SVG.encode(), mapping),
    ):
        root = etree.fromstring(data)
        assert all(
            t.get("direction") in (None, "ltr")
            for t in root.iter("{http://www.w3.org/2000/svg}text")
        )
        assert _svg_text(root, "a").get("text-anchor") is None
    # gate: an empty self-closing <text/> before a translated label must not take its edits;
    # a middle-anchored, rotated label with tspans keeps middle and its transform
    svg2 = (
        b'<svg xmlns="http://www.w3.org/2000/svg"><text id="z" x="1" y="1"/>'
        b'<text id="a" x="10" y="20">Hello</text>'
        b'<text id="m" x="50" y="40" transform="rotate(30 50 40)" text-anchor="middle">'
        b'<tspan x="50">One</tspan><tspan x="50" dy="12" text-anchor="start">Two</tspan></text></svg>'
    )
    root = etree.fromstring(filt.build_localized(svg2, {"0": "שלום", "1": "א", "2": "ב"}, HE))
    z, a, m = (_svg_text(root, x) for x in "zam")
    assert z.get("direction") is None and z.get("text-anchor") is None, z.attrib
    assert a.get("direction") == "rtl" and a.get("text-anchor") == "end", a.attrib
    assert m.get("text-anchor") == "middle" and m.get("transform") == "rotate(30 50 40)", m.attrib
    spans = list(m)
    assert spans[0].get("text-anchor") is None and spans[1].get("text-anchor") == "end", [
        s.attrib for s in spans
    ]
    print("  [svg] direction=rtl + mirrored anchors on translated text only OK")


def main():
    test_xlsx_rtl()
    test_ods_rtl()
    test_idml_rtl()
    test_icml_rtl()
    test_svg_rtl()
    print("RTL LAYOUT OK")


if __name__ == "__main__":
    main()
