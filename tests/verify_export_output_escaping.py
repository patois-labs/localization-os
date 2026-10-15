"""Security train regressions for stored URLs, HTML and spreadsheet output."""

import ast
import csv
import importlib
import io
import os
import sys
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
sys.path.insert(0, str(ROOT))
from tests._bootstrap import init_env

init_env("loctms_output_escaping_")

from openpyxl import Workbook, load_workbook

from app import output_csv
from app.db import init_db
from app.filters.html_filter import HtmlFilter
from app.filters.html_guard import UnsafeMarkupRefused
from app.filters.xlsx_filter import XlsxFilter
from app.output_safety import InvalidOutputURL, csv_original_cell, csv_safe_cell, safe_http_url
from app.services import data_export
from app.services.connectors import store_part2 as store


def _event_attrs(markup) -> list:
    """Every on* attribute a browser-grade parser sees in ``markup`` (#672: a quote in a
    text node is plain text, so the injection checks look at the parsed attributes rather
    than at quote entities)."""
    from html.parser import HTMLParser

    if isinstance(markup, bytes):
        markup = markup.decode("utf-8")
    found: list = []

    class _P(HTMLParser):
        def handle_starttag(self, tag, attrs):
            found.extend(name for name, _v in attrs if name.lower().startswith("on"))

        handle_startendtag = handle_starttag

    _P().feed(markup)
    return found


def test_html():
    from app.services.qa import check_segment

    for target in ("Use < 5", "Use <img src=x>"):
        findings = [f for f in check_segment("Use five", target) if f["type"] == "tags"]
        assert len(findings) == 1 and findings[0]["severity"] == "major"
    filt = HtmlFilter()
    source = b"<p>Type &lt;script&gt; &amp; &#60; &#X3C; now.</p><p>Hello <b>world</b></p>"
    units = filt.parse(source).units
    assert filt.build(source, {str(u.p_index): u.text for u in units}) == source
    payload = 'Hola <img src=x onerror="alert(document.domain)"> mundo'
    for build in (filt.build, filt.build_localized):
        result = build(b"<p>Hello world</p>", {"0": payload})
        assert b"<img" not in result and b"&lt;img" in result
    translated = units[1].text.replace("world", "<script>alert(1)</script>")
    result = filt.build(source, {"1": translated})
    assert b"<b>&lt;script&gt;" in result
    assert b"<script>" not in result

    # Real connector payload builders, before transport, with the reviewer's target.
    extra = dict(  # noqa: C408
        html=filt,
        body_filter=filt,
        body_bytes=b"<p>Hello world</p>",
        masked_body_bytes=b"<p>Hello world</p>",
        shielded_body_bytes=b"<p>Hello world</p>",
        title="Title",
        raw_title="Title",
        source_title="Title",
        shielded_title="Title",
        author_id=1,
        body_key="answer",
        title_codes={},
        body_codes={},
        editor="html",
        shortcodes={},
        post_type="post",
        source_collection="collection",
        slug="hello",
        related=[],
    )
    for module, name in (
        ("zendesk", "_zd_build_payload"),
        ("wordpress", "_wp_build_payload"),
        ("intercom", "_ic_build_payload"),
        ("freshdesk", "_fd_build_payload"),
        ("helpscout", "_hs_build_payload"),
        ("zoho", "_zo_build_payload"),
        ("document360", "_d360_build_payload"),
    ):
        builder = getattr(importlib.import_module("app.services.connectors." + module), name)
        # Intercom's builder also reads prep.body_units (held-body rule, 9e80057bf); the other
        # builders ignore it. [] = "no held unit", so the escaping under test is unchanged.
        prep_fields = {"body_units": []}
        prep = SimpleNamespace(extra=extra, **prep_fields)
        body = builder(prep, "de", "Title", {"0": payload})["body"]
        assert "<img" not in body and "&lt;img" in body, module
        quoted = b'<div title="a > b">Hello world</div>'
        quoted_extra = dict(extra)
        for key in ("body_bytes", "masked_body_bytes", "shielded_body_bytes"):
            quoted_extra[key] = quoted
        body = builder(
            SimpleNamespace(extra=quoted_extra, **prep_fields),
            "de",
            "Title",
            {"0": 'b" onfocus=alert(1) x="'},
        )["body"]
        assert body == '<div title="a > b">b" onfocus=alert(1) x="</div>', module
        assert _event_attrs(body) == [], (module, body)

    from app.filters import get_filter
    from app.services.connectors.fieldmap import extract_record
    from app.services.connectors.hubspot import HubSpotVariantStrategy
    from app.services.connectors.salesforce_common import _sf_build_payload

    extraction = extract_record(
        {"Title": "Title", "Body": '<div title="a > b">Hello world</div>'},
        [{"name": "Body", "type": "html"}],
        get_filter,
    )
    sf_payload = _sf_build_payload(
        SimpleNamespace(extra={"extraction": extraction}),
        "de",
        "Title",
        {"Body#0": 'b" onfocus=alert(1) x="'},
    )
    assert sf_payload["Body"] == '<div title="a > b">b" onfocus=alert(1) x="</div>'
    assert _event_attrs(sf_payload["Body"]) == []

    patch, _, _, _ = HubSpotVariantStrategy._build_patch(
        None,
        {},
        {
            "body": {
                "kind": "html",
                "filter": filt,
                "mapping": {"0": payload},
                "source_bytes": b"<p>Hello world</p>",
                "complete": True,
            }
        },
    )
    assert "<img" not in patch["postBody"] and "&lt;img" in patch["postBody"]
    patch, _, _, _ = HubSpotVariantStrategy._build_patch(
        None,
        {},
        {
            "body": {
                "kind": "html",
                "filter": filt,
                "mapping": {"0": 'b" onfocus=alert(1) x="'},
                "source_bytes": b'<div title="a > b">Hello world</div>',
                "complete": True,
            }
        },
    )
    assert patch["postBody"] == '<div title="a > b">b" onfocus=alert(1) x="</div>'
    assert _event_attrs(patch["postBody"]) == []
    patch, _, _, _ = HubSpotVariantStrategy._build_patch(
        None,
        {"layoutSections": {"html": "<p>Hello world</p>"}},
        {"body": {"kind": "layout", "sets": [(["html"], payload)]}},
    )
    leaf = patch["layoutSections"]["html"]
    assert "<img" not in leaf and "&lt;img" in leaf
    from app.filters.html_filter import escape_fragment_translation
    from app.services.connectors.wordpress import _wp_build_payload

    extra.update(has_excerpt=True, source_excerpt="A summary")
    payload_out = _wp_build_payload(
        SimpleNamespace(extra=extra), "de", "Title", {"excerpt": payload}
    )
    assert "<img" not in payload_out["excerpt"] and "&lt;img" in payload_out["excerpt"]
    assert escape_fragment_translation("<p>Hello</p>", "<p>Hola</p>") == "<p>Hola</p>"
    assert escape_fragment_translation("<p>Hello</p>", "[fr] <p>Hola</p>") == ("[fr] <p>Hola</p>")
    original = "<script>safe()</script><p>Hello</p>"
    assert escape_fragment_translation(original, "<script>evil()</script><p>Hola</p>") == (
        "<script>safe()</script><p>Hola</p>"
    )
    for script in ("<script>safe();</style>keep();</script>", "<script/>safe();</script>"):
        original = (script + "<p>Hello</p>").encode()
        assert [u.text for u in filt.parse(original).units] == ["Hello"]
        assert filt.build(original, {"0": "Hola"}) == (script + "<p>Hola</p>").encode()
        assert (
            escape_fragment_translation(
                script + "<p>Hello</p>", script.replace("safe()", "evil()") + "<p>Hola</p>"
            )
            == script + "<p>Hola</p>"
        )


def test_script_comparison():
    _check_raw_body("<script>if (a < b) { el.innerHTML = '<b>x</b>'; }</script>")


def test_script_list():
    _check_raw_body("<script>for(i=0;i<n;i++){h+='<li>'+x[i]+'</li>';}</script>")


def _check_raw_body(prefix):
    filt = HtmlFilter()
    source = (prefix + "<h1>Hello</h1><p>World</p>").encode()
    assert [u.text for u in filt.parse(source).units] == ["Hello", "World"]
    assert filt.build(source, {"0": "Hello", "1": "World"}) == source
    assert (
        filt.build(source, {"0": "Hola", "1": "Mundo"})
        == (prefix + "<h1>Hola</h1><p>Mundo</p>").encode()
    )


def test_script_double_escaped():
    from app.services.connectors.wordpress import _mask_shortcodes

    for prefix in (
        "<script><!--<script>x</script>alert(0)</script>",
        "<SCRIPT><!--<ScRiPt >x</sCrIpT >alert(0)</SCRIPT >",
        "<script><!--<script>x</script>--></script>",
        "<script><!--<script>x-->alert(0)</script>",
        "<script><!--x</script>",
        "<script><!--><script>x</script>",
        "<script><!---><script>x</script>",
        "<script>var x = '</scripture>'; keep()</script>",
        "<style>p::before { content: '<b title=\"x>y\">'; }</STYLE >",
    ):
        _check_raw_body(prefix)
    prefix = "<script><!--<script>x[i]</script>alert(0)</script>"
    masked, codes = _mask_shortcodes(prefix + "<p>Hello</p>")
    assert not codes and masked == prefix + "<p>Hello</p>"
    _check_raw_body(prefix)


def test_tokenizer_linear():
    # Bound a regressed tokenizer in a subprocess, rather than hanging the suite.
    import subprocess

    failures = []
    for name, source in (
        ("less-than", "<" * 200_000),
        ("double-quote", "<b x" * 50_000 + '">'),
        ("single-quote", "<b x" * 50_000 + "'>"),
        ("quoted-value", "<b x" * 50_000 + " y='>"),
        ("unclosed-comment", "<!--" + "<!--x" * 40_000),
    ):
        code = (
            "import sys, time\n"
            "from app.filters.html_filter import _tokenize\n"
            "text = sys.stdin.read()\n"
            "start = time.perf_counter()\n"
            "tokens = _tokenize(text)\n"
            "elapsed = time.perf_counter() - start\n"
            "print(f'{elapsed:.6f}', flush=True)\n"
            "assert ''.join(tokens) == text\n"
            "assert elapsed < 1.0, elapsed\n"
        )
        try:
            result = subprocess.run(
                [sys.executable, "-c", code],
                input=source,
                capture_output=True,
                text=True,
                timeout=5,
                cwd=ROOT,
            )
        except subprocess.TimeoutExpired:
            failures.append(name + " exceeded 5 s subprocess deadline")
        else:
            measurement = f"tokenizer {name}: {result.stdout.strip()} s"
            print(measurement, flush=True)
            if os.environ.get("OUTPUT_ESCAPING_TIMING_LOG"):
                with open(os.environ["OUTPUT_ESCAPING_TIMING_LOG"], "a", encoding="utf-8") as log:
                    log.write(measurement + "\n")
            if result.returncode:
                failures.append(name + ": " + result.stderr)
    assert not failures, failures


def test_scanner_boundaries():
    from app.filters.html_filter import _tokenize

    filt = HtmlFilter()
    for prefix in ("<!-- unclosed <b>hidden", "<![CDATA[unclosed <b>hidden"):
        assert not filt.parse(prefix.encode()).units
        assert filt.build(prefix.encode(), {}) == prefix.encode()
    for source in ("Use < 5", '<div title="unclosed >text', "<div title='unclosed >text"):
        assert [u.text for u in filt.parse(source.encode()).units] == [source]
        assert filt.build(source.encode(), {"0": 'Safe " < text'}) == b'Safe " &lt; text'
    source = '<!-- <b title="x>y"> --><![CDATA[<i>opaque</i>]]><p>Hello</p>'
    assert len(_tokenize(source)) == 9
    assert [u.text for u in filt.parse(source.encode()).units] == ["Hello"]


def test_source_hubl_masks():
    from app.services.connectors.hubspot import (
        HubSpotConnector,
        HubSpotVariantStrategy,
        _hubspot_build_payload,
    )

    source = "Hi {% if n > 1 %}friends{% endif %}"
    connector = HubSpotConnector({"content_types": "site-pages"}, {})
    title, units, _, extra = connector.extract(
        {"_ctype": "site-pages", "name": "Title", "layoutSections": {"html": source}}
    )
    extra["connector"] = connector
    mapping = {str(u.p_index): u.text.replace("Hi", "Hola") for u in units}
    prep = SimpleNamespace(extra=extra, body_units=units, aid="site-pages:1", title_text=title)
    payload = _hubspot_build_payload(prep, "de", title, mapping)
    for current in ("Old target", "Old {% if n < 2 %}target{% endif %}"):
        patch, _, _, _ = HubSpotVariantStrategy._build_patch(
            None,
            {"layoutSections": {"html": current}},
            payload,
        )
        assert patch["layoutSections"]["html"] == source.replace("Hi", "Hola")


def test_xlsx():
    wb = Workbook()
    wb.active["A1"] = "Hello world"
    wb.active["B1"] = "=2+2"
    original = io.BytesIO()
    wb.save(original)
    filt = XlsxFilter()
    for prefix in ("=", "+", "-", "@", "\t", "\r"):
        target = prefix + "1+1"
        for build in (filt.build, filt.build_localized):
            result = load_workbook(io.BytesIO(build(original.getvalue(), {"0": target})))
            assert result.active["A1"].value == target
            assert result.active["A1"].data_type == "s"
            assert result.active["B1"].data_type == "f"
        out = Workbook(write_only=True)
        ws = out.create_sheet()
        ws.append([data_export._xlsx_cell(target, ws)])
        buf = io.BytesIO()
        out.save(buf)
        cell = load_workbook(io.BytesIO(buf.getvalue())).active["A1"]
        assert cell.data_type == "s" and cell.value == target
        assert data_export._csv_cell(target) == "'" + target


def test_quoted_attributes():
    from app.filters.html_filter import _INLINE_TAGS, escape_fragment_translation
    from app.services.segment import ph_point

    filt = HtmlFilter()
    tags = _INLINE_TAGS | {"div", "p", "input", "area", "custom-tag", "ac:structured-macro"}
    for tag in sorted(tags):
        for quote in ('"', "'"):
            start = f"<{tag} data-label={quote}Home > Docs{quote}>"
            source = (start + "Hello world" + f"</{tag}>").encode()
            units = filt.parse(source).units
            assert len(units) == 1 and "Docs" not in units[0].text, tag
            mapping = {"0": units[0].text.replace("Hello world", "Hola mundo")}
            assert filt.build(source, mapping) == source.replace(b"Hello world", b"Hola mundo")
            target = units[0].text.replace("Hello world", 'b" onfocus=alert(1) x="')
            result = filt.build(source, {"0": target}).decode()
            assert start in result and 'b" onfocus=alert(1) x="' in result
            assert _event_attrs(result) == [], (tag, result)
    source = b'<p>See <img alt="a > b"> the picture</p>'
    target = ph_point("x", 1) + ' b" onerror=alert(1) src=x y="'
    assert filt.build(source, {"0": target}) == (
        b'<p><img alt="a > b"> b" onerror=alert(1) src=x y="</p>'
    )
    assert _event_attrs(filt.build(source, {"0": target})) == []
    original = '<b title="a > b">Hi</b>'
    attack = '<b title="a >b" onmouseover=alert(1) x="</b>'
    assert "<b" not in escape_fragment_translation(original, attack)
    assert escape_fragment_translation(original, original.replace("Hi", "Hola")) == (
        original.replace("Hi", "Hola")
    )
    assert filt.build(b"<p>Hello</p>", {"0": "\"' &amp; &"}) == (b"<p>\"' &amp; &amp;</p>")


def test_quoted_fidelity():
    filt = HtmlFilter()
    source = b'<div title="a > b">Hello world</div>'
    assert filt.parse(source).units[0].text == "Hello world"
    assert filt.build(source, {"0": "Hola mundo"}) == source.replace(b"Hello world", b"Hola mundo")
    # Explicit accessibility translations use their own escaped attribute channel.
    source = b'<p>See <img alt="a > b"> here</p><p>Next</p>'
    units = filt.parse(source).units
    assert [u.p_index for u in units] == [0, 1, 2]
    assert units[2].style == "a11y:img-alt"
    assert filt.build(source, {"0": units[0].text, "1": "Next"}) == source
    assert filt.build(source, {str(u.p_index): u.text for u in units}) == source
    result = filt.build(source, {"2": 'b" onerror=alert(1) x="'})
    assert b'alt="b&quot; onerror=alert(1) x=&quot;"' in result
    for prefix in (
        '<script data-label="a > b">keep()</script>',
        "<style data-label='a > b'>p { color: red }</style>",
        "<!-- author's > comment -->",
        '<![CDATA[quoted " > content]]>',
        '<!DOCTYPE html SYSTEM "a > b">',
    ):
        source = (prefix + "<p>Hello</p>").encode()
        assert [u.text for u in filt.parse(source).units] == ["Hello"]
        assert filt.build(source, {"0": "Hola"}) == (prefix + "<p>Hola</p>").encode()


def test_template_fidelity():
    from app.services.connectors.hubspot import (
        HubSpotConnector,
        HubSpotVariantStrategy,
        _hubspot_build_payload,
        _set_path,
    )
    from app.services.connectors.wordpress import _wp_build_payload

    for expression in ("{% if n > 1 %}", "{% if n < 1 %}", '{{ "a&b" }}'):
        source = "Hi " + expression + "friends{% endif %}"
        target = "Hola " + expression + "amigos{% endif %} <img src=x> &amp;"
        for container, path in (({"html": source}, ["html"]), ([source], [0])):
            assert _set_path(container, path, target)
            result = container[path[0]]
            assert expression in result and "{% endif %}" in result
            assert "<img" not in result and "&amp;amp;" not in result
        patch, _, _, _ = HubSpotVariantStrategy._build_patch(
            None,
            {"layoutSections": {"html": source}},
            {"body": {"kind": "layout", "sets": [(["html"], target)]}},
        )
        assert expression in patch["layoutSections"]["html"]
        connector = HubSpotConnector({"content_types": "blog-posts"}, {})
        html = "<p>" + source + "</p>"
        title, units, _, extra = connector.extract(
            {"_ctype": "blog-posts", "name": "Title", "postBody": html}
        )
        extra["connector"] = connector
        mapping = {str(u.p_index): u.text.replace("Hi", "Hola") for u in units}
        prep = SimpleNamespace(extra=extra, body_units=units, aid="blog-posts:1", title_text=title)
        payload = _hubspot_build_payload(prep, "de", title, mapping)
        patch, _, _, _ = HubSpotVariantStrategy._build_patch(None, {}, payload)
        assert patch["postBody"] == html.replace("Hi", "Hola")
    excerpt = 'Hi [caption label="a > b & c"]friends[/caption]'
    extra = {
        "source_excerpt": excerpt,
        "source_title": "Title",
        "masked_body_bytes": b"",
        "shortcodes": [],
        "post_type": "post",
        "has_excerpt": True,
    }
    result = _wp_build_payload(
        SimpleNamespace(extra=extra),
        "de",
        "Title",
        {"excerpt": excerpt.replace("Hi", "Hola") + " <img src=x>"},
    )["excerpt"]
    assert '[caption label="a > b & c"]' in result and "<img" not in result


def test_stored_html_indices():
    from sqlmodel import Session, select

    from app.db import engine
    from app.models import Project, Segment, TargetLocale, TransMode
    from app.services.export import ExportUnreadableBlob, build_translated_file
    from app.services.ingest import ingest_document

    init_db()
    with Session(engine) as session:
        project = Project(name="HTML indices", source_locale="en-US", default_provider="mock")
        session.add(project)
        session.commit()
        session.add(TargetLocale(project_id=project.id, locale="de-DE"))
        session.commit()
        for script in (
            "<script>if (a < b) { el.innerHTML = '<b>x</b>'; }</script>",
            "<script>for(i=0;i<n;i++){h+='<li>'+x[i]+'</li>';}</script>",
        ):
            original = (script + "<p>Hello world</p>").encode()
            doc = ingest_document(session, project, "script.html", original, TransMode.SEGMENTED)
            segments = session.exec(select(Segment).where(Segment.document_id == doc.id)).all()
            assert [s.source_text for s in segments] == ["Hello world"]
            segments[0].target_text = "Hola mundo"
            session.add(segments[0])
            session.commit()
            assert build_translated_file(session, doc, "de-DE") == original.replace(
                b"Hello world", b"Hola mundo"
            )
        source = b'<div title="a > b">Hello world</div>'
        for suffix in ("html", "htm", "xhtml"):
            doc = ingest_document(
                session,
                project,
                "indices." + suffix,
                source,
                TransMode.SEGMENTED,
            )
            segments = session.exec(select(Segment).where(Segment.document_id == doc.id)).all()
            for segment in segments:
                segment.target_text = "Hola mundo"
                session.add(segment)
            session.commit()
            assert build_translated_file(session, doc, "de-DE") == source.replace(
                b"Hello world", b"Hola mundo"
            )
            doc.skeleton = {**doc.skeleton, "units": [{"p_index": 0, "text": 'b">Hello world'}]}
            try:
                build_translated_file(session, doc, "de-DE")
            except ExportUnreadableBlob as exc:
                assert "Re-import" in str(exc)
            else:
                raise AssertionError("legacy HTML indices were applied to a changed extraction")


def test_csv_apostrophes():
    from app.services.glossary_common import _csv_rows

    values = ["'=x", "'-ism", "'ordinary", "''=x"]
    buf = io.StringIO(newline="")
    writer = output_csv.writer(buf)
    writer.writerow(["source_term", "target_term"])
    writer.writerows([[value, value] for value in values])
    rows = list(_csv_rows(buf.getvalue().encode()))
    assert [row["source_term"] for row in rows] == values
    # Reports and data export sanitize cells before the guarded writer sees them.
    buf = io.StringIO(newline="")
    output_csv.writer(buf).writerow([csv_safe_cell(value) for value in values])
    assert next(csv.reader(io.StringIO(buf.getvalue()))) == ["'" + value for value in values]


def test_csv_prefixes():
    for value in (" =1+1", "\n=1+1", "\u00a0=1+1", "\uff1d1+1", "\uff0b1", "\uff0d1", "\uff201"):
        assert csv_safe_cell(value) == "'" + value
        assert csv_original_cell(csv_safe_cell(value)) == value


def test_legacy_urls():
    from sqlmodel import Session

    from app.db import engine
    from app.models import Connection

    init_db()
    conn = store.create("wordpress", "Legacy", {"base_url": "https://example.com"}, {})
    for legacy in ("https://x.atlassian.net ", "javascript:legacy()"):
        with Session(engine) as session:
            row = session.get(Connection, conn.id)
            row.config = {"base_url": legacy}
            session.add(row)
            session.commit()
        store.patch_config(conn.id, {"marker": 1})
        store.patch_config_merge(
            conn.id,
            "hashes",
            watermark_key="watermark",
            watermark_at_least=42,
        )
        store.patch_config_nested(conn.id, {"hashes": {"add": {"one": "abc"}}})
        cfg = store.get(conn.id).config
        assert cfg["watermark"] == 42 and cfg["hashes"] == {"one": "abc"}
        assert cfg["base_url"] == legacy
    assert safe_http_url("https://x.atlassian.net ") == "https://x.atlassian.net"
    from app.services.connectors.base import ConnectorError

    # Refused by F6's URL guard (InvalidOutputURL) or, first, by F1's unreadable-origin
    # refusal (ConnectorError): either way the unsafe URL is never saved.
    try:
        store.patch_config(conn.id, {"base_url": "javascript:changed()"})
    except (InvalidOutputURL, ConnectorError):
        pass
    else:
        raise AssertionError("changed unsafe URL saved")


def test_csv():
    values = [prefix + "1+1" for prefix in ("=", "+", "-", "@", "\t", "\r")]
    for value in values:
        assert csv_original_cell(csv_safe_cell(value)) == value
    assert csv_safe_cell(-3) == -3
    from app.services.glossary_common import _csv_rows

    assert list(_csv_rows(b"source_term,target_term\r\n'=1+1,'@example\r\n")) == [
        {"source_term": "=1+1", "target_term": "@example"}
    ]
    for dictionary in (False, True):
        buf = io.StringIO(newline="")
        if dictionary:
            writer = output_csv.DictWriter(buf, fieldnames=values)
            writer.writeheader()
            writer.writerows([dict(zip(values, values, strict=True))])
        else:
            writer = output_csv.writer(buf)
            writer.writerows([values])
        for row in csv.reader(io.StringIO(buf.getvalue(), newline="")):
            assert row == ["'" + value for value in values]


def test_urls():
    from app.services.connectors.base import ConnectorError

    init_db()
    conn = store.create("wordpress", "Blog", {"base_url": "https://blog.example.com"}, {})
    for value in (
        "javascript:alert(document.domain)//",
        "data:text/html,Hi",
        "//example.com",
        "https:///missing",
        "https://",
        123,
        "https://host\n/path",
    ):
        assert not safe_http_url(value)
        for write in (
            lambda value=value: store.create("wordpress", "Bad", {"base_url": value}, {}),
            lambda value=value: store.update(conn.id, "Bad", {"base_url": value}, None),
            lambda value=value: store.patch_config(conn.id, {"base_url": value}),
        ):
            try:
                write()
            except (InvalidOutputURL, ConnectorError):  # F6 URL guard or F1 origin refusal
                pass
            else:
                raise AssertionError("unsafe URL saved")
        assert store.get(conn.id).config["base_url"] == "https://blog.example.com"
    assert safe_http_url("https://example.com/path")
    from tests._cli_runner import run_cli

    code, out, err = run_cli(
        [
            "--json",
            "connector",
            "update",
            str(conn.id),
            "--config",
            "base_url=javascript:alert(1)",
        ]
    )
    # Refused cleanly by F6's URL guard or F1's origin refusal (whichever runs first).
    assert code != 0 and "Traceback" not in err, (code, out, err)
    assert "http:// or https://" in out or "https://host" in out, out
    from app.web.deps import templates

    template = templates.env.from_string("{% if url|safe_http_url %}LINK{% endif %}")
    assert template.render(url="javascript:alert(1)") == ""
    assert template.render(url="https://example.com") == "LINK"

    from fastapi.testclient import TestClient
    from sqlmodel import Session, select

    from app.db import engine
    from app.main import app
    from app.models import Connection, User
    from app.services import apikeys, security
    from app.services.connectors import idmap
    from tests._csrf import csrf_for

    security.seed_initial()
    client = TestClient(app)

    def login(who):
        csrf_for(client)
        form = {
            "email": who,
            "password": who,  # pregate: allow-secret seeded fixture
            "next": "/",
        }  # pregate: allow-secret seeded test user
        response = client.post("/auth/login", data=form, follow_redirects=False)
        assert response.status_code == 303
        csrf_for(client)

    login("linguist")
    bad = "javascript:alert(document.domain)//"
    response = client.post(
        f"/connectors/connection/{conn.id}",
        data={"name": "Blog", "base_url": bad},
        follow_redirects=False,
    )
    # Either guard may refuse first: the output-URL handler answers 400, while the
    # connector URL guard (ConnectorError) redirects back to the form with ok=0.
    if response.status_code == 303:
        assert "ok=0" in response.headers["location"], response.headers["location"]
    else:
        assert response.status_code == 400, response.status_code
    assert store.get(conn.id).config["base_url"] == "https://blog.example.com"
    with Session(engine) as session:
        linguist = session.exec(select(User).where(User.email == "linguist")).one()
        _, key = apikeys.create_key(linguist.id, "Output regression")
    response = client.post(
        f"/api/connectors/{conn.id}/update",
        json={"config": {"base_url": bad}},
        headers={"X-Loc-Token": key},
    )
    assert response.status_code == 400, response.text
    assert store.get(conn.id).config["base_url"] == "https://blog.example.com"
    # Legacy rows need a render-time guard even before anyone edits them.
    with Session(engine) as session:
        row = session.get(Connection, conn.id)
        row.config = {"base_url": bad}
        session.add(row)
        session.commit()
    idmap.reserve(conn.id, "post:1", "de", "post:77", status=idmap.NEEDS_LINK)
    login("admin")
    page = client.get(f"/connectors/connection/{conn.id}")
    assert page.status_code == 200
    assert 'href="javascript:' not in page.text


def test_writer_ratchet():
    from collections import Counter

    # The only raw CSV constructors belong to the guarded adapter itself.
    # The XLSX metadata reader never writes cells. All other openpyxl users are
    # enumerated below, so adding a writer requires explicit security review.
    xlsx_allow = {
        "app/filters/xlsx_filter.py": "reads cells; writes go through xlsx_package (#685)",
        "app/filters/xlsx_package.py": "formula Translator only, to unshare a formula group",
        "app/filters/xlsx_guard.py": "read-only package/coordinate/formula helpers for preflight",
        "app/services/data_export.py": "_xlsx_cell wraps every appended value",
    }
    # Exact .value assignment inventory. These are database fields and thread locals,
    # not spreadsheet cells; counts prevent an added write hiding in an allowed file.
    value_allow = {
        ("app/db.py", "marker.value"): 1,
        ("app/db.py", "pb_marker.value"): 1,
        ("app/filters/php_array_filter.py", "self.value"): 1,  # lexer token attribute, not output
        ("app/services/audit.py", "row.value"): 2,
        ("app/services/content_project_rename_delete_sp.py", "self.value"): 1,
        ("app/services/data_export_part2.py", "row.value"): 1,
        ("app/services/keys_common.py", "_actor.value"): 2,
        ("app/services/ledger_controls.py", "anchor.value"): 1,
        ("app/services/ledger_controls.py", "row.value"): 1,
        ("app/services/settings.py", "row.value"): 1,
        ("app/output_safety.py", "cell.value"): 1,  # guarded string write
    }
    value_writes = Counter()
    for path in (ROOT / "app").rglob("*.py"):
        rel = path.relative_to(ROOT).as_posix()
        source = path.read_text(encoding="utf-8")
        tree = ast.parse(source)
        raw_csv = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                raw_csv.update(a.asname or a.name for a in node.names if a.name == "csv")
                if any(a.name.startswith("openpyxl") for a in node.names):
                    assert rel in xlsx_allow, rel
            if isinstance(node, ast.ImportFrom):
                if node.module == "csv":
                    assert not any(a.name in ("writer", "DictWriter") for a in node.names), rel
                if (node.module or "").startswith("openpyxl"):
                    assert rel in xlsx_allow, rel
                    if rel == "app/filters/xlsx_guard.py":
                        allowed = {
                            "openpyxl.utils.cell": {"range_boundaries", "get_column_letter"},
                            "openpyxl.formula.translate": {"Translator"},
                            "openpyxl.packaging.relationship": {"Relationship", "get_rels_path"},
                            "openpyxl.xml.constants": {
                                "SHARED_STRINGS",
                                "XLSM",
                                "XLSX",
                                "XLTM",
                                "XLTX",
                            },
                        }
                        assert node.module in allowed, rel
                        assert {a.name for a in node.names} <= allowed[node.module], rel
                    if rel == "app/filters/xlsx_package.py":
                        assert node.module == "openpyxl.formula.translate", rel
                        assert {a.name for a in node.names} == {"Translator"}, rel
        for node in ast.walk(tree):
            if (
                isinstance(node, ast.Attribute)
                and isinstance(node.ctx, ast.Store)
                and node.attr == "value"
            ):
                value_writes[(rel, ast.unparse(node))] += 1
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute):
                func = node.func
                assert func.attr != "cell", f"Review spreadsheet cell write in {rel}"
                if (
                    isinstance(func.value, ast.Name)
                    and func.value.id in raw_csv
                    and func.attr in ("writer", "DictWriter")
                ):
                    assert rel == "app/output_csv.py", rel
        if rel == "app/filters/xlsx_filter.py":
            # #685: translations are written into the original package XML, never through
            # an openpyxl cell (the workbook is only read); xlsx_package writes a string.
            assert "pkg.set_cell_text(" in source and ".value = " not in source
        if rel == "app/filters/xlsx_package.py":
            assert 'c.set("t", "inlineStr")' in source and 'SubElement(c, _M + "f")' not in source
        if rel == "app/services/data_export.py":
            calls = [
                n
                for n in ast.walk(tree)
                if isinstance(n, ast.Call)
                and isinstance(n.func, ast.Attribute)
                and n.func.attr == "append"
                and ast.unparse(n.func.value) == "self._ws"
            ]
            assert len(calls) == 2
            assert all("_xlsx_cell" in ast.unparse(n) for n in calls)
    assert value_writes == value_allow, value_writes - Counter(value_allow)


def base_regression(issue):
    """Opt-in offline replay of original HEAD seams, without editing the working tree."""
    import subprocess

    from app.filters import html_filter, xlsx_filter
    from app.services import data_export_part2

    module, definition, check = {
        "227": (store, "create", test_urls),
        "234": (html_filter, "_rematerialize", test_html),
        "273": (xlsx_filter, "XlsxFilter", test_xlsx),
        "274": (data_export_part2, "_csv_cell", test_xlsx),
    }[issue]
    path = Path(module.__file__).relative_to(ROOT).as_posix()
    original = subprocess.check_output(["git", "show", "HEAD:" + path], encoding="utf-8")
    node = next(n for n in ast.parse(original).body if getattr(n, "name", None) == definition)
    exec(compile(ast.Module(body=[node], type_ignores=[]), path, "exec"), module.__dict__)
    if issue == "273":
        globals()["XlsxFilter"] = module.XlsxFilter
    if issue == "274":
        data_export._csv_cell = module._csv_cell
    check()  # The unchanged security assertion MUST fail on the original seam.
    print("BASE SECURITY ASSERTIONS OK")
    raise SystemExit(0)


if os.environ.get("OUTPUT_ESCAPING_BASE"):
    base_regression(os.environ["OUTPUT_ESCAPING_BASE"])

round_one = (
    test_quoted_attributes,
    test_quoted_fidelity,
    test_template_fidelity,
    test_csv_apostrophes,
    test_csv_prefixes,
    test_legacy_urls,
)
if os.environ.get("OUTPUT_ESCAPING_ROUND1_BASE"):
    import subprocess

    gate_revision = "86098a58270389d42e1854878ae94b86bd602abe"
    for name in (
        "app.filters.html_filter",
        "app.output_safety",
        "app.services.connectors.store_part2",
        "app.services.connectors.hubspot",
        "app.services.connectors.wordpress",
    ):
        module = importlib.import_module(name)
        path = Path(module.__file__).relative_to(ROOT).as_posix()
        source = subprocess.check_output(
            ["git", "show", gate_revision + ":" + path], encoding="utf-8"
        )
        exec(compile(source, path, "exec"), module.__dict__)
    from app.filters.html_filter import HtmlFilter
    from app.output_safety import csv_original_cell, csv_safe_cell, safe_http_url

    output_csv.csv_safe_cell = csv_safe_cell
    round_one[int(os.environ["OUTPUT_ESCAPING_ROUND1_BASE"]) - 1]()
    raise AssertionError("gate reproduction unexpectedly passed on HEAD")


def test_text_quotes_browser_desync():
    """#672 gate: raw quotes in text are safe only where this scanner and a browser agree.

    Each page below parses differently in Chrome: the browser is INSIDE a quoted
    attribute value where the scanner sees the text "Hello there." (an abrupt <!--> or
    <!---> comment, --!>, `</ x>`, a raw-text <xmp>/<textarea>/<noscript>, a quote after
    `= =`, <style> inside SVG, a CDATA or bogus comment holding `>`). Confirmed in real
    Chrome (DOMParser): a raw `"` in the translation closed the value and added onfocus=.
    Such pages keep quotes escaped; an ordinary page writes them raw."""
    filt = HtmlFilter()
    attack = "Z\" onfocus=alert(1) w='"
    pages = (
        '<!--><a title="--> Hello there. <b title="x">bold</b>',
        '<!---><a title="--> Hello there. <b title="x">bold</b>',
        '<!-- a --!><a title="--> Hello there. <b title="x">bold</b>',
        '<p></ x <a title="> <b title=">Hello there. <i title="y">i</i></p>',
        '<xmp><a title="</xmp><a title=">Hello there. <b title="x">b</b>',
        '<textarea><a title="</textarea><a title=">Hello there. <b title="x">b</b>',
        '<noscript><a title="</noscript><a title=">Hello there. <b title="x">b</b>',
        '<a title= = "x><b title=">Hello there. <i title="y">i</i>',
        '<svg><style><a title="</style>Hello there. <b title="x">b</b></svg>',
        '<![CDATA[><a title="]]>Hello there. <b title="x">b</b>',
        '<?x ><a title="?>Hello there. <b title="x">b</b>',
    )
    # #676: these two pages cannot be rebuilt the way a browser reads the source (the text
    # lands in an SVG <style>, or escaping a stray `</ x` turns a bogus comment back into a
    # tag), so ANY translation of them, benign too, is refused instead of exported
    refused_pages = {pages[3], pages[8]}
    refused = set()
    for page in pages:
        units = filt.parse(page.encode()).units
        hits = [u for u in units if "Hello there." in u.text]
        if not hits:
            continue  # the scanner keeps that text as markup: nothing is written there
        for value in (attack, "Bonjour."):
            mapping = {str(u.p_index): u.text.replace("Hello there.", value) for u in hits}
            try:
                out = filt.build(page.encode(), mapping).decode()
            except UnsafeMarkupRefused:
                refused.add(page)
                continue
            if value == attack:
                assert attack not in out and "Z&quot; onfocus=alert(1) w=&#x27;" in out, out
    assert refused == refused_pages, refused ^ refused_pages
    ok = b'<p>Say "hi", don\'t.</p><noscript><img src=x></noscript><title>T</title>'
    units = filt.parse(ok).units
    assert filt.build(ok, {str(u.p_index): u.text for u in units}) == ok
    print("  text quotes escaped on browser-desync pages, raw on ordinary pages")


round_two = (
    test_script_comparison,
    test_script_list,
    test_script_double_escaped,
    test_tokenizer_linear,
    test_source_hubl_masks,
)
if os.environ.get("OUTPUT_ESCAPING_ROUND2_CASE"):
    round_two[int(os.environ["OUTPUT_ESCAPING_ROUND2_CASE"]) - 1]()
    print("ROUND TWO REGRESSION OK")
    raise SystemExit(0)

for test in (
    *round_two,
    test_scanner_boundaries,
    *round_one,
    test_stored_html_indices,
    test_html,
    test_text_quotes_browser_desync,
    test_xlsx,
    test_csv,
    test_urls,
    test_writer_ratchet,
):
    test()
    print(test.__name__ + " passed")
print("EXPORT OUTPUT ESCAPING OK")
