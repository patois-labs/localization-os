"""F7 final-review regressions; hostile work is stopped at library/render sentinels.

Fixtures adapt %TEMP%/f7-final/{xlsx_probe,normal,docx_mem,pdf_probe}.py.
Run only through tests/test_verify_suites.py. Each case has a 30-second hard deadline.
"""

import asyncio
import datetime
import io
import os
import re
import sys
import threading
import time
import zipfile
from unittest.mock import patch

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from tests import verify_request_limits as fixtures

# Importing the existing suite establishes its offline, throwaway environment first.
_Override, _docx, _pdf, _xlsx = fixtures._Override, fixtures._docx, fixtures._pdf, fixtures._xlsx

from app.filters.base import InputLimitError
from app.filters.xlsx_filter import XlsxFilter


def rewrite(blob, change):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(blob)) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, change(info.filename, src.read(info)))
    return out.getvalue()


def linked_book():
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.worksheet.table import Table

    wb = Workbook()
    ws = wb.active
    ws.append(["Source", "Target"])
    ws.append(["Hello", "World"])
    for coord in ("A1", "B2"):
        ws[coord].hyperlink = "https://example.invalid/"
        ws[coord].comment = Comment("Reviewer note", "QA")
    ws.add_table(Table(displayName="Words", ref="A1:B2"))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def threaded_book():
    from tests.verify_request_limits import _with_members

    def change(name, xml):
        if name == "xl/worksheets/_rels/sheet1.xml.rels":
            return xml.replace(
                b"</Relationships>",
                b'<Relationship Type="http://schemas.microsoft.com/office/2017/10/'
                b'relationships/threadedComment" Target="../threadedComments/thread1.xml" '
                b'Id="thread1"/></Relationships>',
            )
        if name == "xl/_rels/workbook.xml.rels":
            return xml.replace(
                b"</Relationships>",
                b'<Relationship Type="http://schemas.microsoft.com/office/2017/10/'
                b'relationships/person" Target="persons/person.xml" Id="person1"/>'
                b"</Relationships>",
            )
        if name == "[Content_Types].xml":
            return xml.replace(
                b"</Types>",
                b'<Override PartName="/xl/threadedComments/thread1.xml" '
                b'ContentType="application/vnd.ms-excel.threadedcomments+xml"/>'
                b'<Override PartName="/xl/persons/person.xml" '
                b'ContentType="application/vnd.ms-excel.person+xml"/></Types>',
            )
        return xml

    blob = rewrite(linked_book(), change)
    return _with_members(
        blob,
        {
            "xl/threadedComments/thread1.xml": (
                b'<ThreadedComments xmlns="http://schemas.microsoft.com/office/'
                b'spreadsheetml/2018/threadedcomments"><threadedComment ref="B2" '
                b'dT="2026-01-01T00:00:00Z" personId="{11111111-1111-1111-1111-111111111111}" '
                b'id="{22222222-2222-2222-2222-222222222222}"><text>Review</text>'
                b"</threadedComment></ThreadedComments>"
            ),
            "xl/persons/person.xml": (
                b'<personList xmlns="http://schemas.microsoft.com/office/'
                b'spreadsheetml/2018/threadedcomments"><person displayName="QA" '
                b'id="{11111111-1111-1111-1111-111111111111}" userId="qa@example.invalid" '
                b'providerId="None"/></personList>'
            ),
        },
    )


def assert_refused(action):
    start = time.perf_counter()
    try:
        action()
    except InputLimitError:
        assert time.perf_counter() - start < 3
        return
    raise AssertionError("missing InputLimitError")


def regate_part(blob, kind, ref="A1", encoding=None):
    """Reuse the mutations in %TEMP%/f7-regate/bypass.py and e2e_sjis.py."""

    def change(name, xml):
        selected = (
            "comments/comment" in name
            if kind == "comment"
            else name == "xl/tables/table1.xml"
            if kind == "table"
            else name == "docProps/app.xml"
        )
        if not selected:
            return xml
        if kind == "table":
            xml = (
                f'<foo xmlns="urn:x" id="1" name="Words" displayName="Words" '
                f'ref="{ref}" headerRowCount="1"/>'
            ).encode()
        elif kind == "comment":
            xml = xml.replace(b'ref="A1"', f'ref="{ref}"'.encode())
            xml = xml.replace(
                b'<comments xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"',
                b'<foo xmlns="urn:x"',
            ).replace(b"</comments>", b"</foo>")
        if encoding:
            xml = re.sub(rb"<\?xml[^>]*\?>", b"", xml)
            text = xml.decode().replace("Reviewer note", "\u65e5\u672c\u8a9e")
            xml = (f'<?xml version="1.0" encoding="{encoding}"?>' + text).encode(encoding)
        return xml

    return rewrite(blob, change)


def test_regate_relationship_parts():
    failures = []
    for kind, ref, encoding in (
        ("comment", "A1:J100000", None),
        ("comment", "XFD1048576", None),
        ("table", "A1:CV100000", None),
        ("comment", "A1:J100000", "UTF-32"),
        ("comment", "XFD1048576", "Shift_JIS"),
        ("table", "A1:CV100000", "Shift_JIS"),
    ):
        candidate = regate_part(linked_book(), kind, ref, encoding)
        start = time.perf_counter()
        try:
            with patch(
                "app.filters.xlsx_filter.load_workbook",
                side_effect=AssertionError("unsafe related part reached openpyxl"),
            ) as loader:
                assert_refused(lambda candidate=candidate: XlsxFilter().parse(candidate))
                assert_refused(lambda candidate=candidate: XlsxFilter().build(candidate, {}))
                loader.assert_not_called()
        except Exception as exc:
            failures.append(f"{kind} {ref} {encoding}: {exc}")
        print(f"regate {kind} {ref} {encoding}: {time.perf_counter() - start:.3f}s")
    assert not failures, failures


def test_regate_benign_encodings():
    from openpyxl import load_workbook

    for kind, encoding in (
        ("app", "Shift_JIS"),
        ("comment", "Shift_JIS"),
        ("comment", "UTF-32"),
        ("comment", "UTF-16"),
        ("comment", "UTF-32BE"),
    ):
        blob = regate_part(linked_book(), kind, encoding=encoding)
        load_workbook(io.BytesIO(blob)).close()
        assert XlsxFilter().parse(blob).units
        assert XlsxFilter().build(blob, {})


def test_regate_relative_targets():
    from app.filters import xlsx_guard

    for ref in ("A1", "XFD1048576"):
        source = regate_part(linked_book(), "comment", ref, "Shift_JIS")
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(source)) as src, zipfile.ZipFile(out, "w") as dst:
            for info in src.infolist():
                name, data = info.filename, src.read(info)
                if name == "xl/comments/comment1.xml":
                    name = "xl/notes/renamed.part"
                elif name == "xl/worksheets/_rels/sheet1.xml.rels":
                    data = data.replace(b"/xl/comments/comment1.xml", b"../notes/renamed.part")
                dst.writestr(name, data)
        if ref == "A1":
            assert XlsxFilter().parse(out.getvalue()).units
        else:
            assert_refused(lambda out=out: xlsx_guard.check(out.getvalue()))
    hostile = rewrite(
        threaded_book(),
        lambda n, b: b.replace(b'ref="B2"', b'ref="XFD1048576"') if "threadedComments" in n else b,
    )
    assert_refused(lambda: xlsx_guard.check(hostile))


def test_related_part_invalid_value():
    blob = rewrite(
        linked_book(),
        lambda n, b: b.replace(b'id="1"', b'id="invalid"') if n == "xl/tables/table1.xml" else b,
    )
    assert_refused(lambda: XlsxFilter().parse(blob))
    assert_refused(lambda: XlsxFilter().build(blob, {}))


def test_relationship_dispatch_aliases():
    from app.filters import xlsx_guard

    hostile = regate_part(linked_book(), "comment", "XFD1048576")
    failures = []
    for kind in (
        "relationship",
        "relationship_type",
        "sheet",
        "sheet_id",
        "sheet_child_id",
        "comment",
        "comment_child_ref",
    ):

        def change(name, xml, kind=kind):
            if kind == "relationship" and name.endswith(".rels"):
                return xml.replace(b"<Relationship ", b"<other ")
            if kind == "relationship_type" and name.endswith("sheet1.xml.rels"):
                return xml.replace(
                    b'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    b'relationships/comments"',
                    b'type="comments"',
                )
            if kind == "sheet" and name == "xl/workbook.xml":
                return xml.replace(b"<sheet ", b"<other ")
            if kind == "sheet_id" and name == "xl/workbook.xml":
                return xml.replace(b"r:id=", b"id=")
            if kind == "sheet_child_id" and name == "xl/workbook.xml":
                return re.sub(
                    rb"(<sheet\b[^>]*)/>",
                    rb"\1><id>rId1</id></sheet>",
                    xml.replace(b' r:id="rId1"', b""),
                )
            if kind == "comment_child_ref" and "comments/comment" in name:
                return xml.replace(b'ref="XFD1048576"', b'ref="A1"').replace(
                    b"</comment>", b"<ref>XFD1048576</ref></comment>"
                )
            if kind == "comment" and "comments/comment" in name:
                return xml.replace(b"<comment ", b"<other ").replace(b"</comment>", b"</other>")
            return xml

        candidate = rewrite(hostile, change)
        try:
            assert_refused(lambda candidate=candidate: xlsx_guard.check(candidate))
        except Exception as exc:
            failures.append(f"{kind}: {exc}")
    assert not failures, failures


def test_materialising_refs():
    blob = linked_book()
    failures = []
    cases = [
        (kind, ref)
        for kind in ("hyperlink", "comment", "table")
        for ref in ("A1:J100000", "Z200000", "XFD1048576")
    ]
    # Child namespaces are ignored by openpyxl's Serialisable.from_tree.
    cases += [("hyperlink_alias", "XFD1048576"), ("comment_alias", "XFD1048576")]
    for kind, ref in cases:

        def change(name, xml, kind=kind, ref=ref):
            tag = kind.split("_")[0]
            if (
                (tag == "hyperlink" and name == "xl/worksheets/sheet1.xml")
                or (tag == "comment" and "comments/comment" in name)
                or (tag == "table" and name == "xl/tables/table1.xml")
            ):
                xml = re.sub(
                    rb"(<" + tag.encode() + rb'\b[^>]*\bref=")[^"]+',
                    lambda m: m[1] + ref.encode(),
                    xml,
                )
                if kind.endswith("_alias"):
                    xml = xml.replace(
                        b"<" + tag.encode() + b" ", b"<" + tag.encode() + b' xmlns="urn:other" '
                    )
                if tag == "table":
                    xml = re.sub(rb"<tableColumns\b.*?</tableColumns>", b"", xml)
            return xml

        candidate = rewrite(blob, change)
        start = time.perf_counter()
        try:
            with patch(
                "app.filters.xlsx_filter.load_workbook",
                side_effect=AssertionError("unsafe ref reached openpyxl"),
            ) as loader:
                assert_refused(lambda candidate=candidate: XlsxFilter().parse(candidate))
                assert_refused(lambda candidate=candidate: XlsxFilter().build(candidate, {}))
                loader.assert_not_called()
        except Exception as exc:
            failures.append(f"{kind} {ref}: {exc}")
        print(f"ref {kind} {ref}: {time.perf_counter() - start:.3f}s")
    assert not failures, failures


def test_cumulative_refs():
    from app.filters import xlsx_guard

    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    for tag, container in (
        ("hyperlink", "hyperlinks"),
        ("comment", "commentList"),
        ("table", "tables"),
    ):
        xml = (
            f'<worksheet xmlns="{ns}"><{container}>'
            + f'<{tag} ref="A1:D4"/>' * 3
            + f"</{container}></worksheet>"
        ).encode()
        candidate = rewrite(_xlsx(), lambda n, b, xml=xml: xml if n.endswith("sheet1.xml") else b)
        with _Override(XLSX_MAX_CELLS=40):
            assert_refused(lambda candidate=candidate: xlsx_guard.check(candidate))


def test_related_part_rectangle_union():
    from app.filters import xlsx_guard

    # The sheet is wide; the comment part is tall. Each separate rectangle fits,
    # but loading the comment expands this sheet to 26 * 100 = 2600 cells.
    blob = rewrite(
        linked_book(),
        lambda n, b: (
            b.replace(b'ref="A1"', b'ref="A100"')
            if "comments/comment" in n
            else b.replace(b'r="B2"', b'r="Z2"')
            if n.endswith("sheet1.xml")
            else b
        ),
    )
    with _Override(XLSX_MAX_CELLS=500):
        assert_refused(lambda: xlsx_guard.check(blob))


def test_normal_links_comments_identical():
    from app.filters import xlsx_guard

    class FixedDatetime(datetime.datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2026, 1, 1, tzinfo=tz)

    blob = threaded_book()
    filt = XlsxFilter()
    with (
        patch("openpyxl.writer.excel.datetime.datetime", FixedDatetime),
        patch("zipfile.time.localtime", return_value=(2026, 1, 1, 0, 0, 0, 3, 1, -1)),
    ):
        with patch.object(xlsx_guard, "check"):
            expected = filt.parse(blob)
            texts = {str(u.p_index): u.text + " translated" for u in expected.units}
            expected_bytes = filt.build(blob, texts)
        assert filt.parse(blob) == expected
        assert filt.build(blob, texts) == expected_bytes


def external_cache_book(count):
    """Reuse the package and duplicate-cell layout of f7-regate2/extlink.py."""
    from tests.verify_request_limits import _with_members

    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    def change(name, xml):
        if name == "xl/workbook.xml":
            return xml.replace(
                b"</sheets>",
                (
                    f'</sheets><externalReferences><externalReference xmlns:r="{rel}" '
                    'r:id="rIdX"/></externalReferences>'
                ).encode(),
            )
        if name == "xl/_rels/workbook.xml.rels":
            return xml.replace(
                b"</Relationships>",
                (
                    f'<Relationship Id="rIdX" Type="{rel}/externalLink" '
                    'Target="externalLinks/externalLink1.xml"/></Relationships>'
                ).encode(),
            )
        if name == "[Content_Types].xml":
            return xml.replace(
                b"</Types>",
                b'<Override PartName="/xl/externalLinks/externalLink1.xml" '
                b'ContentType="application/vnd.openxmlformats-officedocument.'
                b'spreadsheetml.externalLink+xml"/></Types>',
            )
        return xml

    rows = b"".join(
        f'<row r="{r}">'.encode()
        + f'<cell r="A{r}"><v>1</v></cell>'.encode() * min(20, count - (r - 1) * 20)
        + b"</row>"
        for r in range(1, (count + 19) // 20 + 1)
    )
    return _with_members(
        rewrite(_xlsx(), change),
        {
            "xl/externalLinks/externalLink1.xml": (
                f'<externalLink xmlns="{ns}" xmlns:r="{rel}"><externalBook r:id="rId1">'
                '<sheetNames><sheetName val="S"/></sheetNames><sheetDataSet>'
                '<sheetData sheetId="0">'
            ).encode()
            + rows
            + b"</sheetData></sheetDataSet></externalBook></externalLink>",
            "xl/externalLinks/_rels/externalLink1.xml.rels": (
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
                f'relationships"><Relationship Id="rId1" Type="{rel}/externalLinkPath" '
                'Target="other.xlsx" TargetMode="External"/></Relationships>'
            ).encode(),
        },
    )


def pivot_cache_book(count, sheets=1):
    from openpyxl import Workbook
    from openpyxl.pivot.cache import CacheDefinition, CacheField, CacheSource, WorksheetSource
    from openpyxl.pivot.fields import Text
    from openpyxl.pivot.record import Record, RecordList
    from openpyxl.pivot.table import Location, TableDefinition

    wb = Workbook()
    wb.active.append(["Source"])
    wb.active.append(["Hello"])
    for _ in range(sheets - 1):
        wb.create_sheet()
    pivot = TableDefinition(
        name="Words",
        cacheId=1,
        dataCaption="Values",
        location=Location(ref="C1:C2", firstHeaderRow=1, firstDataRow=1, firstDataCol=1),
    )
    pivot.cache = CacheDefinition(
        cacheSource=CacheSource(
            type="worksheet", worksheetSource=WorksheetSource(ref="A1:A2", sheet="Sheet")
        ),
        cacheFields=[CacheField(name="Source")],
    )
    pivot.cache.records = RecordList(r=[Record(_fields=[Text(v="Hello")])])
    wb.active.add_pivot(pivot)
    out = io.BytesIO()
    wb.save(out)
    return rewrite(
        out.getvalue(),
        lambda n, b: (
            re.sub(
                rb"<r>.*?</r>",
                b"".join(f'<r><s v="Hello {i}"/></r>'.encode() for i in range(count)),
                b,
            )
            if n.endswith("pivotCacheRecords1.xml")
            else b
        ),
    )


def _best_of(fn, samples=3):
    """Fastest of `samples` wall-clock runs: the least-disturbed estimate on a busy box."""
    best = None
    for _ in range(samples):
        start = time.perf_counter()
        fn()
        took = time.perf_counter() - start
        best = took if best is None else min(best, took)
    return best


def _refusal(blob):
    from app.filters import xlsx_guard

    def run():
        try:
            xlsx_guard.check(blob)
        except InputLimitError:
            return
        raise AssertionError("missing InputLimitError")

    return run


def _bare_cache_pass(blob):
    """The floor any guard must pay: one lxml pass with a do-nothing Python target over the
    same externalLink part (same parser, same per-event callbacks, no accounting)."""
    from lxml import etree

    class Noop:
        def start(self, name, attrs):
            pass

        def end(self, name):
            pass

        def data(self, text):
            pass

        def close(self):
            pass

    def run():
        with (
            zipfile.ZipFile(io.BytesIO(blob)) as zf,
            zf.open("xl/externalLinks/externalLink1.xml") as member,
        ):
            parser = etree.XMLParser(target=Noop(), resolve_entities=False, no_network=True)
            etree.parse(member, parser)

    return run


def test_external_cache_budget():
    """A hostile external-link cache is refused, and the refusal is CHEAP.

    "Cheap" used to be an absolute `< 3 s` for 600,000 cached cells at the default 500,000
    cap. That is wall-clock on whatever box runs the suite: the guard must count cells up to
    the cap through lxml Python callbacks, and a bare do-nothing lxml pass over 500,000
    cells alone costs about 1.5 s here, so the budget left ~2x headroom and a loaded machine
    measured 2.7-4.2 s with the guard unchanged. Cheapness is now proven RELATIVE to the
    same box at the same moment (fastest of several samples), which is what the budget
    was standing in for:
      * the refusal happens at the default cap (correctness, no timer);
      * the cost is set by the CAP, not by what the attacker ships: 10x the cached cells
        costs no more than ~the same refusal (early exit; no inflate-then-count);
      * the cost is a small constant multiple of merely reading the cells the cap admits
        (the guard measures ~2x a bare lxml pass; a per-cell regression breaches 5x).
    """
    from app import config

    _refusal(external_cache_book(config.XLSX_MAX_CELLS + 100000))()

    cap = 100000
    with _Override(XLSX_MAX_CELLS=cap):
        small = _best_of(_refusal(external_cache_book(cap + cap // 5)))
        large = _best_of(_refusal(external_cache_book(cap * 12)))
        floor = _best_of(_bare_cache_pass(external_cache_book(cap)))
    assert large <= 2 * small, (
        f"refusal cost grows with input beyond the cap: {large:.3f}s for 12x vs {small:.3f}s"
    )
    assert small <= 5 * floor, (
        f"refusal costs {small:.3f}s, over 5x a bare lxml pass over the capped cells ({floor:.3f}s)"
    )


def test_pivot_cache_budget():
    from app.filters import xlsx_guard

    blob = pivot_cache_book(260000, sheets=2)
    assert_refused(lambda: xlsx_guard.check(blob))


def test_cache_selection_and_text():
    from app.filters import xlsx_guard

    failures = []
    for kind, blob, cell_limit, text_limit in (
        ("external", external_cache_book(20), 25, 70),
        (
            "external_attribute",
            rewrite(
                external_cache_book(20),
                lambda n, b: (
                    b.replace(b"><v>1</v></cell>", b' v="1"/>')
                    if n.endswith("externalLink1.xml")
                    else b
                ),
            ),
            25,
            70,
        ),
        ("pivot", pivot_cache_book(20, sheets=2), 30, 200),
    ):
        # Workbook selection uses ids, irrespective of relationship type or filename.
        blob = rewrite(
            blob,
            lambda n, b: (
                b.replace(b'/externalLink"', b'/other"').replace(
                    b'/pivotCacheDefinition"', b'/other"'
                )
                if n.endswith(".rels")
                else b
            ),
        )
        for budget in (
            {"XLSX_MAX_CELLS": cell_limit},
            {"XLSX_MAX_TOTAL_TEXT_CHARS": text_limit},
        ):
            try:
                with _Override(**budget):
                    assert_refused(lambda blob=blob: xlsx_guard.check(blob))
            except AssertionError:
                failures.append(f"{kind} {budget}")
    assert not failures, failures


def test_small_caches_identical():
    from openpyxl import load_workbook

    for blob in (external_cache_book(20), pivot_cache_book(2, sheets=2)):
        # Reuse the existing byte-for-byte baseline comparison, including translated exports.
        with patch(__name__ + ".threaded_book", return_value=blob):
            test_normal_links_comments_identical()
        result = load_workbook(io.BytesIO(XlsxFilter().build(blob, {})))
        assert result._external_links or result.active._pivots
        result.close()


def test_float_row_numbers():
    from app.filters import xlsx_guard

    def book(value):
        return rewrite(
            _xlsx(),
            lambda n, b: (
                b.replace(b'<row r="1"', f'<row r="{value}"'.encode())
                if n.endswith("sheet1.xml")
                else b
            ),
        )

    blob = book("3.0")
    with patch(__name__ + ".threaded_book", return_value=blob):
        test_normal_links_comments_identical()
    for value in ("3.5", "0.0", "1048577.0", "nan", "inf"):
        assert_refused(lambda value=value: xlsx_guard.check(book(value)))


def test_ordinary_280k_cells():
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    for row in range(20000):
        ws.append([row * 14 + col for col in range(14)])
    out = io.BytesIO()
    wb.save(out)
    assert not XlsxFilter().parse(out.getvalue()).units


def test_unrelated_broken_xml():
    from tests.verify_request_limits import _with_members

    blob = _with_members(_xlsx(), {"customXml/item9.xml": b"<custom><broken"})
    assert XlsxFilter().parse(blob).units
    assert XlsxFilter().build(blob, {})


def test_spreadsheet_broken_xml_safe():
    blob = rewrite(_xlsx(), lambda n, b: b[:-20] if n.endswith("sheet1.xml") else b)
    assert_refused(lambda: XlsxFilter().parse(blob))
    assert_refused(lambda: XlsxFilter().build(blob, {}))


def test_limit_http_envelopes():
    from fastapi.testclient import TestClient
    from sqlmodel import Session, select

    from app.db import engine
    from app.main import app
    from app.models import Document, Segment
    from app.ports import blobstore
    from app.services import apikeys
    from tests.verify_request_limits import _auth, _project, _user

    _auth(False)
    client = TestClient(app, raise_server_exceptions=False)
    user, _ = _user()
    _, key = apikeys.create_key(user.id, "fix2-limits", scopes="api docs")
    headers = {"X-Loc-Token": key}
    pid = _project()
    blob = linked_book()
    hostile = rewrite(
        blob,
        lambda n, b: b.replace(b'ref="A1"', b'ref="XFD1048576"') if n.endswith("sheet1.xml") else b,
    )
    response = client.post(
        "/api/documents",
        headers=headers,
        data={"project_id": str(pid)},
        files={"file": ("hostile.xlsx", hostile)},
    )
    assert response.status_code == 413 and "limit" in response.json()["error"], response.text
    response = client.post(
        f"/projects/{pid}/documents", files={"file": ("safe.xlsx", blob)}, follow_redirects=False
    )
    doc_id = int(response.headers["location"].rsplit("/", 1)[-1].split("?")[0])
    with Session(engine) as session:
        for segment in session.exec(select(Segment).where(Segment.document_id == doc_id)):
            segment.target_text = "Translated text"
            segment.status = "mt"
            session.add(segment)
        session.commit()
    with Session(engine) as session:
        # #399: the upload blob key is the document's persisted identity, not "<id>.xlsx".
        blob_key = blobstore.document_key(session.get(Document, doc_id))
    broken = rewrite(blob, lambda n, b: b[:-20] if n.endswith("sheet1.xml") else b)
    # This BlobStore belongs to the suite's throwaway data directory.
    blobstore.get().put(blob_key, broken)
    for url in (
        f"/documents/{doc_id}/export?locale=de-DE",
        f"/api/documents/{doc_id}/download?locale=de-DE",
    ):
        response = client.get(url, headers=headers)
        assert response.status_code == 413 and "spreadsheet XML" in response.text, response.text
    failures = []
    for candidate in (
        regate_part(blob, "comment", "XFD1048576", "Shift_JIS"),
        regate_part(blob, "comment", "A1:J100000", "UTF-32"),
        rewrite(blob, lambda n, b: b[:-20] if "comments/comment" in n else b),
    ):
        with patch(
            "app.filters.xlsx_filter.load_workbook",
            side_effect=AssertionError("unsafe stored workbook reached openpyxl"),
        ) as loader:
            response = client.post(
                "/api/documents",
                headers=headers,
                data={"project_id": str(pid)},
                files={"file": ("hostile.xlsx", candidate)},
            )
            if response.status_code != 413:
                failures.append(f"upload: {response.status_code}")
            blobstore.get().put(blob_key, candidate)
            for url in (
                f"/documents/{doc_id}/export?locale=de-DE",
                f"/api/documents/{doc_id}/download?locale=de-DE",
            ):
                start = time.perf_counter()
                response = client.get(url, headers=headers)
                if response.status_code != 413:
                    failures.append(f"{url}: {response.status_code}")
                assert time.perf_counter() - start < 3
            loader.assert_not_called()
    invalid_value = rewrite(
        blob,
        lambda n, b: b.replace(b'id="1"', b'id="invalid"') if n == "xl/tables/table1.xml" else b,
    )
    blobstore.get().put(blob_key, invalid_value)
    for url in (
        f"/documents/{doc_id}/export?locale=de-DE",
        f"/api/documents/{doc_id}/download?locale=de-DE",
    ):
        response = client.get(url, headers=headers)
        if response.status_code != 413:
            failures.append(f"invalid table value {url}: {response.status_code}")
    blobstore.get().put(blob_key, regate_part(blob, "app", encoding="Shift_JIS"))
    for url in (
        f"/documents/{doc_id}/export?locale=de-DE",
        f"/api/documents/{doc_id}/download?locale=de-DE",
    ):
        response = client.get(url, headers=headers)
        if response.status_code != 200:
            failures.append(f"benign Shift_JIS {url}: {response.status_code}")
    for kind in ("app", "comment"):
        response = client.post(
            "/api/documents",
            headers=headers,
            data={"project_id": str(pid)},
            files={"file": ("safe.xlsx", regate_part(blob, kind, encoding="Shift_JIS"))},
        )
        if response.status_code not in (200, 201):
            failures.append(f"benign Shift_JIS {kind} upload: {response.status_code}")
    assert not failures, failures


def test_limit_http_envelopes_legacy_key():
    """Pre-#399 documents (NULL storage key) read "<id>.xlsx"; limits must still hold."""
    from fastapi.testclient import TestClient
    from sqlmodel import Session, select

    from app.db import engine
    from app.main import app
    from app.models import Document, Segment
    from app.ports import blobstore
    from app.services import apikeys
    from tests.verify_request_limits import _auth, _project, _user

    _auth(False)
    client = TestClient(app, raise_server_exceptions=False)
    user, _ = _user()
    _, key = apikeys.create_key(user.id, "fix2-limits-legacy", scopes="api docs")
    headers = {"X-Loc-Token": key}
    pid = _project()
    blob = linked_book()
    response = client.post(
        f"/projects/{pid}/documents", files={"file": ("legacy.xlsx", blob)}, follow_redirects=False
    )
    doc_id = int(response.headers["location"].rsplit("/", 1)[-1].split("?")[0])
    with Session(engine) as session:
        for segment in session.exec(select(Segment).where(Segment.document_id == doc_id)):
            segment.target_text = "Translated text"
            segment.status = "mt"
            session.add(segment)
        doc = session.get(Document, doc_id)
        doc.upload_blob_key = None  # make it a pre-#399 row
        session.add(doc)
        session.commit()
    with Session(engine) as session:
        blob_key = blobstore.document_key(session.get(Document, doc_id))
    assert blob_key == f"{doc_id}.xlsx", blob_key
    urls = (
        f"/documents/{doc_id}/export?locale=de-DE",
        f"/api/documents/{doc_id}/download?locale=de-DE",
    )
    failures = []
    broken = rewrite(blob, lambda n, b: b[:-20] if n.endswith("sheet1.xml") else b)
    blobstore.get().put(blob_key, broken)
    for url in urls:
        response = client.get(url, headers=headers)
        if not (response.status_code == 413 and "spreadsheet XML" in response.text):
            failures.append(f"legacy broken sheet {url}: {response.status_code}")
    for candidate in (
        regate_part(blob, "comment", "XFD1048576", "Shift_JIS"),
        regate_part(blob, "comment", "A1:J100000", "UTF-32"),
    ):
        with patch(
            "app.filters.xlsx_filter.load_workbook",
            side_effect=AssertionError("unsafe stored workbook reached openpyxl"),
        ) as loader:
            blobstore.get().put(blob_key, candidate)
            for url in urls:
                start = time.perf_counter()
                response = client.get(url, headers=headers)
                if response.status_code != 413:
                    failures.append(f"legacy hostile {url}: {response.status_code}")
                assert time.perf_counter() - start < 3
            loader.assert_not_called()
    blobstore.get().put(blob_key, regate_part(blob, "app", encoding="Shift_JIS"))
    for url in urls:
        response = client.get(url, headers=headers)
        if response.status_code != 200:
            failures.append(f"legacy benign Shift_JIS {url}: {response.status_code}")
    assert not failures, failures


def test_ordinary_docx_manual():
    from docx import Document

    from app.filters.docx_filter import DocxFilter

    doc = Document()
    for i in range(12000):
        doc.add_paragraph(f"Manual paragraph {i}." if i % 5 else "")
    out = io.BytesIO()
    doc.save(out)
    assert len(DocxFilter().parse(out.getvalue()).units) == 9600


def test_ordinary_docx_glossary():
    from docx import Document

    from app.filters.docx_filter import DocxFilter

    doc = Document()
    doc.add_paragraph("Glossary")
    table = doc.add_table(rows=2000, cols=5)
    for cell in table._cells:
        cell.text = "Term"
    out = io.BytesIO()
    doc.save(out)
    assert len(DocxFilter().parse(out.getvalue()).units) == 10001


def test_docx_paragraph_budget():
    from docx import Document

    from app.filters import docx_filter

    doc = Document()
    doc.add_paragraph("Hello")
    doc.add_paragraph("")
    doc.add_table(rows=1, cols=2)
    buf = io.BytesIO()
    doc.save(buf)
    with (
        _Override(DOCX_MAX_PARAGRAPHS=3),
        patch.object(
            docx_filter,
            "_all_paragraphs",
            side_effect=AssertionError("paragraph expansion reached"),
        ),
    ):
        assert_refused(lambda: docx_filter.DocxFilter().parse(buf.getvalue()))
    with _Override(DOCX_MAX_PARAGRAPHS=4):
        assert len(docx_filter.DocxFilter().parse(buf.getvalue()).units) == 1
    assert docx_filter.DocxFilter().parse(_docx()).units


def test_pdf_total_pixels():
    import pypdfium2 as pdfium

    from app.filters.pdf_filter import PdfFilter

    blob = _pdf([(100, 100)] * 3)
    # At 150 dpi each page is ceil(100 * 150 / 72)^2 = 43681 pixels.
    with (
        _Override(PDF_OCR_MAX_TOTAL_PIXELS=3 * 43681 - 1),
        patch.object(
            pdfium.PdfPage, "render", side_effect=AssertionError("over-budget PDF rendered")
        ),
    ):
        assert_refused(lambda: PdfFilter().parse(blob))
    with (
        _Override(PDF_OCR_MAX_TOTAL_PIXELS=3 * 43681),
        patch("app.filters.pdf_filter.ocr.ocr_image", return_value=""),
    ):
        assert len(PdfFilter().parse(blob).skeleton["pages"]) == 3


def test_upload_busy_wording():
    from app.web import deps
    from app.web.streaming import stream_slot

    async def exercise():
        async with stream_slot() as admitted:
            assert admitted
            try:
                await deps.offload_parse(lambda: None)
            except deps.WorkerBusy as exc:
                assert "stream" not in str(exc).lower(), str(exc)
            else:
                raise AssertionError("busy worker admitted")
        assert b"stream" not in deps.busy_response().body.lower()

    with _Override(STREAM_WORKERS=1, STREAM_WAIT_S=0.01):
        asyncio.run(exercise())


def main():
    failures = []
    for test in (
        test_external_cache_budget,
        test_pivot_cache_budget,
        test_cache_selection_and_text,
        test_small_caches_identical,
        test_float_row_numbers,
        test_regate_relationship_parts,
        test_regate_benign_encodings,
        test_regate_relative_targets,
        test_related_part_invalid_value,
        test_relationship_dispatch_aliases,
        test_materialising_refs,
        test_cumulative_refs,
        test_related_part_rectangle_union,
        test_normal_links_comments_identical,
        test_ordinary_280k_cells,
        test_unrelated_broken_xml,
        test_spreadsheet_broken_xml_safe,
        test_limit_http_envelopes,
        test_limit_http_envelopes_legacy_key,
        test_docx_paragraph_budget,
        test_ordinary_docx_manual,
        test_ordinary_docx_glossary,
        test_pdf_total_pixels,
        test_upload_busy_wording,
    ):
        timer = threading.Timer(30, lambda: os._exit(124))
        timer.daemon = True
        timer.start()
        start = time.perf_counter()
        try:
            test()
            status = "PASS"
        except Exception as exc:
            status = f"FAIL {type(exc).__name__}: {exc}"
            if exc.__cause__:
                status += f" (cause: {exc.__cause__!r})"
            failures.append(test.__name__)
        finally:
            timer.cancel()
        line = f"{test.__name__}: {status} ({time.perf_counter() - start:.3f}s)"
        line = line.encode("ascii", "backslashreplace").decode("ascii")
        print(line, flush=True)
        if path := os.environ.get("F7_FIX2_EVIDENCE"):
            with open(path, "a", encoding="ascii") as log:
                log.write(line + "\n")
    assert not failures, failures
    print("REQUEST LIMITS FIX2 OK")


if __name__ == "__main__":
    main()
