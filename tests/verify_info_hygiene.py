"""Offline regressions for the remaining #270 information-hygiene findings."""

import asyncio
import importlib.util
import io
import json
import subprocess
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from tests._bootstrap import init_env

init_env("loctms_info_hygiene_")

from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app import config
from app.db import engine, init_db
from app.main import app
from app.models import Customer, Document, Project, Role, User
from app.services import apikeys, audit, pii, security, settings
from app.services import security_client_ip_one_authority_ as clock


def source(path):
    return (ROOT / path).read_text(encoding="utf-8")


def test_bound_upload():
    settings.set_many({"security_auth_required": "1"})
    with Session(engine) as session:
        customer = Customer(name="University", ai_opt_out=True)
        session.add(customer)
        session.commit()
        session.refresh(customer)
        user = User(email="clerk@example.test", role=Role.MANAGER, customer_id=customer.id)
        project = Project(name="University docs", source_locale="en", customer_id=customer.id)
        other = Project(name="Other", source_locale="en")
        session.add_all([user, project, other])
        session.commit()
        uid, pid, other_id = user.id, project.id, other.id
    _, key = apikeys.create_key(uid, "upload test", scopes="docs")
    client = TestClient(app)
    headers = {"X-Loc-Token": key}
    file = {"file": ("grades.txt", b"Private student grades.", "text/plain")}
    response = client.post("/api/documents", headers=headers, files=file)
    assert response.status_code == 422 and "project" in response.json()["error"]
    with Session(engine) as session:
        assert not session.exec(select(Document)).all()
    response = client.post(
        "/api/documents", headers=headers, files=file, data={"project_id": str(other_id)}
    )
    assert response.status_code == 404
    response = client.post(
        "/api/documents", headers=headers, files=file, data={"project_id": str(pid)}
    )
    assert response.status_code == 200, response.text


def test_pii_structured():
    settings.set_many({"pii_ner_enabled": "0"})
    for value in (
        "2001:db8::1",
        "::1",
        "867-5309",
        "Student ID 20231234",
        "DE89 3704 0044 0532 0130 00",
    ):
        redacted, mapping = pii.redact(value)
        assert mapping and value not in redacted, (value, redacted)
        assert value in mapping.values(), mapping
        assert pii.restore(redacted, mapping) == value
    with patch.object(config, "PII_INSTITUTION_ID_DIGITS", 6):
        assert pii.redact("Student ID 123456")[1]
    assert pii.redact("Meeting at 12:34:56")[0] == "Meeting at 12:34:56"


def test_background_clock():
    settings.set_many({"session_idle_minutes": "1", "session_max_hours": "24"})
    request = SimpleNamespace(
        session={"uid": 1, "iat": 1000, "la": 1000}, headers={"x-loc-background": "1"}
    )
    with patch.object(clock, "_now", return_value=1059):
        assert security.enforce_session(request) == ""
        assert request.session["la"] == 1000
    with patch.object(clock, "_now", return_value=1061):
        assert security.enforce_session(request) == "idle"
        assert not request.session
    for path in (
        "app/static/jobs.js",
        "app/static/queuedrun.js",
        "app/static/sse.js",
        "app/templates/batch_view.html",
        "app/templates/export_run.html",
        "app/templates/run_view.html",
    ):
        assert "'X-Loc-Background': '1'" in source(path), path
    assert "new EventSource(" not in source("app/static/sse.js")
    subprocess.run(
        ["node", "tests/_info_browser.js"], cwd=ROOT, check=True, capture_output=True, timeout=20
    )


def test_webhook_oracle():
    from app.web import api_part2

    client = TestClient(app)
    with (
        patch.object(api_part2.conn_store, "get", return_value=SimpleNamespace(kind="disabled")),
        patch.object(api_part2.connectors_registry, "kind_enabled", return_value=False),
    ):
        disabled = client.post("/webhooks/connectors/99999")
    missing = client.post("/webhooks/connectors/99999")
    assert disabled.status_code == missing.status_code == 401
    assert disabled.json() == missing.json()


def test_redirects_local():
    from starlette.requests import Request

    from app.web import connectors_common, projects_projects

    def http_request(referer):
        # A complete HTTP scope, as an ASGI server builds it (path, scheme, server, Host):
        # referer_local_path (#501) compares the Referer host to request.url, which needs them.
        return Request(
            {
                "type": "http",
                "method": "POST",
                "scheme": "http",
                "path": "/connectors/manage/test",
                "query_string": b"",
                "server": ("testserver", 80),
                "headers": [(b"host", b"testserver"), (b"referer", referer)],
                "session": {},
            }
        )

    with (
        patch.object(connectors_common, "_admin_block", return_value=(None, None)),
        patch.dict(connectors_common.connectors_registry.REGISTRY, {"test": object()}),
        patch.object(connectors_common.connectors_registry, "kind_enabled", return_value=True),
    ):
        fn = connectors_common.connector_kind_toggle
        # A foreign, protocol-relative, backslash or non-http Referer never leaves the site.
        for referer in (
            b"https://evil.test",
            b"https://evil.test/connectors/catalog",
            b"//evil.test/connectors/catalog",
            b"/\\evil.test",
            b"javascript:alert(1)",
        ):
            location = fn("test", http_request(referer), "1").headers["location"]
            assert location == "/connectors/manage", (referer, location)
        # A same-host Referer comes back as its local path only, never an absolute URL.
        location = fn(
            "test", http_request(b"http://testserver/connectors/catalog?x=1"), "1"
        ).headers["location"]
        assert location == "/connectors/catalog?x=1", location
    with (
        patch.object(projects_projects, "_my_org_id", return_value=None),
        patch("app.kernel.context_pkg.get_profile", return_value=None),
    ):
        request = http_request(b"https://evil.test")
        response = projects_projects.context_card_delete(99999, request, "https://evil.test", None)
        assert response.headers["location"] == "/"


def test_companion_encryption():
    from app.web.deps import connector_api_token

    settings.set_many({"connector_api_token": ""})
    plain = connector_api_token()
    assert settings.get("connector_api_token") != plain
    assert settings.read_secret_setting("connector_api_token") == plain
    assert connector_api_token() == plain
    assert security._genuine_display_prefix(plain) == "lsk_"
    assert security._genuine_display_prefix(settings.get("connector_api_token")) == ""
    with patch.object(settings, "get", return_value="damaged ciphertext"):
        assert settings.read_secret_setting("connector_api_token") == ""


def test_anonymous_audit_budget():
    with (
        patch.object(audit.time, "monotonic", return_value=1000),
        patch.object(audit, "_refusal_window", 0),
        patch.object(audit, "_refusal_count", 0),
    ):
        assert all(audit.allow_refusal_record() for _ in range(60))
        assert not audit.allow_refusal_record()
        with patch.object(audit.time, "monotonic", return_value=1060):
            assert audit.allow_refusal_record()
    assert "audit.allow_refusal_record()" in source("app/main.py")
    assert "linkage only" in audit.iter_export.__doc__


def test_docx_warning():
    from docx import Document as WordDocument
    from docx.oxml import parse_xml

    from app.filters.docx_filter import DocxFilter

    document = WordDocument()
    document.add_paragraph("Body")
    document.sections[0].header.paragraphs[0].text = "Header"
    document.sections[0].footer.paragraphs[0].text = "Footer"
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Nested"
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    document.element.body.append(
        parse_xml(
            f"<w:sdt {ns}><w:sdtContent><w:p><w:r><w:t>TOC</w:t></w:r></w:p></w:sdtContent></w:sdt>"
        )
    )
    data = io.BytesIO()
    document.save(data)
    res = DocxFilter().parse(data.getvalue())
    text = json.dumps(res.skeleton.get("file_warnings", []))
    # #355: the body's nested table and block content control ARE translated now, so the
    # warning must not call them untranslated. #270: nor the header and footer.
    assert {"Header", "Footer", "Nested", "TOC"} <= {u.text for u in res.units}
    assert "untranslated" not in text.lower(), text
    assert not any(s in text for s in ("header", "footer", "nested tables", "content controls"))
    # #686 (walk 6): a text content control INSIDE a paragraph is translated: not listed
    document.element.body.append(
        parse_xml(
            f"<w:p {ns}><w:r><w:t>Lead </w:t></w:r><w:sdt><w:sdtContent><w:r><w:t>inline"
            "</w:t></w:r></w:sdtContent></w:sdt></w:p>"
        )
    )
    data = io.BytesIO()
    document.save(data)
    res = DocxFilter().parse(data.getvalue())
    assert any("inline" in u.text for u in res.units), [u.text for u in res.units]
    assert "content controls" not in json.dumps(res.skeleton.get("file_warnings", []))
    # a drop-down list control inside a paragraph is still carried untranslated: listed
    document.element.body.append(
        parse_xml(
            f"<w:p {ns}><w:r><w:t>Lead </w:t></w:r><w:sdt><w:sdtPr><w:dropDownList/></w:sdtPr>"
            "<w:sdtContent><w:r><w:t>choice</w:t></w:r></w:sdtContent></w:sdt></w:p>"
        )
    )
    data = io.BytesIO()
    document.save(data)
    text = json.dumps(DocxFilter().parse(data.getvalue()).skeleton["file_warnings"])
    assert "content controls" in text and "nested tables" not in text, text


def test_network_policy():
    spec = importlib.util.spec_from_file_location(
        "network_policy", ROOT / "deploy/aws-hosted/network_policy.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    allowed = module.cidrs("173.245.48.0/20\n2400:cb00::/32")
    for bad in ("1.2.3.4/99", "999.2.3.0/24", "::/129", '1.2.3.0/24"}],"x":1', ""):
        try:
            module.cidrs(bad)
            raise AssertionError(bad)
        except ValueError:
            pass

    def groups(port, ranges, **extra):
        return {
            "SecurityGroups": [
                {
                    "IpPermissions": [
                        {
                            "IpProtocol": "tcp",
                            "FromPort": port,
                            "ToPort": port,
                            "IpRanges": [{"CidrIp": c} for c in ranges],
                            **extra,
                        }
                    ]
                }
            ]
        }

    assert not module.check_rules(groups(443, ["173.245.48.0/20"]), allowed)
    for value in (
        groups(443, ["0.0.0.0/1", "128.0.0.0/1"]),
        groups(8000, ["0.0.0.0/0"]),
        groups(443, ["1.2.3.4/32"]),
        groups(443, [], UserIdGroupPairs=[{"GroupId": "sg"}]),
    ):
        assert module.check_rules(value, allowed)
    for path in ("aws-verify.sh", "aws-instance-setup.sh"):
        text = source("deploy/aws-hosted/" + path)
        assert 'python3 "$policy"' in text


def test_runbook_hygiene():
    text = source("deploy/aws-hosted/LAUNCH.md")
    assert "install -m 600 /dev/null /srv/caddy/origin-key.pem" in text
    assert "unrecorded local terminal" in text
    assert "tailscale.com/install.sh | sh" not in text
    unit = source("deploy/aws-hosted/host/install-audit-alert.sh")
    assert "NoNewPrivileges=true" in unit and "ProtectSystem=strict" in unit


def test_pii_boundary_roundtrip():
    from app.services import pii_boundary

    settings.set_many({"pii_redaction": "1", "pii_ner_enabled": "0"})
    samples = {
        "EMAIL": "person@example.test",
        "SSN": "123-45-6789",
        "IBAN": "DE89 3704 0044 0532 0130 00",
        "CARD": "4111 1111 1111 1111",
        "PHONE": "867-5309",
        "IP": "192.0.2.1",
        "IPV6": "2001:db8::1",
        "STUDENT_ID": "Student ID 20231234",
    }
    assert set(pii.PATTERNS) | set(pii._extra_patterns()) <= samples.keys()
    seen = []

    def echo(text):
        seen.append(text)
        return text

    wrapped = pii_boundary.wrap_callable(echo, {"text"})
    for value in [
        *samples.values(),
        "::1",
        "2001:db8::1 and 2001:db8::2",
        " | ".join(samples.values()),
    ]:
        assert wrapped(value) == value, value
        assert seen[-1] != value and "PII_" in seen[-1], value
    settings.set_many({"pii_redaction": "0"})


def test_pii_false_positives():
    for value in (
        "x :: y",
        "a::b",
        "FY24 REVENUE GREW BY TEN PERCENT THIS YEAR",
        "Order AB12 CDEF GHIJ KLMN shipped",
        "Room 123.4567",
    ):
        assert pii.redact(value) == (value, {}), value
        assert not pii.detect(value), value


def test_webhook_large_oracle():
    from app.web import api_part2

    client = TestClient(app)
    for headers in ({}, {"content-length": "1"}):
        body = b"x" * (config.WEBHOOK_MAX_BODY_BYTES + 1)
        with (
            patch.object(
                api_part2.conn_store, "get", return_value=SimpleNamespace(kind="disabled")
            ),
            patch.object(api_part2.connectors_registry, "kind_enabled", return_value=False),
        ):
            disabled = client.post("/webhooks/connectors/99999", content=body, headers=headers)
        missing = client.post("/webhooks/connectors/99999", content=body, headers=headers)
        assert disabled.status_code == missing.status_code == 413
        assert disabled.json() == missing.json()


def test_completion_navigation():
    for path, statement in (
        ("app/static/jobs.js", "window.location.replace(u.toString())"),
        ("app/static/queuedrun.js", "window.location.reload()"),
        ("app/static/keysrun.js", "window.location.reload()"),
        ("app/templates/run_view.html", "window.location.href = '/runs/' + ev.run_id"),
        ("app/templates/batch_view.html", "BATCH_RELOAD_MAX = 1;"),
    ):
        assert statement in source(path), path


def test_refusal_suppression_summary():
    with (
        patch.object(audit.time, "monotonic", return_value=2000),
        patch.object(audit, "_refusal_window", 2000),
        patch.object(audit, "_refusal_count", 0),
        patch.object(audit, "_refusal_dropped", 0, create=True),
        patch.object(audit, "log") as record,
    ):
        for _ in range(63):
            audit.allow_refusal_record()
        record.assert_not_called()
        with patch.object(audit.time, "monotonic", return_value=2060):
            assert audit.allow_refusal_record()
            audit.allow_refusal_record()
        record.assert_called_once_with(action="request.refused.suppressed", detail="dropped=3")


def test_license_environment_scan():
    from email.message import Message

    from tests import verify_licenses

    metadata = Message()
    metadata["Name"] = "unexpected-transitive"
    metadata["License-Expression"] = "GPL-3.0-only"
    with patch.object(
        verify_licenses.im, "distributions", return_value=[SimpleNamespace(metadata=metadata)]
    ):
        count, violations, _ = verify_licenses._scan_copyleft()
        assert count == 1 and any("unexpected-transitive" in v for v in violations)
    assert "--require-hashes --no-deps -r requirements.lock" in source(".github/workflows/ci.yml")


def test_dev_audit_inventory():
    from scripts.audit_runtime_profiles import dev_tool_pins

    metadata = {"Name": "pip"}
    with patch(
        "importlib.metadata.distributions",
        return_value=[SimpleNamespace(metadata=metadata, version="1.0")],
    ):
        assert dev_tool_pins() == ["pip==1.0"]


def test_spotlight_text():
    text = source("app/static/onboarding.js")
    assert "node.textContent = value" in text
    for field in ("title", "description"):
        assert f"{field}: plainText({field})" in text
    assert "nextBtnText: plainText(btnText)" in text


def test_deploy_hygiene():
    from starlette.requests import Request

    from app.web.health import health

    response = asyncio.run(health(Request({"type": "http", "headers": []})))
    assert "detail" not in json.loads(response.body)
    from app.web.health import health_admin

    request = Request({"type": "http", "headers": [], "session": {"uid": 1}})
    with patch.object(security, "session_user", return_value=User(role=Role.VIEWER)):
        assert health_admin(request).status_code == 403
    with patch.object(security, "session_user", return_value=User(role=Role.ADMIN)):
        assert "detail" in json.loads(health_admin(request).body)
    entrypoint = source("docker-entrypoint.sh")
    assert "--limit-concurrency 256" in entrypoint and "--no-access-log" in entrypoint
    assert '>"$TS_STATE_DIR/tailscaled.log"' not in entrypoint


def test_mfa_rotation_guidance():
    text = source("deploy/aws-hosted/LAUNCH.md")
    assert "Before raising `mfa_required`" in text and "those passwords first" in text


def test_env_jev_host_pin():
    from app.services import engines, jev_inspector

    with (
        patch.object(jev_inspector, "_env_key", return_value="offline-fixture"),
        patch.object(config, "TYPESAFE_BASE_URL", engines.JEV_HOSTED_BASE_URL),
    ):
        profile, _ = jev_inspector.import_env()
    assert profile.base_url == engines.JEV_HOSTED_BASE_URL
    try:
        engines.update_profile(profile.id, base_url="https://attacker.invalid")
        raise AssertionError("Imported environment credential could be repointed")
    except ValueError:
        pass


def main():
    init_db()
    tests = [v for k, v in globals().items() if k.startswith("test_") and callable(v)]
    failures = []
    for test in tests:
        try:
            test()
            print(test.__name__ + ": PASS")
        except Exception as exc:
            print(test.__name__ + ": FAIL " + repr(exc))
            failures.append(test.__name__)
    assert not failures, failures
    print("INFO HYGIENE OK")


if __name__ == "__main__":
    main()
