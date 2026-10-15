"""Glossary adherence verification (SPEC-workflows-redesign Track B / B1).

Hermetic + offline (throwaway LOC_TMS_DATA_DIR, provider keys blanked -> no
network, MockProvider forced). Proves the POSITIVE glossary-adherence gate (the
sibling of the DNT gate) and the REQUIRE rulebook directive that makes term
adherence GATE-ELIGIBLE:

  1. quality.glossary_term_violations: a non-DNT required-term miss is detected;
     present -> clean; DNT entries skipped; empty target_term skipped; whole-token
     (no trivial substring pass); case_sensitive honoured; de-duplicated.
  2. commit_mt_output wiring (the sole write seam): a miss -> committed (text
     WRITTEN) + NEEDS_REVIEW + glossary_violation qa_flag; tm_eligible UNCHANGED
     (mirrors DNT); present -> byte-identical clean MT; composes with QE + DNT
     flags (none lost); the DNT path is unchanged; the refusal path is unchanged.
  3. Both translate paths end-to-end (a fake engine drops the required term): the
     MT path AND the segmented workflow write-back land NEEDS_REVIEW +
     glossary_violation, text still committed.
  4. lqa REQUIRE directive: parsing edge cases (blank, lowercase keyword,
     substring-vs-whole-token); FORBID stays byte-identical; score_text's gate
     flips to failed on a miss (a single MAJOR fails the gate though score>=90).
  5. REQUIRE drives the correction loop end-to-end: translate -> LQA(REQUIRE
     missing) -> fix{run_condition prev_lqa:failed} RUNS the fix step; the control
     (REQUIRE present -> LQA passes) SKIPS it.
  6. Parity: the glossary_violation flag renders in the editor qa cell.
  7. fail-before/pass-after: a git-show HEAD reconstruction proves the check is
     load-bearing (degrades to a git-independent behavioral delta once HEAD carries
     the change).

Final line: "GLOSSARY ADHERENCE OK".
"""

import os
import sys
from io import BytesIO

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env
from tests._piece_patch import set_pieces

init_env("loctms_glossadh_")

from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app.db import engine, init_db
from app.main import app
from app.models import (
    Document,
    GlossaryEntry,
    LQAResult,
    Project,
    RunStatus,
    Segment,
    SegStatus,
    StepType,
    TransMode,
    Workflow,
    WorkflowStep,
)
from app.services import lqa, orchestrator, quality, translate

init_db()

_SEQ = [0]


def _seq() -> int:
    _SEQ[0] += 1
    return _SEQ[0]


# --- term/segment builders --------------------------------------------------
def _gloss(source_term, target_term, *, dnt=False, case_sensitive=False, exact=False):
    """One match_terms-shaped constraint dict (the SAME shape the choke point reads).
    `exact` is the Q5 per-entry opt-out back into strict surface matching."""
    return {
        "source_term": source_term,
        "target_term": target_term,
        "dnt": dnt,
        "sense_label": None,
        "case_sensitive": case_sensitive,
        "exact": exact,
    }


def _mk_seg(session, *, source, status=SegStatus.NEW, qa_flags=None):
    p = Project(name="gap-%d" % _seq())  # noqa: UP031
    session.add(p)
    session.commit()
    d = Document(project_id=p.id, filename="d.docx")
    session.add(d)
    session.commit()
    seg = Segment(
        document_id=d.id,
        target_locale="es-MX",
        order_index=0,
        source_text=source,
        status=status,
        qa_flags=qa_flags,
    )
    session.add(seg)
    session.commit()
    return seg


def _canned(text="TRANSLATED"):
    return {
        "text": text,
        "provider_used": "mock",
        "model_used": "mock",
        "attempts": 1,
        "error": None,
        "usage": {},
    }


class _CaptureRC:
    """Swap a module's run_completion for a recorder returning a fixed text (a model
    that ignores the glossary)."""

    def __init__(self, module, text="TRANSLATED"):
        self.module = module
        self.text = text
        self._orig = None

    def __enter__(self):
        self._orig = self.module.run_completion

        def fake(spec, prompt, system=None, task=None, payload=None, **kw):
            return _canned(self.text)

        set_pieces(self.module.__name__, "run_completion", fake)
        return self

    def __exit__(self, *exc):
        set_pieces(self.module.__name__, "run_completion", self._orig)
        return False


def _mk_project(session):
    p = Project(
        name="gap-proj-%d" % _seq(),  # noqa: UP031
        org_id=1,
        source_locale="en-US",
        default_provider="claude",  # noqa: RUF100, UP031
    )
    session.add(p)
    session.commit()
    session.refresh(p)
    return p


def _mk_doc(session, project, *, sources, mode=TransMode.SEGMENTED):
    doc = Document(project_id=project.id, filename="d.txt", fmt="txt", mode=mode)
    session.add(doc)
    session.commit()
    session.refresh(doc)
    for i, s in enumerate(sources):
        session.add(
            Segment(document_id=doc.id, target_locale="es-ES", order_index=i, source_text=s)
        )
    session.commit()
    return doc


def _mk_wf(session):
    wf = Workflow(name="gap-wf-%d" % _seq(), scope="global")  # noqa: UP031
    session.add(wf)
    session.commit()
    session.refresh(wf)
    session.add(
        WorkflowStep(
            workflow_id=wf.id,
            order_index=0,
            name="Step0",
            step_type=StepType.LLM_AGENT,
            provider="claude",
            prompt_text="",
            chunk_policy="none",
            retry_max=1,
            backoff_base=0.0,
        )
    )
    session.commit()
    return wf


def _add_glossary_entry(session, project_id, source_term, target_term, *, dnt=False):
    session.add(
        GlossaryEntry(
            project_id=project_id,
            source_term=source_term,
            target_term=target_term,
            source_locale="en-US",
            target_locale="es-ES",
            dnt=dnt,
        )
    )
    session.commit()


def _seg_q(doc_id):
    return select(Segment).where(Segment.document_id == doc_id, Segment.target_locale == "es-ES")


# ===========================================================================
# 1. quality.glossary_term_violations (pure function)
# ===========================================================================
def test_glossary_term_violations_unit():
    T = [_gloss("Widget", "Aparato")]
    # miss: source has Widget, target lacks Aparato
    assert quality.glossary_term_violations("Buy the Widget now.", "[es] Compra ya.", T) == [
        {"source_term": "Widget", "target_term": "Aparato"}
    ]
    # present: target has Aparato -> clean
    assert (
        quality.glossary_term_violations("Buy the Widget now.", "[es] Compra el Aparato ya.", T)
        == []
    )
    # source term absent -> not applicable (no flag even though target lacks Aparato)
    assert quality.glossary_term_violations("Buy the gadget now.", "[es] Compra ya.", T) == []
    # DNT entries are NOT this gate's business (dnt_violations owns them)
    assert (
        quality.glossary_term_violations(
            "Keep BrandX.", "[es] Mantener.", [_gloss("BrandX", "BrandX", dnt=True)]
        )
        == []
    )
    # empty target_term -> skipped (nothing to require)
    assert (
        quality.glossary_term_violations("Buy the Widget.", "[es] Compra.", [_gloss("Widget", "")])
        == []
    )
    # falsy guards
    assert quality.glossary_term_violations("", "x", T) == []
    assert quality.glossary_term_violations("x", "", T) == []
    assert quality.glossary_term_violations("x", "y", []) == []
    print(
        "  1a. glossary_term_violations: non-DNT miss detected; present/absent-source/"
        "DNT/empty-target/falsy -> clean OK"
    )

    # whole-token: a required "AI" is NOT satisfied by "email" in the target (the mirror
    # of the DNT substring false-positive) -> still a miss.
    assert quality.glossary_term_violations(
        "Use AI today.", "[es] Usa email hoy.", [_gloss("AI", "AI")]
    ) == [{"source_term": "AI", "target_term": "AI"}]
    # ...but a whole-token "AI" in the target satisfies it.
    assert (
        quality.glossary_term_violations("Use AI today.", "[es] Usa AI hoy.", [_gloss("AI", "AI")])
        == []
    )
    print("  1b. glossary_term_violations: whole-token target match (no substring pass) OK")

    # case_sensitive: source match honours the flag. 'widget' (lowercase) in source does
    # NOT match a case-sensitive 'Widget' entry -> not applicable -> no flag.
    assert (
        quality.glossary_term_violations(
            "buy the widget now.",
            "[es] Compra ya.",
            [_gloss("Widget", "Aparato", case_sensitive=True)],
        )
        == []
    )
    # case-insensitive DOES match the lowercase source -> miss (no Aparato in target)
    assert quality.glossary_term_violations(
        "buy the widget now.",
        "[es] Compra ya.",
        [_gloss("Widget", "Aparato", case_sensitive=False)],
    ) == [{"source_term": "Widget", "target_term": "Aparato"}]
    # case_sensitive target check: required 'Aparato' present only as 'aparato' -> miss
    assert quality.glossary_term_violations(
        "Buy the Widget.",
        "[es] compra aparato.",
        [_gloss("Widget", "Aparato", case_sensitive=True)],
    ) == [{"source_term": "Widget", "target_term": "Aparato"}]
    print("  1c. glossary_term_violations: case_sensitive honoured on source + target legs OK")

    # de-dup by (source, target) pair
    dup = [_gloss("Widget", "Aparato"), _gloss("Widget", "Aparato")]
    assert quality.glossary_term_violations("A Widget here.", "[es] Aqui.", dup) == [
        {"source_term": "Widget", "target_term": "Aparato"}
    ]
    print("  1d. glossary_term_violations: duplicate entries collapse to one miss OK")


# ===========================================================================
# 1m. Q5 — morphology-tolerant target-term matching (core-QA fix train 1)
# ===========================================================================
def test_glossary_morphology_tolerant():
    """The live prod campaign flooded dense-glossary reviews with false MAJOR flags on
    legitimate INFLECTIONS of the required term. The target check is now morphology-
    tolerant (deterministic, LLM-free) while staying PRECISE (whole-token discipline +
    a stem-length floor). DNT stays exact; a per-entry `exact` opts back into strict."""

    def _miss(src, tgt, terms, locale=None):
        return quality.glossary_term_violations(src, tgt, terms, target_locale=locale)

    SESION = "sesión"  # sesión (accented, both sides)

    # --- the FOUR live-prod false positives must NOT flag post-fix -----------
    # (1) es plural agreement: dispositivo -> dispositivos
    assert (
        _miss(
            "Sync every device.",
            "Sincroniza en todos sus dispositivos.",
            [_gloss("device", "dispositivo")],
            "es-ES",
        )
        == []
    ), "es plural"
    # (2) es verb conjugation vs noun: descarga -> Descargue
    assert (
        _miss(
            "Download the installer.",
            "Descargue el instalador.",
            [_gloss("download", "descarga")],
            "es-ES",
        )
        == []
    ), "es conjugation"
    # (3) es conjugated MULTIWORD: iniciar sesión -> Inicie sesión
    assert (
        _miss(
            "Log in to continue.",
            f"Inicie {SESION} para continuar.",
            [_gloss("log in", f"iniciar {SESION}")],
            "es-ES",
        )
        == []
    ), "es multiword"
    # (4) fr plural: appareil -> appareils
    assert (
        _miss(
            "On every device.", "Sur tous les appareils.", [_gloss("device", "appareil")], "fr-FR"
        )
        == []
    ), "fr plural"
    print("  1m-a. the four live-prod inflection false positives no longer flag OK")

    # --- PRECISION: a required 'carga' must NOT be satisfied by 'cargo' (the named
    # short-stem collision -- shared stem 'carg'=4 < the floor of 5).
    assert _miss("Upload the file.", "Cargo el archivo.", [_gloss("upload", "carga")], "es-ES") == [
        {"source_term": "upload", "target_term": "carga"}
    ], "carga != cargo"
    # a TRUE positive: a competitor synonym sharing NO stem still flags
    assert _miss(
        "Buy the Widget.", "Compra el trasto.", [_gloss("Widget", "Aparato")], "es-ES"
    ) == [{"source_term": "Widget", "target_term": "Aparato"}], "true miss still flags"
    print("  1m-b. precision: carga!=cargo (floor); a no-stem synonym still flags OK")

    # --- the AI/email whole-token discipline is RETAINED (substring false-negative) --
    assert _miss("Use AI today.", "Usa email hoy.", [_gloss("AI", "AI")], "es-ES") == [
        {"source_term": "AI", "target_term": "AI"}
    ], "AI not satisfied by email"
    assert _miss("Use AI today.", "Usa AI hoy.", [_gloss("AI", "AI")], "es-ES") == []
    print("  1m-c. whole-token discipline retained (AI not in email) OK")

    # --- case_sensitive: an inflection honours the case rule ------------------
    # a case-SENSITIVE 'Aparato' is NOT satisfied by the lowercase plural 'aparatos'
    assert _miss(
        "Buy the Widget.",
        "Compra los aparatos.",
        [_gloss("Widget", "Aparato", case_sensitive=True)],
        "es-ES",
    ) == [{"source_term": "Widget", "target_term": "Aparato"}], "cs inflection strict"
    # ...but case-INSENSITIVE accepts the same plural
    assert (
        _miss("Buy the Widget.", "Compra los aparatos.", [_gloss("Widget", "Aparato")], "es-ES")
        == []
    ), "ci inflection accepted"
    print("  1m-d. case_sensitive honoured on the inflection path OK")

    # --- per-entry `exact` opt-out forces strict SURFACE matching ------------
    ex = _gloss("device", "dispositivo", exact=True)
    assert _miss("Sync every device.", "Sincroniza los dispositivos.", [ex], "es-ES") == [
        {"source_term": "device", "target_term": "dispositivo"}
    ], "exact: plural rejected"
    assert _miss("Sync every device.", "Ver el dispositivo aqui.", [ex], "es-ES") == [], (
        "exact: the exact surface still passes"
    )
    print("  1m-e. per-entry exact opt-out forces strict surface matching OK")

    # --- the GENERIC family (no locale) still gives append/plural tolerance,
    # but NOT full conjugation (documented locale-dependence) -----------------
    assert (
        _miss(
            "Sync every device.",
            "Sincroniza los dispositivos.",
            [_gloss("device", "dispositivo")],
            None,
        )
        == []
    ), "generic plural OK"
    assert _miss(
        "Download it.", "Descargue el instalador.", [_gloss("download", "descarga")], None
    ) == [{"source_term": "download", "target_term": "descarga"}], "generic: no conjugation"
    print("  1m-f. generic family: plural tolerant; conjugation is locale-gated OK")

    # --- through the SEAM: segment.target_locale drives the family; a correct
    # inflected term commits CLEAN (no false NEEDS_REVIEW) --------------------
    with Session(engine) as s:
        seg = _mk_seg(s, source="Sync every device.")  # _mk_seg sets es-MX -> 'es'
        out = quality.commit_mt_output(
            s,
            seg,
            "Sincroniza en todos sus dispositivos.",
            source="mt",
            glossary_terms=[_gloss("device", "dispositivo")],
        )
        s.commit()
        assert out.glossary_violations == () and seg.status == SegStatus.MT, out
        assert not [f for f in (seg.qa_flags or []) if f.get("type") == "glossary_violation"], (
            seg.qa_flags
        )
    print("  1m-g. write seam: an inflected required term commits clean (no false flag) OK")


# ===========================================================================
# 1n. Q5 fix round — well-formedness guard (F1), append rung (F5), accents (F8)
# ===========================================================================
def test_glossary_fix_round():
    """The adversarial review's fix round for the morphology matcher:
    F1  a punctuation-bearing required term is matched STRICTLY (the tolerant rungs
        strip punctuation, which would let a bare 'c' satisfy 'C++');
    F5  the APPEND rung recognises a plural the STEM floor alone would reject;
    F8  the APPEND/STEM rungs fold accents (the exact rung stays accent-sensitive)."""

    def _miss(src, tgt, terms, locale=None):
        return quality.glossary_term_violations(src, tgt, terms, target_locale=locale)

    # --- F1: the SIX exploit pairs. Each required term carries punctuation the tolerant
    # rungs would strip; a decoy target that satisfied the pre-fix matcher must now FLAG.
    exploits = [
        ("cpp", "C++", "Use cpp here.", "Usa c aqui."),  # 'C++' vs bare 'c'
        ("format", "%s", "Insert the format token.", "Inserta s aqui."),  # '%s' vs 's'
        ("telco", "AT&T", "Call telco support.", "Llama a at t."),  # 'AT&T' vs 'at t'
        ("runtime", "Node.js", "Install the runtime.", "Instala node js."),  # vs 'node js'
        ("framework", ".NET", "Use the framework.", "Usa net hoy."),  # '.NET' vs 'net'
        ("mail", "e-mail", "Send the mail.", "Envia el email."),  # 'e-mail' vs 'email'
    ]
    for src_t, tgt_t, src, tgt in exploits:
        assert _miss(src, tgt, [_gloss(src_t, tgt_t)], "es-ES") == [
            {"source_term": src_t, "target_term": tgt_t}
        ], ("F1 exploit must flag", tgt_t, tgt)
    # ...and a punctuation-bearing term whose EXACT surface is present still passes clean
    # (the strict rung matches 'Wi-Fi' as 'Wi-Fi').
    assert (
        _miss("Enable wifi now.", "Activa Wi-Fi ahora.", [_gloss("wifi", "Wi-Fi")], "es-ES") == []
    ), "F1: exact 'Wi-Fi' matches"
    # a punctuation-bearing term genuinely MISSING still flags (strict miss, not vacuous):
    assert _miss(
        "Enable wifi now.", "Activa la conexion ahora.", [_gloss("wifi", "Wi-Fi")], "es-ES"
    ) == [{"source_term": "wifi", "target_term": "Wi-Fi"}], "F1: 'Wi-Fi' true miss flags"
    print("  1n-a. F1: six punctuation-stripped exploits flag; exact 'Wi-Fi' still matches OK")

    # --- F5: the APPEND rung. es 'red'->'redes' and fr 'reseau'->'reseaux' are plurals
    # the 5-char STEM floor alone would REJECT (stem 'red'=3), so passing them PROVES the
    # append rung is live (review mutation iv — dropping the append rung — must die).
    assert (
        _miss("Join the network.", "Une las redes.", [_gloss("network", "red")], "es-ES") == []
    ), "F5: es red->redes (append)"
    assert (
        _miss("Join the network.", "Rejoins les réseaux.", [_gloss("network", "réseau")], "fr-FR")
        == []
    ), "F5: fr réseau->réseaux"
    # a NON-suffix tail is still a miss (the append rung is closed-suffix, not any-tail):
    assert _miss("Join the network.", "Une la redada.", [_gloss("network", "red")], "es-ES") == [
        {"source_term": "network", "target_term": "red"}
    ], "F5: 'redada' is not 'red'+s/es"
    print("  1n-b. F5: append rung recognises red->redes / réseau->réseaux; mutation iv dies OK")

    # --- F8: accent folding on the APPEND/STEM rungs (the dominant remaining FP class).
    assert (
        _miss("Press the button.", "Pulsa los botones.", [_gloss("button", "botón")], "es-ES") == []
    ), "F8: botón->botones"
    assert (
        _miss("Start a session.", "Inicia las sesiones.", [_gloss("session", "sesión")], "es-ES")
        == []
    ), "F8: sesión->sesiones"
    assert (
        _miss("Take an action.", "Realiza las acciones.", [_gloss("action", "acción")], "es-ES")
        == []
    ), "F8: acción->acciones"
    # ...but the `exact` OPT-OUT is accent-strict: an entry marked exact returns before the
    # tolerant rungs (`if exact: return False`), so an accent-different surface still
    # requires its exact accents. NOTE the honest label: this pins the `exact` OPT-OUT
    # path, NOT "rung 1 inside the tolerant matcher" -- the non-exact pin just below is what
    # exercises rung 1's accent-strictness.
    ex = _gloss("button", "botón", exact=True)
    assert _miss("Press the button.", "Pulsa el boton aqui.", [ex], "es-ES") == [
        {"source_term": "button", "target_term": "botón"}
    ], "F8: the `exact` opt-out is accent-strict (returns before the tolerant rungs)"
    assert _miss("Press the button.", "Pulsa el botón aqui.", [ex], "es-ES") == [], (
        "F8: the `exact` opt-out matches the exact accented surface"
    )

    # ...and, WITHOUT `exact`, rung 1 (the whole-token exact check) is STILL accent-strict:
    # a required 'botón' against a bare, accent-DROPPED 'boton' FLAGS. Rung 1 does not fold
    # (`if a == b` on the raw surfaces); the tolerant rungs then also miss because the append
    # rung sees an EMPTY suffix ('boton'->'boton') and the stem floor rejects the 4-char
    # 'boto'. NON-MONOTONICITY, PINNED DELIBERATELY (R6): 'botón' matches its PLURAL
    # 'botones' (append rung, accent-folded) yet FLAGS the accent-dropped SINGULAR 'boton'.
    # This is the intended verdict, not a bug: the missing accent on 'boton' IS a spelling
    # error worth a reviewer's glance, while 'botones' is legitimate morphology. Folding
    # rung 1 (the mutation) would let 'boton' satisfy 'botón' and this assertion would fail
    # -- so the fold-rung-1 mutation dies here.
    assert _miss(
        "Press the button.", "Pulsa el boton aqui.", [_gloss("button", "botón")], "es-ES"
    ) == [{"source_term": "button", "target_term": "botón"}], (
        "F8: rung 1 is accent-strict -- bare 'boton' flags required 'botón'"
    )
    assert (
        _miss("Press the button.", "Pulsa los botones.", [_gloss("button", "botón")], "es-ES") == []
    ), "R6: required 'botón' matches its plural 'botones' (append rung)"
    print(
        "  1n-c. F8: accents folded on append/stem (botón/sesión/acción); exact opt-out "
        "AND rung 1 stay accent-strict (bare 'boton' flags; plural 'botones' passes) OK"
    )


# ===========================================================================
# 2. commit_mt_output wiring (the sole write seam)
# ===========================================================================
def test_commit_wiring():
    with Session(engine) as s:
        # (a) miss -> text WRITTEN + NEEDS_REVIEW + glossary_violation flag; tm_eligible unchanged
        seg = _mk_seg(s, source="Buy the Widget now.")
        out = quality.commit_mt_output(
            s, seg, "[es] Compra ya.", source="mt", glossary_terms=[_gloss("Widget", "Aparato")]
        )
        s.commit()
        assert out.committed and not out.flagged, out
        assert seg.target_text == "[es] Compra ya.", ("text still committed", seg.target_text)
        assert seg.status == SegStatus.NEEDS_REVIEW, seg.status
        assert out.glossary_violations == ({"source_term": "Widget", "target_term": "Aparato"},), (
            out.glossary_violations
        )
        gv = [f for f in (seg.qa_flags or []) if f.get("type") == "glossary_violation"]
        assert gv and gv[0]["source_term"] == "Widget" and gv[0]["target_term"] == "Aparato", (
            seg.qa_flags
        )
        assert out.tm_eligible is True, ("tm_eligible unchanged on a miss (mirrors DNT)", out)
        print(
            "  2a. commit: required-term miss -> committed + NEEDS_REVIEW + glossary_violation, "
            "tm_eligible unchanged OK"
        )

        # tm_eligible=False stays False on a miss too (caller's value passed through)
        seg = _mk_seg(s, source="Buy the Widget now.")
        out = quality.commit_mt_output(
            s,
            seg,
            "[es] Compra ya.",
            source="mt",
            tm_eligible=False,
            glossary_terms=[_gloss("Widget", "Aparato")],
        )
        s.commit()
        assert out.tm_eligible is False and out.glossary_violations, out
        print(
            "  2b. commit: a miss passes the caller's tm_eligible through unchanged (soft gate) OK"
        )

        # (b) present -> clean MT, no flag, byte-identical outcome
        seg = _mk_seg(s, source="Buy the Widget now.")
        out = quality.commit_mt_output(
            s,
            seg,
            "[es] Compra el Aparato ya.",
            source="mt",
            glossary_terms=[_gloss("Widget", "Aparato")],
        )
        s.commit()
        assert out.committed and out.glossary_violations == () and seg.status == SegStatus.MT, out
        assert not [f for f in (seg.qa_flags or []) if f.get("type") == "glossary_violation"], (
            seg.qa_flags
        )
        print("  2c. commit: required term PRESENT -> clean MT, no flag, byte-identical OK")

        # (d) same-text re-check composes with its QE flag AND a DNT miss.
        qe_flag = {"type": "qe", "severity": "minor", "score": 80, "rationale": "ok"}
        seg = _mk_seg(s, source="Keep BrandX and buy the Widget.", qa_flags=[qe_flag])
        seg.target_text = "[es] Mantener y comprar."
        out = quality.commit_mt_output(
            s,
            seg,
            "[es] Mantener y comprar.",
            source="mt",
            glossary_terms=[_gloss("BrandX", "BrandX", dnt=True), _gloss("Widget", "Aparato")],
        )
        s.commit()
        types = {f.get("type") for f in (seg.qa_flags or [])}
        assert types == {"qe", "dnt_violation", "glossary_violation"}, (
            "all three flags",
            seg.qa_flags,
        )
        assert out.dnt_violations == ("BrandX",) and out.glossary_violations == (
            {"source_term": "Widget", "target_term": "Aparato"},
        ), out
        kept = [f for f in seg.qa_flags if f.get("type") == "qe"][0]  # noqa: RUF015
        assert kept["rationale"] == "ok", kept
        print("  2d. commit: glossary flag composes with QE + DNT flags (none lost) OK")

        # (e) refusal path UNAFFECTED (a withheld hedge gets no glossary flag)
        HEDGE = (
            "I'm ready to translate. However, I don't see any source text. "
            "Please provide the actual source text."
        )
        seg = _mk_seg(s, source="Buy the Widget now.")
        out = quality.commit_mt_output(
            s, seg, HEDGE, source="mt", glossary_terms=[_gloss("Widget", "Aparato")]
        )
        s.commit()
        assert out.flagged and not out.committed and out.glossary_violations == (), out
        assert not (seg.target_text or "").strip(), "hedge withheld"
        assert not [f for f in (seg.qa_flags or []) if f.get("type") == "glossary_violation"], (
            seg.qa_flags
        )
        assert [f for f in seg.qa_flags if f.get("type") == "refusal"], seg.qa_flags
        print("  2e. commit: refusal path unaffected (withheld hedge -> no glossary flag) OK")


# ===========================================================================
# 3. DNT regression — the DNT gate is byte-identical alongside the new gate
# ===========================================================================
def test_dnt_regression():
    with Session(engine) as s:
        terms = [_gloss("BrandX", "BrandX", dnt=True), _gloss("Widget", "Aparato")]
        # DNT dropped, required term PRESENT -> ONLY dnt_violation (glossary clean)
        seg = _mk_seg(s, source="Keep BrandX; buy the Widget.")
        out = quality.commit_mt_output(
            s, seg, "[es] Compra el Aparato.", source="mt", glossary_terms=terms
        )
        s.commit()
        assert out.dnt_violations == ("BrandX",) and out.glossary_violations == (), out
        assert seg.status == SegStatus.NEEDS_REVIEW, seg.status
        assert {f["type"] for f in seg.qa_flags} == {"dnt_violation"}, seg.qa_flags
        print("  3a. regression: DNT dropped + required term present -> only dnt_violation OK")

        # DNT preserved, required term MISSING -> ONLY glossary_violation
        seg = _mk_seg(s, source="Keep BrandX; buy the Widget.")
        out = quality.commit_mt_output(
            s, seg, "[es] Mantener BrandX.", source="mt", glossary_terms=terms
        )
        s.commit()
        assert out.dnt_violations == () and out.glossary_violations, out
        assert {f["type"] for f in seg.qa_flags} == {"glossary_violation"}, seg.qa_flags
        print(
            "  3b. regression: DNT preserved + required term missing -> only glossary_violation OK"
        )

    # F3: DNT stays EXACT — it must NOT be routed through the morphology-tolerant matcher.
    # An INFLECTED form of a DNT term is a real breach (a do-not-translate token is defined
    # as surviving VERBATIM). These pin that dnt_violations never gains inflection tolerance
    # (review mutation iii — routing DNT through _term_present_tolerant — must die).
    # a brand DNT term altered by a trailing 's' is a violation, not an inflection:
    assert quality.dnt_violations(
        "Keep BrandX visible.", "Mantener BrandXs visible.", [_gloss("BrandX", "BrandX", dnt=True)]
    ) == ["BrandX"], "F3: DNT 'BrandX' dropped when target has only 'BrandXs'"
    # an es plural of a DNT term is STILL a breach — dnt has no morphology, no locale:
    assert quality.dnt_violations(
        "El aparato aqui.", "Los aparatos aqui.", [_gloss("aparato", "aparato", dnt=True)]
    ) == ["aparato"], "F3: DNT 'aparato' dropped when target has only the plural 'aparatos'"
    # ...and the verbatim DNT term of course still passes clean.
    assert (
        quality.dnt_violations(
            "El aparato aqui.", "El aparato aqui.", [_gloss("aparato", "aparato", dnt=True)]
        )
        == []
    )
    print("  3c. F3: DNT stays EXACT (inflected/plural DNT term still flags; mutation iii dies) OK")


# ===========================================================================
# 4. Both translate paths end-to-end (fake engine drops the required term)
# ===========================================================================
def test_end_to_end():
    # MT path
    with Session(engine) as s:
        p = _mk_project(s)
        _add_glossary_entry(s, p.id, "Widget", "Aparato")
        doc = _mk_doc(s, p, sources=("Buy the Widget now.",))
        doc_id = doc.id
    with Session(engine) as s, _CaptureRC(translate, text="Compra ahora."):
        for _ in translate.pretranslate_stream(s, doc_id, "es-ES"):
            pass
    with Session(engine) as s:
        seg = s.exec(_seg_q(doc_id)).first()
        assert seg.target_text == "Compra ahora.", ("committed, not blocked", seg.target_text)
        assert seg.status == SegStatus.NEEDS_REVIEW, seg.status
        assert [f for f in (seg.qa_flags or []) if f.get("type") == "glossary_violation"], (
            seg.qa_flags
        )
    print("  4a. MT path end-to-end: required term dropped -> NEEDS_REVIEW + glossary_violation OK")

    # Segmented workflow path (exercises the orchestrator write-back re-application)
    with Session(engine) as s:
        p = _mk_project(s)
        _add_glossary_entry(s, p.id, "Widget", "Aparato")
        doc = _mk_doc(s, p, sources=("Buy the Widget now.",))
        wf = _mk_wf(s)
        wf_id, doc_id = wf.id, doc.id
    with Session(engine) as s, _CaptureRC(orchestrator, text="Compra ahora."):
        for _ in orchestrator.run_workflow_stream(s, wf_id, doc_id, "es-ES"):
            pass
    with Session(engine) as s:
        seg = s.exec(_seg_q(doc_id)).first()
        assert seg.target_text == "Compra ahora.", ("committed, not blocked", seg.target_text)
        assert seg.status == SegStatus.NEEDS_REVIEW, seg.status
        assert [f for f in (seg.qa_flags or []) if f.get("type") == "glossary_violation"], (
            seg.qa_flags
        )
    print(
        "  4b. workflow path end-to-end: required term dropped -> NEEDS_REVIEW + glossary_violation OK"  # noqa: E501
    )


# ===========================================================================
# 5. REQUIRE rulebook directive (parsing + gate mechanics)
# ===========================================================================
def test_require_directive_unit():
    # missing -> a MAJOR Terminology error
    errs = lqa._rulebook_errors("the quick fox", "REQUIRE Kubernetes")
    assert errs == [
        {
            "dimension": "Terminology",
            "severity": "major",
            "note": "required term 'Kubernetes' missing",
            "source": "rulebook",
        }
    ], errs
    # present (whole token) -> no error; lowercase keyword accepted
    assert lqa._rulebook_errors("we run kubernetes here", "require Kubernetes") == []
    # substring is NOT enough: REQUIRE cat is a miss when the only 'cat' is in 'category'
    assert lqa._rulebook_errors("a category only", "REQUIRE cat") == [
        {
            "dimension": "Terminology",
            "severity": "major",
            "note": "required term 'cat' missing",
            "source": "rulebook",
        }
    ]
    assert lqa._rulebook_errors("the cat sat", "REQUIRE cat") == []
    # blank term -> ignored
    assert lqa._rulebook_errors("anything", "REQUIRE   ") == []
    assert lqa._rulebook_errors("anything", "REQUIRE") == []
    # FORBID stays byte-identical (minor Style/Awkward, lowercase substring)
    assert lqa._rulebook_errors("please utilisez foo", "FORBID utilisez") == [
        {
            "dimension": "Style/Awkward",
            "severity": "minor",
            "note": "forbidden term 'utilisez'",
            "source": "rulebook",
        }
    ]
    print(
        "  5a. _rulebook_errors: REQUIRE missing->major, present/substring/blank handled; "
        "FORBID byte-identical OK"
    )

    # score_text: a single REQUIRE miss fails the GATE though the score clears the >=90
    # floor (100-5=95) -> gate-eligible (a MAJOR fails regardless of score).
    card = lqa.score_text(
        "Deploy on Kubernetes.", "[fr] Deployez.", "en-US", "fr-FR", "REQUIRE Kubernetes"
    )
    assert card["score"] == 95 and card["passed"] is False and card["gate"] == "needs_review", card
    # present -> passes
    ok = lqa.score_text(
        "Deploy on Kubernetes.",
        "[fr] Deployez sur Kubernetes.",
        "en-US",
        "fr-FR",
        "REQUIRE Kubernetes",
    )
    assert ok["passed"] is True and ok["gate"] == "auto_approved", ok
    print(
        "  5b. score_text: REQUIRE miss -> score 95 but gate FAILED (gate-eligible); present->pass OK"  # noqa: E501
    )


# ===========================================================================
# 6. REQUIRE drives the correction loop end-to-end (translate -> LQA -> fix)
# ===========================================================================
def _add_step(
    session,
    wf_id,
    order,
    name,
    *,
    step_type=StepType.LLM_AGENT,
    provider="mock",
    prompt_text="",
    run_condition=None,
):
    session.add(
        WorkflowStep(
            workflow_id=wf_id,
            order_index=order,
            name=name,
            step_type=step_type,
            provider=provider,
            prompt_text=prompt_text,
            chunk_policy="auto",
            retry_max=1,
            backoff_base=0.0,
            run_condition=run_condition,
        )
    )
    session.commit()


def _docx_bytes(paragraphs):
    d = Docx()
    for p in paragraphs:
        d.add_paragraph(p)
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def _new_project(client, name):
    r = client.post(
        "/projects",
        data={
            "name": name,
            "source_locale": "en-US",
            "target_locales": "fr-FR",
            "default_provider": "mock",
        },
        follow_redirects=False,
    )
    return int(r.headers["location"].rsplit("/", 1)[-1])


def _upload(client, pid, paragraphs):
    r = client.post(
        f"/projects/{pid}/documents",
        data={"mode": "segmented"},
        files={
            "file": (
                "d.docx",
                _docx_bytes(paragraphs),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    return int(r.headers["location"].rsplit("/", 1)[-1])


def _fix_loop_workflow(name, guard):
    """translate -> LQA(mock, deterministic rulebook) -> fix{guard}."""
    with Session(engine) as s:
        wf = Workflow(name=name, scope="global")
        s.add(wf)
        s.commit()
        s.refresh(wf)
        _add_step(s, wf.id, 0, "Translator", prompt_text="Translate: {text}")
        _add_step(s, wf.id, 1, "LQA", step_type=StepType.LQA, provider="mock")
        _add_step(s, wf.id, 2, "Fix", prompt_text="Fix: {text}", run_condition=guard)
        return wf.id


def _stepruns(run_id):
    with Session(engine) as s:
        return s.exec(
            select(orchestrator.StepRun)
            .where(orchestrator.StepRun.workflow_run_id == run_id)
            .order_by(orchestrator.StepRun.id)
        ).all()


def test_require_gate_fix_loop():
    client = TestClient(app)
    guard = {"ref": "prev_lqa", "op": "failed", "value": None}

    # FAIL case: REQUIRE a term absent from the source (mock echoes source) -> LQA fails
    # -> the Fix step RUNS. Structural checks are clean, so the failure is attributable
    # solely to the REQUIRE directive (asserted via the persisted LQAResult note).
    lqa.save_rulebook("en-US", "fr-FR", "REQUIRE Kubernetes")
    wf_id = _fix_loop_workflow("B1 fix-loop FAIL", guard)
    pid = _new_project(client, "B1 fix-loop FAIL")
    doc = _upload(client, pid, ["Buy the widget today.", "Ship it now."])
    with Session(engine) as s:
        summary = orchestrator.run_workflow(s, wf_id, doc, "fr-FR")
    assert summary["status"] == "completed", summary
    srs = _stepruns(summary["run_id"])
    assert srs[1].gate == "failed", (
        "precondition: LQA REQUIRE miss must FAIL the gate",
        srs[1].gate,
    )
    assert srs[2].status == RunStatus.COMPLETED, (
        "Fix must RUN when the LQA failed",
        [sr.status for sr in srs],
    )
    with Session(engine) as s:
        notes = [
            r.note for r in s.exec(select(LQAResult).where(LQAResult.document_id == doc)).all()
        ]
    assert "required term 'Kubernetes' missing" in notes, (
        "causation: REQUIRE drove the fail",
        notes,
    )
    print("  6a. REQUIRE miss -> LQA gate FAILED -> Fix step RUNS (the correction loop) OK")

    # PASS control: REQUIRE a term that IS present (echoed from source) -> LQA passes ->
    # the Fix step is SKIPPED. Proves the REQUIRE directive is what flips the gate.
    lqa.save_rulebook("en-US", "fr-FR", "REQUIRE widget")
    wf_id = _fix_loop_workflow("B1 fix-loop PASS", guard)
    pid = _new_project(client, "B1 fix-loop PASS")
    doc = _upload(client, pid, ["Buy the widget today.", "Ship the widget now."])
    with Session(engine) as s:
        summary = orchestrator.run_workflow(s, wf_id, doc, "fr-FR")
    assert summary["status"] == "completed", summary
    srs = _stepruns(summary["run_id"])
    assert srs[1].gate == "passed", ("control: REQUIRE present must PASS", srs[1].gate)
    assert srs[2].status == RunStatus.SKIPPED, (
        "Fix must be SKIPPED when the LQA passed",
        [sr.status for sr in srs],
    )
    print("  6b. REQUIRE present -> LQA gate PASSED -> Fix step SKIPPED (control) OK")


# ===========================================================================
# 7. Parity — the glossary_violation flag renders in the editor qa cell
# ===========================================================================
def test_editor_surfacing():
    client = TestClient(app)
    pid = _new_project(client, "B1 editor surfacing")
    doc = _upload(client, pid, ["Buy the Widget today."])
    with Session(engine) as s:
        seg = s.exec(
            select(Segment).where(Segment.document_id == doc, Segment.target_locale == "fr-FR")
        ).first()
        # Persist a glossary_violation the way the choke point does (miss: no Aparato).
        quality.commit_mt_output(
            s, seg, "[fr] Compra hoy.", source="mt", glossary_terms=[_gloss("Widget", "Aparato")]
        )
        s.commit()
        seg_id = seg.id
    page = client.get(f"/documents/{doc}/editor?locale=fr-FR").text
    assert f'id="qa-{seg_id}"' in page, "qa cell anchor missing"
    assert "glossary_violation" in page, (
        "glossary_violation chip did not render in the editor qa cell"
    )
    print("  7. parity: glossary_violation flag renders in the editor qa cell OK")


# ===========================================================================
# 8. fail-before / pass-after — git-show HEAD reconstruction
# ===========================================================================
def _load_module_from_source(src, name):
    import importlib.util

    spec = importlib.util.spec_from_loader(name, loader=None)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod  # dataclasses needs the module registered for KW_ONLY resolution
    exec(compile(src, f"{name}.py", "exec"), mod.__dict__)
    return mod


def test_fail_before_pass_after():
    import subprocess

    TERMS = [_gloss("Kubernetes", "Kubernetes")]
    SRC, TGT = "Deploy on Kubernetes now.", "[es] Despliega ahora."  # required term MISSING

    old_src = None
    try:
        old_src = subprocess.check_output(
            ["git", "show", "HEAD:app/services/quality.py"],
            text=True,
            cwd=_ROOT,
            stderr=subprocess.DEVNULL,
        )
    except Exception:
        old_src = None

    # SRP split 2026-09-15: the checker may live in a sibling piece of quality.py at HEAD
    head_quality_tree = ""
    try:
        listed = subprocess.check_output(
            ["git", "ls-tree", "--name-only", "HEAD", "app/services/"],
            text=True,
            cwd=_ROOT,
            stderr=subprocess.DEVNULL,
        ).split()
        for rel in listed:
            if rel.startswith("app/services/quality") and rel.endswith(".py"):
                head_quality_tree += subprocess.check_output(
                    ["git", "show", f"HEAD:{rel}"],
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    cwd=_ROOT,
                    stderr=subprocess.DEVNULL,
                )
    except Exception:
        head_quality_tree = old_src or ""

    if old_src is not None and "def glossary_term_violations" not in head_quality_tree:
        # HEAD predates the change: reconstruct it and prove the miss goes UNFLAGGED.
        oldq = _load_module_from_source(old_src, "old_quality_b1")
        with Session(engine) as s:
            seg = _mk_seg(s, source=SRC)
            out = oldq.commit_mt_output(s, seg, TGT, source="mt", glossary_terms=TERMS)
            s.commit()
            assert seg.target_text == TGT and seg.status == SegStatus.MT, (
                "fail-before: pre-change commit leaves a required-term miss as clean MT",
                seg.status,
            )
            assert not [f for f in (seg.qa_flags or []) if f.get("type") == "glossary_violation"], (
                seg.qa_flags
            )
            assert not hasattr(oldq, "glossary_term_violations")
        print(
            "  8a. fail-before (git HEAD reconstruction): pre-change commit does NOT flag the miss OK"  # noqa: E501
        )
    else:
        print(
            "  8a. fail-before: HEAD already carries the change (or git unavailable) -> "
            "behavioral-delta proof below stands in OK"
        )

    # pass-after (always): the current module DOES flag the same miss...
    with Session(engine) as s:
        seg = _mk_seg(s, source=SRC)
        out = quality.commit_mt_output(s, seg, TGT, source="mt", glossary_terms=TERMS)
        s.commit()
        assert seg.status == SegStatus.NEEDS_REVIEW and out.glossary_violations, out
        assert [f for f in (seg.qa_flags or []) if f.get("type") == "glossary_violation"], (
            seg.qa_flags
        )
    # ...and the git-independent delta: the OLD dnt-only gate returns [] for this non-DNT
    # miss, while the new gate catches it — the change is load-bearing either way.
    assert quality.dnt_violations(SRC, TGT, TERMS) == []
    assert quality.glossary_term_violations(SRC, TGT, TERMS) == [
        {"source_term": "Kubernetes", "target_term": "Kubernetes"}
    ]
    print("  8b. pass-after: new gate flags the miss; the old dnt-only gate returns [] (delta) OK")


class _CapturePayloadRC:
    """Swap a module's run_completion for a recorder that captures the provider
    PAYLOAD (its `glossary` block is what we assert on) and returns a fixed text."""

    def __init__(self, module, bucket, text="TRANSLATED"):
        self.module, self.bucket, self.text = module, bucket, text
        self._orig = None

    def __enter__(self):
        self._orig = self.module.run_completion

        def fake(spec, prompt, system=None, task=None, payload=None, **kw):
            self.bucket.append({"task": task, "glossary": (payload or {}).get("glossary")})
            return _canned(self.text)

        set_pieces(self.module.__name__, "run_completion", fake)
        return self

    def __exit__(self, *exc):
        set_pieces(self.module.__name__, "run_completion", self._orig)
        return False


# ===========================================================================
# 9. Plain-MT vs workflow glossary PARITY (issue #57 guard)
#    Both surfaces must resolve the SAME glossary set for the same project/segment
#    through the shared build_glossary_plan + match_terms_merged seam. This LOCKS the
#    parity so a future divergence (e.g. re-adding a typed-only gate on one path) is
#    caught. On the same UNTYPED project with a legacy per-project term, the plain
#    doc-translate path and a 1-step llm_agent workflow must both inject the term.
# ===========================================================================
def test_plain_workflow_glossary_parity():

    def _setup(session, name):
        p = _mk_project(session)
        p.name = name
        session.add(p)
        session.commit()
        # A LEGACY per-project entry (glossary_id=None) — the issue #57 scenario.
        _add_glossary_entry(session, p.id, "activation code", "codigo de activacion")
        doc = _mk_doc(session, p, sources=("Enter your activation code now.",))
        return p, doc

    plain_calls, wf_calls = [], []

    # Plain doc-translate MT path.
    with Session(engine) as s:
        _, doc = _setup(s, "gap57-plain")
        doc_id = doc.id
    with Session(engine) as s, _CapturePayloadRC(translate, plain_calls):
        for _ in translate.pretranslate_stream(s, doc_id, "es-ES"):
            pass

    # 1-step llm_agent workflow on an identical project/segment.
    with Session(engine) as s:
        _, doc = _setup(s, "gap57-wf")
        wf = _mk_wf(s)
        wf_id, doc_id = wf.id, doc.id
    with Session(engine) as s, _CapturePayloadRC(orchestrator, wf_calls):
        for _ in orchestrator.run_workflow_stream(s, wf_id, doc_id, "es-ES"):
            pass

    def _terms(calls):
        # Fresh provider calls only (task="translate"); normalize to comparable tuples.
        for c in calls:
            if c["task"] == "translate":
                return sorted(
                    (t["source_term"], t["target_term"], bool(t.get("dnt")))
                    for t in (c["glossary"] or [])
                )
        return None

    plain_terms, wf_terms = _terms(plain_calls), _terms(wf_calls)
    assert plain_terms is not None, ("plain path made no translate call", plain_calls)
    assert wf_terms is not None, ("workflow path made no translate call", wf_calls)
    # NON-EMPTY on both (the term must actually be injected, not silently dropped).
    assert plain_terms, ("plain path sent an EMPTY glossary (issue #57)", plain_calls)
    assert wf_terms, ("workflow path sent an EMPTY glossary", wf_calls)
    # IDENTICAL across surfaces.
    assert plain_terms == wf_terms, (
        "glossary parity broke between the plain-MT and workflow paths (issue #57): "
        f"plain={plain_terms!r} workflow={wf_terms!r}"
    )
    assert ("activation code", "codigo de activacion", False) in plain_terms, plain_terms
    print("  9. plain-MT vs workflow glossary parity: same non-empty term set on both surfaces OK")


def main():
    test_glossary_term_violations_unit()
    test_glossary_morphology_tolerant()
    test_glossary_fix_round()
    test_commit_wiring()
    test_dnt_regression()
    test_end_to_end()
    test_require_directive_unit()
    test_require_gate_fix_loop()
    test_editor_surfacing()
    test_fail_before_pass_after()
    test_plain_workflow_glossary_parity()
    print("GLOSSARY ADHERENCE OK")


if __name__ == "__main__":
    main()
