"""PowerPoint filter + Office macro/template aliases (format Batch 4).

  [A] registry: .pptx .pptm .potx .potm .ppsx .ppsm -> pptx filter; .docm .dotx .dotm ->
      docx; .xlsm .xltx .xltm -> xlsx; every surface list names them.
  [B] parse walks slides in sldIdLst order: shapes, groups, tables (merged cells skipped),
      SmartArt, chart titles, notes (marked), hidden slides (marked), then alt text;
      fields, empty paragraphs, mc:Fallback, layouts and masters are not units.
  [C] build with no translations keeps every member byte-identical; a translated build
      rewrites only the touched parts ([Content_Types].xml and vbaProject.bin untouched).
  [D] inline fidelity: run rPr (bold, hyperlink) survives; a:br and a:fld points come back;
      a dropped point is appended, never deleted; a newline becomes a:br.
  [E] SmartArt cached drawing mirrors the translated data points.
  [F] build_localized stamps a:rPr lang on translated runs only, a:pPr rtl for RTL.
  [G] macro / template variants keep their main content type and the VBA bytes.
  [H] hostile packages: not a zip, zip bomb, DOCTYPE/entity parts, rels that escape the
      package, markup budget, segment cap, no presentation part.
  [I] Word .docm / .dotx / .dotm round trip through the docx filter with the same walk
      stamp; content types and word/vbaProject.bin kept; .docx unchanged.
  [J] Excel .xlsm / .xltm keep VBA and macro content type; .xltx stays a template;
      .xlsx never becomes macro-enabled.
  [K] ingest + export end to end for a .pptm and a .docm (doc.fmt keeps the extension).

Fixtures: tests/fixtures/formats/pptx/make_pptx.py (hand-written XML, stdlib zip only).

Hermetic + offline. Final line: "PPTX FILTER OK".
"""

import io
import os
import re
import sys
import zipfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_pptx_")

from app.filters import get_filter, supported_exts
from app.filters.base import InputLimitError
from tests.fixtures.formats.pptx.make_pptx import (
    MAIN_TYPES,
    SLIDE1,
    VBA_BYTES,
    make_pptx,
)

PPT_EXTS = ("pptx", "pptm", "potx", "potm", "ppsx", "ppsm")
WORD_EXTS = ("docm", "dotx", "dotm")
XL_EXTS = ("xlsm", "xltx", "xltm")


def _members(data: bytes) -> dict:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return {i.filename: zf.read(i) for i in zf.infolist()}


def _units(data: bytes, fmt="pptx"):
    return get_filter(fmt).parse(data).units


def _refused(fn, *args) -> str:
    try:
        fn(*args)
    except InputLimitError as exc:
        return str(exc)
    raise AssertionError(f"expected a refusal from {fn}")


def test_a_registry_and_surfaces():
    exts = set(supported_exts())
    for e in PPT_EXTS:
        assert e in exts and get_filter(e).fmt == "pptx", e
    for e in WORD_EXTS:
        assert e in exts and get_filter(e).fmt == "docx", e
    for e in XL_EXTS:
        assert e in exts and get_filter(e).fmt == "xlsx", e
    with open("app/templates/project.html", encoding="utf-8") as fh:
        page = fh.read()
    # anchored on the upload input id: other formats may come first in the list
    accept = re.search(r'id="doc-upload-file"[^>]*?accept="([^"{]*)', page, re.S).group(1)
    for e in PPT_EXTS + WORD_EXTS + XL_EXTS:
        assert f".{e}," in accept + ",", ("upload accept list", e)
    from app.services import mime
    from app.services.connectors import msgraph

    for e in PPT_EXTS + WORD_EXTS + XL_EXTS:
        assert mime.mime_for_fmt(e) != mime.DEFAULT_MIME, ("mime", e)
        assert e in msgraph.FORMATS, ("msgraph", e)
    assert "pptx" in msgraph.FORMATS and msgraph.FORMATS["pptm"] == "pptx"
    assert not any("PowerPoint" in v for v in msgraph._NO_FILTER_NOTE.values())
    assert "ppt" in msgraph._NO_FILTER_NOTE  # legacy binary .ppt is still skipped with a reason
    from app.services import export

    with open(export.__file__, encoding="utf-8") as fh:
        src = fh.read()
    assert '"pptx"' in src, "export maps pptx rebuild failures like docx/xlsx"
    for doc in ("app/capabilities/documents.md", "llms.txt", "docs/CLI.md"):
        with open(doc, encoding="utf-8") as fh:
            body = fh.read()
        assert "pptx" in body and "xlsm" in body and "docm" in body, doc
    print("  [A] 6 PowerPoint + 3 Word + 3 Excel extensions registered on every surface OK")


def test_b_parse_walk():
    units = _units(make_pptx())
    got = [(u.style, re.sub(r"⟦/?[a-z]\d+⟧", "", u.text)) for u in units]
    assert got == [
        ("slide 1", "Quarterly results"),
        ("slide 1", "Sales grew twelve percent see the reportPage "),
        ("slide 1", "Grouped caption"),
        ("slide 1", "Region"),
        ("slide 1", "North"),
        ("slide 1", "42"),
        ("slide 1", "Choice text"),
        ("slide 1 SmartArt", "Plan"),
        ("slide 1 SmartArt", "Build"),
        ("slide 1 chart", "Revenue by region"),
        ("slide 1 notes", "Mention the new office."),
        ("slide 2 hidden", "Backup slide"),
        ("a11y:img-alt", "A rising bar chart"),
        ("a11y:img-title", "Growth"),
    ], got
    texts = " ".join(t for _s, t in got)
    for absent in (
        "hidden merged",
        "Fallback text",
        "Click to edit",
        "Series cache",
        "10/10/2026",
        "Fallback alt",
    ):
        assert absent not in texts, absent
    multi = units[1].text
    assert multi.startswith("⟦r1⟧Sales grew ⟦/r1⟧⟦r2⟧twelve percent⟦/r2⟧"), multi
    assert multi.count("⟦x") == 2, ("a:br and a:fld are points", multi)
    skel = get_filter("pptx").parse(make_pptx()).skeleton
    assert skel["format"] == "pptx" and skel["pptx_walk"] == 1
    assert skel["units"][-3]["part"] == "ppt/slides/slide2.xml"
    print("  [B] slides, groups, tables, SmartArt, charts, notes, hidden, alt text in order OK")


def _translate_all(data, fn=lambda t: t):
    return {str(u.p_index): fn(u.text) for u in _units(data)}


def _tx(text):
    """Translate the words, keep every token."""
    return re.sub(r"(⟦/?[a-z]\d+⟧)|([^⟦⟧]+)", lambda m: m.group(1) or "FR " + m.group(2), text)


def test_c_byte_identity():
    for kind in ("pptx", "pptm"):
        src = make_pptx(kind, vba=kind == "pptm")
        out = get_filter("pptx").build(src, {})
        assert _members(out) == _members(src), "no translation: every member keeps its bytes"
        built = get_filter("pptx").build(src, _translate_all(src, _tx))
        a, b = _members(src), _members(built)
        assert list(a) == list(b), "member order kept"
        changed = sorted(n for n in a if a[n] != b[n])
        assert changed == [
            "ppt/charts/chart1.xml",
            "ppt/diagrams/data1.xml",
            "ppt/diagrams/drawing1.xml",
            "ppt/notesSlides/notesSlide1.xml",
            "ppt/slides/slide1.xml",
            "ppt/slides/slide2.xml",
        ], changed
        assert b["[Content_Types].xml"] == a["[Content_Types].xml"]
        if kind == "pptm":
            assert b["ppt/vbaProject.bin"] == VBA_BYTES
        assert b["ppt/slides/slide1.xml"].startswith(
            b"<?xml version='1.0' encoding='UTF-8' standalone='yes'?>"
        ), b["ppt/slides/slide1.xml"][:80]
    print("  [C] untranslated build byte-identical; translated build touches only text parts OK")


def test_d_inline_fidelity():
    src = make_pptx()
    f = get_filter("pptx")
    units = _units(src)
    body = units[1]
    tr = _tx(body.text)
    out = f.build(src, {str(body.p_index): tr})
    again = _units(out)
    assert again[1].text == tr, (again[1].text, tr)
    slide = _members(out)["ppt/slides/slide1.xml"].decode()
    assert '<a:rPr lang="en-US" b="1"/><a:t>FR twelve percent</a:t>' in slide, slide
    assert '<a:hlinkClick r:id="rId9"/></a:rPr><a:t>FR the report</a:t>' in slide
    assert 'type="slidenum"' in slide and "<a:br>" in slide
    # the translator dropped every point and every run token: nothing is deleted
    out2 = f.build(src, {str(body.p_index): "Tout en un\nDeux lignes"})
    s2 = _members(out2)["ppt/slides/slide1.xml"].decode()
    para = s2[s2.index("Tout en un") - 200 : s2.index("a:endParaRPr", s2.index("Tout en un"))]
    assert para.count("<a:br>") == 2, ("the newline became a:br and the old a:br kept", para)
    assert 'type="slidenum"' in para, "a dropped field is appended, not deleted"
    assert len(_units(out2)) == len(units)
    # escaping: markup characters in a target are text, not XML
    out3 = f.build(src, {"0": "R&D <b>"})
    assert "R&amp;D &lt;b&gt;" in _members(out3)["ppt/slides/slide1.xml"].decode()
    assert _units(out3)[0].text == "R&D <b>"
    print("  [D] run formatting, hyperlink, break and field survive; dropped points kept OK")


def test_e_smartart_cache():
    src = make_pptx()
    units = _units(src)
    plan = next(u for u in units if u.text == "Plan")
    out = get_filter("pptx").build(src, {str(plan.p_index): "Planifier"})
    drawing = _members(out)["ppt/diagrams/drawing1.xml"].decode()
    assert "Planifier" in drawing and ">Plan<" not in drawing and ">Build<" in drawing, drawing
    # a cache that no longer reads like the source is left alone
    m = _members(src)
    stale = m["ppt/diagrams/drawing1.xml"].replace(b">Plan<", b">Edited by hand<")
    src2 = make_pptx(extra_parts={"ppt/diagrams/drawing1.xml": stale})
    out2 = get_filter("pptx").build(src2, {str(plan.p_index): "Planifier"})
    assert b"Edited by hand" in _members(out2)["ppt/diagrams/drawing1.xml"]
    print("  [E] SmartArt drawing cache mirrors translated points (stale cache untouched) OK")


def test_f_language():
    src = make_pptx()
    f = get_filter("pptx")
    out = f.build_localized(src, {"0": "Resultats"}, "fr-FR")
    slide = _members(out)["ppt/slides/slide1.xml"].decode()
    assert '<a:rPr lang="fr-FR"/><a:t>Resultats</a:t>' in slide, slide[:1500]
    assert slide.count('lang="fr-FR"') == 1, "untranslated paragraphs keep their language"
    out_ar = f.build_localized(src, {"0": "نتائج"}, "ar")
    s_ar = _members(out_ar)["ppt/slides/slide1.xml"].decode()
    assert '<a:pPr rtl="1"/><a:r><a:rPr lang="ar"/>' in s_ar, s_ar[:1500]
    assert f.build_localized(src, {}, "fr-FR") == f.build(src, {})
    print("  [F] build_localized stamps lang on translated runs and rtl for Arabic OK")


def test_g_variants_keep_types():
    for kind in PPT_EXTS:
        vba = kind.endswith("m")
        src = make_pptx(kind, vba=vba)
        out = get_filter(kind).build_localized(src, _translate_all(src, _tx), "de-DE")
        m = _members(out)
        assert MAIN_TYPES[kind].encode() in m["[Content_Types].xml"], kind
        assert m["[Content_Types].xml"] == _members(src)["[Content_Types].xml"]
        if vba:
            assert m["ppt/vbaProject.bin"] == VBA_BYTES, kind
        assert len(_units(out, kind)) == len(_units(src, kind)), kind
    print("  [G] pptx/pptm/potx/potm/ppsx/ppsm keep content type and VBA bytes OK")


def test_h_hostile():
    from app import config

    f = get_filter("pptx")
    assert "not a valid PowerPoint" in _refused(f.parse, b"not a zip at all")
    # zip bomb: one part inflating far past the ratio limit
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, body in _members(make_pptx()).items():
            zf.writestr(name, body)
        zf.writestr("ppt/media/zeros.bin", b"\0" * (config.ZIP_RATIO_MIN_BYTES * 1000 + 1))
    assert "compression bomb" in _refused(f.parse, bomb.getvalue())
    # DOCTYPE / entity expansion in a slide
    lol = (
        '<?xml version="1.0"?><!DOCTYPE p:sld [<!ENTITY a "aaaaaaaaaa">'
        '<!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">]>' + SLIDE1.split("?>", 1)[1]
    )
    assert "document type" in _refused(f.parse, make_pptx(slide1=lol))
    xxe = (
        '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]>'
        + SLIDE1.split("?>", 1)[1]
    )
    assert "document type" in _refused(f.parse, make_pptx(slide1=xxe))
    # a relationship that escapes the package root is ignored, never resolved
    esc = make_pptx(
        extra_slide_rels=[
            (
                "rId20",
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide",
                "../../../../outside.xml",
                "",
            )
        ]
    )
    assert len(_units(esc)) == len(_units(make_pptx()))
    # no presentation part (a Word package renamed .pptx)
    from docx import Document

    buf = io.BytesIO()
    Document().save(buf)
    assert "not a PowerPoint presentation" in _refused(f.parse, buf.getvalue())
    # markup budget and segment cap
    old = config.MAX_DOCUMENT_SEGMENTS
    try:
        config.MAX_DOCUMENT_SEGMENTS = 5
        assert "formatting runs or markup" in _refused(f.parse, make_pptx())
        config.MAX_DOCUMENT_SEGMENTS = 18
        many = SLIDE1.replace(
            "</p:spTree>",
            "<!-- -->"
            + "".join(
                f"<p:sp><p:txBody><a:p><a:r><a:t>x{i}</a:t></a:r></a:p></p:txBody></p:sp>"
                for i in range(3)
            )
            + "</p:spTree>",
            1,
        )
        msg = _refused(f.parse, make_pptx(slide1=many))
        assert "segments" in msg, msg
    finally:
        config.MAX_DOCUMENT_SEGMENTS = old
    print("  [H] not-a-zip, bomb, DOCTYPE/XXE, escaping rels, wrong package, caps refused OK")


# ----------------------------------------------------------------- Word aliases

WORD_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
    "dotx": "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
    "docm": "application/vnd.ms-word.document.macroEnabled.main+xml",
    "dotm": "application/vnd.ms-word.template.macroEnabledTemplate.main+xml",
}


def _content_map(ct_xml: bytes) -> dict:
    """Part name -> content type, resolving Defaults by extension (what a reader sees)."""
    defaults = dict(re.findall(rb'<Default Extension="([^"]+)" ContentType="([^"]+)"', ct_xml))
    overrides = dict(re.findall(rb'<Override PartName="([^"]+)" ContentType="([^"]+)"', ct_xml))
    return {"defaults": {k.lower(): v for k, v in defaults.items()}, "overrides": overrides}


def _resolve_types(data: bytes) -> dict:
    m = _members(data)
    cm = _content_map(m["[Content_Types].xml"])
    out = {}
    for name in m:
        if name == "[Content_Types].xml":
            continue
        key = ("/" + name).encode()
        ext = name.rsplit(".", 1)[-1].lower().encode()
        out[name] = cm["overrides"].get(key) or cm["defaults"].get(ext)
    return out


def make_word(kind: str) -> bytes:
    """A python-docx document re-typed as `kind`, with a VBA project for the macro types."""
    from docx import Document

    d = Document()
    d.add_paragraph("Hello macro world")
    p = d.add_paragraph("Plain then ")
    p.add_run("bold").bold = True
    buf = io.BytesIO()
    d.save(buf)
    m = _members(buf.getvalue())
    ct = m["[Content_Types].xml"].replace(WORD_TYPES["docx"].encode(), WORD_TYPES[kind].encode())
    if kind.endswith("m"):
        ct = ct.replace(
            b"</Types>",
            b'<Override PartName="/word/vbaProject.bin" '
            b'ContentType="application/vnd.ms-office.vbaProject"/></Types>',
        )
        m["word/vbaProject.bin"] = VBA_BYTES
        rels = m["word/_rels/document.xml.rels"]
        m["word/_rels/document.xml.rels"] = rels.replace(
            b"</Relationships>",
            b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/'
            b'relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
        )
    m["[Content_Types].xml"] = ct
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, body in m.items():
            zf.writestr(name, body)
    return out.getvalue()


def test_i_word_aliases():
    from docx import Document

    plain = make_word("docx")
    base_units = [u.text for u in _units(plain, "docx")]
    for kind in WORD_EXTS:
        src = make_word(kind)
        f = get_filter(kind)
        parsed = f.parse(src)
        assert [u.text for u in parsed.units] == base_units, (kind, parsed.units)
        assert parsed.skeleton["docx_walk"] == get_filter("docx").parse(plain).skeleton["docx_walk"]
        tr = {str(u.p_index): _tx(u.text) for u in parsed.units}
        out = f.build_localized(src, tr, "fr-FR")
        types_in, types_out = _resolve_types(src), _resolve_types(out)
        assert types_out == types_in, (kind, types_in, types_out)
        assert any(v == WORD_TYPES[kind].encode() for v in types_out.values()), kind
        if kind.endswith("m"):
            assert _members(out)["word/vbaProject.bin"] == VBA_BYTES, kind
            assert b"vbaProject" in _members(out)["word/_rels/document.xml.rels"]
        assert [u.text for u in f.parse(out).units] == [tr[str(u.p_index)] for u in parsed.units]
    # the plain .docx path is unchanged, and python-docx's refusal text for a non-Word main
    # part is kept
    assert Document(io.BytesIO(plain)).paragraphs[0].text == "Hello macro world"
    from app.filters.docx_filter import DocxDocument

    try:
        DocxDocument(io.BytesIO(make_pptx()))
    except ValueError as exc:
        assert "is not a Word file" in str(exc), exc
    else:
        raise AssertionError("a presentation is not a Word file")
    print("  [I] .docm/.dotx/.dotm round-trip through the docx walk; types and VBA kept OK")


# ----------------------------------------------------------------- Excel aliases

XL_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
    "xltx": "application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml",
    "xlsm": "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
    "xltm": "application/vnd.ms-excel.template.macroEnabled.main+xml",
}


def make_excel(kind: str) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = "Hello sheet"
    wb.active["A2"] = 42
    buf = io.BytesIO()
    wb.save(buf)
    m = _members(buf.getvalue())
    ct = m["[Content_Types].xml"].replace(XL_TYPES["xlsx"].encode(), XL_TYPES[kind].encode())
    if kind.endswith("m"):
        ct = ct.replace(
            b"</Types>",
            b'<Override PartName="/xl/vbaProject.bin" '
            b'ContentType="application/vnd.ms-office.vbaProject"/></Types>',
        )
        m["xl/vbaProject.bin"] = VBA_BYTES
        rels = m["xl/_rels/workbook.xml.rels"]
        m["xl/_rels/workbook.xml.rels"] = rels.replace(
            b"</Relationships>",
            b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/'
            b'relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
        )
    m["[Content_Types].xml"] = ct
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, body in m.items():
            zf.writestr(name, body)
    return out.getvalue()


def test_j_excel_aliases():
    for kind in ("xlsx", *XL_EXTS):
        src = make_excel(kind)
        f = get_filter(kind)
        units = f.parse(src).units
        assert [u.text for u in units] == ["Hello sheet"], (kind, units)
        out = f.build_localized(src, {"0": "Bonjour"}, "fr-FR")
        m = _members(out)
        ct = m["[Content_Types].xml"]
        assert XL_TYPES[kind].encode() in ct, (kind, ct)
        for other, t in XL_TYPES.items():
            if other != kind:
                assert t.encode() not in ct, (kind, "became", other)
        if kind.endswith("m"):
            assert m["xl/vbaProject.bin"] == VBA_BYTES, kind
            assert b"vbaProject" in m["xl/_rels/workbook.xml.rels"], kind
        else:
            assert "xl/vbaProject.bin" not in m, kind
        assert [u.text for u in f.parse(out).units] == ["Bonjour"]
    print("  [J] .xlsm/.xltm keep VBA + macro type; .xltx stays template; .xlsx stays xlsx OK")


def test_k_ingest_export():
    from sqlmodel import Session, select

    from app.db import engine, init_db
    from app.models import Customer, Project, Segment, SegStatus, TargetLocale, TransMode
    from app.services import export as export_svc
    from app.services.ingest import ingest_document

    init_db()
    with Session(engine) as s:
        c = Customer(name="B4 Cust")
        s.add(c)
        s.commit()
        s.refresh(c)
        p = Project(
            name="b4", org_id=1, source_locale="en-US", default_provider="mock", customer_id=c.id
        )
        s.add(p)
        s.commit()
        s.refresh(p)
        s.add(TargetLocale(project_id=p.id, locale="fr-FR"))
        s.commit()
        cases = [("deck.pptm", make_pptx("pptm", vba=True)), ("memo.docm", make_word("docm"))]
        for name, data in cases:
            doc = ingest_document(s, p, name, data, TransMode.SEGMENTED)
            s.commit()
            s.refresh(doc)
            assert doc.fmt == name.rsplit(".", 1)[1], doc.fmt
            segs = s.exec(
                select(Segment).where(
                    Segment.document_id == doc.id, Segment.target_locale == "fr-FR"
                )
            ).all()
            assert segs, name
            for seg in segs:
                seg.target_text = _tx(seg.source_text)
                seg.status = SegStatus.MT
                s.add(seg)
            s.commit()
            out = export_svc.build_translated_file(s, doc, "fr-FR")
            m = _members(out)
            if name.endswith("pptm"):
                assert m["ppt/vbaProject.bin"] == VBA_BYTES
                assert m["[Content_Types].xml"] == _members(data)["[Content_Types].xml"]
                assert b'lang="fr-FR"' in m["ppt/slides/slide1.xml"]
                assert b"FR Quarterly results" in m["ppt/slides/slide1.xml"]
                # a stored walk stamp this build does not know is refused, never guessed
                skel = dict(doc.skeleton)
                skel["pptx_walk"] = 99
                doc.skeleton = skel
                s.add(doc)
                s.commit()
                try:
                    export_svc.build_translated_file(s, doc, "fr-FR")
                except export_svc.ExportUnreadableBlob as exc:
                    assert "walk 99" in str(exc), exc
                else:
                    raise AssertionError("unknown pptx walk stamp must be refused")
            else:
                assert m["word/vbaProject.bin"] == VBA_BYTES
                assert _resolve_types(out) == _resolve_types(data)
                assert b"FR Hello macro world" in m["word/document.xml"]
    print("  [K] .pptm and .docm ingest keep their extension as fmt; export keeps VBA OK")


def test_l_disguised_macro_refused():
    """Review gate (Batch 4 decision): a macro-enabled package under a NON-macro name
    (.docx/.xlsx/.pptx, .dotx/.xltx/.potx/.ppsx) is refused at the ingest chokepoint and in
    the Microsoft 365 connector; the same bytes under the macro name are accepted. A
    template under the document name (no macros) is accepted."""
    from sqlmodel import Session

    from app.db import engine, init_db
    from app.filters.zip_guard import MacroExtensionMismatch, refuse_disguised_macro
    from app.models import Customer, Project, TargetLocale, TransMode
    from app.services.ingest import UnsupportedFormat, ingest_document

    word_m, xl_m, ppt_m = make_word("docm"), make_excel("xlsm"), make_pptx("pptm", vba=True)
    for data, bad, good in (
        (word_m, ("docx", "dotx"), "docm"),
        (xl_m, ("xlsx", "xltx"), "xlsm"),
        (ppt_m, ("pptx", "potx", "ppsx"), "pptm"),
    ):
        for ext in bad:
            try:
                refuse_disguised_macro(data, ext)
            except MacroExtensionMismatch as exc:
                assert "macros" in str(exc) and f".{ext}" in str(exc), exc
            else:
                raise AssertionError(f"macro package named .{ext} must be refused")
        refuse_disguised_macro(data, good)
    # a VBA part alone (main part re-typed as plain) is still a macro package
    m = _members(word_m)
    m["[Content_Types].xml"] = m["[Content_Types].xml"].replace(
        WORD_TYPES["docm"].encode(), WORD_TYPES["docx"].encode()
    )
    sneaky = io.BytesIO()
    with zipfile.ZipFile(sneaky, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, body in m.items():
            zf.writestr(name, body)
    try:
        refuse_disguised_macro(sneaky.getvalue(), "docx")
    except MacroExtensionMismatch:
        pass
    else:
        raise AssertionError("a .docx carrying vbaProject.bin must be refused")
    refuse_disguised_macro(make_word("dotx"), "docx")  # template, no macros: harmless
    refuse_disguised_macro(make_word("docx"), "docx")

    init_db()
    with Session(engine) as s:
        c = Customer(name="B4 Gate Cust")
        s.add(c)
        s.commit()
        s.refresh(c)
        p = Project(
            name="b4-gate",
            org_id=1,
            source_locale="en-US",
            default_provider="mock",
            customer_id=c.id,
        )
        s.add(p)
        s.commit()
        s.refresh(p)
        s.add(TargetLocale(project_id=p.id, locale="fr-FR"))
        s.commit()
        for name, data in (("memo.docx", word_m), ("book.xlsx", xl_m), ("deck.pptx", ppt_m)):
            try:
                ingest_document(s, p, name, data, TransMode.SEGMENTED)
            except UnsupportedFormat as exc:
                assert "macros" in str(exc), exc
                s.rollback()
            else:
                raise AssertionError(f"{name} with macros must be refused at ingest")
        doc = ingest_document(s, p, "plan.docx", make_word("dotx"), TransMode.SEGMENTED)
        assert doc.fmt == "docx"
        s.rollback()

    from app.services.connectors import msgraph

    ConnectorError = msgraph.ConnectorError
    load = msgraph.MsGraphConnector.load

    class _Client:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    class _Conn:
        def _client(self):
            return _Client()

        def download(self, c, iid, cap):
            return word_m

        def max_bytes(self):
            return 10**8

    try:
        load(_Conn(), {"fmt": "docx", "id": "1", "name": "memo.docx"})
    except ConnectorError as exc:
        assert "macros" in str(exc), exc
    else:
        raise AssertionError("msgraph must refuse a macro package named .docx")
    units, _data, _f = load(_Conn(), {"fmt": "docx", "id": "1", "name": "memo.docm"})
    assert units
    print("  [L] macro package under a non-macro name refused (ingest + M365); template OK")


if __name__ == "__main__":
    test_a_registry_and_surfaces()
    test_b_parse_walk()
    test_c_byte_identity()
    test_d_inline_fidelity()
    test_e_smartart_cache()
    test_f_language()
    test_g_variants_keep_types()
    test_h_hostile()
    test_i_word_aliases()
    test_j_excel_aliases()
    test_k_ingest_export()
    test_l_disguised_macro_refused()
    print("PPTX FILTER OK")
