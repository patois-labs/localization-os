"""Bold / italic / size survive a docx translation into a complex-script language (#722).

Format soak round 9: Word keeps separate run properties for complex-script text (Arabic,
Hebrew, Persian, Thai, Hindi, ...): `w:b` / `w:i` / `w:sz` apply only to non-complex-script
characters, `w:bCs` / `w:iCs` / `w:szCs` to complex-script ones (ECMA-376 Part 1, 17.3.2).
An English run with only `w:b` exported its Arabic translation regular weight (LibreOffice read
the run as bold for Latin and East Asian text only).

  twins     build_localized into ar-SA / he-IL / fa-IR / th-TH / hi-IN gives each translated
            run with w:b / w:i / w:sz the missing w:bCs / w:iCs / w:szCs, same value;
  kept      a run that already has its twin keeps it (an explicit w:bCs w:val="0" stays);
  latin     fr-FR / de-DE / ja-JP exports add no twin.

Hermetic (no engine, no server). ASCII output; final line "DOCX CS FORMAT 722 OK".
"""

import io
import os
import sys

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_docxcs722_")

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _fail(msg):
    print("FAIL:", str(msg).encode("ascii", "backslashreplace").decode("ascii"))
    sys.exit(1)


def _source() -> bytes:
    from docx import Document
    from docx.oxml.parser import OxmlElement
    from docx.shared import Pt

    d = Document()
    p = d.add_paragraph()
    p.add_run("Click the ")
    p.add_run("file").bold = True
    r = p.add_run(" to open it, ")
    r.italic = True
    big = p.add_run("today")
    big.font.size = Pt(16)
    # a run that already has its twin, explicitly off: must stay as it is
    p2 = d.add_paragraph()
    keep = p2.add_run("Keep")
    keep.bold = True
    off = OxmlElement("w:bCs")
    off.set(f"{W}val", "0")
    keep._r.get_or_add_rPr().find(f"{W}b").addnext(off)
    p2.add_run(" this as it is.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


TARGETS = {
    0: "⟦r1⟧انقر على ⟦/r1⟧⟦r2⟧الملف⟦/r2⟧⟦r3⟧ لفتحه، ⟦/r3⟧⟦r4⟧اليوم⟦/r4⟧",
    1: "⟦r1⟧احتفظ⟦/r1⟧⟦r2⟧ بهذا كما هو.⟦/r2⟧",
}


def _runs(out: bytes):
    from docx import Document

    doc = Document(io.BytesIO(out))
    return [[r._r for r in par.runs] for par in doc.paragraphs]


def _prop(r, name):
    rpr = r.find(f"{W}rPr")
    el = None if rpr is None else rpr.find(f"{W}{name}")
    if el is None:
        return None
    return el.get(f"{W}val", "on")


def test_twins_and_kept():
    from app.filters import get_filter

    f = get_filter("docx")
    src = _source()
    units = f.parse(src).units
    idx = {u.p_index for u in units}
    if idx != {0, 1}:
        _fail(f"unexpected units: {[u.text for u in units]!r}")
    for loc in ("ar-SA", "he-IL", "fa-IR", "th-TH", "hi-IN"):
        out = f.build_localized(src, dict(TARGETS), loc)
        p0, p1 = _runs(out)
        bold, ital, big = p0[1], p0[2], p0[3]
        if _prop(bold, "bCs") != "on":
            _fail(f"[twins] {loc}: bold run has no w:bCs (complex-script text not bold)")
        if _prop(ital, "iCs") != "on":
            _fail(f"[twins] {loc}: italic run has no w:iCs")
        if _prop(big, "szCs") != _prop(big, "sz") or _prop(big, "sz") != "32":
            _fail(f"[twins] {loc}: sized run szCs {_prop(big, 'szCs')} != sz {_prop(big, 'sz')}")
        if _prop(p0[0], "bCs") is not None:
            _fail(f"[twins] {loc}: a plain run must not gain properties")
        if _prop(p1[0], "bCs") != "0":
            _fail(f"[kept] {loc}: an explicit w:bCs w:val=0 was changed: {_prop(p1[0], 'bCs')}")
        names = [c.tag.replace(W, "") for c in bold.find(f"{W}rPr")]
        if names.index("bCs") != names.index("b") + 1:
            _fail(f"[twins] {loc}: twin out of schema order: {names}")
    print("  [twins] ar / he / fa / th / hi: bCs, iCs, szCs added; explicit twin kept OK")


def test_latin_targets_unchanged():
    from app.filters import get_filter

    f = get_filter("docx")
    src = _source()
    fr = {
        0: "⟦r1⟧Cliquez sur le ⟦/r1⟧⟦r2⟧fichier⟦/r2⟧⟦r3⟧ pour l'ouvrir, ⟦/r3⟧⟦r4⟧aujourd'hui⟦/r4⟧"
    }
    for loc in ("fr-FR", "de-DE", "ja-JP"):
        out = f.build_localized(src, fr, loc)
        p0 = _runs(out)[0]
        if any(_prop(r, n) is not None for r in p0 for n in ("bCs", "iCs", "szCs")):
            _fail(f"[latin] {loc}: complex-script twins added to a non-complex-script target")
    print("  [latin] fr / de / ja: no twin added OK")


def main():
    test_twins_and_kept()
    test_latin_targets_unchanged()
    print("DOCX CS FORMAT 722 OK")


if __name__ == "__main__":
    main()
