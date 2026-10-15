"""Chunk 1 verification — runs offline against the app via TestClient.

Builds a sample .docx in memory, uploads it in both modes, and asserts:
  - DOCX parsing extracts paragraph units
  - segmented mode splits into sentence segments
  - full-text mode produces a single whole-document segment
  - a chunk plan is produced and flags an oversized paragraph
  - inline tags are detected
  - the DOCX filter round-trips (parse -> build) preserving structure
"""

import os
import sys
from io import BytesIO

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)

# Run against a throwaway data dir so we never touch the real loc_tms.db
from tests._bootstrap import init_env

init_env("loctms_test_")

from docx import Document as Docx
from fastapi.testclient import TestClient

from app.db import init_db
from app.filters import get_filter
from app.main import app


def make_docx() -> bytes:
    d = Docx()
    d.add_heading("Spring Sale", level=1)
    d.add_paragraph("Adventure awaits. Save up to 50% this weekend only.")
    d.add_paragraph("Use code {PROMO} at checkout. Visit https://example.com for details.")
    d.add_paragraph("Empty-ish.")
    # one very large paragraph to trip the oversized-chunk detector
    d.add_paragraph(("This is a long sentence about localization. " * 600).strip())
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def main() -> None:
    init_db()
    client = TestClient(app)

    # project with two target locales
    r = client.post(
        "/projects",
        data={
            "name": "Chunk1 Test",
            "source_locale": "en-US",
            "target_locales": "fr-FR, de-DE",
            "default_provider": "claude",
        },
        follow_redirects=False,
    )
    assert r.status_code == 303, r.status_code
    pid = int(r.headers["location"].rsplit("/", 1)[-1])

    docx_bytes = make_docx()

    # --- segmented upload ---
    r = client.post(
        f"/projects/{pid}/documents",
        data={"mode": "segmented"},
        files={
            "file": (
                "spring.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    assert r.status_code == 303, r.text
    seg_doc_id = int(r.headers["location"].rsplit("/", 1)[-1])

    page = client.get(f"/documents/{seg_doc_id}?locale=fr-FR").text
    assert "Adventure awaits." in page, "source segment not shown"
    assert "Save up to 50% this weekend only." in page, "sentence split failed"
    assert "{PROMO}" in page, "inline tag not detected/shown"
    assert "oversized" in page, "oversized chunk not flagged"

    # inspect DB directly
    from sqlmodel import Session, select

    from app.db import engine
    from app.models import Segment

    with Session(engine) as s:
        seg_fr = s.exec(
            select(Segment).where(
                Segment.document_id == seg_doc_id, Segment.target_locale == "fr-FR"
            )
        ).all()
        seg_de = s.exec(
            select(Segment).where(
                Segment.document_id == seg_doc_id, Segment.target_locale == "de-DE"
            )
        ).all()
    # "Adventure awaits." + "Save up to 50%..." => the first paragraph alone is 2 sentences
    assert len(seg_fr) >= 5, f"expected several sentence segments, got {len(seg_fr)}"
    assert len(seg_fr) == len(seg_de), "segments should exist per target locale"
    print(f"  segmented: {len(seg_fr)} segments x 2 locales")

    # --- full-text upload ---
    r = client.post(
        f"/projects/{pid}/documents",
        data={"mode": "full_text"},
        files={
            "file": (
                "spring.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    assert r.status_code == 303, r.text
    ft_doc_id = int(r.headers["location"].rsplit("/", 1)[-1])

    with Session(engine) as s:
        ft_fr = s.exec(
            select(Segment).where(
                Segment.document_id == ft_doc_id, Segment.target_locale == "fr-FR"
            )
        ).all()
    assert len(ft_fr) == 1, f"full-text should be ONE segment, got {len(ft_fr)}"
    assert "Adventure awaits." in ft_fr[0].source_text
    assert "Spring Sale" in ft_fr[0].source_text
    print(f"  full-text: {len(ft_fr)} segment, {len(ft_fr[0].source_text)} chars")

    # --- filter round-trip (parse -> build) ---
    filt = get_filter("docx")
    parsed = filt.parse(docx_bytes)
    assert len(parsed.units) >= 4, "expected >=4 non-empty paragraphs"
    # translate the first unit and rebuild
    target = {str(parsed.units[0].p_index): "Vente de printemps"}
    rebuilt = filt.build(docx_bytes, target)
    re_doc = Docx(BytesIO(rebuilt))
    texts = [p.text for p in re_doc.paragraphs]
    assert "Vente de printemps" in texts, "round-trip build did not inject translation"
    print(f"  round-trip: rebuilt docx, {len(texts)} paragraphs, injection OK")

    # --- multi-run inline-formatting fidelity (Phase 5 chunk 6) ---
    md = Docx()
    md.add_paragraph("A plain single-run paragraph.")
    mp = md.add_paragraph("Please click ")
    mp.add_run("Submit").bold = True
    mp.add_run(" to continue.")
    mbuf = BytesIO()
    md.save(mbuf)
    mdocx = mbuf.getvalue()

    mparsed = filt.parse(mdocx)
    # single-run paragraph stays plain text (no tokens, no regression); the
    # multi-run paragraph carries paired run tokens + rPr side-channel.
    single = next(u for u in mparsed.units if u.text == "A plain single-run paragraph.")
    assert single.inline_map == {}, "single-run paragraph should carry no inline_map"
    multi = next(u for u in mparsed.units if "Submit" in u.text)
    assert multi.inline_map, "multi-run paragraph should carry an inline_map"
    # token-stripped text is the original readable sentence
    from app.services.segment import find_placeholders, strip_placeholders

    assert strip_placeholders(multi.text) == "Please click Submit to continue."
    assert len(find_placeholders(multi.text)) == 6, "expected 3 paired run tokens (6 tokens)"

    # translate (mock keeps tokens verbatim) and rebuild
    mmap = {str(u.p_index): "[FR] " + u.text for u in mparsed.units}
    mrebuilt = filt.build(mdocx, mmap)
    rmd = Docx(BytesIO(mrebuilt))
    mpara = next(p for p in rmd.paragraphs if "Submit" in p.text)
    # the bold run survived as its own run (formatting NOT collapsed to first run)
    bolds = [r.text for r in mpara.runs if r.bold]
    assert bolds == ["Submit"], f"bold run not preserved: {[(r.text, r.bold) for r in mpara.runs]}"
    assert mpara.text == "[FR] Please click Submit to continue.", repr(mpara.text)
    # more than one run => mid-paragraph formatting kept (old code collapsed to 1)
    assert len([r for r in mpara.runs if r.text]) >= 3, "runs collapsed"
    print(f"  inline fidelity: multi-run docx kept {len(mpara.runs)} runs, bold='{bolds[0]}'")

    # --- hyperlink paragraph: a paragraph whose text lives partly OUTSIDE its direct
    # runs (here in a <w:hyperlink>) must lose no text and duplicate none. It used to
    # fall back to untokenized whole-paragraph handling, which deleted the hyperlink on
    # export (docx-fidelity2, 2026-09-26); it is now tokenized per run like any
    # multi-run paragraph and the hyperlink is re-created around its translated run
    # (verify_docx_po_fidelity pins the link targets).
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    hd = Docx()
    hp = hd.add_paragraph("Please click ")
    hlink_run = hp.add_run("here")  # this run will live in a <w:hyperlink>
    hl = OxmlElement("w:hyperlink")
    hp._p.remove(hlink_run._r)
    hl.append(hlink_run._r)
    hp._p.append(hl)
    hp.add_run(" to continue.")
    hbuf = BytesIO()
    hd.save(hbuf)
    hdocx = hbuf.getvalue()

    hparsed = filt.parse(hdocx)
    hu = next(u for u in hparsed.units if "click" in u.text)
    # one run token pair per run, the hyperlink's run included
    assert hu.text == ("⟦r1⟧Please click ⟦/r1⟧⟦r2⟧here⟦/r2⟧⟦r3⟧ to continue.⟦/r3⟧"), repr(hu.text)
    assert len(find_placeholders(hu.text)) == 6, hu.text

    hrebuilt = filt.build(hdocx, {str(hu.p_index): "[FR] " + hu.text})
    rhd = Docx(BytesIO(hrebuilt))
    hpara = next(p for p in rhd.paragraphs if "click" in p.text)
    # the hyperlink's "here" must NOT be duplicated or dropped, and stays in the link
    assert hpara.text == "[FR] Please click here to continue.", repr(hpara.text)
    hls = hpara._p.findall(qn("w:hyperlink"))
    assert len(hls) == 1 and "".join(hls[0].xpath(".//w:t/text()")) == "here", [
        "".join(h.xpath(".//w:t/text()")) for h in hls
    ]
    print("  hyperlink paragraph: text survives round-trip (no loss/dup), link kept")

    # --- unsupported format rejected ---
    r = client.post(
        f"/projects/{pid}/documents",
        data={"mode": "segmented"},
        files={"file": ("archive.zip", b"PK\x03\x04hello", "application/zip")},
        follow_redirects=False,
    )
    assert r.status_code == 303 and "err=format" in r.headers["location"], "zip should be rejected"
    print("  unsupported format rejected OK")

    print("CHUNK 1 OK")


if __name__ == "__main__":
    main()
