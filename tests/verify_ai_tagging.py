"""AI Tagging (Phase 1) - deterministic inline-tag integrity verification.

Hermetic + offline (throwaway LOC_TMS_DATA_DIR, provider keys blanked -> no network,
MockProvider forced). Proves every Phase-1 component + its fail-pre-fix behavior, and
the review-round fixes:

  1. mask/restore round-trip: paired tokens never split; colliding per-paragraph ids
     across segments safe; feedback-buffer seam never leaks a sentinel; residual.
  1H. COLLISION-PROOFING (HIGH, reviewer #2): source prose that legitimately contains a
     fixed-shape sentinel is NEVER rewritten by restore, both reviewers' laundering
     repros end clean/quarantined, and the gate/editor guard never false-positive.
  2. RESTORED-forms-only (HIGH-1): a masked structural compare is blind; the restored
     gate catches the drop.
  3. gate fires on drop/add/dup/reorder/nesting/residual; byte-identical when no tags.
  4. repair rung: de-dup (clean, whitespace collapsed) + point-order (BEST-GUESS, flagged)
     before quarantine; dropped-span & paired-reorder -> quarantine; repairs re-verify.
  5. commit wiring: unrepairable -> NEEDS_REVIEW + tag_integrity + TM-ineligible; de-dup ->
     clean; point-order -> committed-but-FLAGGED; composes; byte-identical when intact.
  6. lqa 1.3 remap (+ git-show fail-pre-fix proof).
  7. bounded fix-loop (RunGuard total-step cap).
  8. FULL_TEXT export guard: TAGGED desync -> whole export quarantined; TAGLESS
     parts<units -> REFUSED loud (honest-failure, 2026-08-13; superseded M2's silent
     positional fallback).
  8M. project BUNDLE export skip-and-report (M1): a quarantined doc is excluded, the clean
     doc still ships, the skip is reported.
  9. editor save guard.
 10. batch verify CLI (flags-only).
 11. end-to-end masking: plain-MT + both orchestrator paths.

Final line: "AI TAGGING OK".
"""

import io
import os
import subprocess
import sys

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
# Defense in depth (B1): never die on a non-ASCII glyph under a cp1252 pipe. The prints
# below are ALSO written pure-ASCII; this belt covers assertion messages that embed the
# real tokens on a failure.
for _s in (sys.stdout, sys.stderr):
    try:  # noqa: SIM105
        _s.reconfigure(encoding="utf-8", errors="backslashreplace")
    except Exception:
        pass
from tests._bootstrap import init_env
from tests._piece_patch import set_pieces

init_env("loctms_aitag_")

from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app.db import engine, init_db
from app.main import app
from app.models import (
    Document,
    Project,
    RunStatus,
    Segment,
    SegStatus,
    StepType,
    TransMode,
    Workflow,
    WorkflowRun,
    WorkflowStep,
)
from app.services import (
    ai_tagging,
    delivery,
    export,
    lqa,
    orchestrator,
    qa,
    quality,
    segment,
    settings,
    tag_repair,
    translate,
)
from app.services.segment import _PLACEHOLDER_RE

init_db()

_SEQ = [0]
_BASE_COMMIT = "c2b3f49"  # main this branch was cut from (git-show proofs pin to it)

R1 = segment.ph_open("r", 1)
R1C = segment.ph_close("r", 1)
R2 = segment.ph_open("r", 2)
R2C = segment.ph_close("r", 2)
G1 = segment.ph_open("g", 1)
G1C = segment.ph_close("g", 1)
G2 = segment.ph_open("g", 2)
G2C = segment.ph_close("g", 2)
X1 = segment.ph_point("x", 1)
X2 = segment.ph_point("x", 2)
LAQ = chr(0xAB)  # <<  (a legitimate French guillemet in real prose)
RAQ = chr(0xBB)  # >>


def _seq() -> int:
    _SEQ[0] += 1
    return _SEQ[0]


def _git_show(path: str):
    try:
        out = subprocess.run(
            ["git", "show", f"{_BASE_COMMIT}:{path}"],
            cwd=_ROOT,
            capture_output=True,
            text=True,
            timeout=30,
        )
        return out.stdout if out.returncode == 0 else None
    except Exception:
        return None


# --- builders ---------------------------------------------------------------
def _mk_seg(session, *, source, target=None, status=SegStatus.NEW, qa_flags=None, locale="es-MX"):
    p = Project(name="aitag-%d" % _seq())  # noqa: UP031
    session.add(p)
    session.commit()
    d = Document(project_id=p.id, filename="d.docx")
    session.add(d)
    session.commit()
    seg = Segment(
        document_id=d.id,
        target_locale=locale,
        order_index=0,
        source_text=source,
        target_text=target,
        status=status,
        qa_flags=qa_flags,
    )
    session.add(seg)
    session.commit()
    return seg


def _mk_project(session):
    p = Project(
        name="aitag-proj-%d" % _seq(),  # noqa: UP031
        org_id=1,
        source_locale="en-US",
        default_provider="mock",  # noqa: RUF100, UP031
    )
    session.add(p)
    session.commit()
    session.refresh(p)
    return p


def _mk_doc(session, project, *, sources, mode=TransMode.SEGMENTED, locale="es-ES"):
    doc = Document(project_id=project.id, filename="d.txt", fmt="txt", mode=mode)
    session.add(doc)
    session.commit()
    session.refresh(doc)
    for i, s in enumerate(sources):
        session.add(Segment(document_id=doc.id, target_locale=locale, order_index=i, source_text=s))
    session.commit()
    return doc


def _mk_wf(session):
    wf = Workflow(name="aitag-wf-%d" % _seq(), scope="global")  # noqa: UP031
    session.add(wf)
    session.commit()
    session.refresh(wf)
    session.add(
        WorkflowStep(
            workflow_id=wf.id,
            order_index=0,
            name="Step0",
            step_type=StepType.LLM_AGENT,
            provider="mock",
            prompt_text="Translate: {text}",
            chunk_policy="none",
            retry_max=1,
            backoff_base=0.0,
        )
    )
    session.commit()
    return wf


def _mk_ft_doc(session, units, source_text):
    p = _mk_project(session)
    doc = Document(
        project_id=p.id,
        filename="d.docx",
        fmt="docx",
        mode=TransMode.FULL_TEXT,
        skeleton={"units": units, "seg_locale": "en-US"},
    )
    session.add(doc)
    session.commit()
    session.refresh(doc)
    session.add(
        Segment(document_id=doc.id, target_locale="es-ES", order_index=0, source_text=source_text)
    )
    session.commit()
    return doc.id


class _Engine:
    """Swap a module's run_completion for a fake whose output is fn(payload['text']) -
    the REAL mask/restore path runs (the engine sees generated sentinels; we control
    exactly how they are damaged/preserved)."""

    def __init__(self, module, fn):
        self.module = module
        self.fn = fn
        self._orig = None

    def __enter__(self):
        self._orig = self.module.run_completion

        def fake(spec, prompt, system=None, task=None, payload=None, **kw):
            txt = (payload or {}).get("text", "")
            return {
                "text": self.fn(txt),
                "provider_used": "mock",
                "model_used": "mock",
                "attempts": 1,
                "error": None,
                "usage": {},
            }

        set_pieces(self.module.__name__, "run_completion", fake)
        return self

    def __exit__(self, *exc):
        set_pieces(self.module.__name__, "run_completion", self._orig)
        return False


def _drop_sentinels(t):
    return ai_tagging._GEN_SENTINEL_RE.sub("", t)


def _dup_first_sentinel(t):
    m = ai_tagging._GEN_SENTINEL_RE.search(t)
    return t if not m else t[: m.start()] + m.group(0) + t[m.start() :]


# ===========================================================================
# 1. mask / restore round-trip
# ===========================================================================
def test_mask_restore():
    src = "Click " + R1 + "here" + R1C + " then " + X1 + " done."
    masked, mp = ai_tagging.mask(src)
    assert not _PLACEHOLDER_RE.search(masked), "no raw token reaches the engine"
    assert len(mp) == 3, mp
    restored, resid = ai_tagging.restore(masked, mp)
    assert restored == src and resid is False, (restored, resid)
    print(
        "  1a. mask -> opaque sentinels (no raw token reaches the engine); restore round-trips OK"
    )

    assert mp[[k for k, v in mp.items() if v == R1][0]] == R1  # noqa: RUF015
    assert mp[[k for k, v in mp.items() if v == R1C][0]] == R1C  # noqa: RUF015
    print("  1b. paired open/close each get their own sentinel (never split / renumbered) OK")

    s1, m1 = ai_tagging.mask("A " + R1 + "x" + R1C)
    s2, m2 = ai_tagging.mask("B " + R1 + "y" + R1C)
    assert ai_tagging.restore(s1, m1)[0].endswith(R1 + "x" + R1C)
    assert ai_tagging.restore(s2, m2)[0].endswith(R1 + "y" + R1C)
    print("  1c. colliding r1 across segments safe (segment-granular mapping) OK")

    fb_src, _ = ai_tagging.restore(s1, m1)
    assert not ai_tagging.has_residual_sentinel(fb_src)
    print("  1d. feedback-buffer seam: restored form carries no sentinel OK")

    # a mangled GENERATED sentinel (nonce-bearing, truncated) -> residual detected.
    _, m_r1 = ai_tagging.mask(R1)
    key = list(m_r1.keys())[0]  # noqa: RUF015
    # #726 gate: a dropped closing guillemet leaves a ONE-sided frame (the `«` could be the
    # author's): not guessed at, a residual. With BOTH guillemets dropped it restores.
    assert ai_tagging.restore("x " + key[:-1] + " y", m_r1)[1] is True
    assert ai_tagging.restore("x " + key[1:-1] + " y", m_r1) == ("x " + R1 + " y", False)
    mangled = "x " + key[:-2] + " y"  # truncated mid-id -> generated prefix survives
    _, resid2 = ai_tagging.restore(mangled, m_r1)
    assert resid2 is True and ai_tagging.has_residual_sentinel(mangled)
    assert ai_tagging.mask("no tokens here") == ("no tokens here", {})
    print("  1e. mangled generated sentinel detected as residual; tagless mask byte-identical OK")


# ===========================================================================
# 1H. COLLISION-PROOFING (HIGH, reviewer #2) + prose-mention (reviewer #1)
# ===========================================================================
def test_collision_proofing():
    # Source PROSE legitimately containing a fixed-shape sentinel AND a real token.
    prose = "Doc mentions " + LAQ + "TAG_1" + RAQ + " and formats " + R1 + "this" + R1C + "."
    masked, mp = ai_tagging.mask(prose)
    assert LAQ + "TAG_1" + RAQ in masked, "fixed-shape prose must be left intact by mask"
    restored, resid = ai_tagging.restore(masked, mp)
    assert restored == prose and resid is False, ("restore must not rewrite the prose", restored)
    print("  1H-a. fixed-shape sentinel in source prose survives mask+restore verbatim OK")

    # the gate / residual check do NOT false-positive on a fixed-shape sentinel as content.
    assert (
        quality.tag_damage(LAQ + "TAG_5" + RAQ + " is code", LAQ + "TAG_5" + RAQ + " es codigo")
        == ""
    )
    assert not ai_tagging.has_residual_sentinel(LAQ + "TAG_5" + RAQ + " literal prose")
    print("  1H-b. fixed-shape sentinel as content -> gate/residual do NOT false-positive OK")

    # Reviewer #2 laundering repro: real-token DROP is caught/quarantined, prose untouched
    # (WITHOUT the nonce this laundered to certified-clean-with-wrong-text).
    with Session(engine) as s:
        seg = _mk_seg(s, source=prose)
        restored_out, _ = ai_tagging.restore(_drop_sentinels(masked), mp)
        out = quality.commit_mt_output(s, seg, restored_out, source="mt")
        s.commit()
        assert out.tag_violation and seg.status == SegStatus.NEEDS_REVIEW, out
        assert LAQ + "TAG_1" + RAQ in seg.target_text, "prose mention preserved"
    print(
        "  1H-c. laundering repro: real-token drop -> quarantined (never certified-clean-wrong) OK"
    )

    # a CLEAN preserve of prose + token commits clean, unflagged.
    with Session(engine) as s:
        seg = _mk_seg(s, source=prose)
        clean_out, _ = ai_tagging.restore(masked, mp)
        out = quality.commit_mt_output(s, seg, clean_out, source="mt")
        s.commit()
        assert out.tag_violation == "" and seg.status == SegStatus.MT, out
        assert not [f for f in (seg.qa_flags or []) if f["type"] == "tag_integrity"]
    print("  1H-d. prose-mention + preserved token commits CLEAN, unflagged OK")


# ===========================================================================
# 2. RESTORED-forms-only (HIGH-1)
# ===========================================================================
def test_restored_forms_only():
    src = "Read " + R1 + "the title" + R1C + "."
    masked_src, mp = ai_tagging.mask(src)
    masked_tgt = _drop_sentinels(masked_src)
    assert not any(f["type"] == "placeholder" for f in qa.check_segment(masked_src, masked_tgt)), (
        "a MASKED compare must not see tokens (that is why we restore first)"
    )
    restored_tgt, _ = ai_tagging.restore(masked_tgt, mp)
    assert quality.tag_damage(src, restored_tgt), "restored-form gate MUST catch the drop"
    print("  2. HIGH-1: masked compare blind; restored-form gate catches the drop OK")


# ===========================================================================
# 3. the deterministic gate (quality.tag_damage)
# ===========================================================================
def test_tag_damage_unit():
    good = R1 + "Titulo" + R1C + " y " + X1
    assert quality.tag_damage(R1 + "Title" + R1C + " and " + X1, good) == ""
    assert quality.tag_damage("plain source", "[es] traduccion") == ""
    assert "dropped" in quality.tag_damage(R1 + "a" + R1C, "a")
    assert "added" in quality.tag_damage("a", "a" + X1)
    assert "duplicated" in quality.tag_damage(R1 + "a" + R1C, R1 + R1 + "a" + R1C)
    assert "reordered" in quality.tag_damage(X1 + "a" + X2, X2 + "a" + X1)
    crossed = quality.tag_damage(G1 + G2 + "a" + G2C + G1C, G1 + G2 + "a" + G1C + G2C)
    assert crossed, ("crossed pairs must be damage", crossed)
    # a GENERATED-shape residual sentinel is unconditional damage
    _, mp = ai_tagging.mask(R1 + "a" + R1C)
    gen_key = [k for k, v in mp.items() if v == R1][0]  # noqa: RUF015
    assert quality.tag_damage(R1 + "a" + R1C, gen_key + "a" + R1C) != ""
    print("  3. tag_damage: drop/add/dup/reorder/crossed/residual caught; no-tags -> '' OK")


# ===========================================================================
# 4. the repair rung (tag_repair) - returns (text, clean)
# ===========================================================================
def test_repair_unit():
    cand, clean = tag_repair.repair(R1 + "a" + R1C, R1 + R1 + "a" + R1C)
    assert cand == R1 + "a" + R1C and clean is True, (cand, clean)
    # de-dup collapses the doubled whitespace a removed surplus token leaves (LOW-1)
    cand, clean = tag_repair.repair("a " + X1 + " b", "a " + X1 + " " + X1 + " b")
    assert cand == "a " + X1 + " b" and clean is True, repr(cand)
    # point-order-restore: BEST-GUESS, NOT clean (MEDIUM)
    cand, clean = tag_repair.repair(X1 + "a" + X2, X2 + "a" + X1)
    assert cand == X1 + "a" + X2 and clean is False, (cand, clean)
    # dedup THEN reorder compose -> best-guess
    cand, clean = tag_repair.repair(X1 + "a" + X2, X2 + X2 + "a" + X1)
    assert cand == X1 + "a" + X2 and clean is False, (cand, clean)
    # UNRECOVERABLE: dropped paired span -> None (quarantine)
    assert tag_repair.repair(R1 + "a" + R1C, "a") is None
    # a PAIRED-token reorder is NEVER auto-fixed -> None
    assert (
        tag_repair.repair(R1 + "a" + R1C + R2 + "b" + R2C, R2 + "b" + R2C + R1 + "a" + R1C) is None
    )
    # a clean target needs no repair -> None
    assert tag_repair.repair(R1 + "a" + R1C, R1 + "b" + R1C) is None
    # every repair re-verifies structurally clean
    cand, _ = tag_repair.repair(X1 + "a" + X2, X2 + "a" + X1)
    assert not quality.tag_damage(X1 + "a" + X2, cand)
    print(
        "  4. repair: de-dup(clean, ws-collapsed) + point-order(best-guess); "
        "dropped-span/paired-reorder -> quarantine; repairs re-verify OK"
    )


# ===========================================================================
# 5. commit_mt_output wiring
# ===========================================================================
def test_commit_wiring():
    with Session(engine) as s:
        seg = _mk_seg(s, source="Read " + R1 + "this" + R1C + ".")
        out = quality.commit_mt_output(s, seg, "[es] Lee esto.", source="mt")
        s.commit()
        assert out.committed and not out.flagged
        assert seg.target_text == "[es] Lee esto."
        assert seg.status == SegStatus.NEEDS_REVIEW
        assert out.tag_violation and out.tm_eligible is False, out
        assert [f for f in seg.qa_flags if f["type"] == "tag_integrity"], seg.qa_flags
        print("  5a. commit: unrepairable drop -> NEEDS_REVIEW + tag_integrity + TM-INELIGIBLE OK")

        seg = _mk_seg(s, source=R1 + "a" + R1C)
        out = quality.commit_mt_output(s, seg, R1 + R1 + "hola" + R1C, source="mt")
        s.commit()
        assert out.tag_repaired and out.tag_violation == "", out
        assert seg.target_text == R1 + "hola" + R1C
        assert seg.status == SegStatus.MT and out.tm_eligible is True
        assert not [f for f in (seg.qa_flags or []) if f["type"] == "tag_integrity"]
        print("  5b. commit: duplicated token DE-DUP'd -> clean MT, no flag, TM-eligible OK")

        seg = _mk_seg(s, source=X1 + "a" + X2)
        out = quality.commit_mt_output(s, seg, X2 + "a" + X1, source="mt")
        s.commit()
        assert out.tag_repaired and out.tag_violation, ("best-guess must still flag", out)
        assert seg.target_text == X1 + "a" + X2
        assert seg.status == SegStatus.NEEDS_REVIEW and out.tm_eligible is False
        assert [f for f in seg.qa_flags if f["type"] == "tag_integrity"], seg.qa_flags
        print("  5b2. commit: point-order reorder -> committed BEST-GUESS + FLAGGED + review OK")

        seg = _mk_seg(s, source=R1 + "a" + R1C)
        out = quality.commit_mt_output(s, seg, R1 + "hola" + R1C, source="mt")
        s.commit()
        assert out.tag_violation == "" and not out.tag_repaired and seg.status == SegStatus.MT
        print("  5c. commit: tags intact -> byte-identical clean MT OK")

        seg = _mk_seg(s, source="plain text")
        out = quality.commit_mt_output(s, seg, "[es] texto plano", source="mt")
        s.commit()
        assert out.tag_violation == "" and seg.status == SegStatus.MT and out.tm_eligible
        print("  5d. commit: no tags present -> byte-identical (tagless guarantee) OK")

        seg = _mk_seg(
            s,
            source="Keep " + R1 + "X" + R1C + ".",
            qa_flags=[{"type": "qe", "severity": "minor", "score": 80}],
        )
        # This verdict describes the text being re-checked, not an older target.
        seg.target_text = "[es] sin."
        quality.commit_mt_output(s, seg, "[es] sin.", source="mt")
        s.commit()
        types = {f["type"] for f in seg.qa_flags}
        assert "tag_integrity" in types and "qe" in types, seg.qa_flags
        print("  5e. commit: tag_integrity composes with a pre-existing qe flag OK")


# ===========================================================================
# 6. lqa 1.3 remap (+ git-show fail-pre-fix proof)
# ===========================================================================
def test_lqa_remap():
    assert lqa._QA_TO_MQM["placeholder"] == ("Design/Markup", "major")
    card = lqa.score_text("Read " + R1 + "this" + R1C + ".", "[es] Lee esto.", "en-US", "es-ES")
    assert card["passed"] is False and card["gate"] == "needs_review", card
    assert any(
        e["dimension"] == "Design/Markup" and e["severity"] == "major" for e in card["errors"]
    ), card["errors"]
    ok = lqa.score_text(
        "Read " + R1 + "this" + R1C + ".", "[es] Lee " + R1 + "esto" + R1C + ".", "en-US", "es-ES"
    )
    assert not any(e["dimension"] == "Design/Markup" for e in ok["errors"]), ok
    print("  6a. lqa 1.3: dropped run-token -> Design/Markup/MAJOR -> LQA gate FAILS OK")

    base = _git_show("app/services/lqa.py")
    if base is None:
        print("  6b. (git-show unavailable - behavioral delta above is the guarantee) SKIP")
    else:
        assert '"placeholder": ("Design/Markup"' not in base, (
            "pre-fix lqa._QA_TO_MQM must NOT map 'placeholder' (the gate-bypass bug)"
        )
        print("  6b. fail-pre-fix: base commit's _QA_TO_MQM had NO 'placeholder' key OK")


# ===========================================================================
# 7. bounded fix-loop (1.5c)
# ===========================================================================
def test_bounded_fix_loop():
    prev = settings.get("workflow_total_step_cap")
    settings.set_many({"workflow_total_step_cap": "3"})
    try:
        with Session(engine) as s:
            p = _mk_project(s)
            doc = _mk_doc(s, p, sources=["x"])
            wf = _mk_wf(s)
            step = s.exec(select(WorkflowStep).where(WorkflowStep.workflow_id == wf.id)).first()
            run = WorkflowRun(
                workflow_id=wf.id,
                document_id=doc.id,
                target_locale="es-ES",
                status=RunStatus.RUNNING,
            )
            s.add(run)
            s.commit()
            s.refresh(run)
            guard = orchestrator.RunGuard(s, run, active=True)
            reasons = [guard.enter(step) for _ in range(6)]
        assert reasons[:3] == [None, None, None], reasons
        assert "loop_cap_reached" in reasons[3:], reasons
        print("  7. bounded fix-loop: RunGuard total-step cap halts oscillation OK")
    finally:
        settings.set_many({"workflow_total_step_cap": prev})


# ===========================================================================
# 8. FULL_TEXT export guard - TAGGED quarantines; TAGLESS parts<units REFUSES (honest-failure)
# ===========================================================================
def test_full_text_export_guard():
    units = [{"p_index": 0, "text": R1 + "Title" + R1C}, {"p_index": 1, "text": R1 + "Body" + R1C}]
    with Session(engine) as s:
        doc_id = _mk_ft_doc(s, units, R1 + "Title" + R1C + "\n\n" + R1 + "Body" + R1C)

    def _set(target):
        with Session(engine) as s:
            seg = s.exec(select(Segment).where(Segment.document_id == doc_id)).first()
            seg.target_text = target
            s.add(seg)
            s.commit()

    _set(R1 + "Titulo" + R1C + "\n\n" + R1 + "Cuerpo" + R1C)
    with Session(engine) as s:
        mapping = export.paragraph_texts_for(s, s.get(Document, doc_id), "es-ES")
        assert mapping["0"] == R1 + "Titulo" + R1C and mapping["1"] == R1 + "Cuerpo" + R1C
    print("  8a. FULL_TEXT (tagged): intact 2-paragraph body maps cleanly OK")

    _set(R1 + "Titulo" + R1C + " " + R1 + "Cuerpo" + R1C)  # merged blank line
    with Session(engine) as s:
        try:
            export.paragraph_texts_for(s, s.get(Document, doc_id), "es-ES")
            assert False, "a tagged paragraph-count desync must quarantine"
        except export.TagExportQuarantine as e:
            assert "desync" in str(e)
    print("  8b. FULL_TEXT (tagged): merged blank line -> WHOLE export quarantined OK")

    _set("Titulo\n\n" + R1 + R1 + "Cuerpo" + R1C)  # per-paragraph token-set shift
    with Session(engine) as s:
        try:
            export.paragraph_texts_for(s, s.get(Document, doc_id), "es-ES")
            assert False, "a tagged per-paragraph token-set desync must quarantine"
        except export.TagExportQuarantine:
            pass
    print("  8c. FULL_TEXT (tagged): per-paragraph token-set shift -> WHOLE export quarantined OK")

    # HONEST-FAILURE (2026-08-13 big-doc flood): a TAGLESS full-text doc whose model
    # emitted FEWER paragraphs than the source no longer silently source-fills the missing
    # positions (that served ENGLISH under a target filename). It now REFUSES loud
    # (kind="untranslated"). See verify_ft_export_untranslated.py for full coverage.
    with Session(engine) as s:
        tl_id = _mk_ft_doc(
            s,
            [{"p_index": 0, "text": "Title here"}, {"p_index": 1, "text": "Body here"}],
            "Title here\n\nBody here",
        )
        seg = s.exec(select(Segment).where(Segment.document_id == tl_id)).first()
        seg.target_text = "Titulo y Cuerpo juntos"  # reflowed to ONE paragraph
        s.add(seg)
        s.commit()
        try:
            export.paragraph_texts_for(s, s.get(Document, tl_id), "es-ES")
            assert False, "a tagless parts<units desync must refuse, not source-fill"
        except export.TagExportQuarantine as e:
            assert e.kind == "untranslated", e.kind
    print(
        "  8d. FULL_TEXT (TAGLESS): parts<units REFUSES (honest-failure), no silent source-fill OK"
    )


# ===========================================================================
# 8M. project BUNDLE export skip-and-report (M1)
# ===========================================================================
def test_bundle_skip_and_report():
    import zipfile
    from io import BytesIO

    from docx import Document as Docx

    def _docx(paras):
        d = Docx()
        for p in paras:
            d.add_paragraph(p)
        b = BytesIO()
        d.save(b)
        return b.getvalue()

    client = TestClient(app)
    r = client.post(
        "/projects",
        data={
            "name": "bundle-%d" % _seq(),  # noqa: UP031
            "source_locale": "en-US",
            "target_locales": "es-ES",
            "default_provider": "mock",
        },
        follow_redirects=False,
    )
    pid = int(r.headers["location"].rsplit("/", 1)[-1])
    client.post(
        f"/projects/{pid}/documents",
        data={"mode": "segmented"},
        files={
            "file": (
                "clean.docx",
                _docx(["Hello world."]),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )

    with Session(engine) as s:
        proj = s.get(Project, pid)
        from app.services import translate as _tr

        cdoc = s.exec(select(Document).where(Document.project_id == proj.id)).first()
        for _ in _tr.pretranslate_stream(s, cdoc.id, "es-ES"):
            pass
        # a second FULL_TEXT tagged doc, deliberately desynced -> quarantines on export
        qdoc = Document(
            project_id=proj.id,
            filename="q.docx",
            fmt="docx",
            mode=TransMode.FULL_TEXT,
            skeleton={
                "units": [
                    {"p_index": 0, "text": R1 + "A" + R1C},
                    {"p_index": 1, "text": R1 + "B" + R1C},
                ],
                "seg_locale": "en-US",
            },
        )
        s.add(qdoc)
        s.commit()
        s.refresh(qdoc)
        s.add(
            Segment(
                document_id=qdoc.id,
                target_locale="es-ES",
                order_index=0,
                source_text=R1 + "A" + R1C + "\n\n" + R1 + "B" + R1C,
                target_text=R1 + "A" + R1C + " " + R1 + "B" + R1C,
            )
        )
        s.commit()
        data, _name, _ctype, skipped = delivery.export_project(s, proj.id, zip=True)

    zf = zipfile.ZipFile(BytesIO(data))
    names = zf.namelist()
    assert any("clean" in n for n in names), names
    assert not any("q.es-ES" in n for n in names), ("quarantined doc excluded", names)
    assert "_QUARANTINED.txt" in names, ("manifest present", names)
    assert len(skipped) == 1 and skipped[0]["filename"] == "q.docx", skipped
    manifest = zf.read("_QUARANTINED.txt").decode()
    assert "q.docx" in manifest and "quarantined" in manifest.lower()
    print("  8M. bundle export: quarantined doc EXCLUDED + reported; clean doc still ships OK")


# ===========================================================================
# 9. editor save guard (1.7)
# ===========================================================================
def test_editor_save_guard():
    client = TestClient(app)
    with Session(engine) as s:
        seg = _mk_seg(s, source="Click " + R1 + "here" + R1C + ".", locale="es-ES")
        seg_id = seg.id

    r = client.post(
        f"/segments/{seg_id}/save", data={"target_text": "Haz clic aqui.", "tte_ms": "0"}
    )
    body = r.json()
    assert body.get("needs_confirm") is True and "formatting" in body["message"].lower(), body
    with Session(engine) as s:
        assert not (s.get(Segment, seg_id).target_text or "")
    print("  9a. save guard: token-dropping save -> blocking needs_confirm, nothing written OK")

    r = client.post(
        f"/segments/{seg_id}/save",
        data={"target_text": "Haz clic aqui.", "tte_ms": "0", "confirm": "1"},
    )
    assert r.json().get("ok") is True
    with Session(engine) as s:
        assert s.get(Segment, seg_id).target_text == "Haz clic aqui."
    print("  9b. save guard: confirm=1 -> save proceeds (non-destructive) OK")

    r = client.post(
        f"/segments/{seg_id}/save",
        data={"target_text": "Haz clic " + R1 + "aqui" + R1C + ".", "tte_ms": "0"},
    )
    assert r.json().get("ok") is True and "needs_confirm" not in r.json()
    print("  9c. save guard: a token-preserving save is byte-identical (no prompt) OK")

    with Session(engine) as s:
        seg2 = _mk_seg(s, source="Code is " + LAQ + "TAG_9" + RAQ, locale="es-ES")
        sid2 = seg2.id
    r = client.post(
        f"/segments/{sid2}/save",
        data={"target_text": "El codigo es " + LAQ + "TAG_9" + RAQ, "tte_ms": "0"},
    )
    assert r.json().get("ok") is True and "needs_confirm" not in r.json()
    print("  9d. save guard: fixed-shape sentinel as content -> no false confirm OK")


# ===========================================================================
# 10. batch verify CLI (flags-only)
# ===========================================================================
def test_batch_verify_cli():
    from app.cli.__main__ import main as cli_main

    with Session(engine) as s:
        p = _mk_project(s)
        doc = _mk_doc(
            s, p, sources=["Read " + R1 + "this" + R1C + ".", "no tags here"], locale="es-ES"
        )
        segs = s.exec(
            select(Segment).where(Segment.document_id == doc.id).order_by(Segment.order_index)
        ).all()
        segs[0].target_text = "[es] Lee esto."
        segs[0].status = SegStatus.MT
        segs[1].target_text = "[es] sin etiquetas"
        segs[1].status = SegStatus.MT
        s.add(segs[0])
        s.add(segs[1])
        s.commit()
        seg0_id, seg1_id, doc_id = segs[0].id, segs[1].id, doc.id

    rc = cli_main(["--json", "tagging", "verify", "--doc", str(doc_id)])
    assert rc == 0
    with Session(engine) as s:
        seg0 = s.get(Segment, seg0_id)
        assert seg0.status == SegStatus.MT
        assert not [f for f in (seg0.qa_flags or []) if f.get("type") == "tag_integrity"]
    print("  10a. tagging verify (dry-run): flags-only report, byte-identical (no writes) OK")

    rc = cli_main(["tagging", "verify", "--doc", str(doc_id), "--apply"])
    assert rc == 0
    with Session(engine) as s:
        seg0 = s.get(Segment, seg0_id)
        seg1 = s.get(Segment, seg1_id)
        assert seg0.status == SegStatus.NEEDS_REVIEW
        assert [f for f in seg0.qa_flags if f.get("type") == "tag_integrity"], seg0.qa_flags
        assert seg1.status == SegStatus.MT and not (seg1.qa_flags or [])
    print(
        "  10b. tagging verify --apply: damaged -> NEEDS_REVIEW + tag_integrity; clean untouched OK"
    )


# ===========================================================================
# 11. end-to-end masking
# ===========================================================================
def _seg_of(doc_id, order=0, locale="es-ES"):
    with Session(engine) as s:
        return s.exec(
            select(Segment).where(
                Segment.document_id == doc_id,
                Segment.order_index == order,
                Segment.target_locale == locale,
            )
        ).first()


def test_end_to_end_plain_mt():
    with Session(engine) as s:
        p = _mk_project(s)
        doc = _mk_doc(s, p, sources=["Read " + R1 + "this" + R1C + " now."])
        doc_id = doc.id
    with Session(engine) as s, _Engine(translate, lambda t: t):
        for _ in translate.pretranslate_stream(s, doc_id, "es-ES"):
            pass
    seg = _seg_of(doc_id)
    assert seg.status == SegStatus.MT and not ai_tagging.has_residual_sentinel(seg.target_text)
    assert quality.tag_damage(seg.source_text, seg.target_text) == "", seg.target_text
    print("  11a. plain-MT: preserving engine -> tokens restored, clean MT, no leak OK")

    with Session(engine) as s:
        p = _mk_project(s)
        doc = _mk_doc(s, p, sources=["Read " + R1 + "this" + R1C + " now."])
        doc_id = doc.id
    with Session(engine) as s, _Engine(translate, _drop_sentinels):
        for _ in translate.pretranslate_stream(s, doc_id, "es-ES"):
            pass
    seg = _seg_of(doc_id)
    assert seg.status == SegStatus.NEEDS_REVIEW
    assert [f for f in (seg.qa_flags or []) if f.get("type") == "tag_integrity"], seg.qa_flags
    print("  11b. plain-MT: token-dropping engine -> NEEDS_REVIEW + tag_integrity OK")


def test_end_to_end_orchestrator():
    with Session(engine) as s:
        p = _mk_project(s)
        doc = _mk_doc(s, p, sources=["Read " + R1 + "this" + R1C + "."])
        wf = _mk_wf(s)
        wf_id, doc_id = wf.id, doc.id
    with Session(engine) as s, _Engine(orchestrator, _drop_sentinels):
        orchestrator.run_workflow(s, wf_id, doc_id, "es-ES")
    seg = _seg_of(doc_id)
    assert seg.status == SegStatus.NEEDS_REVIEW
    assert [f for f in (seg.qa_flags or []) if f.get("type") == "tag_integrity"], seg.qa_flags
    print("  11c. orchestrator (segmented): dropping engine -> NEEDS_REVIEW + tag_integrity OK")

    with Session(engine) as s:
        p = _mk_project(s)
        doc = _mk_doc(s, p, sources=["Read " + R1 + "this" + R1C + "."])
        wf = _mk_wf(s)
        wf_id, doc_id = wf.id, doc.id
    with Session(engine) as s, _Engine(orchestrator, _dup_first_sentinel):
        orchestrator.run_workflow(s, wf_id, doc_id, "es-ES")
    seg = _seg_of(doc_id)
    assert not ai_tagging.has_residual_sentinel(seg.target_text)
    assert quality.tag_damage(seg.source_text, seg.target_text) == "", seg.target_text
    assert not [f for f in (seg.qa_flags or []) if f.get("type") == "tag_integrity"]
    print("  11d. orchestrator (segmented): duplicating engine -> DE-DUP'd clean, no flag OK")

    with Session(engine) as s:
        p = _mk_project(s)
        doc_id = _mk_ft_doc(
            s,
            [{"p_index": 0, "text": "Read " + R1 + "this" + R1C + "."}],
            "Read " + R1 + "this" + R1C + ".",
        )
        wf = _mk_wf(s)
        wf_id = wf.id
    with Session(engine) as s, _Engine(orchestrator, _drop_sentinels):
        for _ in orchestrator.run_workflow_stream(s, wf_id, doc_id, "es-ES"):
            pass
    seg = _seg_of(doc_id)
    assert seg.status == SegStatus.NEEDS_REVIEW
    assert [f for f in (seg.qa_flags or []) if f.get("type") == "tag_integrity"], seg.qa_flags
    print("  11e. orchestrator (FULL_TEXT): dropping engine -> NEEDS_REVIEW + tag_integrity OK")


def main():
    print("AI TAGGING - deterministic inline-tag integrity (Phase 1)")
    test_mask_restore()
    test_collision_proofing()
    test_restored_forms_only()
    test_tag_damage_unit()
    test_repair_unit()
    test_commit_wiring()
    test_lqa_remap()
    test_bounded_fix_loop()
    test_full_text_export_guard()
    test_bundle_skip_and_report()
    test_editor_save_guard()
    test_batch_verify_cli()
    test_end_to_end_plain_mt()
    test_end_to_end_orchestrator()
    test_segmented_export_tag_guard()
    test_emptied_span_and_empty_run()
    test_rewritten_delimiters_restore()
    test_bracket_and_brace_frames_restore()
    test_mangled_delimiters_726()
    print("AI TAGGING OK")


# --- 14. engine rewrites the sentinel delimiters (GitHub #673, format soak) ---------
def test_rewritten_delimiters_restore():
    """gemma-4-12b into zh-CN turned every «TAG<nonce>_N» into 「TAG<nonce>_N」: the exact
    replace restored nothing and every formatted segment was held for review."""
    src = "This is " + R1 + "preformatted" + R1C + " text " + X1 + "."
    _masked, mp = ai_tagging.mask(src)
    keys = list(mp)  # in document order: _1, _2, _3
    bodies = [k[1:-1] for k in keys]
    # (a) the live zh-CN shape: corner brackets around every sentinel.
    zh = "这是「{}」预格式「{}」文本「{}」。".format(*bodies)
    out, resid = ai_tagging.restore(zh, mp)
    assert out == "这是" + R1 + "预格式" + R1C + "文本" + X1 + "。", out
    assert resid is False
    # (b) mixed: one exact, one half-rewritten (« ... 」), one in curly quotes.
    mixed = "A " + keys[0] + "b" + LAQ + bodies[1] + "」 c “" + bodies[2] + "”."
    out, resid = ai_tagging.restore(mixed, mp)
    assert out == "A " + R1 + "b" + R1C + " c " + X1 + ".", out
    assert resid is False
    # (b2) a body framed on one side only, or TRUNCATED (`「TAG<nonce>_`), is not guessed
    # at: it stays a residual (review). #726: a body with no delimiter beside it restores.
    out, resid = ai_tagging.restore("A 「" + bodies[0] + " b", mp)
    assert R1 not in out and resid is True, (out, resid)
    assert ai_tagging.restore("A " + bodies[0] + " b", mp) == ("A " + R1 + " b", False)
    out, resid = ai_tagging.restore("A 「" + bodies[0][:-1] + " b", mp)
    assert R1 not in out and resid is True, (out, resid)
    # (c) _1 must not eat _10: a longer id we never issued stays a residual.
    _s10, m10 = ai_tagging.mask(" ".join([X1] * 10))
    k1 = next(iter(m10))
    out, resid = ai_tagging.restore("「" + k1[1:-1] + "0」", {k1: X1})
    assert X1 not in out and resid is True, (out, resid)
    # (d) an unrestored 「TAG<nonce>_ remnant is reported as a residual sentinel now.
    assert ai_tagging.has_residual_sentinel("这是「" + bodies[0] + "」")
    # (e) prose with a fixed-shape 「TAG_1」 / «TAG_1» is still never touched or flagged.
    prose = "见「TAG_1」与 " + LAQ + "TAG_1" + RAQ
    assert ai_tagging.restore(prose, mp) == (prose, False)
    assert not ai_tagging.has_residual_sentinel(prose)
    print("  [14] rewritten sentinel delimiters (zh 「」, mixed, bare) restore; residuals flagged")


def test_bracket_and_brace_frames_restore():
    """#684 (format soak round 4): gemma-4-12b into he-IL / pt-BR also wrote the sentinel
    inside the brackets of the real tokens around it (`⟦TAG<nonce>_1⟧`), half-rewritten
    (`«TAG<nonce>_2⟧`) and in placeholder braces (`{{TAG<nonce>_1}}`). None restored, so
    the segment was held with sentinel garbage; `⟦TAG…` / `{{TAG…` were not even residuals."""
    src = "You edited " + X1 + " this " + R1 + "week" + R1C + "."
    _masked, mp = ai_tagging.mask(src)
    keys = list(mp)
    bodies = [k[1:-1] for k in keys]
    ok = "You edited " + X1 + " this " + R1 + "week" + R1C + "."
    LB, RB = chr(0x27E6), chr(0x27E7)  # the token brackets
    # (a) he-IL live shape: token brackets around the body.
    fr = [LB + b + RB for b in bodies]
    t = "You edited " + fr[0] + " this " + fr[1] + "week" + fr[2] + "."
    assert ai_tagging.restore(t, mp) == (ok, False), ai_tagging.restore(t, mp)
    # (b) pt-BR live shape: opening guillemet kept, closing rewritten to the token bracket.
    t = "You edited " + LAQ + bodies[0] + RB + " this " + keys[1] + "week" + keys[2] + "."
    assert ai_tagging.restore(t, mp) == (ok, False), ai_tagging.restore(t, mp)
    # (c) he-IL live shape: doubled placeholder braces, restored whole (no `{...}` left).
    t = "You edited {{" + bodies[0] + "}} this " + keys[1] + "week" + keys[2] + "."
    assert ai_tagging.restore(t, mp) == (ok, False), ai_tagging.restore(t, mp)
    # (c2) gate: one side doubled alone leaves no stray brace beside the token.
    for o, c in (("{{", "}"), ("{", "}}"), ("[[", "]"), (LB + LB, RB)):
        t = "You edited " + o + bodies[0] + c + " this " + keys[1] + "week" + keys[2] + "."
        assert ai_tagging.restore(t, mp) == (ok, False), (o, c, ai_tagging.restore(t, mp))
    # (d) single ASCII frames.
    for o, c in ("{}", "[]", "()", "<>"):
        t = "You edited " + o + bodies[0] + c + " this " + keys[1] + "week" + keys[2] + "."
        assert ai_tagging.restore(t, mp) == (ok, False), (o, ai_tagging.restore(t, mp))
    # (e) an unframed body after `⟦` / `{{` (no closing frame).
    for o in (LB, "{{", "[", "("):
        # framed on one side only: not guessed at (the frame could be the author's)...
        out, resid = ai_tagging.restore("You edited " + o + bodies[0] + " x", mp)
        assert X1 not in out and resid is True, (o, out, resid)
        # ...nor is a TRUNCATED body: it stays a residual.
        out, resid = ai_tagging.restore("You edited " + o + bodies[0][:-1] + " x", mp)
        assert X1 not in out and resid is True, (o, out, resid)
        assert ai_tagging.has_residual_sentinel(o + bodies[0])
    # (f) _1 never eats _10 inside the new frames either.
    _s10, m10 = ai_tagging.mask(" ".join([X1] * 10))
    k1 = next(iter(m10))
    out, resid = ai_tagging.restore(LB + k1[1:-1] + "0" + RB, {k1: X1})
    assert X1 not in out and resid is True, (out, resid)
    # (g) fixed-shape prose (`{TAG_1}`, `⟦TAG_1⟧`, `[TAG_1]`) is never touched or flagged.
    prose = "see {TAG_1} and " + LB + "TAG_1" + RB + " and [TAG_1] and {{TAG_2}}"
    assert ai_tagging.restore(prose, mp) == (prose, False)
    assert not ai_tagging.has_residual_sentinel(prose)
    print("  [15] token-bracket / brace frames (#684) restore; unframed remnants flagged")


def test_mangled_delimiters_726():
    """#726 (HubSpot rehearsal, ja-JP, gemma-4-26b): the engine wrote `「TAG<nonce>_1»` /
    `「TAG<nonce>_1」` and other mangled frames. Debris an engine leaves in ANY frame, with
    spaces, or with no frame at all must restore when it is this call's body, and any
    nonce-bearing `TAG<hex>_` leftover must be flagged whatever surrounds it."""
    src = G1 + "Starter" + G1C + " plan " + X1 + "."
    _masked, mp = ai_tagging.mask(src)
    keys = list(mp)
    bodies = [k[1:-1] for k in keys]
    ok = G1 + "スターター" + G1C + "プラン" + X1 + "。"

    def ja(a, b, c):
        return a + "スターター" + b + "プラン" + c + "。"

    # (a) the live shapes: half-rewritten and fully rewritten corner brackets.
    t = ja("「" + bodies[0] + RAQ, keys[1], "「" + bodies[2] + "」")
    assert ai_tagging.restore(t, mp) == (ok, False), ai_tagging.restore(t, mp)
    # (b) fullwidth angle brackets / quotes / braces / parens, halfwidth corner brackets,
    # double-prime quotes, white lenticular brackets, with whitespace inside the frame.
    for o, c in (
        ("＜", "＞"),
        ("＂", "＂"),
        ("｛", "｝"),
        ("（", "）"),
        ("｢", "｣"),
        ("〝", "〞"),
        ("〖", "〗"),
        ("「 ", " 」"),
        ("【 ", "】"),
    ):
        t = ja(o + bodies[0] + c, keys[1], keys[2])
        assert ai_tagging.restore(t, mp) == (ok, False), (o, c, ai_tagging.restore(t, mp))
    # (c) no delimiter beside any body (both guillemets dropped): restores, nothing eaten.
    t = ja(bodies[0], bodies[1], bodies[2])
    assert ai_tagging.restore(t, mp) == (ok, False), ai_tagging.restore(t, mp)
    # (c2) gate: a delimiter on ONE side only is ambiguous. Here the 「」 are the AUTHOR's
    # (`「«TAG_1»Starter«TAG_2»」` with the guillemets dropped); consuming them would ship
    # `⟦g1⟧...⟦/g1⟧` without them. Never guessed at: a residual, held for review.
    t = "「" + bodies[0] + "スターター" + bodies[1] + "」プラン" + keys[2] + "。"
    out, resid = ai_tagging.restore(t, mp)
    assert resid is True and "「" in out and "」" in out and X1 in out, (out, resid)
    assert G1 not in out and G1C not in out, out
    # (d) a bare body next to the next sentinel keeps the space and the next frame.
    t = bodies[0] + " 「" + bodies[1] + "」"
    assert ai_tagging.restore(t, mp) == (G1 + " " + G1C, False), ai_tagging.restore(t, mp)
    # (e) a restored token is never taken as the frame of the next sentinel.
    t = keys[0] + bodies[1]
    assert ai_tagging.restore(t, mp) == (G1 + G1C, False), ai_tagging.restore(t, mp)
    out, resid = ai_tagging.restore(keys[0] + bodies[1] + RAQ, mp)
    assert out.startswith(G1 + bodies[1]) and resid is True, (out, resid)
    # (f) residual detection: nonce-bearing debris in any frame, or none, is flagged...
    for o in ("＜", "〔 ", "“", "", "x"):
        assert ai_tagging.has_residual_sentinel(o + "TAG88005b59_1"), o
        out, resid = ai_tagging.restore("A " + o + "TAG88005b59_1 b", mp)
        assert resid is True, (o, out)
    # ...and a truncated body of THIS call stays a residual, never a guessed token.
    out, resid = ai_tagging.restore("A ＜" + bodies[0][:-1] + " b", mp)
    assert G1 not in out and resid is True, (out, resid)
    # (g) _1 never eats _10 in the new forms.
    _s10, m10 = ai_tagging.mask(" ".join([X1] * 10))
    k1 = next(iter(m10))
    for t in ("＜" + k1[1:-1] + "0＞", k1[1:-1] + "0 x"):
        out, resid = ai_tagging.restore(t, {k1: X1})
        assert X1 not in out and resid is True, (t, out, resid)
    # (h) fixed-shape prose is still never touched or flagged.
    prose = "「TAG_1」 ＜TAG_2＞ " + LAQ + "TAG_5" + RAQ + " TAG_3"
    assert ai_tagging.restore(prose, mp) == (prose, False)
    assert not ai_tagging.has_residual_sentinel(prose)
    # (i) an empty mapping leaves the text byte-identical.
    for t in (prose, "「TAG88005b59_1」", ""):
        assert ai_tagging.restore(t, {})[0] == t
    # (j) the comparison helper strips CJK / fullwidth-framed debris as one unit.
    assert ai_tagging.strip_generated_sentinels("A「TAG88005b59_1」B") == "AB"
    assert ai_tagging.strip_generated_sentinels("A＜TAG88005b59_1＞B") == "AB"
    # (k) gate: a nonce-SHAPED identifier the source itself carries is content, not
    # debris: the write gate does not hold it; debris absent from the source still is.
    from app.services import quality

    ident = "Build BUILD_TAG20241011_rc1 " + X1
    assert ai_tagging.has_residual_sentinel(ident)
    assert not ai_tagging.has_residual_sentinel(ident, ident)
    assert quality.tag_damage(ident, "Version BUILD_TAG20241011_rc1 " + X1) == ""
    bad = "Version BUILD_TAG20241011_rc1 TAG88005b59_1 " + X1
    assert ai_tagging.has_residual_sentinel(bad, ident)
    assert quality.tag_damage(ident, bad) == "residual masking sentinel in target"
    print("  [16] mangled / missing sentinel delimiters (#726) restore; debris flagged")


# --- 12. SEGMENTED export tag guard (GitHub #41) -----------------------------
def test_segmented_export_tag_guard():
    """A tag-damaged SEGMENTED segment must NOT be written into the deliverable.

    Pre-1.7 `paragraph_texts_for` consulted neither status nor qa_flags nor tag_damage,
    so the gate flagged the segment, made it TM-ineligible and connector-withheld -- and
    export then put it in the CUSTOMER'S FILE anyway, silently collapsing the formatting
    (measured live on a realistic DOCX: 6 of 16 formatted paragraphs lost their runs; on
    an XLIFF a unit lost its inline markup entirely). A dropped tag does NOT leak a
    visible token, it MERGES RUNS -- which is why every leak/readability check passed.
    Uses a REAL ingested document so the skeleton, units and segment split are genuine."""
    from app.models import TargetLocale
    from app.services import export
    from app.services.ingest import ingest_document

    HTML = (
        b"<html><body><p>Read <b>the report</b> now.</p><p>Second unit is clean.</p></body></html>"
    )
    with Session(engine) as s:
        p = _mk_project(s)
        s.add(TargetLocale(project_id=p.id, locale="es-ES"))
        s.commit()
        d = ingest_document(s, p, "guard.html", HTML, TransMode.SEGMENTED)
        s.commit()
        s.refresh(d)
        row = s.exec(
            select(Segment).where(Segment.document_id == d.id, Segment.order_index == 0)
        ).first()
        assert segment.find_placeholders(row.source_text), "fixture must carry tokens"
        row.target_text = "Lea el informe ahora."  # BOTH tokens dropped
        s.add(row)
        s.commit()

        settings.set_many({"export_tag_damage_policy": "source"})
        rep = []
        m = export.paragraph_texts_for(s, d, "es-ES", report=rep)
        # GitHub #42: the substitution channel is now self-describing records naming the
        # UNIT (segment order_index) and KIND (a routine #41 automatic fallback), not a
        # bare list[int] that read ambiguously between SEGMENTED and FULL_TEXT.
        assert rep == [{"unit": "segment", "index": 0, "kind": export.SUB_AUTOMATIC}], (
            rep
        )  # damage detected
        joined = " ".join(m.values())
        assert "⟦g1⟧" in joined and "⟦/g1⟧" in joined, joined
        assert "Second unit is clean." in joined, "clean units must be untouched"

        settings.set_many({"export_tag_damage_policy": "quarantine"})
        try:
            export.paragraph_texts_for(s, d, "es-ES")
            raise AssertionError("quarantine policy must refuse the export")
        except export.TagExportQuarantine:
            pass

        settings.set_many({"export_tag_damage_policy": "ship"})
        joined = " ".join(export.paragraph_texts_for(s, d, "es-ES").values())
        assert "⟦g1⟧" not in joined  # explicit pre-1.7 opt-out
        settings.set_many({"export_tag_damage_policy": "source"})

        # A TAGLESS document is byte-identical under every policy.
        p2 = _mk_project(s)
        s.add(TargetLocale(project_id=p2.id, locale="es-ES"))
        s.commit()
        d2 = ingest_document(
            s,
            p2,
            "plain.html",
            b"<html><body><p>No tokens at all.</p></body></html>",
            TransMode.SEGMENTED,
        )
        s.commit()
        s.refresh(d2)
        base = export.paragraph_texts_for(s, d2, "es-ES")
        for pol in ("source", "quarantine", "ship"):
            settings.set_many({"export_tag_damage_policy": pol})
            assert export.paragraph_texts_for(s, d2, "es-ES") == base, pol
        settings.set_many({"export_tag_damage_policy": "source"})
    print("  [12] SEGMENTED export guard: a damaged segment never reaches the deliverable")
    print("       (source fallback restores tokens / quarantine refuses / ship opts out)")


# --- 13. emptied paired span + the docx empty-run round trip -----------------
def test_emptied_span_and_empty_run():
    """Two damage classes the gate was structurally blind to, both measured live.

    (A) EMPTIED SPAN: identity, multiset, ORDER and nesting are ALL preserved by a
    target that moves the words OUT of a pair. The empty span is dropped at build time
    and the neighbour absorbs the words wearing the NEIGHBOUR's formatting. Live
    (llama-3-8b, fr-CA): the model moved "export pipeline" out of its bold span leaving
    the pair empty; the delivered DOCX lost the bold while `tag_damage` said "".

    (B) EMPTY `<w:r>`: real Word emits `[<w:r/>, <w:r>bold text]` constantly (rsid
    splits, revision merges). `_tokenize_paragraph` counts CONTENT runs and so does not
    tokenize it, but the build branch wrote into `para.runs[0]` -- the EMPTY run -- and
    blanked the bold one. Formatting lost with the target BYTE-IDENTICAL to the source:
    a pure round-trip bug, no model involved."""
    from docx import Document as Docx

    from app.filters.docx_filter import DocxFilter

    SRC = "⟦r1⟧The ⟦/r1⟧⟦r2⟧Acme Pro⟦/r2⟧"
    # (A) emptied -> damage; a blank SOURCE body is legitimate and must NOT flag.
    assert "emptied span r2" in quality.tag_damage(SRC, "⟦r1⟧El ⟦/r1⟧⟦r2⟧⟦/r2⟧conector Acme Pro")
    assert quality.tag_damage(SRC, SRC) == ""
    assert quality.tag_damage("⟦r1⟧⟦/r1⟧a", "⟦r1⟧⟦/r1⟧b") == ""
    # ...and the export guard inherits it through that one verdict.
    assert quality.tag_damage(SRC, "⟦r1⟧El ⟦/r1⟧⟦r2⟧ ⟦/r2⟧x"), (
        "a whitespace-only target body is still an emptied span"
    )

    # (B) a leading EMPTY run must not receive the write, on a PERFECT translation.
    d = Docx()
    para = d.add_paragraph()
    para.add_run("")  # empty, no rPr -- real Word does this
    r = para.add_run("Section 3. Review the policy.")
    r.bold = True
    buf = io.BytesIO()
    d.save(buf)
    data = buf.getvalue()
    f = DocxFilter()
    units = f.parse(data).skeleton["units"]
    out = f.build(data, {str(u["p_index"]): u["text"] for u in units})  # target == source
    runs = [(x.text, bool(x.bold)) for x in Docx(io.BytesIO(out)).paragraphs[0].runs if x.text]
    assert runs and all(b for _, b in runs), runs  # bold survived the round trip
    print("  [13] emptied paired span flagged; docx empty-run round trip preserves formatting")


if __name__ == "__main__":
    main()
