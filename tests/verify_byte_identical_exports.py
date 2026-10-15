"""Untouched files come back byte-identical (B-bytes, 2026-10-11 decision 1).

The format soak WARNed "same content, different bytes" on untouched PO, XLIFF, JSON, SRT,
INI/CSV, InCopy, text, HTML, DOCX and XLSX files (quote style, attribute order, CRLF to
LF, BOM dropped, Windows-1252 to UTF-8). Each case below is one of those WARNs.

  noop        build(original, {}) is the original bytes for every format; build(original,
              identity) too for monolingual formats. (build_localized may still stamp the
              target language or plural forms; the untouched EXPORT is checked below.)
  bilingual   PO / XLIFF still WRITE a target equal to the source (identity is a real write);
              XMB still exports XTB (a different file).
  po          a translated PO keeps every untouched line (wrapped header, quoting, blank
              lines), CRLF, a UTF-16 BOM and its Windows-1252 code page (when the text fits).
  xliff       a translated XLIFF keeps its XML declaration, CRLF and trailing newline.
  json        a translated JSON keeps its layout (compact stays compact), BOM and CRLF.
  srt         a translated SRT keeps CRLF, BOM, its blank-line spacing and untouched cues.
  legacy      INI / CSV / TXT / HTML in Windows-1252 go back in it when the text fits,
              UTF-8 (HTML: meta charset UTF-8) when it does not.
  icml        a rebuilt InCopy story keeps the file's CRLF.
  export      build_translated_file: an untranslated document, and a plain-text document
              whose every target equals its source, export as the original bytes; a
              translated SRT keeps CRLF and untouched cues; HTML still declares the target.

Hermetic + offline. Final line: "BYTE IDENTICAL EXPORTS OK".
"""

import os
import sys

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_bytes_")

from io import BytesIO

from sqlmodel import Session, select

from app.db import engine, init_db
from app.filters import get_filter
from app.models import Document, Project, Segment, SegStatus, TargetLocale, TransMode
from app.services import export
from app.services.ingest import ingest_document

init_db()

BOM = b"\xef\xbb\xbf"


def ident(fmt: str, data: bytes) -> dict:
    return {str(u.p_index): u.text for u in get_filter(fmt).parse(data).units}


def check_noop(fmt: str, data: bytes, monolingual: bool = True) -> None:
    f = get_filter(fmt)
    assert f.build(data, {}) == data, f"{fmt}: build(original, {{}}) is not the original"
    if monolingual:
        assert f.build(data, ident(fmt, data)) == data, f"{fmt}: identity build differs"


# --- fixtures: each is a soak WARN shape -------------------------------------------
PO_CRLF = (
    b'msgid ""\r\n'
    b'msgstr ""\r\n'
    b'"Content-Type: text/plain; charset=UTF-8\\n"\r\n'
    b'"Language: \\n"\r\n'
    b'"Plural-Forms: nplurals=2; plural=(n != 1);\\n"\r\n'
    b"\r\n"
    b"#: src/app.c:10\r\n"
    b'msgid "Hello {0}, welcome back!"\r\n'
    b'msgstr ""\r\n'
    b"\r\n"
    b"\r\n"
    b"#, c-format\r\n"
    b'msgid ""\r\n'
    b'"You have %d new files "\r\n'
    b'"in your inbox."\r\n'
    b'msgstr ""\r\n'
)
PO_1252 = (
    b'msgid ""\n'
    b'msgstr ""\n'
    b'"Content-Type: text/plain; charset=windows-1252\\n"\n'
    b'"Plural-Forms: nplurals=2; plural=n!=1;\\n"\n'
    b"\n"
    b'msgid "The caf\xe9 menu"\n'
    b'msgstr ""\n'
    b"\n"
    b'msgid "Open"\n'
    b'msgstr ""\n'
)
PO_BARE_1252 = (
    b'#: a.c:1\r\nmsgid "The caf\xe9 menu"\r\nmsgstr ""\r\n\r\nmsgid "Open"\r\nmsgstr ""\r\n'
)
PO_UTF16 = '\ufeff# Test\r\n\r\nmsgid "Close"\r\nmsgstr ""\r\n'.encode("utf-16-be")

XLIFF_CRLF = (
    b'<?xml version="1.0" encoding="UTF-8"?>\r\n'
    b'<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">\r\n'
    b'  <file original="a.txt" source-language="en-US" datatype="plaintext">\r\n'
    b"    <body>\r\n"
    b'      <trans-unit id="1">\r\n'
    b"        <source>Hello world</source>\r\n"
    b"      </trans-unit>\r\n"
    b'      <trans-unit id="2">\r\n'
    b"        <source>Good night</source>\r\n"
    b"      </trans-unit>\r\n"
    b"    </body>\r\n"
    b"  </file>\r\n"
    b"</xliff>\r\n"
)

JSON_COMPACT = b'{"app":{"greeting":"Hello {0}, welcome back!","count":3,"bye":"Good\\u0020bye"}}'
JSON_BOM_CRLF = BOM + b'{\r\n  "a": "First line",\r\n  "b": "Second line"\r\n}\r\n'

SRT_CRLF_BOM = BOM + (
    b"1\r\n00:00:01,000 --> 00:00:03,500\r\nHello {0}, welcome back!\r\n"
    b" \t\r\n"
    b"2\r\n00:00:04,000 --> 00:00:06,000\r\nYou have new files.\r\n"
    b"\r\n"
    b"3\r\n00:00:07,000 --> 00:00:09,000\r\nGood night.  \r\n"
)

INI_1252 = b"[Strings]\ncafe=The caf\xe9 menu costs \x805\nbye=Good bye\n"
CSV_1252 = b'key,source\ngreeting,"Hello, caf\xe9"\nbye,Good bye\n'
TXT_BOM = BOM + b"Hello {0}, welcome back!\n\nGood night.\n"
TXT_1252 = b"The caf\xe9 menu.\nGood night.\n"
HTML_1252 = (
    b'<!DOCTYPE html>\r\n<html lang="en">\r\n<head><meta charset="windows-1252">'
    b"<title>Caf\xe9</title></head>\r\n<body>\r\n<p>Don't forget the caf\xe9.</p>\r\n"
    b"<p>Good night.</p>\r\n</body>\r\n</html>\r\n"
)
HTML_APOS = b"<html><body><p>Don't forget to restart the application.</p></body></html>\n"


def _docx() -> bytes:
    from docx import Document as Docx

    d = Docx()
    d.add_paragraph("Hello world.")
    d.add_paragraph("Good night.")
    b = BytesIO()
    d.save(b)
    return b.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = "Hello world."
    wb.active["A2"] = "Good night."
    b = BytesIO()
    wb.save(b)
    return b.getvalue()


def _icml_crlf() -> bytes:
    lines = [
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        b'<?aid style="50" type="snippet" readerVersion="6.0" featureSet="513" ?>',
        b'<?aid SnippetType="InCopyInterchange"?>',
        b'<Document DOMVersion="20.0" Self="d">',
        b'\t<Story Self="u1" TrackChanges="false" StoryTitle="$ID/">',
        b'\t\t<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/$ID/Normal">',
        b'\t\t\t<CharacterStyleRange AppliedCharacterStyle="CharacterStyle/$ID/None">'
        b"<Content>Hello world.</Content></CharacterStyleRange>",
        b"\t\t</ParagraphStyleRange>",
        b"\t</Story>",
        b"</Document>",
        b"",
    ]
    return b"\r\n".join(lines)


# --- noop -------------------------------------------------------------------------
def test_noop():
    cases = [
        ("po", PO_CRLF, False),
        ("po", PO_1252, False),
        ("po", PO_UTF16, False),
        ("xliff", XLIFF_CRLF, False),
        ("json", JSON_COMPACT, True),
        ("json", JSON_BOM_CRLF, True),
        ("srt", SRT_CRLF_BOM, True),
        ("ini", INI_1252, True),
        ("csv", CSV_1252, True),
        ("txt", TXT_BOM, True),
        ("txt", TXT_1252, True),
        ("html", HTML_1252, True),
        ("html", HTML_APOS, True),
        ("icml", _icml_crlf(), True),
        ("docx", _docx(), True),
        ("xlsx", _xlsx(), True),
    ]
    for fmt, data, mono in cases:
        check_noop(fmt, data, mono)
    print(f"  noop: {len(cases)} soak shapes come back byte-identical OK")


def test_bilingual():
    po = get_filter("po")
    out = po.build(PO_CRLF, ident("po", PO_CRLF))
    assert out != PO_CRLF and b'msgstr "Hello {0}, welcome back!"' in out, out
    xl = get_filter("xliff")
    out = xl.build(XLIFF_CRLF, ident("xliff", XLIFF_CRLF))
    assert b"<target>Hello world</target>" in out, out
    xmb = (
        b'<?xml version="1.0" encoding="UTF-8"?>\n<messagebundle>\n'
        b'  <msg id="1">Hello world</msg>\n</messagebundle>\n'
    )
    assert get_filter("xmb").build(xmb, {}).find(b"translationbundle") != -1
    print("  bilingual: PO/XLIFF identity writes targets; XMB still exports XTB OK")


# --- translated files -----------------------------------------------------------
def test_po():
    po = get_filter("po")
    out = po.build_localized(PO_CRLF, {"0": "Bonjour {0}, bon retour !"}, "fr-FR")
    assert out.count(b"\n") == out.count(b"\r\n"), out  # CRLF throughout
    expected = PO_CRLF.replace(b'"Language: \\n"', b'"Language: fr_FR\\n"').replace(
        b"plural=(n != 1)", b"plural=(n > 1)"
    )
    expected = expected.replace(
        b'welcome back!"\r\nmsgstr ""', b'welcome back!"\r\nmsgstr "Bonjour {0}, bon retour !"'
    )
    assert out == expected, out
    # build() of a Windows-1252 catalog: same code page, same label, untouched lines kept
    out = po.build(PO_1252, {"0": "Le menu du caf\xe9"})
    assert b"charset=windows-1252" in out and b"Le menu du caf\xe9" in out, out
    assert b'msgid "Open"\nmsgstr ""\n' in out, out
    # the localized export declares UTF-8 (decision 9, B-intake), so its bytes are UTF-8
    for text, loc in (("Le menu du caf\xe9", "fr-FR"), ("\u30ab\u30d5\u30a7", "ja-JP")):
        out = po.build_localized(PO_1252, {"0": text}, loc)
        assert b"charset=UTF-8" in out and b"windows-1252" not in out, out
        body = out.decode("utf-8")  # the bytes match the declared charset
        assert 'msgid "The caf\xe9 menu"' in body and text in body, body
    # a HEADERLESS Windows-1252 catalog: build() keeps its bytes; the localized export
    # gains B-intake's UTF-8 header and is UTF-8 throughout (no mojibake)
    bare = PO_BARE_1252
    out = po.build(bare, {"1": "Ouvrir"})
    assert out == bare.replace(b'"Open"\r\nmsgstr ""', b'"Open"\r\nmsgstr "Ouvrir"'), out
    out = po.build_localized(bare, {"0": "Le menu du caf\xe9"}, "fr-FR")
    body = out.decode("utf-8")
    assert "charset=UTF-8" in body and "Language: fr_FR" in body, body
    assert 'msgid "The caf\xe9 menu"\r\nmsgstr "Le menu du caf\xe9"' in body, body
    assert out.count(b"\n") == out.count(b"\r\n"), out
    assert po.build(bare, {}) == bare
    # UTF-16 with its BOM stays UTF-16
    out = po.build(PO_UTF16, {"0": "Fermer"})
    assert out.startswith(b"\xfe\xff"), out[:4]
    assert 'msgstr "Fermer"\r\n'.encode("utf-16-be") in out, out
    print("  po: untouched lines, CRLF, Windows-1252 (when it fits) and UTF-16 kept OK")


def test_xliff():
    out = get_filter("xliff").build_localized(XLIFF_CRLF, {"0": "Bonjour le monde"}, "fr-FR")
    assert out.startswith(b'<?xml version="1.0" encoding="UTF-8"?>\r\n'), out[:60]
    assert out.endswith(b"</xliff>\r\n"), out[-20:]
    assert out.count(b"\n") == out.count(b"\r\n"), out
    assert b"<target>Bonjour le monde</target>" in out and b'target-language="fr-FR"' in out
    print("  xliff: declaration, CRLF and trailing newline kept OK")


def test_json():
    js = get_filter("json")
    out = js.build(JSON_COMPACT, {"0": "Bonjour {0}, bon retour !"})
    assert out == JSON_COMPACT.replace(b"Hello {0}, welcome back!", b"Bonjour {0}, bon retour !"), (
        out
    )
    out = js.build_localized(JSON_BOM_CRLF, {"1": "Deuxi\u00e8me ligne"}, "fr-FR")
    assert out == JSON_BOM_CRLF.replace(b"Second line", "Deuxi\u00e8me ligne".encode()), out
    print("  json: layout, escapes, BOM and CRLF kept; only the changed string moves OK")


def test_srt():
    srt = get_filter("srt")
    out = srt.build_localized(SRT_CRLF_BOM, {"1": "Vous avez de nouveaux fichiers."}, "fr-FR")
    assert out == SRT_CRLF_BOM.replace(
        b"You have new files.", b"Vous avez de nouveaux fichiers."
    ), out
    print("  srt: BOM, CRLF, blank-line spacing and untouched cues kept OK")


def test_legacy():
    out = get_filter("ini").build(INI_1252, {"1": "Au revoir \xe0 bient\xf4t"})
    assert out == INI_1252.replace(b"Good bye", b"Au revoir \xe0 bient\xf4t"), out
    out = get_filter("ini").build(INI_1252, {"1": "\u3055\u3088\u3046\u306a\u3089"})
    assert "caf\u00e9".encode() in out and "\u3055\u3088".encode() in out, out
    out = get_filter("csv").build(CSV_1252, {"1": "Au revoir \xe0 bient\xf4t"})
    assert out == CSV_1252.replace(b"Good bye", b"Au revoir \xe0 bient\xf4t"), out
    out = get_filter("csv").build(CSV_1252, {"1": "\u3055\u3088\u3046\u306a\u3089"})
    assert out.startswith(BOM) and "caf\u00e9".encode() in out, out  # Excel-readable UTF-8
    out = get_filter("txt").build(TXT_BOM, {"1": "Bonne nuit."})
    assert out == TXT_BOM.replace(b"Good night.", b"Bonne nuit."), out
    out = get_filter("txt").build(TXT_1252, {"1": "Bonne nuit \xe0 tous."})
    assert out == TXT_1252.replace(b"Good night.", b"Bonne nuit \xe0 tous."), out
    html = get_filter("html")
    out = html.build_localized(
        HTML_1252, {str(len(ident("html", HTML_1252)) - 1): "Bonne nuit."}, "fr-FR"
    )
    assert b'<meta charset="windows-1252">' in out and b"caf\xe9" in out, out
    assert b"Bonne nuit." in out and b'lang="fr' in out and out.count(b"\r\n") >= 6, out
    out = html.build_localized(HTML_1252, {"1": "\u304a\u3084\u3059\u307f"}, "ja-JP")
    assert b'charset="UTF-8"' in out, out
    out.decode("utf-8")
    print("  legacy: INI/CSV/TXT/HTML keep Windows-1252 when it fits, else UTF-8 OK")


def test_icml():
    data = _icml_crlf()
    out = get_filter("icml").build(data, {"0": "Bonjour le monde."})
    assert b"Bonjour le monde." in out and out.count(b"\n") == out.count(b"\r\n"), out
    print("  icml: a rebuilt story keeps CRLF OK")


# --- the export path ----------------------------------------------------------------
_N = [0]


def _ingest(name: str, data: bytes):
    _N[0] += 1
    with Session(engine) as s:
        p = Project(name=f"bytes-{_N[0]}", org_id=1, source_locale="en-US", default_provider="mock")
        s.add(p)
        s.commit()
        s.refresh(p)
        s.add(TargetLocale(project_id=p.id, locale="fr-FR"))
        s.commit()
        doc = ingest_document(s, p, name, data, TransMode.SEGMENTED)
        s.commit()
        return doc.id


def _set_targets(doc_id: int, fn) -> None:
    with Session(engine) as s:
        for seg in s.exec(select(Segment).where(Segment.document_id == doc_id)).all():
            new = fn(seg)
            if new is not None:
                seg.target_text = new
                seg.status = SegStatus.MT
                s.add(seg)
        s.commit()


def _export(doc_id: int) -> bytes:
    with Session(engine) as s:
        return export.build_translated_file(s, s.get(Document, doc_id), "fr-FR")


def test_export():
    for name, data in (
        ("a.srt", SRT_CRLF_BOM),
        ("a.po", PO_CRLF),
        ("bare.po", PO_BARE_1252),  # no header: an untouched export gains none
        ("a.xlf", XLIFF_CRLF),
        ("a.json", JSON_COMPACT),
        ("a.txt", TXT_1252),
        ("a.html", HTML_1252),
        ("a.docx", _docx()),
    ):
        did = _ingest(name, data)
        assert _export(did) == data, f"{name}: untouched export is not the original"
    # a plain-text file whose every target equals its source: nothing written
    did = _ingest("same.txt", TXT_BOM)
    _set_targets(did, lambda seg: seg.source_text)
    assert _export(did) == TXT_BOM
    # HTML declares the target language, so a source-equal target is still a change
    did = _ingest("same.html", HTML_APOS)
    _set_targets(did, lambda seg: seg.source_text)
    out = _export(did)
    assert out != HTML_APOS and b'lang="fr' in out, out
    # a translated headerless Windows-1252 PO: B-intake's UTF-8 header, UTF-8 bytes
    did = _ingest("t.po", PO_BARE_1252)
    _set_targets(did, lambda seg: "Ouvrir" if seg.source_text == "Open" else None)
    body = _export(did).decode("utf-8")
    assert "charset=UTF-8" in body and 'msgid "The caf\xe9 menu"' in body, body
    assert 'msgstr "Ouvrir"' in body, body
    # a translated SRT keeps CRLF, BOM and every untouched cue
    did = _ingest("t.srt", SRT_CRLF_BOM)
    _set_targets(
        did, lambda seg: "Bonne nuit." if seg.source_text.startswith("Good night") else None
    )
    assert _export(did) == SRT_CRLF_BOM.replace(b"Good night.  ", b"Bonne nuit."), _export(did)
    print("  export: untouched and source-equal plain files are the original; SRT kept OK")


def _decls(data: bytes) -> list:
    import re

    text = data.decode("latin-1")
    return re.findall(
        r"""(?:encoding\s*=\s*["']|charset\s*=\s*["']?)([A-Za-z0-9_.:-]+)""", text, re.I
    )


def _agrees(data: bytes) -> bool:
    """Every charset the file declares decodes its bytes (bytes and labels agree)."""
    labels = _decls(data)
    for label in labels:
        try:
            data.decode(label)
        except (UnicodeDecodeError, LookupError):
            return False
    return True


def test_gate_charset_labels():
    """B-bytes gate: a legacy code page is kept only when the text fits AND every
    declaration in the file names the codec written; otherwise UTF-8 under UTF-8 labels.
    Never Windows-1252 bytes under a UTF-8 or ISO-8859-1 label (mojibake)."""
    html = get_filter("html")
    po = get_filter("po")
    iso_html = (
        '<html><head><meta charset="iso-8859-1"><title>Caf\xe9</title></head>'
        "<body><p>Caf\xe9 cr\xe8me</p></body></html>"
    ).encode("latin-1")
    two_html = (
        '<html><head><meta http-equiv="Content-Type" content="text/html; charset=windows-1252">'
        '<meta charset="windows-1252"></head><body><p>Caf\xe9 cr\xe8me</p></body></html>'
    ).encode("cp1252")
    late_html = (
        "<html><head><script>/*" + "x" * 5000 + '*/</script><meta charset="windows-1252">'
        "</head><body><p>Caf\xe9 cr\xe8me</p></body></html>"
    ).encode("cp1252")
    iso_po = (
        b'msgid ""\nmsgstr ""\n"Content-Type: text/plain; charset=ISO-8859-1\n"\n\n'
        b'msgid "Caf\xe9"\nmsgstr ""\n'
    )
    for name, filt, data in (
        ("iso html", html, iso_html),
        ("two-meta html", html, two_html),
        ("late-meta html", html, late_html),
        ("iso po", po, iso_po),
    ):
        units = filt.parse(data).units
        fits = {str(u.p_index): u.text + " \xe0 la cr\xe8me" for u in units}
        nofit = {str(u.p_index): u.text + " \u0436" for u in units}
        for mapping, want in ((fits, "\xe0 la cr\xe8me"), (nofit, "\u0436")):
            builds = [filt.build(data, mapping)]
            if filt is html:
                builds.append(filt.build_localized(data, mapping, "fr-FR"))
            for out in builds:
                assert _agrees(out), (name, _decls(out), out[:200])
                labels = _decls(out)
                assert labels, (name, out[:200])
                assert want in out.decode(labels[0]), (name, labels, out[:300])
        assert filt.build(data, {}) == data, name
    # the ISO-8859-1 page that fits stays ISO-8859-1 with its own label
    out = html.build(iso_html, {"0": "Caf\xe9!"})
    assert _decls(out) == ["iso-8859-1"] and "Caf\xe9!".encode("latin-1") in out, out
    print("  gate: legacy code page only when bytes and every declaration agree OK")


def test_gate_json_duplicate_keys():
    """B-bytes gate: the generic JSON splice with a duplicate key (walk order is not file
    order) still writes valid JSON with every translation in place."""
    import json

    f = get_filter("json")
    data = b'{"a": "x one", "b": "y two", "a": "z three", "c": ["q r", 1]}'
    assert [u.text for u in f.parse(data).units] == ["z three", "y two", "q r"]
    out = f.build(data, {"0": "ZZ", "1": "YY", "2": "Q"})
    assert json.loads(out) == {"a": "ZZ", "b": "YY", "c": ["Q", 1]}, out
    assert out == b'{"a": "x one", "b": "YY", "a": "ZZ", "c": ["Q", 1]}', out
    # the localized export (B-plurals' i18next fill) keeps the splice when it adds nothing,
    # and still fills a BOM file's missing plural keys (BOM and CRLF kept)
    compact = b'{"title":"Hello","body":"Some text"}'
    out = f.build_localized(compact, {"0": "Bonjour"}, "fr-FR")
    assert out == b'{"title":"Bonjour","body":"Some text"}', out
    i18n = b'\xef\xbb\xbf{\r\n  "item_one": "one item",\r\n  "item_other": "many items"\r\n}\r\n'
    out = f.build_localized(i18n, {"0": "jeden", "1": "wiele"}, "pl-PL")
    assert out.startswith(b"\xef\xbb\xbf") and b"\r\n" in out and b"\n\n" not in out, out
    got = json.loads(out[3:])
    assert got["item_few"] == got["item_many"] == "wiele" and got["item_one"] == "jeden", got
    print("  gate: JSON splice with duplicate keys stays valid; localized splice kept OK")


if __name__ == "__main__":
    test_noop()
    test_bilingual()
    test_po()
    test_xliff()
    test_json()
    test_srt()
    test_legacy()
    test_icml()
    test_export()
    test_gate_charset_labels()
    test_gate_json_duplicate_keys()
    print("BYTE IDENTICAL EXPORTS OK")
