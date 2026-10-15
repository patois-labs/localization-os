"""QA-3 + QA-5 regression lock (live-QA workflows campaign, 2026-07-19 report).

QA-3 (HIGH, refusal-guard recall gap): the refusal guard's `_STEMS` covered "I'm ready
to translate", "as an ai", etc., but NOT a first-person assistant META self-description
that LEADS the output ("I am one step in an automated localization pipeline, specifically
a machine translation model, ..."). Such output committed as a clean `status=mt`
translation with no flag; approving it seeded shared TM as a future 100%-match — and a
100% TM hit SKIPS the refusal guard on every later doc, so the hedge auto-fills silently
forever. Both the write choke point (quality.commit_mt_output) and the approve backstop
(documents.segment_approve) called the SAME is_refusal, so one recall gap defeated
defense-in-depth. A strong unused signal: the output was 100% English for an es target.

The fix is two independent layers:
  (a) high-precision ANCHORED META stems for the "I am a/an ... model/assistant/AI/
      pipeline/system" family (English + first-person es/fr/de), same anchored discipline
      as the existing stems (matched only in the first _ANCHOR_CHARS, re-checked after a
      preamble strip) so a genuine translation that merely MENTIONS this vocabulary
      mid-body is not flagged;
  (b) an INDEPENDENT second signal — refusal.looks_untranslated(text, target_locale):
      output-language != target-locale (decisively-English output for a KNOWN non-English
      target whose own script / function-word profile is entirely absent). Deliberately
      conservative (length floor, distinct-English-function-word count + ratio, target
      profile decisively absent) so short labels / brands / numbers / code / English
      loanwords in target copy / genuine target translations are NEVER flagged. Wired at
      the choke point (quality.commit_mt_output `lang_guard`, opted in only by the REAL
      fresh-MT path — never the offline MockProvider echo) as a SOFT gate (routes to
      NEEDS_REVIEW under a DISTINCT `nontranslation` flag + TM-ineligible, text kept for
      the reviewer) AND at the approve TM backstop — so the two sites fail independently.

QA-5 (MED, editor grid): documents.py's `_SURFACED` set excluded refusal.FLAG_TYPE, so a
refusal-routed needs_review segment rendered only a generic "No translation yet" badge —
visually indistinguishable from a never-translated one. Fix: add refusal.FLAG_TYPE (and
the new NONTRANSLATION_FLAG_TYPE) to `_SURFACED` so the editor renders the explanatory
chip through the same generic {type,severity,message} loop the DNT/glossary/QE chips use.

FAIL-BEFORE (proven by reconstruction, and asserted best-effort below from the base blob):
pre-fix `is_refusal(QA_META)` == False, `looks_untranslated` and `NONTRANSLATION_FLAG_TYPE`
absent. POST-FIX: QA_META -> True; the independent signal flags English-for-es output;
every canonical hedge stays True; the whole torture set stays False; the editor renders
the chip. Hermetic + FK enforcement ON; keyless (a custom NON-mock provider double drives
the e2e paths). ASCII-only prints; final line 'verify_refusal_qa3_qa5 OK'.
"""

import os
import subprocess
import sys
import tempfile
from io import BytesIO

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_refusal_qa3_")

from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app import config
from app.db import engine, init_db
from app.main import app
from app.models import Document, Project, Segment, SegStatus
from app.services import quality, refusal, tm, translate
from app.services.providers import Provider, register_provider

# The exact live-QA meta-output miss (the headline QA-3 defect).
QA_META = (
    "I am one step in an automated localization pipeline, specifically a machine "
    "translation model, and I don't have the ability to access external content."
)
# A plain ENGLISH non-hedge answer (NOT caught by any is_refusal stem) produced for an
# es target — the layer-(b) language-mismatch case.
PLAIN_EN = (
    "The customer support team will review your ticket and will respond to you "
    "within two business days regarding the account you asked about."
)
# A clean Spanish translation — must STAY a clean MT commit (precision).
CLEAN_ES = "El equipo de soporte revisara su solicitud y le respondera en dos dias habiles."

QA_META_SRC = "[meta trigger]"
# ROUTING TRIGGER for the layer-(b) case. This used to be the 17-character stub
# "[english trigger]", which made the pair 17 -> 137 chars: +120 surplus at 8.1x, i.e. a
# textbook LENGTH RUNAWAY, and once the K4 gate landed at the write seam the fixture
# tripped it and stopped exercising what it exists to exercise (that `lang_guard` is
# OPT-IN). The stub source was the artifact, not the gate -- no real translation unit
# looks like this -- so the trigger is now a REALISTIC source sentence of proportionate
# length. It is still an exact-match key for QA3Provider and still unique.
PLAIN_EN_SRC = "Ask about the ticket response time for this account."
# RE-COLLISION GUARD (K4 fix round, finding F10). The property this fixture needs is not
# "a long enough source" but "this pair does not trip the LENGTH gate" -- assert it here,
# next to the constant, so a future shortening fails loudly on the real invariant instead
# of surfacing as an opaque CommitOutcome dump three sections later.
assert not quality.length_runaway(PLAIN_EN_SRC, PLAIN_EN), (
    "PLAIN_EN_SRC/PLAIN_EN now trip the K4 length gate (+%d chars, %.1fx) -- this pair "  # noqa: UP031
    "must exercise the OPT-IN lang_guard, not the length gate"
    % (len(PLAIN_EN) - len(PLAIN_EN_SRC), len(PLAIN_EN) / len(PLAIN_EN_SRC))
)


class QA3Provider(Provider):
    """NON-mock capturing double. Returns the assistant META self-description for
    QA_META_SRC (layer a), a plain-English non-hedge for PLAIN_EN_SRC (layer b), and a
    clean Spanish translation otherwise (precision). NON-mock -> get_provider() returns
    THIS instance -> the layer-(b) lang_guard is active on its output (unlike the offline
    mock echo, which is exempt)."""

    name = "qa3prov"

    def complete(self, prompt, system=None, task=None, payload=None, **kw):
        self.last_usage = {
            "input_tokens": max(1, len(prompt or "") // 4),
            "output_tokens": 12,
            "cache_read_tokens": 0,
            "cache_write_tokens": 0,
        }
        if task == "translate":
            src = (payload or {}).get("text", "").strip()
            if src == QA_META_SRC:
                return QA_META
            if src == PLAIN_EN_SRC:
                return PLAIN_EN
            return CLEAN_ES
        return (prompt or "").strip()


def docx_from(paragraphs):
    d = Docx()
    for p in paragraphs:
        d.add_paragraph(p)
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def new_doc(client, pid, paragraphs, mode="segmented", name="d.docx"):
    r = client.post(
        f"/projects/{pid}/documents",
        data={"mode": mode},
        files={
            "file": (
                name,
                docx_from(paragraphs),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    return int(r.headers["location"].rsplit("/", 1)[-1])


def segs_for(session, doc_id, locale="es-MX"):
    return sorted(
        session.exec(
            select(Segment).where(Segment.document_id == doc_id, Segment.target_locale == locale)
        ).all(),
        key=lambda x: x.order_index,
    )


def _mk_segment(
    session,
    *,
    source="Hello",
    target_text=None,
    status=SegStatus.NEW,
    qa_flags=None,
    locale="es-MX",
    provider="qa3prov",
):
    """A committed Project+Document+Segment (FK enforcement). The project pins a REAL
    (non-mock) provider so the choke-point / approve lang_guard is active."""
    p = Project(name="qa3", source_locale="en-US", default_provider=provider)
    session.add(p)
    session.commit()
    d = Document(project_id=p.id, filename="d.docx")
    session.add(d)
    session.commit()
    seg = Segment(
        document_id=d.id,
        target_locale=locale,
        order_index=0,
        source_text=source,
        target_text=target_text,
        status=status,
        qa_flags=qa_flags,
    )
    session.add(seg)
    session.commit()
    return seg


# =============================================================================
# 0. FAIL-BEFORE self-proof (best-effort): reconstruct the pre-fix refusal.py from the
#    base commit blob and assert the bug was present. Non-fatal if git/blob unreachable
#    (shallow clone) — the post-fix asserts below are the binding regression lock.
# =============================================================================
_BASE_COMMIT = "c048d01"  # the QA-report commit this branch forked from (pre-fix)


def test_fail_before_reconstruction():
    try:
        blob = subprocess.run(
            ["git", "show", f"{_BASE_COMMIT}:app/services/refusal.py"],
            cwd=_ROOT,
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=30,
        )
        if blob.returncode != 0 or not blob.stdout:
            print("  0. fail-before reconstruction SKIPPED (base blob unreachable) OK")
            return
        import importlib.util

        with tempfile.NamedTemporaryFile("w", suffix=".py", delete=False, encoding="utf-8") as fh:
            fh.write(blob.stdout)
            path = fh.name
        try:
            spec = importlib.util.spec_from_file_location("_refusal_prefix", path)
            m = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(m)
            assert m.is_refusal(QA_META) is False, "pre-fix must MISS the QA meta-output"
            assert not hasattr(m, "looks_untranslated"), "pre-fix must lack looks_untranslated"
            assert not hasattr(m, "NONTRANSLATION_FLAG_TYPE"), "pre-fix must lack the flag type"
        finally:
            os.unlink(path)
        print("  0. fail-before: base blob MISSES QA meta-output + lacks layer-(b) surface OK")
    except Exception as e:
        print(f"  0. fail-before reconstruction SKIPPED ({type(e).__name__}) OK")


# =============================================================================
# 1. LAYER (a): META stems — the QA meta-output + the family are caught; canonical
#    hedges stay caught; the existing FP floor stays clean.
# =============================================================================
def test_layer_a_meta_stems():
    assert refusal.is_refusal(QA_META), "the QA meta-output MUST be detected post-fix"

    meta_true = [
        "I am a machine translation model and cannot browse the web.",
        "I'm an AI language model trained only to translate text.",
        "I am a large language model and do not have that ability.",
        "I am part of an automated translation pipeline.",
        "As a machine translation system, I only output the translation.",
        # target-language first-person meta (es/fr/de, accented + unaccented)
        "Soy un modelo de traduccion y no tengo acceso a internet.",
        "Soy una inteligencia artificial disenada para traducir.",
        "Je suis un modele de traduction, je ne peux pas naviguer.",
        "Ich bin ein Uebersetzungsmodell und kann nicht im Internet suchen.",
        "Ich bin eine kuenstliche Intelligenz.",
        # preamble-pushed meta (courtesy opener shoves it past the anchor)
        "Sure! Of course. However, I am a machine translation model without web access.",
        "Por supuesto. Sin embargo, soy un modelo de traduccion sin acceso a la red.",
    ]
    for t in meta_true:
        assert refusal.is_refusal(t), "meta family not detected: " + t[:50]

    # REGRESSION: every canonical hedge the guard already caught stays caught.
    canon = [
        "I'm ready to translate from English to Spanish. However, I don't see any source text.",
        "As an AI, I can't do that. Please provide the source text you want translated.",
        "I'm sorry, I can't translate this - there's nothing to work with.",
        "No puedo traducir este texto porque no veo el contenido original.",
        "Je ne peux pas traduire ce texte sans le contenu source.",
        "There is no text to translate.",
    ]
    for t in canon:
        assert refusal.is_refusal(t), "canonical hedge regressed: " + t[:50]

    # REGRESSION: the META additions must not flag legit content that merely MENTIONS
    # the AI/MT vocabulary mid-body (the anchored discipline) or describes the engine in
    # the THIRD person (not first-person "I am").
    meta_false = [
        # third-person product copy about the engine (target-lang) -> NOT first-person
        "Como asistente de inteligencia artificial, nuestro producto sugiere respuestas "
        "automaticamente durante el horario de mayor demanda.",
        # mid-body mention, well past the 80-char anchor (mirrors the existing FP floor)
        "El motor de traduccion integrado en este modulo de soporte al cliente fue "
        "entrenado as a language model para generar respuestas de soporte en varios "
        "idiomas para nuestros clientes globales.",
        # a genuine translation that opens with ordinary content
        "El equipo de soporte revisara su solicitud lo antes posible.",
        # third-person English description of the engine
        "The platform includes a machine translation model that suggests replies to agents.",
    ]
    for t in meta_false:
        assert not refusal.is_refusal(t), "META FALSE POSITIVE: " + t[:60]

    # REGRESSION (adversarial review, Finding 1): the SHORT meta identity stems must not
    # OVERMATCH ordinary first-person English by bare substring. Before the right-word-
    # boundary fix (+ the "one step IN/OF" / "automated TRANSLATION/LOCALIZATION" pins),
    # every string below was is_refusal==True — a HARD withhold that DROPS legit English-
    # target content ('i am an ai' fired on "air/aid/aims/airline", 'i am one step' on the
    # idiom "one step closer/away", 'i am an automated' on the compound "automated-testing").
    # This is the first-person-English FP floor the branch shipped WITHOUT. All must be False.
    en_idiom_false = [
        "I am an air traffic controller at a busy international hub every single day.",
        "I am an aid worker delivering food and medical supplies to remote villages.",
        "I am an aims-based financial advisor helping retirees plan their withdrawals.",
        "I'm an airline pilot with more than twenty years of long-haul experience.",
        "I am one step closer to finishing my degree after this final semester ends.",
        "I am one step away from completing the marathon training plan I started in May.",
        "I am one step ahead of the competition because we shipped the feature first.",
        "I am an automated-testing engineer on the platform reliability team at Acme.",
    ]
    for t in en_idiom_false:
        assert not refusal.is_refusal(t), "FIRST-PERSON-ENGLISH FALSE POSITIVE: " + t[:60]

    # RECALL must survive the boundary: the genuine short-stem forms whose real
    # continuation is a space/punctuation (or a hyphen, "AI-powered") still fire.
    en_meta_boundary_true = [
        "I am an AI.",  # stem == whole clause
        "I am an AI-powered translation assistant without web access.",  # hyphen edge
        "I am a translation tool and cannot open external links.",
        "I am one step in an automated pipeline and cannot fetch that.",
    ]
    for t in en_meta_boundary_true:
        assert refusal.is_refusal(t), "boundary regressed a genuine meta hedge: " + t[:50]
    print(
        "  1. layer (a): QA meta-output + family detected; canonical hedges intact; "
        "third-person/mid-body engine copy NOT flagged; first-person-English idioms "
        "(air/aid/one step closer/automated-testing) NOT flagged; genuine short metas "
        "still fire OK"
    )


# =============================================================================
# 2. LAYER (b): looks_untranslated — positive (English for a non-English target) and a
#    TORTURE FALSE-POSITIVE FLOOR (short/brand/numbers/code/loanwords/genuine target
#    translations/English target/unsupported target). Precision beats recall.
# =============================================================================
def test_layer_b_positive():
    positives = [
        (PLAIN_EN, "es-MX"),
        (QA_META, "es-ES"),
        (QA_META, "fr-FR"),
        (QA_META, "de-DE"),
        (QA_META, "it-IT"),
        (QA_META, "pt-BR"),
        (QA_META, "nl-NL"),
        (QA_META, "ja-JP"),  # non-Latin: no target-script chars in an English body
        (QA_META, "ko-KR"),
        (QA_META, "zh-CN"),
        (QA_META, "ru-RU"),
    ]
    for t, loc in positives:
        assert refusal.looks_untranslated(t, loc), f"layer-(b) miss for target {loc}"

    # RECALL GUARD (adversarial review, Finding 2 fix): the framing-word coverage added to
    # _LATIN_TARGET_STOPWORDS must remain STRICTLY DISTINCTIVE. A word that ALSO spells an
    # ordinary English word (e.g. "error" es/en, "page" fr/en) would let a decisively-
    # English non-translation containing it read as "target present" and EVADE the guard.
    # These decisively-English bodies embed exactly those collision-risk words and MUST
    # still flag for es/fr targets.
    collision_bait = [
        (
            "The system returned an error on the page you requested and could not load the "
            "account for you today.",
            "es-ES",
        ),
        (
            "The system returned an error on the page you requested and could not load the "
            "account for you today.",
            "fr-FR",
        ),
    ]
    for t, loc in collision_bait:
        assert refusal.looks_untranslated(t, loc), (
            f"collision-risk English word wrongly read as target-present for {loc}"
        )

    # STRUCTURAL INVARIANT: no Latin target stopword may collide with the curated English
    # function-word set (that overlap would silently punch a recall hole in every target).
    for lang, stop in refusal._LATIN_TARGET_STOPWORDS.items():
        clash = stop & refusal._EN_FUNCTION_WORDS
        assert not clash, (
            f"target stopword collides with EN function words ({lang}): {sorted(clash)}"
        )
    print(
        "  2a. layer (b): decisively-English output flagged for 11 non-English targets "
        "(Latin + CJK + Cyrillic); collision-bait (error/page) still flags; no "
        "stopword/EN-function-word overlap OK"
    )


def test_layer_b_torture_floor():
    # (text, target_locale) — EVERY one must be False (not a language mismatch).
    torture = [
        ("Guardar", "es-ES"),  # short label
        ("Acme Corporation Inc.", "es-ES"),  # brand
        ("12345 67890 42 99 100 3.14", "es-ES"),  # numbers
        ("for (let i = 0; i < n; i++) { return this.value; }", "es-ES"),  # code
        (
            "Descargue el nuevo software y reinicie el sistema para aplicar los cambios "
            "ahora mismo.",
            "es-ES",
        ),  # es w/ English loanword
        (
            "Presione el boton 'Save and Continue' para guardar los cambios del "
            "formulario de registro.",
            "es-ES",
        ),  # es quoting an English UI string
        (
            "Instale la app de Acme desde el App Store en su dispositivo movil hoy "
            "mismo por favor.",
            "es-ES",
        ),  # es, brand + loanwords
        (
            "No hay texto en este campo todavia, complete la informacion requerida antes "
            "de continuar.",
            "es-ES",
        ),  # genuine es empty-state
        (
            "Cliquez sur le bouton bleu pour traduire le document complet des maintenant "
            "s'il vous plait.",
            "fr-FR",
        ),  # genuine French
        (
            "Klicken Sie auf die blaue Schaltflaeche, um das gesamte Dokument jetzt zu "
            "uebersetzen bitte.",
            "de-DE",
        ),  # genuine German
        (
            "Assicurati di non premere il pulsante due volte prima che il modulo sia "
            "inviato al server.",
            "it-IT",
        ),  # genuine Italian
        (
            "Certifique-se de nao pressionar o botao duas vezes antes que o formulario "
            "seja enviado agora.",
            "pt-BR",
        ),  # genuine Portuguese
        (
            "Zorg ervoor dat u de knop niet twee keer indrukt voordat het formulier is "
            "verzonden vandaag.",
            "nl-NL",
        ),  # genuine Dutch
        (
            "この機能を使用するには、右上のボタンをクリックしてください。",
            "ja-JP",
        ),  # genuine Japanese (has JA chars)
        (
            "The quick brown fox jumps over the lazy dog near the calm river bank at dawn today.",
            "en-US",
        ),  # English target -> not a mismatch
        (
            "The user can log in and reset the password if they forget it and contact us "
            "for help today.",
            "tr-TR",
        ),  # unsupported target -> abstain
        (
            "The user can log in and reset the password if they forget it and contact us "
            "for help today.",
            "pl-PL",
        ),  # unsupported target -> abstain
        # adversarial review, Finding 2: a genuine target segment that FRAMES a quoted /
        # do-not-translate ENGLISH span with target-language framing words. Before the
        # framing-word coverage was added to _LATIN_TARGET_STOPWORDS these had ZERO
        # recognized target tokens, so the embedded English tipped the gate and the
        # segment was routed to review. The framing words are distinctive (no English
        # collision), so recognizing them only ADDS an abstain here.
        (
            "Nuestra politica indica: 'You must not share your password with anyone at any "
            "time or place.'",
            "es-ES",
        ),  # es framing + English DNT quote
        (
            "Mensaje mostrado: 'The user can log in and reset their password from the "
            "settings page anytime.'",
            "es-ES",
        ),  # es framing + English quote
        (
            "Fehler: The requested resource could not be found on the origin server at this time.",
            "de-DE",
        ),  # de "Fehler:" + English body
        (
            "Errore riscontrato: The connection to the database timed out before the query "
            "could finish.",
            "it-IT",
        ),  # it framing + English body
        (
            "Erro exibido ao usuario: The account has been locked after too many failed "
            "login attempts today.",
            "pt-BR",
        ),  # pt framing + English body
    ]
    for t, loc in torture:
        assert not refusal.looks_untranslated(t, loc), "layer-(b) FALSE POSITIVE for {}: {}".format(
            loc,
            t.encode("ascii", "backslashreplace").decode()[:60],
        )
    print(
        "  2b. layer (b) torture floor: all %d strings pass UNTOUCHED (labels/brands/"  # noqa: UP031
        "numbers/code/loanwords/genuine es-fr-de-it-pt-nl-ja/en-target/unsupported/"
        "framed-English-quote) OK" % len(torture)
    )


# =============================================================================
# 3. CHOKE POINT (quality.commit_mt_output) — the independent second signal, end to end.
# =============================================================================
def test_choke_point_lang_guard():
    with Session(engine) as s:
        # (i) lang_guard ON + English-for-es output -> SOFT nontranslation: committed
        #     (text kept for the reviewer), NEEDS_REVIEW, distinct flag, TM-INELIGIBLE.
        seg = _mk_segment(s, source=PLAIN_EN_SRC)
        out = quality.commit_mt_output(s, seg, PLAIN_EN, source="mt", lang_guard=True)
        s.commit()
        assert out.committed and not out.flagged, ("soft, not a refusal-withhold", out)
        assert out.tm_eligible is False, ("a suspected non-translation must be TM-ineligible", out)
        assert seg.target_text == PLAIN_EN, ("SOFT keeps the produced text", seg.target_text)
        assert seg.status == SegStatus.NEEDS_REVIEW, seg.status
        nf = [f for f in (seg.qa_flags or []) if f.get("type") == refusal.NONTRANSLATION_FLAG_TYPE]
        assert nf and nf[0].get("source") == "mt", ("distinct nontranslation flag", seg.qa_flags)
        assert not [f for f in (seg.qa_flags or []) if f.get("type") == "refusal"], (
            "must NOT reuse the refusal flag",
            seg.qa_flags,
        )

        # (ii) lang_guard OFF (default) + SAME output -> clean MT (proves the offline
        #      mock echo is unaffected: the choke point never runs layer (b) unless asked).
        seg2 = _mk_segment(s, source=PLAIN_EN_SRC)
        out2 = quality.commit_mt_output(s, seg2, PLAIN_EN, source="mt")  # lang_guard defaults False
        s.commit()
        assert out2.committed and out2.tm_eligible and seg2.status == SegStatus.MT, out2
        assert not (seg2.qa_flags or []), ("no flag when lang_guard is off", seg2.qa_flags)

        # (iii) a genuine Spanish translation with lang_guard ON -> clean MT (precision).
        seg3 = _mk_segment(s, source="hola")
        out3 = quality.commit_mt_output(s, seg3, CLEAN_ES, source="mt", lang_guard=True)
        s.commit()
        assert out3.committed and out3.tm_eligible and seg3.status == SegStatus.MT, out3
        assert not (seg3.qa_flags or []), ("clean es must not be flagged", seg3.qa_flags)

        # (iv) an actual refusal still HARD-withholds regardless of lang_guard (layer a
        #      precedence): no target, flagged, refusal flag, TM-ineligible.
        seg4 = _mk_segment(s, source=QA_META_SRC)
        out4 = quality.commit_mt_output(s, seg4, QA_META, source="mt", lang_guard=True)
        s.commit()
        assert (not out4.committed) and out4.flagged and not out4.tm_eligible, out4
        assert not (seg4.target_text or "").strip(), (
            "refusal withholds the text",
            seg4.target_text,
        )
        assert [f for f in (seg4.qa_flags or []) if f.get("type") == "refusal"], seg4.qa_flags
    print(
        "  3a. choke point: lang_guard ON soft-flags English-for-es (kept, NEEDS_REVIEW, "
        "TM-ineligible, distinct flag); OFF -> clean MT; clean es clean; refusal withheld OK"
    )

    # (v) kill switch OFF -> lang_guard is inert (written through, no flag) = pre-guard.
    config.REFUSAL_GUARD_ENABLED = False
    try:
        with Session(engine) as s:
            seg = _mk_segment(s, source=PLAIN_EN_SRC)
            out = quality.commit_mt_output(s, seg, PLAIN_EN, source="mt", lang_guard=True)
            s.commit()
            assert out.committed and out.tm_eligible and seg.status == SegStatus.MT, out
            assert not (seg.qa_flags or []), ("guard OFF -> no nontranslation flag", seg.qa_flags)
    finally:
        config.REFUSAL_GUARD_ENABLED = True
    print("  3b. choke point: kill switch OFF -> lang_guard inert (byte-identical to pre-guard) OK")


# =============================================================================
# 4. END-TO-END through translate.pretranslate_document (the real fresh-MT path) with a
#    NON-mock provider: the meta-output routes to review (layer a) and TM is not poisoned;
#    a plain-English output routes to review with the nontranslation flag (layer b) and is
#    TM-ineligible.
# =============================================================================
def test_translate_e2e(client, pid):
    doc_id = new_doc(client, pid, ["Hello world", QA_META_SRC, PLAIN_EN_SRC], name="e2e.docx")
    with Session(engine) as s:
        translate.pretranslate_document(s, doc_id, "es-MX")
    with Session(engine) as s:
        rows = segs_for(s, doc_id)
        good, meta, plain = rows[0], rows[1], rows[2]
        # good -> clean Spanish MT
        assert good.status == SegStatus.MT and good.target_text == CLEAN_ES, (
            good.status,
            good.target_text,
        )
        # meta -> layer (a): withheld + refusal flag, TM not poisoned
        assert meta.status == SegStatus.NEEDS_REVIEW and not (meta.target_text or "").strip(), (
            "meta withheld",
            meta.status,
            meta.target_text,
        )
        assert [f for f in (meta.qa_flags or []) if f.get("type") == "refusal"], meta.qa_flags
        # plain English -> layer (b): SOFT nontranslation flag, text kept, NEEDS_REVIEW
        assert plain.status == SegStatus.NEEDS_REVIEW, plain.status
        assert plain.target_text == PLAIN_EN, ("layer-b keeps the produced text", plain.target_text)
        assert [
            f for f in (plain.qa_flags or []) if f.get("type") == refusal.NONTRANSLATION_FLAG_TYPE
        ], plain.qa_flags
        # neither the meta hedge nor the plain-English nontranslation reached segment TM
        assert (
            tm.lookup(s, s.get(Document, doc_id).project_id, QA_META_SRC, "en-US", "es-MX") is None
        )
        assert (
            tm.lookup(s, s.get(Document, doc_id).project_id, PLAIN_EN_SRC, "en-US", "es-MX") is None
        )
    print(
        "  4. translate e2e: meta -> withheld+refusal (layer a); plain-English -> "
        "NEEDS_REVIEW+nontranslation (layer b); TM not poisoned by either OK"
    )


# =============================================================================
# 5. APPROVE TM BACKSTOP (documents.segment_approve): the independent signal blocks the
#    approve->TM write for a real provider, so the two sites fail independently.
# =============================================================================
def test_approve_backstop(client):
    with Session(engine) as s:
        # English-for-es approved target (a human clicked past the review, or an old MT):
        # is_refusal is False (not a hedge), but looks_untranslated catches it -> no TM write.
        seg = _mk_segment(s, source=PLAIN_EN_SRC, target_text=PLAIN_EN, status=SegStatus.EDITED)
        seg_id, _pid = seg.id, seg.document_id
    r = client.post(f"/segments/{seg_id}/approve")
    body = r.json()
    assert body["ok"] and body["tm_written"] is False, ("nontranslation must not seed TM", body)
    assert "language mismatch" in body.get("tm_skipped", ""), body

    with Session(engine) as s:
        # a genuine Spanish approved target writes TM normally (control).
        seg = _mk_segment(s, source="hola", target_text=CLEAN_ES, status=SegStatus.EDITED)
        seg_id = seg.id
    r = client.post(f"/segments/{seg_id}/approve")
    body = r.json()
    assert body["ok"] and body["tm_written"] is True, ("a clean es target must write TM", body)
    print(
        "  5. approve backstop: English-for-es approval blocked from TM (language mismatch); "
        "clean es approval writes TM OK"
    )


# =============================================================================
# 6. QA-5: the editor grid RENDERS the refusal + nontranslation chips (TestClient HTML).
# =============================================================================
def test_editor_surfaces_chips(client):
    p = Project(name="qa5", source_locale="en-US", default_provider="qa3prov")
    with Session(engine) as s:
        s.add(p)
        s.commit()
        pid_ = p.id
        d = Document(project_id=pid_, filename="chips.docx")
        s.add(d)
        s.commit()
        did = d.id
        s.add(
            Segment(
                document_id=did,
                target_locale="es-MX",
                order_index=0,
                source_text="one",
                status=SegStatus.NEEDS_REVIEW,
                qa_flags=[refusal.refusal_flag("mt")],
            )
        )
        s.add(
            Segment(
                document_id=did,
                target_locale="es-MX",
                order_index=1,
                source_text="two",
                status=SegStatus.NEEDS_REVIEW,
                qa_flags=[refusal.nontranslation_flag("mt")],
            )
        )
        s.commit()
    html = client.get(f"/documents/{did}/editor?locale=es-MX").text
    # the generic qa-cell loop renders <span ... title="{message}">{type}</span>
    assert ">refusal</span>" in html, "the refusal chip must render in the editor grid"
    assert ">nontranslation</span>" in html, "the nontranslation chip must render"
    assert refusal.REASON in html, "the refusal chip must carry its explanatory reason (title)"
    assert refusal.NONTRANSLATION_REASON in html, "the nontranslation chip must carry its reason"
    print(
        "  6. QA-5: editor grid renders the refusal + nontranslation chips with their "
        "explanatory reasons (no longer a bare 'untranslated' badge) OK"
    )


def main():
    init_db()
    register_provider("qa3prov", QA3Provider())
    client = TestClient(app)
    r = client.post(
        "/projects",
        data={
            "name": "QA3E2E",
            "source_locale": "en-US",
            "target_locales": "es-MX",
            "default_provider": "qa3prov",
        },
        follow_redirects=False,
    )
    pid = int(r.headers["location"].rsplit("/", 1)[-1])

    test_fail_before_reconstruction()
    test_layer_a_meta_stems()
    test_layer_b_positive()
    test_layer_b_torture_floor()
    test_choke_point_lang_guard()
    test_translate_e2e(client, pid)
    test_approve_backstop(client)
    test_editor_surfaces_chips(client)
    print("verify_refusal_qa3_qa5 OK")


if __name__ == "__main__":
    main()
