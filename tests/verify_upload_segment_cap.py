"""#571: one big plain-text upload must not take the instance down.

Hermetic: throwaway LOC_TMS_DATA_DIR, provider keys cleared, ASCII-only prints.

Covers:
  * an over-cap .txt (too many lines) gets a fast 413 on POST /api/documents and NO
    Document or Segment row is stored;
  * the cap counts SEGMENTS, not only lines: a few long lines of many sentences each are
    refused the same way (checked while splitting, before rows are written);
  * a single huge line is refused (MAX_SEGMENT_CHARS);
  * the cap is generic: markdown, JSON and CSV-free formats go through the same guard,
    and the in-process CLI ingest raises the same refusal;
  * the web upload answers with the refusal text and stores nothing;
  * an at-cap file still ingests (the cap is inclusive), and the setting is the one
    config.MAX_DOCUMENT_SEGMENTS;
  * while a large under-cap file ingests on a live server, /health answers within a short
    timeout and a concurrent writer is never stalled by one giant write transaction.
Final line: "UPLOAD SEGMENT CAP OK".
"""

import os
import sys
import threading
import time

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_segcap_")

import httpx
from fastapi.testclient import TestClient
from sqlmodel import Session, func, select

from app import config
from app.db import engine, init_db
from app.models import Document, Project, Segment
from app.services import apikeys, security, settings

init_db()
settings.set_many({"security_auth_required": "1"})
settings.invalidate_cache()
_user = security.create_user("cap@x.test", "cap", "admin", "cap-password-123")
_row, RAW = apikeys.create_key(_user.id, "cap-test", scopes="api docs")
H = {"X-Loc-Token": RAW}

with Session(engine) as _s:
    _p = Project(name="cap", source_locale="en", default_provider="mock")
    _s.add(_p)
    _s.commit()
    _s.refresh(_p)
    PID = _p.id

from app.main import app

client = TestClient(app)
CAP = 1000
config.MAX_DOCUMENT_SEGMENTS = CAP


def counts():
    with Session(engine) as s:
        return (
            s.exec(select(func.count()).select_from(Document)).one(),
            s.exec(select(func.count()).select_from(Segment)).one(),
        )


def post(name, body):
    t = time.time()
    r = client.post(
        "/api/documents",
        headers=H,
        files={"file": (name, body)},
        data={"project_id": str(PID)},
    )
    return r, time.time() - t


# 1. too many lines: fast 413, nothing stored
before = counts()
r, dt = post("lines.txt", ("This is a short sentence.\n" * (CAP + 500)).encode())
assert r.status_code == 413, (r.status_code, r.text[:200])
assert "segments" in r.json()["error"] and "MAX_DOCUMENT_SEGMENTS" in r.json()["error"], r.text
assert dt < 5, dt
assert counts() == before, ("rows stored for a refused upload", before, counts())
print(f"ok over-cap lines refused 413 in {dt:.2f}s, nothing stored")

# 2. few lines, many sentences: refused while splitting, before any row
before = counts()
line = "This is a short sentence. " * 400
r, dt = post("sentences.txt", (line + "\n") * 5)  # 5 units, 2000 sentences
assert r.status_code == 413, (r.status_code, r.text[:200])
assert counts() == before, ("rows stored for a refused upload", before, counts())
print("ok over-cap sentences refused 413, nothing stored")

# 3. one huge line
before = counts()
r, dt = post("line.txt", b"word " * (config.MAX_SEGMENT_CHARS // 5 + 100))
assert r.status_code == 413, (r.status_code, r.text[:200])
assert "MAX_SEGMENT_CHARS" in r.json()["error"], r.text
assert counts() == before
print("ok one huge line refused 413")

# 4. other formats share the guard
before = counts()
md = "".join(f"Paragraph number {i}.\n\n" for i in range(CAP + 50)).encode()
r, _ = post("many.md", md)
assert r.status_code == 413, (r.status_code, r.text[:200])
js = ("{" + ",".join(f'"k{i}": "Value {i}."' for i in range(CAP + 50)) + "}").encode()
r, _ = post("many.json", js)
assert r.status_code == 413, (r.status_code, r.text[:200])
assert counts() == before
print("ok markdown and JSON refused through the same guard")

# 5. at the cap still works (inclusive) and rows land
r, dt = post("exact.txt", ("Line number one.\n" * CAP).encode())
assert r.status_code == 200, (r.status_code, r.text[:200])
assert r.json()["segments"] == CAP, r.json()
print(f"ok an at-cap file ingests ({CAP} segments, {dt:.2f}s)")

# 6. in-process (CLI) ingest raises the same refusal and stores nothing
from app.filters.base import InputLimitError
from app.services.ingest import ingest_document

before = counts()
with Session(engine) as s:
    proj = s.get(Project, PID)
    try:
        ingest_document(s, proj, "cli.txt", ("Another line.\n" * (CAP + 1)).encode(), "segmented")
        raise AssertionError("over-cap ingest was accepted")
    except InputLimitError as exc:
        assert "MAX_DOCUMENT_SEGMENTS" in str(exc)
assert counts() == before
print("ok in-process ingest refused, nothing stored")

# 7. web upload: a redirect carrying the refusal text, nothing stored
settings.set_many({"security_auth_required": "0"})
settings.invalidate_cache()
before = counts()
r = client.post(
    f"/projects/{PID}/documents",
    files={"file": ("web.txt", ("Web line.\n" * (CAP + 5)).encode())},
    data={"mode": "segmented"},
    follow_redirects=False,
)
assert r.status_code == 303, r.status_code
loc = r.headers["location"]
assert "err=" in loc and "segments" in loc.lower(), loc
assert counts() == before
print("ok web upload refused with the reason")
settings.set_many({"security_auth_required": "1"})
settings.invalidate_cache()

# 7b. #571 review: every format stops EARLY (parse memory and time are bounded), not only
# after a full parse. Each hostile file below has far more units than the cap; the filter
# must raise before it has built much more than CAP units, and the upload is a 413.
from unittest.mock import patch

N = CAP * 5
hostile = {
    "x.md": "".join(f"Paragraph {i}.\n\n" for i in range(N)).encode(),
    "x.srt": "".join(
        f"{i}\n00:00:01,000 --> 00:00:02,000\nCue {i}.\n\n" for i in range(N)
    ).encode(),
    "x.po": "".join(f'msgid "Text {i}"\nmsgstr ""\n\n' for i in range(N)).encode(),
    "x.xlf": (
        '<?xml version="1.0"?><xliff version="1.2"><file source-language="en" '
        'target-language="de" original="x" datatype="plaintext"><body>'
        + "".join(f'<trans-unit id="{i}"><source>Text {i}</source></trans-unit>' for i in range(N))
        + "</body></file></xliff>"
    ).encode(),
    "x.html": (
        "<html><body>" + "".join(f"<p>Para {i}</p>" for i in range(N)) + "</body></html>"
    ).encode(),
    "x.json": ("[" + ",".join(f'"Value {i}"' for i in range(N)) + "]").encode(),
}
before = counts()
for name, body in hostile.items():
    r, dt = post(name, body)
    assert r.status_code == 413, (name, r.status_code, r.text[:200])
    assert "MAX_DOCUMENT_SEGMENTS" in r.json()["error"], (name, r.text[:200])
assert counts() == before
print("ok md, srt, po, xliff, html and json over the cap are refused 413")

import app.filters.markdown_filter as _mdf
import app.filters.po_filter as _pof
import app.filters.srt_filter as _srtf
import app.filters.xliff_filter as _xlf
from app.filters import get_filter

for mod, name in ((_mdf, "x.md"), (_srtf, "x.srt"), (_pof, "x.po"), (_xlf, "x.xlf")):
    made = [0]
    real = mod.TextUnit

    def counting(*a, _real=real, _made=made, **k):
        _made[0] += 1
        return _real(*a, **k)

    with patch.object(mod, "TextUnit", counting):
        try:
            get_filter(name.rsplit(".", 1)[1]).parse(hostile[name])
            raise AssertionError(f"{name}: over-cap file parsed")
        except InputLimitError:
            pass
    assert made[0] <= CAP + 2, (name, "built units past the cap", made[0])
print("ok md, srt, po and xliff stop building units at the cap")

# a file of nothing but newlines is refused by the line budget before any per-line work
t = time.time()
r, _ = post("blank.md", b"\n" * (20 * CAP + 10))
assert r.status_code == 413 and "lines" in r.json()["error"], (r.status_code, r.text[:200])
r, _ = post("blank.txt", b"\n" * (20 * CAP + 10))
assert r.status_code == 413, r.status_code
assert time.time() - t < 5
# deeply nested JSON is a clean refusal, not a 500
r, _ = post("deep.json", b"[" * 200000 + b"]" * 200000)
assert r.status_code == 413, (r.status_code, r.text[:200])
# many tiny JSON strings are refused by the quote budget before json.loads builds them
r, _ = post("tiny.json", b"[" + b",".join([b'"a"'] * (21 * CAP)) + b"]")
assert r.status_code == 413, (r.status_code, r.text[:200])
assert counts() == before
print("ok line budget, deep JSON and tiny-string JSON refused 413")

# a tag-heavy HTML file is refused before tokenizing; export re-reads a legacy over-cap
# document with enforce_caps=False and still works
r, _ = post("tags.html", b"<html><body>" + b"<b>" * (50 * CAP + 10) + b"</body></html>")
assert r.status_code == 413, (r.status_code, r.text[:200])
legacy = hostile["x.html"]
try:
    get_filter("html").parse(legacy)
    raise AssertionError("over-cap html parsed")
except InputLimitError:
    pass
assert len(get_filter("html").parse(legacy, enforce_caps=False).units) == N
assert counts() == before
print("ok tag-heavy html refused; export re-read of an over-cap legacy document still parses")

# 7d. #571 gate: the remaining upload formats and hostile shapes. Every budget below is a
# multiple of MAX_DOCUMENT_SEGMENTS (CAP here), so each file is far past it.
import io
import random
import zipfile
import zlib

from docx import Document as _DocxDocument
from openpyxl import Workbook


def _docx(body_xml: bytes) -> bytes:
    buf = io.BytesIO()
    _DocxDocument().save(buf)
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zin.infolist():
            b = zin.read(item.filename)
            if item.filename == "word/document.xml":
                at = b.index(b"<w:body>") + len(b"<w:body>")
                b = b[:at] + body_xml + b[at:]
            zo.writestr(item, b)
    return out.getvalue()


def _pdf(ops_per_page: int, pages: int) -> bytes:
    """`pages` pages that all share ONE compressed content stream of text objects, all at
    one spot so each page reads as one short paragraph (few units, many objects)."""
    ops = b"BT /F1 2 Tf 10 10 Td (Hi) Tj ET\n" * ops_per_page
    st = zlib.compress(ops, 9)
    kids = b" ".join(b"%d 0 R" % (5 + i) for i in range(pages))
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % pages,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(st) + st + b"\nendstream",
    ]
    page = (
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 3 0 R >> >> /Contents 4 0 R >>"
    )
    objs += [page] * pages
    out = b"%PDF-1.4\n"
    offs = []
    for n, o in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % n + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offs)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return out


XLF_HEAD = (
    b'<?xml version="1.0"?><xliff version="1.2"><file source-language="en" '
    b'original="x" datatype="plaintext"><body><trans-unit id="1"><source>a'
)
XLF_TAIL = b"</source></trans-unit></body></file></xliff>"
random.seed(571)
gate_hostile = {
    # JSON values without quotes (numbers, {} and []) used to build every object first
    "values.json": b"[" + b"{}," * (20 * CAP + 10) + b"{}]",
    # one <source> with a markup flood, and nested inline tags that multiply deep copies
    "flood.xlf": XLF_HEAD + b"<x/>" * (40 * CAP + 10) + XLF_TAIL,
    "nested.xlf": XLF_HEAD + b"<g>" * 50 + b"<x/>" * CAP + b"</g>" * 50 + XLF_TAIL,
    # one tag of unquoted attribute soup: a scan step per run, budgeted
    "soup.html": b"<p><img " + b"x " * (60 * CAP) + b">Hi</p>",
    # escapes are one replacement each: budgeted before the per-entry work
    "escapes.po": b'msgid "' + b"\\n" * (40 * CAP + 10) + b'"\nmsgstr ""\n',
    # ONE paragraph of many runs (the paragraph cap never saw it), and a markup flood
    "runs.docx": _docx(
        b"<w:p>"
        + b"".join(
            b"<w:r><w:t>%06x</w:t></w:r>" % random.getrandbits(24) for _ in range(2 * CAP + 10)
        )
        + b"</w:p>"
    ),
    "marks.docx": _docx(
        b"<w:p>"
        + b"".join(
            b'<w:bookmarkStart w:id="%d" w:name="b%x"/>' % (i, random.getrandbits(24))
            for i in range(40 * CAP + 10)
        )
        + b"</w:p>"
    ),
    # 50 pages sharing one stream of 2 x CAP text objects: 100 x CAP objects in all
    "shared.pdf": _pdf(2 * CAP, 50),
}
before = counts()
for name, body in gate_hostile.items():
    t = time.time()
    r, _ = post(name, body)
    assert r.status_code == 413, (name, r.status_code, r.text[:200])
    assert "MAX_DOCUMENT_SEGMENTS" in r.json()["error"], (name, r.text[:200])
    assert time.time() - t < 10, (name, time.time() - t)
assert counts() == before
print("ok json values, xliff markup and nesting, html tag soup, po escapes, docx runs and")
print("   markup, and pdf objects are refused 413 before the heavy work, nothing stored")

# XLSX stops building units at the cap (its cell guard bounds the load itself)
import app.filters.xlsx_filter as _xlsxf

wb = Workbook(write_only=True)
ws = wb.create_sheet()
for row in range(CAP // 2 + 10):
    ws.append([f"Cell text {row}a", f"Cell text {row}b"])
xbuf = io.BytesIO()
wb.save(xbuf)
made = [0]
real_unit = _xlsxf.TextUnit


def _counting_xlsx(*a, **k):
    made[0] += 1
    return real_unit(*a, **k)


with patch.object(_xlsxf, "TextUnit", _counting_xlsx):
    try:
        get_filter("xlsx").parse(xbuf.getvalue())
        raise AssertionError("over-cap xlsx parsed")
    except InputLimitError:
        pass
assert made[0] <= CAP + 2, ("xlsx built units past the cap", made[0])
r, _ = post("cells.xlsx", xbuf.getvalue())
assert r.status_code == 413, (r.status_code, r.text[:200])
assert counts() == before
print("ok xlsx stops building units at the cap")

# SVG (image localization): a tag flood and too many text nodes are refused in the filter
from app.filters import image_filter as _imgf

svg_head = b'<svg xmlns="http://www.w3.org/2000/svg">'
for body in (
    svg_head + b"<a/>" * (20 * CAP + 10) + b"</svg>",
    svg_head + b"<text>Hi</text>" * (CAP + 5) + b"</svg>",
):
    try:
        _imgf.ImageFilter()._parse_svg(body)
        raise AssertionError("over-cap svg parsed")
    except InputLimitError as exc:
        assert "MAX_DOCUMENT_SEGMENTS" in str(exc), str(exc)
at_cap = _imgf.ImageFilter()._parse_svg(svg_head + b"<text>Hi</text>" * CAP + b"</svg>")
assert len(at_cap.units) == CAP
print("ok svg tag flood and over-cap text nodes refused; at-cap svg parses")

# PO continuation lines are joined once, not re-copied per line (a 4 MB msgid took 49 s).
# Line budget raised for this check only, so the time measured is the join itself.
from app.filters.po_filter import _parse_entries

config.MAX_DOCUMENT_SEGMENTS = 50000
po_body = b'msgid ""\n' + (b'"' + b"a" * 60 + b'"\n') * 60000 + b'msgstr ""\n'
t = time.time()
entries = _parse_entries(po_body)
dt = time.time() - t
config.MAX_DOCUMENT_SEGMENTS = CAP
assert len(entries[0]["msgid"]) == 60 * 60000, len(entries[0]["msgid"])
assert dt < 5, f"PO continuation join took {dt:.1f}s (quadratic)"
print(f"ok a 3.7 MB PO msgid of continuation lines parses in {dt:.2f}s")

# 7c. segments x target languages: the row total is capped too, before anything is stored
from app.models import TargetLocale

with Session(engine) as s:
    wide = Project(name="wide", source_locale="en", default_provider="mock")
    s.add(wide)
    s.commit()
    s.refresh(wide)
    for loc in ("de", "fr", "es"):
        s.add(TargetLocale(project_id=wide.id, locale=loc))
    s.commit()
    WIDE = wide.id
old_rows = config.MAX_DOCUMENT_ROWS
config.MAX_DOCUMENT_ROWS = 2 * CAP
before = counts()
r = client.post(
    "/api/documents",
    headers=H,
    files={"file": ("rows.txt", ("Row line.\n" * CAP).encode())},
    data={"project_id": str(WIDE)},
)
assert r.status_code == 413, (r.status_code, r.text[:200])
assert "MAX_DOCUMENT_ROWS" in r.json()["error"], r.text
assert counts() == before, "rows stored for a refused (segments x locales) upload"
r = client.post(
    "/api/documents",
    headers=H,
    files={"file": ("rows.txt", ("Row line.\n" * (CAP // 2)).encode())},
    data={"project_id": str(WIDE)},
)
assert r.status_code == 200, (r.status_code, r.text[:200])  # 500 x 3 = 1500 <= 2000
config.MAX_DOCUMENT_ROWS = old_rows
print("ok segments x languages over MAX_DOCUMENT_ROWS refused 413, under it ingests")

# 8. a live server stays responsive while a large under-cap file ingests
import socket

import uvicorn

BIG = 40000
config.MAX_DOCUMENT_SEGMENTS = 50000
with socket.socket() as sk:
    sk.bind(("127.0.0.1", 0))
    PORT = sk.getsockname()[1]
srv = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=PORT, log_level="error"))
threading.Thread(target=srv.run, daemon=True).start()
base = f"http://127.0.0.1:{PORT}"
for _ in range(100):
    try:
        httpx.get(base + "/health", timeout=2)
        break
    except Exception:
        time.sleep(0.2)

result = {}


def upload():
    try:
        r = httpx.post(
            base + "/api/documents",
            headers=H,
            files={"file": ("big.txt", ("Big line number.\n" * BIG).encode())},
            data={"project_id": str(PID)},
            timeout=600,
        )
        result["status"] = r.status_code
    except Exception as exc:  # pragma: no cover
        result["status"] = repr(exc)


stall = {"max": 0.0, "n": 0}


def writer():
    """A second writer (another request's tiny insert) must not wait out the ingest."""
    while not result:
        t = time.time()
        with Session(engine) as s:
            s.add(Project(name=f"w{stall['n']}", source_locale="en", default_provider="mock"))
            s.commit()
        stall["max"] = max(stall["max"], time.time() - t)
        stall["n"] += 1
        time.sleep(0.05)


th = threading.Thread(target=upload)
wt = threading.Thread(target=writer)
th.start()
time.sleep(0.3)
wt.start()
worst = 0.0
polls = 0
while th.is_alive():
    t = time.time()
    hr = httpx.get(base + "/health", timeout=10)
    assert hr.status_code == 200
    worst = max(worst, time.time() - t)
    polls += 1
    time.sleep(0.2)
th.join()
wt.join()
assert result.get("status") == 200, result
assert polls >= 3, f"ingest finished too fast to prove anything: {polls} polls"
assert worst < 4.0, f"health stalled {worst:.2f}s during ingest"
assert stall["max"] < 1.5, f"a concurrent writer stalled {stall['max']:.2f}s behind the ingest"
print(
    f"ok live ingest of {BIG} segments: worst /health {worst:.2f}s, "
    f"worst writer stall {stall['max']:.2f}s"
)

print("UPLOAD SEGMENT CAP OK", flush=True)
sys.stdout.flush()
os._exit(0)
