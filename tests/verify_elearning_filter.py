"""E-learning round trips (Articulate Storyline 360 / Rise 360 translation exports).

Loc OS does not read .story projects or published SCORM packages; it handles the files the
authoring tools export FOR translation:
- Storyline XLIFF 1.2 / 2.0 (datatype x-DocumentState units whose formatting travels as
  <bpt>/<ph>/<ept> codes carrying escaped <Style .../> markup, plus plaintext ".Name" units),
- Rise 360 XLIFF 1.2 (nested <g ctype="x-html-*"> groups with an xhtml: attribute),
- Storyline's "Export to Word" document: tables with an ID column, a source column and a
  "Translation" column; only the Translation column may change (app/filters/docx_bilingual).

Hand-written fixtures shaped after the vendor docs (no copied customer or GPL data).
Hermetic + offline. Final line: "ELEARNING FILTER OK".
"""

import os
import re
import sys
from io import BytesIO

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(_ROOT)
sys.path.insert(0, _ROOT)
from tests._bootstrap import init_env

init_env("loctms_elearning_")

import zipfile

import docx
from lxml import etree

from app.filters import get_filter
from app.filters.docx_bilingual import NOTICE_CODE

XLIFF = get_filter("xliff")
DOCX = get_filter("docx")
NS12 = "urn:oasis:names:tc:xliff:document:1.2"


def _mock(units, loc="fr-FR"):
    """Mock translation: upper-case every word of 3+ letters; tokens and {{vars}} untouched."""
    return {
        str(u.p_index): re.sub(
            r"\{\{[^}]*\}\}|[A-Za-z]{3,}",
            lambda m: m.group(0) if m.group(0).startswith("{") else m.group(0).upper(),
            u.text,
        )
        for u in units
    }


# ---------------------------------------------------------------- Storyline XLIFF 1.2

STORYLINE_12 = (
    '\ufeff<?xml version="1.0" encoding="utf-8"?>'
    '<xliff xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" '
    'xmlns:xsd="http://www.w3.org/2001/XMLSchema" version="1.2" '
    'xmlns="urn:oasis:names:tc:xliff:document:1.2">'
    '<file original="C:\\Courses\\Safety Basics.story" source-language="en-US" datatype="xml">'
    "<body>"
    '<trans-unit id="5aBcDeFgHiJ.Name" datatype="plaintext"><source>Welcome</source></trans-unit>'
    '<trans-unit id="6kLmNoPqRsT" xml:space="preserve" datatype="x-DocumentState"><source>'
    '<bpt ctype="x-block" id="block_0" />'
    '<ph id="generic_1">&lt;Style Justification="Center" /&gt;</ph>'
    '<bpt ctype="x-style" id="span_2">&lt;Style FontFamily="Open Sans" FontSize="12" /&gt;</bpt>'
    '<g ctype="x-text" id="text_2">   Hello %LearnerName%,&#xD;&#xD;press Next to start.</g>'
    '<ept id="span_2" /><ept id="block_0" /></source></trans-unit>'
    '<trans-unit id="6uVwXyZaBcD" xml:space="preserve" datatype="x-DocumentState"><source>'
    '<bpt ctype="x-block" id="block_0" /><ph id="generic_1">&lt;Style /&gt;</ph>'
    '<g ctype="x-text" id="text_2">Read the </g>'
    '<bpt ctype="x-style" id="span_3">&lt;Style FontIsBold="True" /&gt;</bpt>'
    '<g ctype="x-text" id="text_3">safety rules</g><ept id="span_3" />'
    '<g ctype="x-text" id="text_4"> carefully.</g><ept id="block_0" /></source></trans-unit>'
    '<trans-unit id="6eMpTyBlOcK" xml:space="preserve" datatype="x-DocumentState"><source>'
    '<bpt ctype="x-block" id="block_0" /><ph id="generic_1">&lt;Style /&gt;</ph>'
    '<ept id="block_0" />'
    "</source></trans-unit>"
    '<trans-unit id="6nOtRaNsLaT" translate="no" datatype="plaintext"><source>ACME LMS</source>'
    "</trans-unit>"
    "</body></file></xliff>"
)


def test_storyline_xliff_12():
    data = STORYLINE_12.encode("utf-8")
    units = XLIFF.parse(data).units
    texts = [u.text for u in units]
    # .Name unit, two text units; the formatting-only block and translate="no" are skipped
    assert len(units) == 3, texts
    assert texts[0] == "Welcome", texts
    assert "   Hello %LearnerName%,\r\rpress Next to start." in texts[1], ascii(texts[1])
    assert "Style" not in re.sub(r"⟦/?[a-z]?\d+⟧", "", texts[1]), ascii(texts[1])
    assert [u.p_index for u in units] == [0, 1, 2], [u.p_index for u in units]

    # identity: nothing translated -> no <target> anywhere, same units on re-parse
    same = XLIFF.build(data, {})
    assert b"<target" not in same, same[:200]
    assert [u.text for u in XLIFF.parse(same).units] == texts

    out = XLIFF.build_localized(data, _mock(units), "fr-FR")
    root = etree.fromstring(out)
    assert root.nsmap.get(None) == NS12
    assert b'xmlns=""' not in out, "a stray empty default namespace breaks Storyline import"
    f = root.find(f"{{{NS12}}}file")
    assert f.get("target-language") == "fr-FR" and f.get("source-language") == "en-US"
    tus = {tu.get("id"): tu for tu in root.iter(f"{{{NS12}}}trans-unit")}
    tgt = tus["5aBcDeFgHiJ.Name"].find(f"{{{NS12}}}target")
    assert tgt.text == "WELCOME", tgt.text
    tu = tus["6kLmNoPqRsT"]
    src, tgt = tu[0], tu[1]
    assert etree.QName(tgt).localname == "target", "target must follow its source"
    # every code element is reproduced with its id, ctype and escaped style payload
    src_codes = [(etree.QName(c).localname, c.get("id"), c.get("ctype")) for c in src]
    tgt_codes = [(etree.QName(c).localname, c.get("id"), c.get("ctype")) for c in tgt]
    assert src_codes == tgt_codes, (src_codes, tgt_codes)
    ph = tgt.find(f"{{{NS12}}}ph")
    assert ph.text == '<Style Justification="Center" />', ph.text
    g = tgt.find(f"{{{NS12}}}g")
    assert g.text == "   HELLO %LEARNERNAME%,\r\rPRESS NEXT to START.", ascii(g.text)
    bpt = tgt.find(f"{{{NS12}}}bpt[@id='span_2']")
    assert bpt.text == '<Style FontFamily="Open Sans" FontSize="12" />', bpt.text
    assert b"&#13;&#13;" in out, "carriage returns (Storyline line breaks) must survive"
    # the formatting-only block and the translate="no" unit get no target
    assert tus["6eMpTyBlOcK"].find(f"{{{NS12}}}target") is None
    assert tus["6nOtRaNsLaT"].find(f"{{{NS12}}}target") is None
    # bold run kept as its own code pair around the translated words
    tu = tus["6uVwXyZaBcD"]
    tgt = tu.find(f"{{{NS12}}}target")
    gs = [c.text for c in tgt if etree.QName(c).localname == "g"]
    assert gs == ["READ THE ", "SAFETY RULES", " CAREFULLY."], gs
    print("  [storyline 1.2] codes, CR, spaces, .Name, skip rules, target-language OK")


# ---------------------------------------------------------------- Storyline XLIFF 2.0

STORYLINE_20 = (
    '<?xml version="1.0" encoding="utf-8"?>'
    '<xliff xmlns="urn:oasis:names:tc:xliff:document:2.0" version="2.0" srcLang="en-US">'
    '<file id="f1" original="Safety Basics.story">'
    '<unit id="5aBcDeFgHiJ.Name"><segment><source>Quiz</source></segment></unit>'
    '<unit id="6kLmNoPqRsT"><segment><source>'
    '<ph id="generic_1" dataRef="d1"/><pc id="text_2">Choose the </pc>'
    '<pc id="text_3">correct</pc><pc id="text_4"> answer.</pc></source></segment></unit>'
    '<unit id="6only"><segment><source><ph id="generic_1"/></source></segment></unit>'
    '<unit id="6no" translate="no"><segment><source>Skip me</source></segment></unit>'
    "</file></xliff>"
)


def test_storyline_xliff_20():
    data = STORYLINE_20.encode("utf-8")
    units = XLIFF.parse(data).units
    assert [u.p_index for u in units] == [0, 1], [(u.p_index, u.text) for u in units]
    out = XLIFF.build_localized(data, _mock(units), "de-DE")
    root = etree.fromstring(out)
    assert root.get("trgLang") == "de-DE" and root.get("srcLang") == "en-US"
    ns = "urn:oasis:names:tc:xliff:document:2.0"
    tgts = list(root.iter(f"{{{ns}}}target"))
    assert len(tgts) == 2, out
    assert [c.get("id") for c in tgts[1]] == ["generic_1", "text_2", "text_3", "text_4"]
    assert "".join(tgts[1].itertext()) == "CHOOSE THE CORRECT ANSWER.", "".join(tgts[1].itertext())
    print("  [storyline 2.0] pc/ph codes, tag-only and translate=no skipped, trgLang OK")


# ---------------------------------------------------------------- Rise 360 XLIFF 1.2

RISE = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    '<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2" '
    'xmlns:xhtml="http://www.w3.org/1999/xhtml">\n'
    '  <file original="course" datatype="html" source-language="en-us">\n'
    "    <body>\n"
    '      <trans-unit id="ck1|title"><source>Customer onboarding</source></trans-unit>\n'
    '      <trans-unit id="ck1|blocks|0|items|0|paragraph"><source>'
    '<g ctype="x-html-DIV" id="3kRw9m2X" xhtml:content-editor-id="ed-1">'
    '<g ctype="x-html-P" id="yHPQ4rWK"><g ctype="x-html-STRONG" id="9F7NK43_">Best practices</g>'
    " for {{company}} teams</g>"
    '<g ctype="x-html-P" id="p2">Visit <g ctype="x-html-A" id="a1" xhtml:href="https://example.com/x">'
    "our site</g>.</g></g></source></trans-unit>\n"
    '      <trans-unit id="ck1|blocks|1|items|0|paragraph"><source>'
    '<g ctype="x-html-DIV" id="d2"><g ctype="x-html-P" id="p3"> </g></g></source></trans-unit>\n'
    "    </body>\n"
    "  </file>\n"
    "</xliff>\n"
)


def test_rise_xliff():
    data = RISE.encode("utf-8")
    units = XLIFF.parse(data).units
    assert [u.p_index for u in units] == [0, 1], [(u.p_index, u.text) for u in units]
    assert "{{company}}" in units[1].text
    out = XLIFF.build_localized(data, _mock(units), "es-ES")
    text = out.decode("utf-8")
    assert 'target-language="es-ES"' in text
    root = etree.fromstring(out)
    tu = list(root.iter(f"{{{NS12}}}trans-unit"))[1]
    tgt = tu.find(f"{{{NS12}}}target")
    div = tgt[0]
    xh = "{http://www.w3.org/1999/xhtml}"
    assert div.get("ctype") == "x-html-DIV" and div.get(xh + "content-editor-id") == "ed-1"
    a = tgt.find(f".//{{{NS12}}}g[@id='a1']")
    assert a.get(xh + "href") == "https://example.com/x" and a.text == "OUR SITE"
    assert "{{company}}" in "".join(tgt.itertext())
    assert "".join(tgt.itertext()) == "BEST PRACTICES FOR {{company}} TEAMSVISIT OUR SITE."
    print("  [rise] nested x-html groups, xhtml attributes, {{var}}, blank block OK")


# ---------------------------------------------------------------- Storyline Word export


def _table(d, header, rows):
    t = d.add_table(rows=1 + len(rows), cols=len(header))
    for c, h in enumerate(header):
        t.cell(0, c).text = h
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            if isinstance(val, list):  # runs: [(text, bold)]
                p = t.cell(r, c).paragraphs[0]
                for text, bold in val:
                    p.add_run(text).bold = bold
            else:
                t.cell(r, c).text = val
    return t


def storyline_word(enhanced: bool, empty_translation: bool = False) -> bytes:
    d = docx.Document()
    bold = [("Click ", False), ("Next", True), (" to continue.", False)]
    if enhanced:
        d.add_heading("Project Information", 1)
        _table(
            d,
            ["Project Title", "Safety Basics"],
            [["Source Language", "English (United States)"], ["Word Count", "11"]],
        )
        d.add_heading("Scene 1: Introduction", 1)
        d.add_heading("Slide 1.1: Welcome", 2)
        t = _table(
            d,
            ["ID", "Object", "Source Text", "Translation"],
            [
                ["5aBcDeFgHiJ", "Slide Title", "Welcome", "Welcome"],
                ["6kLmNoPqRsT", "Text Box", bold, bold],
                ["6pQrStUvWxY", "Button", "Next", "" if empty_translation else "Next"],
            ],
        )
        # a merged layer label row (cell count differs from the header): never a unit
        row = t.add_row()
        merged = row.cells[0].merge(row.cells[3])
        merged.text = "Layer: Feedback"
        r = t.add_row()
        for c, val in enumerate(["6zZzZzZzZzZ", "Text Box", "Well done!", "Well done!"]):
            r.cells[c].text = val
    else:
        _table(
            d,
            ["ID", "Original Text", "Translation"],
            [
                ["5aBcDeFgHiJ.Name", "Welcome", "Welcome"],
                ["6kLmNoPqRsT", bold, bold],
                ["6zZzZzZzZzZ", "Well done!", "Well done!"],
            ],
        )
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def _cells(data: bytes):
    d = docx.Document(BytesIO(data))
    return [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]


def test_storyline_word_legacy():
    data = storyline_word(enhanced=False)
    res = DOCX.parse(data)
    texts = [u.text for u in res.units]
    assert texts == [
        "Welcome",
        "⟦r1⟧Click ⟦/r1⟧⟦r2⟧Next⟦/r2⟧⟦r3⟧ to continue.⟦/r3⟧",
        "Well done!",
    ], [ascii(t) for t in texts]
    notes = [w for w in res.skeleton.get("file_warnings", []) if w.get("code") == NOTICE_CODE]
    assert notes and notes[0]["rows"] == 3 and notes[0]["tables"] == 1, notes

    same = DOCX.build(data, {})
    assert _cells(same) == _cells(data)
    out = DOCX.build_localized(data, _mock(res.units), "fr-FR")
    before, after = _cells(data)[0], _cells(out)[0]
    assert after[0] == ["ID", "Original Text", "Translation"]
    for b, a in zip(before[1:], after[1:], strict=True):
        assert a[0] == b[0] and a[1] == b[1], (a, b)  # ID and source untouched
    assert [a[2] for a in after[1:]] == ["WELCOME", "CLICK NEXT to CONTINUE.", "WELL DONE!"]
    # the bold run stays bold in the Translation cell
    tbl = docx.Document(BytesIO(out)).tables[0]
    runs = tbl.cell(2, 2).paragraphs[0].runs
    assert [(r.text, bool(r.bold)) for r in runs] == [
        ("CLICK ", False),
        ("NEXT", True),
        (" to CONTINUE.", False),
    ], [(r.text, r.bold) for r in runs]
    print("  [word legacy] ID | Original Text | Translation: only Translation changes OK")


def test_storyline_word_enhanced():
    data = storyline_word(enhanced=True)
    res = DOCX.parse(data)
    texts = [re.sub(r"⟦/?[a-z]?\d+⟧", "", u.text) for u in res.units]
    assert texts == ["Welcome", "Click Next to continue.", "Next", "Well done!"], texts
    out = DOCX.build_localized(data, _mock(res.units), "de-DE")
    b_tabs, a_tabs = _cells(data), _cells(out)
    assert a_tabs[0] == b_tabs[0], "project information table must not change"
    for b, a in zip(b_tabs[1], a_tabs[1], strict=True):
        if len(set(b)) == 1:  # merged label row
            assert a == b, (a, b)
            continue
        assert a[:3] == b[:3], (a, b)
    assert [r[3] for r in a_tabs[1][1:]] == [
        "WELCOME",
        "CLICK NEXT to CONTINUE.",
        "NEXT",
        "Layer: Feedback",
        "WELL DONE!",
    ], a_tabs[1]
    d = docx.Document(BytesIO(out))
    heads = [p.text for p in d.paragraphs if p.text]
    assert heads == ["Project Information", "Scene 1: Introduction", "Slide 1.1: Welcome"], heads

    # an empty Translation cell stays empty and is counted in the notice
    data = storyline_word(enhanced=True, empty_translation=True)
    res = DOCX.parse(data)
    assert len(res.units) == 3
    note = next(w for w in res.skeleton["file_warnings"] if w.get("code") == NOTICE_CODE)
    assert note["empty_rows"] == 1 and "empty" in note["message"], note
    print("  [word enhanced] project table, headings, label rows, empty cells OK")


def test_ordinary_docx_unchanged():
    """No bilingual header -> every paragraph is still a unit (old behaviour)."""
    d = docx.Document()
    d.add_paragraph("Intro paragraph.")
    _table(d, ["Term", "Meaning"], [["Hazard", "Something that can hurt you"]])
    # a Translation header whose column is all empty is not a translation table either
    _table(d, ["Source", "Translation"], [["Exit", ""]])
    out = BytesIO()
    d.save(out)
    res = DOCX.parse(out.getvalue())
    texts = [u.text for u in res.units]
    assert texts == [
        "Intro paragraph.",
        "Term",
        "Meaning",
        "Hazard",
        "Something that can hurt you",
        "Source",
        "Translation",
        "Exit",
    ], texts
    assert not any(w.get("code") == NOTICE_CODE for w in res.skeleton.get("file_warnings", []))
    # gate train-11: a style guide with prose that merely contains a FILLED Source |
    # Translation term table keeps the whole-document behaviour (its prose is not lost)
    d = docx.Document()
    d.add_heading("House style", 1)
    d.add_paragraph("Write short sentences and use the approved terms below in every product.")
    _table(d, ["Source", "Translation"], [["Exit", "Sortie"], ["Save", "Enregistrer"]])
    out = BytesIO()
    d.save(out)
    res = DOCX.parse(out.getvalue())
    texts = [u.text for u in res.units]
    assert "Write short sentences and use the approved terms below in every product." in texts
    assert "Exit" in texts and "House style" in texts, texts
    assert not any(w.get("code") == NOTICE_CODE for w in res.skeleton.get("file_warnings", []))
    print("  [docx] ordinary tables keep the whole-document behaviour OK")


def test_hostile_inputs():
    from app.filters.base import InputLimitError
    from app.services.vendor_xliff import XliffParseError

    laughs = (
        b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;">]>'
        b'<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2"><file><body>'
        b"<trans-unit id='1' datatype='x-DocumentState'><source>&b;</source></trans-unit>"
        b"</body></file></xliff>"
    )
    for bad in (laughs, b"<xliff><file><body><trans-unit", b"\x00\x01PK"):
        try:
            XLIFF.parse(bad)
        except (XliffParseError, ValueError):
            pass
        else:
            raise AssertionError(bad[:40])
    # a hostile code-only flood is bounded by the markup budget before lxml runs
    from app import config

    flood = (
        b'<xliff version="1.2"><file><body><trans-unit id="1"><source>'
        + b"<ph/>" * (40 * config.MAX_DOCUMENT_SEGMENTS + 1)
        + b"</source></trans-unit></body></file></xliff>"
    )
    try:
        XLIFF.parse(flood)
    except InputLimitError:
        pass
    else:
        raise AssertionError("code flood not capped")
    # a DOCX that is not a zip is refused politely
    try:
        DOCX.parse(b"PK\x03\x04not really a zip")
    except (ValueError, zipfile.BadZipFile):
        pass
    else:
        raise AssertionError("broken docx accepted")
    print("  [hostile] DTD/entities, broken XML, code flood, broken zip refused OK")


def test_api_and_cli_round_trip():
    """Upload the Storyline Word export through the API and the CLI, mock-translate, export:
    only the Translation column carries the target text; the notice reaches both surfaces."""
    import tempfile

    from fastapi.testclient import TestClient
    from sqlmodel import Session

    from app.db import engine, init_db
    from app.main import app
    from app.models import Project, Role, TargetLocale
    from app.services import apikeys, security
    from tests._cli_runner import run_cli

    loc = "de-DE"
    init_db()
    with Session(engine) as s:
        p = Project(name="elearning", org_id=1, source_locale="en-US", default_provider="mock")
        s.add(p)
        s.commit()
        s.refresh(p)
        s.add(TargetLocale(project_id=p.id, locale=loc))
        s.commit()
        pid = p.id
    user = security.create_user("elearn@x.test", "elearn", Role.ADMIN, "elearn-pw-123456")
    _row, raw_key = apikeys.create_key(user.id, "elearn", scopes="api docs admin")
    client = TestClient(app, raise_server_exceptions=False, follow_redirects=False)
    hdr = {"X-Loc-Token": raw_key}  # pregate: allow-secret throwaway test key
    data = storyline_word(enhanced=False)

    def check(out: bytes, surface: str):
        rows = _cells(out)[0]
        assert rows[0] == ["ID", "Original Text", "Translation"], (surface, rows)
        assert [r[0] for r in rows[1:]] == ["5aBcDeFgHiJ.Name", "6kLmNoPqRsT", "6zZzZzZzZzZ"]
        assert [r[1] for r in rows[1:]] == ["Welcome", "Click Next to continue.", "Well done!"]
        assert all(r[2].startswith(f"[{loc}]") for r in rows[1:]), (surface, rows)

    r = client.post(
        "/api/documents",
        headers=hdr,
        data={"project_id": str(pid)},
        files={"file": ("Safety Basics.docx", data, "application/octet-stream")},
    )
    assert r.status_code in (200, 201), (r.status_code, r.text)
    body = r.json()
    assert any(w.get("code") == NOTICE_CODE for w in body.get("file_warnings", [])), body
    doc = body["document_id"]
    code, out, err = run_cli(["doc", "translate", str(doc), "--locale", loc])
    assert code == 0, (code, out, err)
    r = client.get(f"/api/documents/{doc}/download", headers=hdr, params={"locale": loc})
    assert r.status_code == 200, (r.status_code, r.text)
    check(r.content, "api")

    tmp = tempfile.mkdtemp(prefix="loctms_elearning_cli_")
    src = os.path.join(tmp, "course.docx")
    with open(src, "wb") as fh:
        fh.write(data)
    code, out, err = run_cli(["doc", "upload", str(pid), src])
    assert code == 0, (code, out, err)
    assert "Translation column" in out, out
    m = re.search(r"Uploaded #(\d+): course\.docx", out)
    assert m, out
    cli_doc = int(m.group(1))
    code, out, err = run_cli(["doc", "translate", str(cli_doc), "--locale", loc])
    assert code == 0, (code, out, err)
    dest = os.path.join(tmp, "out.docx")
    code, out, err = run_cli(["doc", "export", str(cli_doc), "--locale", loc, "-o", dest])
    assert code == 0, (code, out, err)
    with open(dest, "rb") as fh:
        check(fh.read(), "cli")

    # Storyline XLIFF through the same pipeline: codes and CR survive, target-language set
    r = client.post(
        "/api/documents",
        headers=hdr,
        data={"project_id": str(pid)},
        files={"file": ("course.xlf", STORYLINE_12.encode("utf-8"), "application/xliff+xml")},
    )
    assert r.status_code in (200, 201), (r.status_code, r.text)
    xdoc = r.json()["document_id"]
    code, out, err = run_cli(["doc", "translate", str(xdoc), "--locale", loc])
    assert code == 0, (code, out, err)
    r = client.get(f"/api/documents/{xdoc}/download", headers=hdr, params={"locale": loc})
    assert r.status_code == 200, (r.status_code, r.text)
    root = etree.fromstring(r.content)
    assert root.find(f"{{{NS12}}}file").get("target-language") == loc
    tgts = list(root.iter(f"{{{NS12}}}target"))
    assert len(tgts) == 3, r.content
    assert b"&#13;&#13;" in r.content, r.content
    assert r.content.count(b'&lt;Style FontIsBold="True" /&gt;') == 2, r.content
    print("  [surfaces] API + CLI upload/translate/export, notice on both OK")


def main():
    test_storyline_xliff_12()
    test_storyline_xliff_20()
    test_rise_xliff()
    test_storyline_word_legacy()
    test_storyline_word_enhanced()
    test_ordinary_docx_unchanged()
    test_hostile_inputs()
    test_api_and_cli_round_trip()
    print("ELEARNING FILTER OK")


if __name__ == "__main__":
    main()
