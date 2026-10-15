"""Bound worksheet expansion before openpyxl allocates cells or iterates empty rows.

Documented in: docs/architecture/translation-pipeline.md
"""

import io
import posixpath
import zipfile

from lxml import etree

from app import config
from app.filters import zip_guard
from app.filters.base import InputLimitError

_SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _parse(zf, path, label, start, end=lambda name: None, characters=lambda text: None):
    """Incremental lxml target: identical encoding detection, without building a tree."""

    class Target:
        def start(self, name, attrs):
            start(name.removeprefix("{"), attrs)

        def end(self, name):
            end(name.removeprefix("{"))

        def data(self, text):
            characters(text)

        def doctype(self, *args):
            raise InputLimitError(f"This {label} contains an unsupported XML document type.")

        def close(self):
            pass

    try:
        with zf.open(path) as member:
            # libxml's IO/push entry point mistakes UTF-32LE's BOM for UTF-16 on
            # this build, unlike fromstring (openpyxl). Supply the XML byte-order
            # signature in that case; libxml still does all decoding/declarations.
            signature = member.read(4)
            member.seek(0)
            encoding = None
            if signature in (b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff"):
                encoding = "UTF-32"
            elif signature == b"<\x00\x00\x00":
                encoding = "UTF-32LE"
            parser = etree.XMLParser(
                target=Target(),
                resolve_entities=False,
                no_network=True,
                huge_tree=False,
                encoding=encoding,
            )
            etree.parse(member, parser)
    except InputLimitError:
        raise
    except Exception as exc:
        # Encoding, parser, ZIP and callback errors must be safe at stored export too.
        raise InputLimitError(
            f"This {label} has invalid spreadsheet XML. Save it again in Excel."
        ) from exc


def _records(zf, path, label, tags, parent=None):
    records = []
    stack = []
    record = None
    record_depth = 0
    value = []

    def start(name, attrs):
        nonlocal record, record_depth, value
        name = name.rsplit("}", 1)[-1]
        at_parent = len(stack) == 1 if parent is None else (len(stack) == 2 and stack[-1] == parent)
        if at_parent and (tags is None or name in tags):
            if len(records) >= config.ZIP_MAX_MEMBERS:
                raise InputLimitError(f"This {label} exceeds the package relationship limit.")
            records.append((name, dict(attrs)))
            record = records[-1][1]
            if parent in ("sheets", "externalReferences", "pivotCaches") and (
                "{" + _REL_NS + "}id" in record
            ):
                record["id"] = record.pop("{" + _REL_NS + "}id")
            record_depth = len(stack) + 1
        if record is not None and len(stack) == record_depth:
            value = []
        stack.append(name)

    def characters(text):
        if record is not None and len(stack) == record_depth + 1:
            value.append(text)

    def end(name):
        nonlocal record
        # Serialisable.from_tree also accepts primitive fields as child elements.
        # Restrict these to the fields used for package selection and resolution.
        if (
            record is not None
            and len(stack) == record_depth + 1
            and stack[-1]
            in ("id", "Id", "Target", "TargetMode", "Type", "type", "PartName", "ContentType")
        ):
            record[stack[-1]] = "".join(value)
        if len(stack) == record_depth:
            record = None
        stack.pop()

    _parse(zf, path, label, start, end, characters)
    return records


def _relationships(zf, path, label):
    """Mirror openpyxl.packaging.relationship.get_dependents path resolution."""
    from openpyxl.packaging.relationship import Relationship, get_rels_path

    rels_path = get_rels_path(path)
    if rels_path not in zf.namelist():
        return []
    parent = posixpath.split(posixpath.dirname(rels_path))[0]
    # RelationshipList is an ElementList: every child is a relationship,
    # irrespective of its tag/namespace, just like NestedSequence sheet entries.
    records = _records(zf, rels_path, label, None)
    for _, attrs in records:
        try:
            rel = Relationship.from_tree(etree.Element("Relationship", attrib=attrs))
        except Exception as exc:
            raise InputLimitError(
                f"This {label} has invalid spreadsheet XML relationships."
            ) from exc
        attrs.clear()
        attrs.update(dict(rel))
        target = rel.Target
        if attrs.get("TargetMode") != "External":
            target = (
                target[1:]
                if target.startswith("/")
                else posixpath.normpath(posixpath.join(parent, target))
            )
        attrs["Target"] = target
    return [attrs for _, attrs in records]


def check(data: bytes, label: str = "XLSX file") -> None:
    # openpyxl is imported on first use (test speed, 2026-10-08): importing it costs about
    # 0.5 s, and every app and CLI process imports this guard through the filter registry.
    from openpyxl.formula.translate import Translator
    from openpyxl.utils.cell import get_column_letter, range_boundaries
    from openpyxl.xml.constants import SHARED_STRINGS, XLSM, XLSX, XLTM, XLTX

    zip_guard.check(data, label)
    cells = strings = area = materialised_cells = 0
    inline_chars = stored_chars = 0
    string_lengths = {}
    references = {}

    def check_text(size):
        if size > config.XLSX_MAX_TOTAL_TEXT_CHARS:
            raise InputLimitError(
                f"This {label} exceeds the total cell text limit "
                f"({config.XLSX_MAX_TOTAL_TEXT_CHARS} characters, including repeated shared "
                "strings and formulas). Split it into smaller files."
            )

    def refuse():
        return InputLimitError(
            f"This {label} exceeds a worksheet safety limit: "
            f"{config.XLSX_MAX_CELLS} cells across worksheet rectangles, "
            f"{config.XLSX_MAX_ROWS} rows per sheet, or "
            f"{config.XLSX_MAX_CELL_CHARS} characters per cell. Split it into smaller files."
        )

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        validated = set()
        sheet_paths = []

        def scan_part(path, mode, extents=(0, 0)):
            row = col = max_row = max_col = text_chars = 0
            max_row, max_col = extents
            text_depth = 0
            text_node = False
            worksheet = mode in ("worksheet", "comments", "threadedComment", "table")
            stack = []
            string_index = 0
            shared_cell = False
            reference_text = []
            cell_coordinate = "A1"
            formula_kind = formula_key = None
            formula_text = []
            shared_formulas = {}
            table_ids = []
            reference_value = []

            def bounds(ref, prior_area=area):
                nonlocal max_row, max_col
                try:
                    if len(ref) > 32:
                        raise ValueError("coordinate too long")
                    min_c, min_r, max_c, max_r = range_boundaries(ref)
                    if not all(v is not None and v > 0 for v in (min_c, min_r, max_c, max_r)):
                        raise ValueError("invalid coordinate")
                    if max_c < min_c or max_r < min_r:
                        raise ValueError("reversed coordinate")
                except ValueError as exc:
                    raise refuse() from exc
                max_row, max_col = max(max_row, max_r), max(max_col, max_c)
                if (
                    max_row > config.XLSX_MAX_ROWS
                    or prior_area + max_row * max_col > config.XLSX_MAX_CELLS
                ):
                    raise refuse()
                return (max_r - min_r + 1) * (max_c - min_c + 1)

            def start(name, attrs):
                nonlocal row, col, cells, strings, worksheet, text_depth, text_chars
                nonlocal materialised_cells
                nonlocal text_node
                nonlocal shared_cell, reference_text
                nonlocal cell_coordinate, formula_kind, formula_key, formula_text
                nonlocal reference_value
                # openpyxl dispatches on child tags, not the document root, and treats
                # EVERY direct child of a spreadsheet row as a cell. Mirror that seam.
                raw_name = name
                parent = stack[-1][0] if stack else None
                relevant = name.startswith(_SHEET_NS) or text_depth > 0
                name = name.rsplit("}", 1)[-1]
                if name == "ref" and mode in ("table", "comments", "threadedComment"):
                    reference_value = []
                if mode == "worksheet" and name == "tablePart":
                    table_ids.append(attrs.get("{" + _REL_NS + "}id", attrs.get("id")))
                if parent == _SHEET_NS + "row":
                    relevant, name = True, "c"
                elif parent == _SHEET_NS + "mergeCells" and "ref" in attrs:
                    relevant, name = True, "mergeCell"
                elif stack and (stack[-1][1], name) in (
                    ("hyperlinks", "hyperlink"),
                    ("comments", "commentList"),
                    ("commentList", "comment"),
                ):
                    # Serialisable.from_tree dispatches these children by local name.
                    relevant = True
                elif mode == "table" and not stack:
                    relevant, name = True, "table"
                elif mode == "comments" and stack and stack[-1][1] == "commentList":
                    relevant, name = True, "comment"
                elif mode == "comments" and name == "commentList":
                    relevant = True
                elif mode == "threadedComment" and name == "threadedComment":
                    relevant, name = True, "comment"
                stack.append((raw_name, name if relevant else ""))
                if not relevant:
                    return
                if name in ("t", "v", "f"):
                    text_node = True
                if name in (
                    "worksheet",
                    "row",
                    "dimension",
                    "mergeCells",
                    "mergeCell",
                    "hyperlink",
                    "comment",
                    "table",
                ):
                    worksheet = True
                if name == "si":
                    strings += 1
                    if strings > config.XLSX_MAX_CELLS:
                        raise refuse()
                if name in ("si", "c"):
                    text_depth += 1
                    text_chars = 0
                if name == "c":
                    shared_cell = attrs.get("t") == "s"
                    reference_text = []
                    formula_kind = formula_key = None
                    formula_text = []
                elif raw_name == _SHEET_NS + "f" and text_depth:
                    formula_kind = attrs.get("t", "normal")
                    formula_key = attrs.get("si")
                if not worksheet:
                    return
                if name in (
                    "dimension",
                    "mergeCell",
                    "hyperlink",
                    "comment",
                    "table",
                ) and attrs.get("ref"):
                    size = bounds(attrs["ref"])
                    if name != "dimension":
                        # Hyperlinks/comments create cells on load. Tables without
                        # column definitions do so on save (write_tables uses ws[ref]).
                        # Charge cumulative work as well as this sheet's union rectangle.
                        materialised_cells += size
                        if materialised_cells + cells > config.XLSX_MAX_CELLS:
                            raise refuse()
                elif name == "row":
                    try:
                        row = int(attrs.get("r", row + 1))
                    except ValueError:
                        value = float(attrs["r"])
                        if not value.is_integer():
                            raise refuse() from None
                        row = int(value)
                    col = 0
                    if row < 1 or row > config.XLSX_MAX_ROWS:
                        raise refuse()
                elif name == "c":
                    cells += 1
                    col += 1
                    if cells + materialised_cells > config.XLSX_MAX_CELLS:
                        raise refuse()
                    ref = attrs.get("r")
                    cell_coordinate = ref or f"{get_column_letter(col)}{max(row, 1)}"
                    if ref:
                        bounds(ref)
                        col = range_boundaries(ref)[0]
                    else:
                        bounds(cell_coordinate)

            def end(name):
                nonlocal text_depth, text_node, text_chars
                nonlocal materialised_cells
                nonlocal string_index, inline_chars, stored_chars, shared_cell
                nonlocal formula_kind, formula_key, formula_text, cell_coordinate, shared_formulas
                raw_name, name = stack.pop()
                if raw_name.rsplit("}", 1)[-1] == "ref" and mode in (
                    "table",
                    "comments",
                    "threadedComment",
                ):
                    # Table/CommentRecord.from_tree allow ref as a primitive child.
                    materialised_cells += bounds("".join(reference_value))
                    if cells + materialised_cells > config.XLSX_MAX_CELLS:
                        raise refuse()
                if name == "si":
                    stored_chars += text_chars
                    check_text(stored_chars)
                    string_lengths[string_index] = max(
                        string_lengths.get(string_index, 0), text_chars
                    )
                    string_index += 1
                elif name == "c":
                    if formula_kind is not None:
                        value = "=" + "".join(formula_text)
                        if formula_kind == "shared":
                            if formula_text:
                                shared_formulas[formula_key] = Translator(value, cell_coordinate)
                            elif formula_key in shared_formulas:
                                value = shared_formulas[formula_key].translate_formula(
                                    cell_coordinate
                                )
                        if len(value) > config.XLSX_MAX_CELL_CHARS:
                            raise refuse()
                        inline_chars += len(value)
                        check_text(inline_chars)
                    elif shared_cell and reference_text:
                        index = int("".join(reference_text))
                        if index < 0:
                            raise InputLimitError(
                                f"This {label} has an invalid shared-string index."
                            )
                        references[index] = references.get(index, 0) + 1
                    else:
                        inline_chars += text_chars
                        check_text(inline_chars)
                    shared_cell = False
                if name in ("t", "v", "f"):
                    text_node = False
                if name in ("si", "c"):
                    text_depth -= 1

            def characters(text):
                nonlocal text_chars
                if (
                    stack
                    and stack[-1][0].rsplit("}", 1)[-1] == "ref"
                    and mode in ("table", "comments", "threadedComment")
                ):
                    reference_value.append(text)
                    if sum(map(len, reference_value)) > 32:
                        raise refuse()
                # Callbacks run before the next part is opened. XML indentation is not text.
                if text_depth and text_node:
                    text_chars += len(text)
                    if shared_cell and stack[-1][1] == "v":
                        reference_text.append(text)
                    if stack[-1][1] == "f":
                        formula_text.append(text)
                    if text_chars > config.XLSX_MAX_CELL_CHARS:
                        raise refuse()

            _parse(zf, path, label, start, end, characters)
            validated.add(path)
            return (max_row, max_col), table_ids

        manifest = _records(zf, "[Content_Types].xml", label, {"Override", "Default"})
        overrides = [a for tag, a in manifest if tag == "Override"]
        workbook = None
        for content_type in (XLTM, XLTX, XLSM, XLSX):
            workbook = next(
                (
                    a.get("PartName", "")[1:]
                    for a in overrides
                    if a.get("ContentType") == content_type
                ),
                None,
            )
            if workbook is not None:
                break
        if workbook is None:
            if any(
                a.get("ContentType") in (XLTM, XLTX, XLSM, XLSX)
                for tag, a in manifest
                if tag == "Default"
            ):
                workbook = "xl/workbook.xml"
            else:
                raise InputLimitError(f"This {label} has invalid spreadsheet XML: no workbook.")
        shared = next(
            (
                a.get("PartName", "")[1:]
                for a in overrides
                if a.get("ContentType") == SHARED_STRINGS
            ),
            None,
        )
        if shared is not None:
            scan_part(shared, "strings")
        sheets = _records(zf, workbook, label, None, parent="sheets")
        rels = _relationships(zf, workbook, label)
        by_id = {r.get("Id"): r for r in rels}
        for _, sheet in sheets:
            rid = sheet.get("{" + _REL_NS + "}id", sheet.get("id"))
            if not rid:
                continue
            if rid not in by_id:
                raise InputLimitError(f"This {label} has invalid spreadsheet XML: missing sheet.")
            rel = by_id[rid]
            path = rel["Target"]
            if path not in zf.namelist():
                continue  # ExcelReader.read_worksheets skips missing worksheet targets.
            sheet_paths.append(path)
            extents, table_ids = scan_part(path, "worksheet")
            sheet_rels = _relationships(zf, path, label)
            sheet_rel_ids = {r.get("Id") for r in sheet_rels}
            if any(rid not in sheet_rel_ids for rid in table_ids):
                raise InputLimitError(f"This {label} has invalid spreadsheet XML: missing table.")
            for related in sheet_rels:
                kind = related.get("Type", "").rsplit("/", 1)[-1]
                if kind in ("comments", "threadedComment", "table") or (
                    related.get("Id") in table_ids
                ):
                    mode = kind if kind in ("comments", "threadedComment") else "table"
                    extents, _ = scan_part(related["Target"], mode, extents)
            area += extents[0] * extents[1]
            if area > config.XLSX_MAX_CELLS:
                raise refuse()

        def scan_cache(path, mode, multiplier=1):
            """Charge cached objects and text without materialising their XML tree."""
            depth = dataset_depth = cell_depth = 0
            cache_id = None
            id_text = []
            in_id = False

            def characters(text):
                nonlocal inline_chars
                if in_id:
                    id_text.append(text)
                if cell_depth:
                    inline_chars += len(text) * multiplier
                    check_text(inline_chars)

            def start(name, attrs):
                nonlocal depth, dataset_depth, cell_depth, cells, cache_id, in_id
                name = name.rsplit("}", 1)[-1]
                depth += 1
                if depth == 1:
                    cache_id = attrs.get("{" + _REL_NS + "}id", attrs.get("id"))
                if mode == "definition" and depth == 2 and name == "id":
                    in_id = True
                if mode == "external" and name == "sheetDataSet" and not dataset_depth:
                    dataset_depth = depth
                if (mode == "external" and dataset_depth and name == "cell") or (
                    mode == "records" and depth == 2
                ):
                    cells += multiplier
                    if cells + materialised_cells > config.XLSX_MAX_CELLS:
                        raise refuse()
                    if not cell_depth:
                        cell_depth = depth
                # Pivot strings normally use v attributes; ExternalCell accepts them too.
                if cell_depth and "v" in attrs:
                    characters(attrs["v"])

            def end(name):
                nonlocal depth, dataset_depth, cell_depth, cache_id, in_id
                if in_id and depth == 2:
                    cache_id = "".join(id_text)
                    in_id = False
                if depth == cell_depth:
                    cell_depth = 0
                if depth == dataset_depth:
                    dataset_depth = 0
                depth -= 1

            _parse(zf, path, label, start, end, characters)
            validated.add(path)
            return cache_id

        def selected_target(record, relationships):
            rid = record.get("id")
            if rid not in relationships:
                raise InputLimitError(f"This {label} has invalid spreadsheet XML relationships.")
            return relationships[rid]["Target"]

        for _, reference in _records(zf, workbook, label, None, parent="externalReferences"):
            scan_cache(selected_target(reference, by_id), "external")
        # ExcelReader.read_worksheets evaluates parser.pivot_caches once for every
        # existing non-chart worksheet, including sheets with no pivot table.
        worksheet_count = sum(
            bool(sheet.get("id"))
            and sheet["id"] in by_id
            and by_id[sheet["id"]]["Target"] in zf.namelist()
            and "chartsheet" not in by_id[sheet["id"]].get("Type", "")
            for _, sheet in sheets
        )
        if worksheet_count:
            for _, cache in _records(zf, workbook, label, None, parent="pivotCaches"):
                path = selected_target(cache, by_id)
                rid = scan_cache(path, "definition")
                deps = _relationships(zf, path, label)
                if deps:
                    if rid is None:
                        # get_rel falls back to RecordList.rel_type when id is absent.
                        selected = next(
                            (r for r in deps if r.get("Type") == (_REL_NS + "/pivotCacheRecords")),
                            None,
                        )
                        if selected is None:
                            continue
                        records_path = selected["Target"]
                    else:
                        # RelationshipList.get selects the first matching id; workbook
                        # relationships above use to_dict and select the last instead.
                        records_path = selected_target(
                            {"id": rid}, {r.get("Id"): r for r in reversed(deps)}
                        )
                    scan_cache(records_path, "records", worksheet_count)

        # These are read by fixed path, or by related XML readers, without creating
        # worksheet cells. Validate incrementally too; unrelated custom XML is ignored.
        pending = [workbook, *sheet_paths]
        pending += [
            p
            for p in ("xl/styles.xml", "docProps/core.xml", "docProps/custom.xml")
            if p in zf.namelist()
        ]
        visited = set()
        xml_kinds = {
            "worksheet",
            "chartsheet",
            "comments",
            "threadedComment",
            "table",
            "drawing",
            "chart",
            "pivotTable",
            "pivotCacheDefinition",
            "pivotCacheRecords",
            "externalLink",
        }
        while pending:
            path = pending.pop()
            if path in visited:
                continue
            visited.add(path)
            if path not in validated:
                _parse(zf, path, label, lambda name, attrs: None)
            for rel in _relationships(zf, path, label):
                if (
                    rel.get("TargetMode") != "External"
                    and rel.get("Type", "").rsplit("/", 1)[-1] in xml_kinds
                ):
                    if path == workbook and rel.get("Type", "").rsplit("/", 1)[-1] in (
                        "worksheet",
                        "chartsheet",
                    ):
                        continue  # Selected above by workbook sheet id, not Type alone.
                    pending.append(rel["Target"])
    # ZIP order is arbitrary: resolve references only after all shared-string parts.
    # If several parts declare an index, the largest length is the safe upper bound.
    check_text(
        inline_chars + sum(count * string_lengths.get(i, 0) for i, count in references.items())
    )
