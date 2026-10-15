"""XLSX filter: Excel workbooks (.xlsx, .xlsm, .xltx, .xltm), fully offline.

Only STRING cell *values* are translated. Numbers, booleans, dates, and
formulas (any string starting with "=") are left completely untouched; a
formula cell isn't even counted as a translatable unit, matching how other
filters skip non-text content while keeping unit indices aligned.

Cells whose string has no alphanumerics (pure ids/symbols) are skipped as
translatable units, exactly like json_filter._is_translatable.

UNITS are enumerated with openpyxl (read only), sheets in workbook order, rows
top-down, cells left-to-right; the skeleton carries (sheet index, row, col) per
unit. That order IS the segment numbering, so it never changes for a stored
project.

WALKS (skeleton["xlsx_walk"], #685): 1 (no stamp, every import before #685) =
the cells only; 2 = the cells, then the text-box and shape paragraphs of each
worksheet's drawing (DrawingML a:p, tokenized with the PowerPoint filter's run
tokens), numbered after the last cell so cell ids are the same under both walks.
Column mode (Batch 6) walks cells only.

EXPORT (#685) writes the translations straight into the ORIGINAL package
(xlsx_package): only the worksheets, shared strings, drawings and
docProps/core.xml that change are re-serialized; every other member (text
boxes, shapes, EMF/WMF pictures, charts, printer settings, data validation,
comments, the calc chain, a VBA project, the template / macro content type)
keeps its exact bytes. openpyxl is never asked to save the workbook: its
rebuild dropped every part it does not model. A Strict Open XML workbook comes
back Strict, as Word and PowerPoint files do.
"""

import io

from app.filters import ooxml_strict, xlsx_guard
from app.filters.base import (
    Filter,
    InputLimitError,
    ParseResult,
    TextUnit,
    check_segment_cap,
    register,
)

XLSX_WALK_VERSION = 2
_KNOWN_WALKS = (1, 2)
_DRAWING_STYLE = "text box"


def load_workbook(*args, **kwargs):
    """openpyxl.load_workbook, imported on first use.

    openpyxl (and the numpy it pulls in) costs about 0.5 s to import, and every app and
    CLI process imports this filter through the registry (test speed, 2026-10-08). The
    module-level name stays so callers and patch targets are unchanged."""
    from openpyxl import load_workbook as _load_workbook

    return _load_workbook(*args, **kwargs)


def walk_version(skeleton) -> int:
    """The unit walk a stored workbook was numbered by (module docstring). A stamp this
    build does not know is refused rather than guessed."""
    stamp = (skeleton or {}).get("xlsx_walk", 1)
    if type(stamp) is int and stamp in _KNOWN_WALKS:
        return stamp
    raise ValueError(
        f"This Excel file was imported with walk {stamp!r}, which this version does "
        "not know. Re-import the document to export it."
    )


def _is_translatable(s: str) -> bool:
    return any(c.isalnum() for c in s)


def _is_formula(value) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _open(data: bytes):
    """(workbook, normalized package bytes, strict). #669: a Strict Open XML workbook is
    read through its Transitional names; strict is (original, normalized) so a rebuild
    maps the parts it rewrote back to Strict."""
    xlsx_guard.check(data)
    normalized, is_strict = ooxml_strict.to_transitional(data, "XLSX file")
    try:
        wb = load_workbook(io.BytesIO(normalized), data_only=False)
    except Exception as exc:
        # Well-formed XML can still fail openpyxl's typed OOXML conversion.
        # Uploads and stored-file exports share this user-safe error boundary.
        raise InputLimitError(
            "This XLSX file has invalid spreadsheet XML. Save it again in Excel."
        ) from exc
    # #669: openpyxl reads a workbook part it does not understand as zero sheets. Refuse
    # it here, at upload, with a message, instead of reading it as empty.
    sheets = list(wb.worksheets) + list(wb.chartsheets)
    if not sheets or all(ws.sheet_state != "visible" for ws in sheets):
        raise InputLimitError(
            "This XLSX file has no visible sheet this version can read. Open it in Excel, "
            "save it as an Excel Workbook (.xlsx) and upload it again."
        )
    return wb, normalized, ((data, normalized) if is_strict else None)


def _load_checked(data: bytes):
    """The guarded openpyxl workbook, for reading (tabular_import uses it too)."""
    return _open(data)[0]


def _iter_cells(wb):
    """Yield (sheet_index, cell) for every translatable string cell, in the
    deterministic order shared by parse() and build()."""
    for sheet_idx, ws in enumerate(wb.worksheets):
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str) or _is_formula(value):
                    continue
                if not _is_translatable(value):
                    continue
                yield sheet_idx, cell


def _cell_str(value) -> str:
    return value if isinstance(value, str) else ("" if value is None else str(value))


def _iter_table(wb, table: dict, source_locale: str = "", targets=()):
    """Column mode (Batch 6): yield (sheet_idx, ws, row_cells, source_cell, plan) for
    every source cell the column mapping makes a unit, in sheet / row order. Each sheet
    is planned from its own first non-empty row; a stored plan (table["plans"]) is
    replayed so export walks exactly what ingest walked. A sheet the mapping does not
    fit (a named column it lacks) is left untouched."""
    from app.filters import csv_filter_common as tabular

    spec = table.get("spec") or {}
    stored = table.get("plans")
    for sheet_idx, ws in enumerate(wb.worksheets):
        plan = None
        header_seen = False
        for row in ws.iter_rows():
            values = [_cell_str(c.value) for c in row]
            if plan is None:
                if not any(v.strip() for v in values):
                    continue
                if stored is not None:
                    plan = stored.get(str(sheet_idx))
                else:
                    try:
                        plan = tabular.resolve(
                            spec, values, len(values), source_locale, list(targets)
                        )
                    except tabular.ColumnMappingError:
                        plan = None
                    table.setdefault("_plans", {})[str(sheet_idx)] = plan
                if plan is None or plan.get("inplace"):
                    break
            if plan["header"] and not header_seen:
                header_seen = True
                continue
            s = plan["source"]
            if s < len(row):
                cell = row[s]
                if isinstance(cell.value, str) and tabular.is_translatable(cell.value):
                    yield sheet_idx, ws, row, cell, plan


def _write_drawing_paragraph(p, text: str, n: int, tag, rtl) -> None:
    from app.filters.pptx_filter import _stamp_language, write_paragraph
    from app.filters.xml_text import XmlTargetError

    try:
        write_paragraph(p, text)
    except ValueError as exc:
        raise XmlTargetError(
            f"Excel text box unit {n + 1} could not be written. Review its target text."
        ) from exc
    if tag:
        _stamp_language(p, tag, rtl)


@register
class XlsxFilter(Filter):
    fmt = "xlsx"
    exts = ("xlsm", "xltx", "xltm")  # macro / template workbooks, content type kept
    takes_columns = True  # Batch 6: a column mapping switches on column mode
    takes_comments = True  # #690: cell comments + threaded comments, per-project option

    def _parse_table(self, data, columns, source_locale, target_locales) -> ParseResult:
        from app.filters import csv_filter_common as tabular

        wb = _load_checked(data)
        spec = columns if isinstance(columns, dict) else tabular.parse_spec(columns)
        if spec.get("delimiter"):
            raise tabular.ColumnMappingError("delimiter= applies to CSV/TSV files only.")
        table: dict = {"spec": spec}
        units: list[TextUnit] = []
        out: list[dict] = []
        for sheet_idx, _ws, row, cell, plan in _iter_table(
            wb, table, source_locale, target_locales or []
        ):
            idx = len(units)
            units.append(TextUnit(p_index=idx, text=cell.value))
            unit = {
                "p_index": idx,
                "text": cell.value,
                "style": "",
                "sheet": sheet_idx,
                "row": cell.row,
                "col": cell.column,
            }
            note = tabular.unit_note(plan, [_cell_str(c.value) for c in row])
            if note:
                unit["note"] = note
            out.append(unit)
            check_segment_cap(idx + 1)
        plans = table.get("_plans") or {}
        if not any(p and not p.get("inplace") for p in plans.values()):
            raise tabular.ColumnMappingError(
                "The column mapping matches no sheet in this workbook (check the source= column)."
            )
        skeleton = {
            "format": "xlsx",
            "units": out,
            "table": {"spec": spec, "plans": plans},
        }
        return ParseResult(skeleton=skeleton, units=units)

    def _rebuild_table(self, wb, pkg, paragraph_texts, target_locale, table, fallback_keys=None):
        from app.filters import csv_filter_common as tabular
        from app.filters.xml_text import xml_target_text

        for idx, (si, _ws, _row, cell, plan) in enumerate(_iter_table(wb, dict(table))):
            tr = paragraph_texts.get(str(idx))
            if tr is None:
                continue
            tr = xml_target_text(tr)
            tcol = tabular.target_col(plan, target_locale)
            if tcol is None:
                if tr != cell.value:
                    pkg.set_cell_text(si, cell.row, cell.column, tr)
                continue
            if fallback_keys is not None and str(idx) in fallback_keys:
                continue  # untranslated: never copy the source into the target column
            target = cell.offset(column=tcol + 1 - cell.column)
            existing = _cell_str(target.value)
            if tr == cell.value and existing.strip():
                continue  # untranslated: keep a target someone already filled
            if tr != existing:
                pkg.set_cell_text(si, target.row, target.column, tr, create=True)

    def parse(
        self,
        data: bytes,
        provider: str | None = None,
        columns: str | dict | None = None,
        source_locale: str = "",
        target_locales: list[str] | None = None,
        comments: bool = False,
    ) -> ParseResult:
        if columns:
            # column mode translates cells into target columns; comments are not offered
            return self._parse_table(data, columns, source_locale, target_locales)
        from app.filters import xlsx_package
        from app.filters.pptx_filter import tokenize

        wb, normalized, _strict = _open(data)
        units: list[TextUnit] = []
        rows: list[dict] = []
        for idx, (sheet_idx, cell) in enumerate(_iter_cells(wb)):
            units.append(TextUnit(p_index=idx, text=cell.value))
            rows.append({"sheet": sheet_idx, "row": cell.row, "col": cell.column})
            check_segment_cap(idx + 1)  # #571 gate: stop at the per-document cap
        # walk 2 (#685): text boxes and shapes, numbered after the last cell
        pkg = xlsx_package.Package(normalized)
        for sheet_idx, part, p in pkg.drawing_paragraphs():
            text, inline_map = tokenize(p)
            idx = len(units)
            units.append(
                TextUnit(p_index=idx, text=text, style=_DRAWING_STYLE, inline_map=inline_map)
            )
            rows.append({"sheet": sheet_idx, "part": part})
            check_segment_cap(idx + 1)
        skeleton = {
            "format": "xlsx",
            "xlsx_walk": XLSX_WALK_VERSION,
            "units": [
                {"p_index": u.p_index, "text": u.text, "style": u.style, **rows[i]}
                for i, u in enumerate(units)
            ],
        }
        if comments:
            # #690: comments and threaded comments, appended at office_comments indices
            from app.filters import office_comments

            for p_index, text, meta in office_comments.xlsx_comment_slots(data):
                units.append(TextUnit(p_index=p_index, text=text, style="comment"))
                skeleton["units"].append(
                    {"p_index": p_index, "text": text, "style": "comment", **meta}
                )
                check_segment_cap(len(units))
            skeleton["comments"] = True
        return ParseResult(skeleton=skeleton, units=units)

    def _rebuild(
        self,
        original: bytes,
        paragraph_texts: dict,
        target_locale=None,
        table=None,
        fallback_keys=None,
        walk: int = XLSX_WALK_VERSION,
    ) -> bytes:
        from app.filters import a11y, xlsx_package
        from app.filters.xml_text import xml_target_text

        if walk not in _KNOWN_WALKS:
            raise ValueError(f"Unknown Excel walk {walk!r}.")
        wb, normalized, strict = _open(original)
        pkg = xlsx_package.Package(normalized)
        tag = a11y.target_tag(target_locale) if target_locale else None
        if table and table.get("plans") is not None:
            self._rebuild_table(wb, pkg, paragraph_texts, target_locale, table, fallback_keys)
        else:
            cells = list(_iter_cells(wb))
            for idx, (sheet_idx, cell) in enumerate(cells):
                tr = paragraph_texts.get(str(idx))
                if tr is None:
                    continue
                tr = xml_target_text(tr)
                if tr != cell.value:  # an unchanged cell keeps its exact XML
                    pkg.set_cell_text(sheet_idx, cell.row, cell.column, tr)
            if walk >= 2:
                rtl = bool(tag) and a11y.target_direction(target_locale) == "rtl"
                base = len(cells)
                for j, (_si, part, p) in enumerate(pkg.drawing_paragraphs()):
                    tr = paragraph_texts.get(str(base + j))
                    if tr is None:
                        continue
                    _write_drawing_paragraph(p, xml_target_text(str(tr)), base + j, tag, rtl)
                    pkg.touch(part)
        if tag:
            pkg.stamp_language(tag)
            if a11y.target_direction(tag) == "rtl":
                set_sheets_rtl(pkg)
        out = pkg.save()
        if strict is not None:  # #669: Strict in, Strict out
            out = ooxml_strict.to_strict(strict[0], strict[1], out)
        return out

    def build(self, original: bytes, paragraph_texts: dict) -> bytes:
        return self._comments(self._rebuild(original, paragraph_texts), original, paragraph_texts)

    @staticmethod
    def _comments(built: bytes, original: bytes, paragraph_texts: dict) -> bytes:
        """#690 export hook, kept apart from the cell rebuild: write comment translations
        (keys at office_comments.COMMENT_BASE and up) into the finished package. The
        package-patching export (#685) keeps every comment part byte-identical, so this
        patch finds notes by sheet + cell and threaded comments by id. A no-op (the
        package returned as is) for a document imported without comments."""
        from app.filters import office_comments

        return office_comments.xlsx_apply(built, original, paragraph_texts)

    def build_localized(
        self,
        original: bytes,
        paragraph_texts: dict,
        target_locale: str | None = None,
        table: dict | None = None,
        fallback_keys: set | None = None,
        walk: int = XLSX_WALK_VERSION,
    ) -> bytes:
        """build() plus the TARGET language in docProps/core.xml (dc:language), the only
        workbook-level language declaration SpreadsheetML has, and on the runs of every
        translated text-box paragraph (a:rPr lang). A right-to-left target also turns
        every worksheet right to left (set_sheets_rtl, decision 17). Comment translations
        (#690) are patched in after, except in column mode (no comment units there)."""
        out = self._rebuild(
            original,
            paragraph_texts,
            target_locale or None,
            table=table,
            fallback_keys=fallback_keys,
            walk=walk,
        )
        if table and table.get("plans") is not None:
            return out
        return self._comments(out, original, paragraph_texts)


def set_sheets_rtl(pkg) -> None:
    """Right-to-left target (B-rtl, decision 17): every worksheet view gets
    rightToLeft="1" (ECMA-376 Part 1, 18.3.1.87 sheetView), so column A sits on the right
    and the sheet reads right to left in Excel and LibreOffice. Chart sheets have no
    rightToLeft and are left alone. Since #685 the export patches the original package,
    so the attribute is set on each worksheet part's sheetView elements."""
    pkg.set_sheets_rtl()
